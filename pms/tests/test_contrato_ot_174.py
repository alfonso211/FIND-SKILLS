"""1.74 · OT sin foto obligatoria desde recepción, tipo «apto. en revisión», contrato turístico con motivo,
«¿Dónde nos conoció?» y domicilio en la calle de los apartamentos (cliente antiguo), documento «empadronados» e
informe de dirección. Datos ZZ-PRUEBA."""
import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select

from app import automatismos, empadronados
from app.database import SessionLocal
from app.models import Asset
from test_pms import d

Z = "ZZ-PRUEBA-174"


def _unidad(client, admin, asset_id, n):
    us = [u for u in client.get(f"/api/unidades?asset_id={asset_id}", headers=admin).json() if u.get("uso") != "garaje"]
    return us[-n]


def _reserva(client, admin, asset_id, n, entrada, salida, doc, direccion=None, municipio="Madrid"):
    u = _unidad(client, admin, asset_id, n)
    r = client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": u["id"], "localizador": f"{Z}-{doc}", "fecha_entrada": d(entrada), "fecha_salida": d(salida),
        "importe_total": 300, "guest": {"nombre": "Prueba", "apellidos": Z, "documento_tipo": "PAS",
                                        "documento_num": doc, "nacionalidad": "Española", "direccion": direccion,
                                        "municipio": municipio}})
    assert r.status_code == 201, r.text
    return r.json()


def test_calle_de_los_apartamentos():
    assert empadronados.calle("C/ de Alcalá, 35 3ºB") == "alcala"
    assert empadronados.calle("Avda. de la Constitución nº 12") == "constitucion"
    a = Asset(nombre="X", direccion="Calle Campezo 2", municipio="Madrid")
    assert empadronados.misma_calle("C/ Campezo, 8 - 2ºA", "Madrid", a)
    assert empadronados.misma_calle("calle campezo 15", None, a)
    assert not empadronados.misma_calle("Calle Campezo 2", "Sevilla", a)  # otra ciudad
    assert not empadronados.misma_calle("Calle Mayor 2", "Madrid", a)
    assert not empadronados.misma_calle(None, "Madrid", a)


def test_ot_sin_foto_desde_recepcion_y_tipo_revision(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    w = client.post("/api/mantenimiento/ordenes", headers=admin, json={
        "asset_id": sfl, "titulo": f"{Z} grifo", "tipo": "correctivo"}).json()
    assert w["exige_foto_antes"] is False and w["falta_foto_antes"] is False
    rev = client.post("/api/mantenimiento/ordenes", headers=admin, json={
        "asset_id": sfl, "titulo": f"{Z} revisión", "tipo": "revision"})
    assert rev.status_code in (200, 201) and rev.json()["tipo"] == "revision"
    lista = client.get("/api/mantenimiento/ordenes", headers=admin, params={"id": rev.json()["id"]}).json()
    assert [x["tipo"] for x in lista] == ["revision"]


def test_contrato_motivo_conocio_y_domicilio(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    r = _reserva(client, admin, sfl, 1, 300, 303, "ZZ174A")
    url = f"/api/turistico/reservas/{r['id']}/contrato"
    info = client.get(url, headers=admin).json()
    assert info["conocio"]["booking"] == "Booking" and "otro" in info["conocio"]
    base = {**info["datos"], "permitir_huecos": True, "sin_garaje": True, "acreditacion": ["dni"],
            "cliente_domicilio": "Calle Mayor 3", "cliente_municipio": "Toledo"}

    def imprimir(**k):
        return client.post(url, headers=admin, json={**base, **k})

    # sin motivo no se continúa, ni con «Imprimir aunque falten datos»
    res = imprimir(motivo=[], conocio="booking")
    assert res.status_code == 400 and "Motivo de la estancia" in res.json()["detail"]
    res = imprimir(motivo=["otro"], conocio="booking")
    assert res.status_code == 400 and "indique cuál" in res.json()["detail"]
    # ¿Dónde nos conoció? obligatorio; «Otro» con su detalle
    res = imprimir(motivo=["turismo"])
    assert res.status_code == 400 and "¿Dónde nos conoció?" in res.json()["detail"]
    res = imprimir(motivo=["turismo"], conocio="otro")
    assert res.status_code == 400 and "«Otro»: indique cuál" in res.json()["detail"]
    # «solo guardar» guarda y dice lo que falta
    g = client.post(url, headers=admin, json={**base, "solo_guardar": True, "motivo": []}).json()
    assert g["guardado"] and any("Motivo" in f for f in g["faltan"]) and any("conoció" in f for f in g["faltan"])

    # domicilio en la calle de los apartamentos: solo continúa si es cliente antiguo con su fecha
    calle = {"motivo": ["laboral"], "conocio": "otro", "conocio_otro": "feria", "cliente_domicilio": "C/ Campezo 8, 3ºA",
             "cliente_municipio": "Madrid"}
    res = imprimir(**calle)
    assert res.status_code == 400 and "calle de los apartamentos" in res.json()["detail"]
    res = imprimir(**calle, cliente_antiguo=True)
    assert res.status_code == 400 and "cliente antiguo" in res.json()["detail"]
    res = imprimir(**calle, cliente_antiguo=True, cliente_antiguo_desde=d(5))
    assert res.status_code == 400 and "no puede ser futura" in res.json()["detail"]
    res = imprimir(**calle, cliente_antiguo=True, cliente_antiguo_desde="2019-03-01")
    assert res.status_code == 200, res.text
    guardado = client.get(url, headers=admin).json()["datos"]
    assert guardado["cliente_antiguo"] is True and guardado["cliente_antiguo_desde"] == "2019-03-01"

    # informe de dirección: contratos del mes por «¿Dónde nos conoció?» (una vez por reserva)
    r2 = _reserva(client, admin, sfl, 2, 310, 312, "ZZ174B")
    url2 = f"/api/turistico/reservas/{r2['id']}/contrato"
    b2 = {**client.get(url2, headers=admin).json()["datos"], "permitir_huecos": True, "sin_garaje": True,
          "acreditacion": ["dni"], "motivo": ["turismo"], "conocio": "booking", "cliente_domicilio": "Calle Real 1",
          "cliente_municipio": "Segovia"}
    assert client.post(url2, headers=admin, json=b2).status_code == 200
    assert client.post(url2, headers=admin, json=b2).status_code == 200  # reimpresión: no suma
    with SessionLocal() as db:
        a = db.get(Asset, sfl)
        o = automatismos.donde_nos_conocio(db, a, date.today())
        assert o["por_origen"].get("Booking", 0) >= 1 and o["por_origen"].get("Otro: feria", 0) == 1
        assert o["contratos"] >= 2


def test_documento_empadronados(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    _reserva(client, admin, sae, 1, -1, 3, "ZZ174C", "Calle Campezo 2, 1ºB")
    _reserva(client, admin, sae, 2, -1, 3, "ZZ174D", "Avenida de América 5")
    with SessionLocal() as db:
        datos = empadronados.datos(db)
    sae_res = next(x for x in datos["resumen"] if x["calle"] == "Calle Campezo 8")
    assert sae_res["personas"] >= 1 and sae_res["apartamentos"] >= 1
    assert any(f["reserva"] == f"{Z}-ZZ174C" for f in datos["filas"])
    assert not any(f["reserva"] == f"{Z}-ZZ174D" for f in datos["filas"])
    res = client.get("/api/turistico/empadronados.xlsx", headers=admin)
    assert res.status_code == 200 and "empadronados_" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Resumen", "Detalle"]
    assert wb["Resumen"]["A3"].value == "Activo"
    with SessionLocal() as db:
        assert db.scalar(select(Asset.id).where(Asset.codigo == "SAE")) == sae
