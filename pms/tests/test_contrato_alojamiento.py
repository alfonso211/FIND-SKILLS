import io


import docx

from conftest import login  # noqa: F401
from test_pms import _new_user, d


def _texto(contenido: bytes) -> str:
    doc = docx.Document(io.BytesIO(contenido))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        vistas = set()
        for fila in t.rows:
            for c in fila.cells:
                if c._tc not in vistas:
                    vistas.add(c._tc)
                    partes.append(c.text)
    return "\n".join(partes)


def _reserva(client, admin, asset_id, codigo, entrada, salida):
    u = client.get(f"/api/unidades?asset_id={asset_id}&q={codigo}", headers=admin).json()[0]
    r = client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": u["id"], "localizador": f"LOC-{codigo}", "fecha_entrada": d(entrada), "fecha_salida": d(salida),
        "importe_total": 1234.5, "guest": {"nombre": "Lucía", "apellidos": "Martín Sanz", "documento_tipo": "DNI",
                                           "documento_num": "12345678Z", "nacionalidad": "Española"}})
    assert r.status_code == 201, r.text
    return r.json()


def test_contrato_suite_florida(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    client.put(f"/api/activos/{sfl}", headers=admin, json={
        "contrato_representante": "Recepción Suite Florida", "contrato_representante_dni": "00000000T",
        "contrato_email": "info@apartamentossuitesflorida.es"})
    r = _reserva(client, admin, sfl, "P1-1A", 40, 43)
    info = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()
    assert info["datos"]["representante"] == "Recepción Suite Florida"
    assert info["datos"]["cliente_documento"] == "12345678Z" and info["historial"] == []

    datos = {**info["datos"], "cliente_domicilio": "C/ Mayor 1", "cliente_cp": "41001", "cliente_municipio": "Sevilla",
             "cliente_pais": "España", "fianza": 300, "tarjeta_titular": "Lucía Martín", "tarjeta_terminacion": "4242",
             "tarjeta_caducidad": "08/28", "motivo": ["laboral", "transito"], "acreditacion": "dni",
             "garaje_sotano": "1", "garaje_plaza": "15"}
    # faltan capacidad, dormitorios, correo y móvil: no se imprime y se dice qué falta
    res = client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json=datos)
    assert res.status_code == 400
    for falta in ("Capacidad máxima", "Correo del cliente", "Móvil del cliente"):  # dormitorios: de la tipología
        assert falta in res.json()["detail"]
    assert "Motivo" not in res.json()["detail"]
    # ...pero lo tecleado queda guardado en la reserva
    guardado = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()["datos"]
    assert guardado["fianza"] == 300 and guardado["motivo"] == ["laboral", "transito"]
    assert guardado["acreditacion"] == ["dni"]
    datos.update(capacidad=2, dormitorios=1, cliente_email="lucia@example.com", cliente_movil="600000000")
    res = client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json=datos)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "LOC-P1-1A" in res.headers["content-disposition"]
    t = _texto(res.content)
    for esperado in ("“Apartamentos Suites Florida”", "nº AM 265", "D./Dª Recepción Suite Florida",
                     "D./Dª Lucía Martín Sanz", "nacionalidad Española", "nº 12345678Z", "en C/ Mayor 1 (C.P. 41001",
                     "Portal 1 Planta 1 Nº A", "Nº de noches: 3", "Precio total: 1.234,50 €", "Fianza: 300,00 €",
                     "terminación nº 4242 · caducidad 08 / 28", "☒ desplazamiento laboral temporal",
                     "☒ tránsito aeroportuario", "☐ turismo / ocio", "☒ DNI o equivalente",
                     "sótano -1 plaza nº 15", "Capacidad máxima: 2 plazas", "Móvil: 600000000",
                     ):
        assert esperado in t, esperado
    # «otro» sin marcar: sin puntos ni dos puntos colgando (motivo y acreditación)
    assert (t + "\n").count("☐ otro\n") == 2 and "otro:" not in t
    # el cliente solo firma: no queda ningún hueco con puntos
    assert "…" not in t
    assert res.headers["x-huecos-pendientes"] == "0"
    # la unidad aprende su capacidad y dormitorios
    u = client.get(f"/api/unidades?asset_id={sfl}&q=P1-1A", headers=admin).json()[0]
    assert (u["capacidad"], u["dormitorios"]) == (2, 1)
    # la ficha del huésped se completa con lo tecleado
    g = client.get("/api/terceros?tipo=huesped&q=12345678Z", headers=admin).json()[0]
    assert (g["direccion"], g["cp"], g["municipio"], g["pais"]) == ("C/ Mayor 1", "41001", "Sevilla", "España")
    # historial, reimpresión idéntica y datos recordados para la siguiente vez
    info2 = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()
    assert len(info2["historial"]) == 1 and info2["datos"]["fianza"] == 300
    rep = client.get(f"/api/turistico/reservas/{r['id']}/contrato/{info2['historial'][0]['id']}", headers=admin)
    assert _texto(rep.content) == t


def test_contrato_suite_aeropuerto_y_seguridad(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    r = _reserva(client, admin, sae, "B-101", 44, 45)
    base = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()["datos"]
    # al hacer la reserva se puede solo guardar, con lo que falte
    g = client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin,
                    json={**base, "solo_guardar": True, "motivo": ["otro"], "motivo_otro": "congreso médico"})
    assert g.status_code == 200 and g.json()["guardado"] and "Fianza" in g.json()["faltan"]
    res = client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin,
                      json={**base, "permitir_huecos": True, "sin_garaje": True, "motivo": ["otro"],
                            "motivo_otro": "congreso médico"})
    assert res.status_code == 200, res.text
    t = _texto(res.content)
    assert "Garaje: No incluido" in t and "☒ otro: congreso médico" in t
    assert "“Apartamentos Suites Aeropuerto”" in t and "nº AM 259" in t and "Bloque B Planta 1 Nº 101" in t
    # los huecos sin dato se dejan con puntos para rellenar a mano
    assert "Fianza: ……" in t
    # nunca se admite el número completo de la tarjeta
    bad = client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin,
                      json={**base, "tarjeta_terminacion": "4111111111111111"})
    assert bad.status_code == 422
    # solo quien gestiona reservas de ese activo
    lim = _new_user(client, admin, "lim.contrato@inversiete.es",
                    [{"role_id": ids["roles"]["Gobernanta / Limpieza"], "asset_id": sae}])
    assert client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=lim, json=base).status_code == 403
    rec_sf = _new_user(client, admin, "rec.sf.contrato@inversiete.es",
                       [{"role_id": ids["roles"]["Recepción"], "asset_id": ids["assets"]["SFL"]["id"]}])
    assert client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=rec_sf).status_code == 403
