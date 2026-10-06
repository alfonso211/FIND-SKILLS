"""Facturas emitidas sin cobrar (renovaciones y reservas que pagan por transferencia): no se factura dos veces,
recepción las ve cada día (panel y resumen por correo), las marca cobradas y queda el control económico."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app import avisos
from app.database import SessionLocal
from conftest import domicilio_fiscal
from test_adjuntos_ot import _usuario

HOY = date.today()
COMPLETO = {"nombre": "Paula", "apellidos": "Transferencia Pendiente", "documento_tipo": "DNI",
            "documento_num": "71234568S", "num_soporte": "BAA555556", "nacionalidad": "España",
            "fecha_nacimiento": "1982-03-03", "sexo": "F", "telefono": "611555446", "email": "paula@example.com",
            "direccion": "C/ Luna 4", "cp": "28013", "municipio": "Madrid", "pais": "España"}


def d(n):
    return (HOY + timedelta(days=n)).isoformat()


def test_facturas_pendientes_de_cobro(client, admin, ids):
    domicilio_fiscal(client, admin)
    sae = ids["assets"]["SAE"]["id"]
    apto = [u for u in client.get(f"/api/unidades?asset_id={sae}", headers=admin).json() if u["uso"] != "garaje"][-9]
    r = client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": apto["id"], "localizador": "TR-PEND-1", "fecha_entrada": d(600), "fecha_salida": d(630),
        "importe_total": 1000, "facturar_pendiente": True, "guest": {"nombre": "Paula"}})
    assert r.status_code == 201, r.text
    r = r.json()
    assert r["factura"]["cobro"] == "pendiente" and r["importe_pagado"] == 0
    fid = r["factura"]["id"]

    # la estancia ya está facturada: ni otro cobro con factura ni otra factura sin cobrar
    cobro = client.post(f"/api/turistico/reservas/{r['id']}/cobro", headers=admin, json={"importe": 1000})
    assert cobro.status_code == 400 and "pendiente de cobro" in cobro.json()["detail"]
    assert client.post(f"/api/turistico/reservas/{r['id']}/facturar", headers=admin, json={}).status_code == 400

    # recepción la ve: listado, panel y resumen diario
    pend = client.get("/api/facturas/pendientes-cobro", headers=admin).json()
    fila = next(f for f in pend["facturas"] if f["id"] == fid)
    assert fila["dias"] == 0 and fila["localizador"] == "TR-PEND-1" and fila["total"] == 1000
    panel = next(a for a in client.get("/api/panel", headers=admin).json()["activos"] if a["id"] == sae)
    assert panel["facturas_pendientes"] >= 1 and panel["facturas_pendientes_importe"] >= 1000
    with SessionLocal() as db:
        datos = avisos._datos_resumen(db, HOY)
    assert any(f[1] == r["factura"]["codigo"] for _, f in datos["facturas_pendientes"])
    assert [x["id"] for x in client.get("/api/facturas", headers=admin, params={"cobro": "pendiente"}).json()
            if x["id"] == fid] == [fid]

    # otro activo no la toca; recepción del activo la marca cobrada (fecha no futura)
    otra = _usuario(client, admin, ids, "recepcion.cobros.sfl@inversiete.com", "Recepción", "SFL")
    assert client.post(f"/api/facturas/{fid}/cobro", headers=otra, json={"fecha": d(0)}).status_code in (403, 404)
    rec = _usuario(client, admin, ids, "recepcion.cobros.sae@inversiete.com", "Recepción", "SAE")
    assert client.post(f"/api/facturas/{fid}/cobro", headers=rec, json={"fecha": d(1)}).status_code == 422
    ok = client.post(f"/api/facturas/{fid}/cobro", headers=rec,
                     json={"fecha": d(0), "forma_pago": "transferencia", "referencia": "TRF 0001"})
    assert ok.status_code == 200, ok.text
    assert ok.json()["cobro"] == "cobrada" and ok.json()["cobro_ref"] == "TRF 0001"
    assert client.post(f"/api/facturas/{fid}/cobro", headers=rec, json={"fecha": d(0)}).status_code == 400
    res = client.get(f"/api/turistico/reservas?asset_id={sae}&q=TR-PEND-1", headers=admin).json()[0]
    assert res["importe_pagado"] == 1000
    assert fid not in [f["id"] for f in client.get("/api/facturas/pendientes-cobro", headers=admin).json()["facturas"]]
    lista = client.get("/api/facturas", headers=admin, params={"q": r["factura"]["codigo"]}).json()[0]
    assert lista["cobro_usuario"] == "recepcion.cobros.sae"

    # un cobro marcado por error se deshace (solo administración)
    assert client.post(f"/api/facturas/{fid}/cobro/deshacer", headers=rec, json={"motivo": "error"}).status_code == 403
    assert client.post(f"/api/facturas/{fid}/cobro/deshacer", headers=admin,
                       json={"motivo": "marcada por error"}).json()["cobro"] == "pendiente"
    assert client.get(f"/api/turistico/reservas?asset_id={sae}&q=TR-PEND-1", headers=admin).json()[0][
        "importe_pagado"] == 0
    client.post(f"/api/facturas/{fid}/cobro", headers=rec, json={"fecha": d(0)})

    # renovación: una parte cobrada al renovar y el resto facturado sin cobrar
    guest = client.get(f"/api/terceros/{r['guest_id']}", headers=admin).json()
    client.put(f"/api/terceros/{r['guest_id']}", headers=admin,
               json={**COMPLETO, "company_id": guest["company_id"], "tipo": "huesped"})
    ren = client.post(f"/api/turistico/reservas/{r['id']}/renovar", headers=admin, json={
        "fecha_salida": d(660), "importe_total": 600, "importe_pagado": 200, "forma_pago": "tarjeta",
        "facturar_pendiente": True})
    assert ren.status_code == 201, ren.text
    ren = ren.json()
    assert [f["cobro"] for f in ren["facturas"]] == ["cobrada", "pendiente"] and ren["importe_pagado"] == 200
    pendiente = ren["facturas"][1]
    assert client.get(f"/api/facturas/{pendiente['id']}", headers=admin).json()["total"] == 400

    # rectificar una factura pendiente: deja de estar pendiente y no toca lo cobrado; se puede volver a facturar
    rect = client.post(f"/api/facturas/{pendiente['id']}/rectificar", headers=admin, json={"motivo": "importe mal"})
    assert rect.status_code == 201, rect.text
    assert client.get(f"/api/facturas/{pendiente['id']}", headers=admin).json()["cobro"] == "anulada"
    assert client.get(f"/api/turistico/reservas?asset_id={sae}&q={ren['localizador']}", headers=admin).json()[0][
        "importe_pagado"] == 200
    nueva = client.post(f"/api/turistico/reservas/{ren['id']}/facturar", headers=admin, json={})
    assert nueva.status_code == 200 and client.get(f"/api/facturas/{nueva.json()['factura']['id']}",
                                                   headers=admin).json()["total"] == 400

    # informe de control económico: pendientes con sus días y cobradas en el periodo
    xl = client.get("/api/informes/cobros", headers=admin, params={"desde": d(-30), "hasta": d(0), "asset_id": sae})
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    pend_xl = [[c.value for c in row] for row in wb["Pendientes de cobro"].iter_rows(min_row=2)]
    assert any(f[1] == nueva.json()["factura"]["codigo"] for f in pend_xl)
    cobradas = [[c.value for c in row] for row in wb["Cobradas en el periodo"].iter_rows(min_row=2)]
    fila = next(f for f in cobradas if f[1] == r["factura"]["codigo"])
    assert fila[6] == 0 and fila[9] == "recepcion.cobros.sae"
