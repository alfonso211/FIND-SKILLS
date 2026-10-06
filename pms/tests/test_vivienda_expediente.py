"""Expediente del contrato de vivienda (LAU): checklist, contrato Word rellenado, cobros con recibo, carpeta y agenda."""
import io

import docx

from app.database import SessionLocal
from app.models import AgendaEvent
from test_adjuntos_ot import _usuario

PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


def _preparar(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    comp = next(c for c in client.get("/api/sociedades", headers=admin).json() if c["id"] == ids["assets"]["BAB35"]["company_id"])
    r = client.put(f"/api/sociedades/{comp['id']}", headers=admin, json={
        **{k: comp[k] for k in ("nombre", "cif", "parent_id", "activa")},
        "contratos": {"rm_tomo": "1", "rm_folio": "2", "rm_hoja": "M-3", "representante": "Representante Ficticio",
                      "representante_dni": "00000001R", "representante_cargo": "Apoderado",
                      "representante_poder": "escritura de prueba", "iban": "ES00 0000 0000 0000 0000 0000",
                      "email_notificaciones": "alquileres@example.com", "telefono_averias": "600000000"}})
    assert r.status_code == 200, r.text
    u = client.post("/api/unidades", headers=admin, json={"asset_id": bab, "codigo": "BAB35-9Z", "uso": "vivienda",
                                                        "planta": "9", "superficie_m2": 80,
                                                        "ref_catastral": "0000000AA0000A0001AA"}).json()
    inv = client.get(f"/api/unidades/{u['id']}/ficha-alquiler", headers=admin).json()["inventario"]
    inv[0]["entrega"] = False  # el frigorífico no se entrega: no debe salir en el Anexo I
    assert client.put(f"/api/unidades/{u['id']}/ficha-alquiler", headers=admin, json={
        "ficha": {"cee_letra": "Z"}}).status_code == 400
    r = client.put(f"/api/unidades/{u['id']}/ficha-alquiler", headers=admin, json={
        "ficha": {"puerta": "Z", "cp": "28001", "superficie_util": 70, "distribucion": "2 dormitorios",
                  "registro_propiedad": "5", "finca": "999", "cee_letra": "E", "cee_registro": "R-9",
                  "cee_vigencia": "2034-01-01", "ibi_anual": 480, "tasa_residuos_anual": 120,
                  "llaves": [{"elemento": "Llave portal", "uds": 2, "obs": ""}]},
        "inventario": inv})
    assert r.status_code == 200, r.text
    lease = client.post("/api/alquiler/contratos", headers=admin, json={
        "unit_id": u["id"], "tenant": {"nombre": "Inquilina", "apellidos": "De Prueba", "documento_num": "00000002W",
                                       "nacionalidad": "española", "telefono": "611000000", "email": "i@example.com",
                                       "direccion": "Calle Inventada 1", "cp": "28002", "municipio": "Madrid"},
        "fecha_inicio": "2031-03-01", "fecha_fin": "2038-02-28", "renta_mensual": 1000, "fianza": 1000,
        "garantia_adicional": 2000, "estado": "borrador"})
    assert lease.status_code == 201, lease.text
    return lease.json()["id"]


def test_expediente_completo(client, admin, ids):
    lid = _preparar(client, admin, ids)
    url = f"/api/alquiler/contratos/{lid}/expediente"
    x = client.get(url, headers=admin).json()
    assert x["referencia"] == "BAB35-9Z-2031"
    ck = {c["clave"]: c for c in x["checklist"]}
    assert ck["cee"]["estado"] == "hecho" and ck["ibi_tasa"]["estado"] == "hecho"  # comprobados con la ficha
    assert ck["aval"]["estado"] == "no_aplica" and ck["cobros"]["estado"] == "pendiente"
    assert "Fecha de firma" in x["faltan"]
    assert not x["inventario"][0]["entrega"]

    # datos del contrato; nunca efectivo para la renta; garantía mediante aval
    assert client.put(f"{url}/datos", headers=admin, json={"forma_pago": "efectivo"}).status_code == 400
    assert client.put(f"{url}/datos", headers=admin, json={"inventado": 1}).status_code == 400
    x = client.put(f"{url}/datos", headers=admin, json={
        "fecha_firma": "2031-02-20", "hora_firma": "10:30", "fecha_entrega": "2031-03-01", "garantia_modalidad": "aval",
        "aval_vencimiento": "2033-02-20", "forma_pago": "transferencia", "importe_inicial": 1050,
        "periodo_desde": "2031-03-01", "periodo_hasta": "2031-03-31", "max_ocupantes": 3, "paginas": 14}).json()
    assert x["faltan"] == [], x["faltan"]
    ck = {c["clave"]: c for c in x["checklist"]}
    assert ck["aval"]["estado"] == "pendiente" and ck["deposito_fianza"]["limite"] == "2031-03-22"

    # cobros: reglas de efectivo, de lo pactado y de la modalidad de garantía
    def entrega(**kw):
        return client.post(f"{url}/entregas", headers=admin, json={"fecha": "2031-02-20", **kw})
    assert entrega(concepto="renta_inicial", importe=1050, forma="efectivo").status_code == 400
    assert entrega(concepto="fianza", importe=1000, forma="efectivo").status_code == 400  # Ley 7/2012
    assert entrega(concepto="fianza", importe=1200, forma="transferencia").status_code == 400  # supera lo pactado
    assert entrega(concepto="garantia", importe=2000, forma="transferencia").status_code == 400  # es con aval
    assert entrega(concepto="fianza", importe=999, forma="efectivo").status_code == 201
    assert entrega(concepto="fianza", importe=1, forma="efectivo").status_code == 400  # suma 1.000 € en efectivo
    assert entrega(concepto="fianza", importe=1, forma="transferencia", referencia="TR-1").status_code == 201
    assert entrega(concepto="renta_inicial", importe=1050, forma="cheque", referencia="123", entidad="Banco X",
                   periodo="marzo 2031").status_code == 201
    r = entrega(concepto="garantia", importe=2000, forma="aval", referencia="AV-9", entidad="Banco Y",
                vencimiento="2033-02-20")
    assert r.status_code == 201, r.text
    x = r.json()
    assert x["pendiente_cobro"] == [] and [e["numero"] for e in x["entregas"]][-1] == "BAB35-9Z-2031-R04"
    rec = client.get(f"{url}/entregas/4/recibo", headers=admin)
    assert rec.status_code == 200 and rec.content[:5] == b"%PDF-"
    assert client.post(f"{url}/entregas/1/anular", headers=admin, json={"motivo": "error de prueba"}).json()[
        "pendiente_cobro"] == ["fianza"]

    # contrato Word: sin hoja de control, con los datos y solo el inventario que se entrega
    r = client.get(f"{url}/contrato.docx", headers=admin)
    assert r.status_code == 200, r.text
    d = docx.Document(io.BytesIO(r.content))
    texto = "\n".join(p.text for p in d.paragraphs)
    assert "HOJA DE CONTROL" not in texto.upper()
    assert "Ref.: BAB35-9Z-2031" in texto and "Inquilina De Prueba" in texto and "Representante Ficticio" in texto
    assert "(mil euros)" in texto and "constan de 14 páginas" in texto and "☒ aval bancario" in texto
    inv = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Estancia")
    assert "Frigorífico" not in [r.cells[1].text for r in inv.rows]

    # carpeta: solo PDF, imágenes o Word; se ve y se borra
    assert client.post(f"{url}/documentos", headers=admin, data={"clave": "contrato_firmado"},
                       files={"fichero": ("x.exe", b"MZ\x90\x00", "application/octet-stream")}).status_code == 400
    x = client.post(f"{url}/documentos", headers=admin, data={"clave": "contrato_firmado"},
                    files={"fichero": ("firmado.pdf", PDF, "application/pdf")}).json()
    doc = x["documentos"][0]
    assert {c["clave"]: c for c in x["checklist"]}["contrato_firmado"]["estado"] == "hecho"
    assert client.get(f"{url}/documentos/{doc['id']}", headers=admin).content == PDF

    # agenda de quien prepara el contrato; al marcar el checklist, la tarea queda hecha
    x = client.post(f"{url}/agenda", headers=admin).json()
    claves = {a["clave"] for a in x["agenda"]}
    assert {"firma", "entrega", "deposito_fianza", "actualizacion", "preaviso", "aval_vencimiento"} <= claves
    assert {c["clave"]: c for c in x["checklist"]}["avisos"]["estado"] == "hecho"
    client.put(f"{url}/checklist/deposito_fianza", headers=admin, json={"estado": "hecho"})
    eid = next(a["id"] for a in x["agenda"] if a["clave"] == "deposito_fianza")
    with SessionLocal() as db:
        e = db.get(AgendaEvent, eid)
        assert e.hecha and e.visibilidad == "privada" and e.inicio.date().isoformat() == "2031-03-17"
    # repetir no duplica
    assert len(client.post(f"{url}/agenda", headers=admin).json()["agenda"]) == len(x["agenda"])

    # otro activo: sin acceso al expediente
    otro = _usuario(client, admin, ids, "recepcion.expediente@inversiete.com", "Recepción", "SAE")
    assert client.get(url, headers=otro).status_code == 403
    assert client.delete(f"{url}/documentos/{doc['id']}", headers=otro).status_code == 403
    assert client.delete(f"{url}/documentos/{doc['id']}", headers=admin).json()["documentos"] == []


def test_importe_en_letra():
    from app.contrato_vivienda import en_letra
    assert en_letra(1) == "un euro"
    assert en_letra(21) == "veintiún euros"
    assert en_letra(21001) == "veintiún mil un euros"
    assert en_letra(1250.5) == "mil doscientos cincuenta euros con cincuenta céntimos"
    assert en_letra(100) == "cien euros" and en_letra(115) == "ciento quince euros"
