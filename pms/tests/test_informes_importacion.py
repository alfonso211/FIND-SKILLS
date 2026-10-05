"""Informes en Excel, importación de reservas desde Excel/CSV y resumen diario por correo."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import avisos
from conftest import domicilio_fiscal, login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _libro(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX
    return load_workbook(BytesIO(resp.content))


def _filas(ws):
    cab = [c.value for c in ws[1]]
    out = []
    for f in ws.iter_rows(min_row=2):
        if f[0].value in (None, "TOTAL"):  # fin de los datos (después vienen el total y la nota)
            break
        out.append(dict(zip(cab, [c.value for c in f])))
    return out


def _xlsx(filas: list[list]) -> bytes:
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_importar_reservas(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    plantilla = client.get("/api/turistico/importar/plantilla", headers=admin)
    ws = _libro(plantilla)["Reservas"]
    assert [c.value for c in ws[1]][:4] == ["Localizador", "Unidad", "Entrada", "Salida"]

    fichero = _xlsx([
        ["Listado de reservas mayo"],  # título encima de la cabecera
        ["Localizador", "Unidad", "Fecha de entrada", "Fecha de salida", "Adultos", "Canal", "Importe total",
         "Nombre", "Apellidos", "DNI", "Columna rara"],
        ["IMP-1", "A-140", date(2025, 5, 10), date(2025, 5, 13), 2, "Booking.com", 300, "Luis", "Sanz", "11111111H", "x"],
        ["IMP-2", "", "12/05/2025", "14/05/2025", 1, "airbnb", "1.234,50", "Eva", "Mora", None, None],
        ["IMP-3", "A-140", "11/05/2025", "12/05/2025", 1, "directo", 50, "Solapa", None, None, None],  # solapa IMP-1
        ["IMP-4", "Z-999", "10/05/2025", "11/05/2025", 1, "", 10, "Nadie", None, None, None],  # unidad inexistente
        ["IMP-5", "A-141", "10/05/2025", "11/05/2025", 1, "", 10, None, None, None, None],  # sin nombre
        ["IMP-6", "A-141", "10/05/2025", "09/05/2025", 1, "", 10, "Al revés", None, None, None],
    ])
    files = {"fichero": ("reservas.xlsx", fichero, XLSX)}
    prev = client.post("/api/turistico/importar", headers=admin, data={"asset_id": sae}, files=files).json()
    assert (prev["validas"], prev["importadas"], prev["errores"]) == (2, 0, 4)
    assert prev["columnas_ignoradas"] == ["Columna rara"]
    por_loc = {f["localizador"]: f for f in prev["filas"]}
    assert por_loc["IMP-1"]["estado"] == "valida" and por_loc["IMP-1"]["unidad"] == "A-140"
    assert por_loc["IMP-2"]["asignada"] is True and por_loc["IMP-2"]["importe_total"] == 1234.5
    assert "ya está reservada" in por_loc["IMP-3"]["motivo"]
    assert "Z-999 no existe" in por_loc["IMP-4"]["motivo"]
    assert "Falta el nombre" in por_loc["IMP-5"]["motivo"]
    assert "posterior" in por_loc["IMP-6"]["motivo"]
    assert client.get("/api/turistico/reservas?q=IMP-1&desde=2025-05-01", headers=admin).json() == []  # nada creado

    r = client.post("/api/turistico/importar", headers=admin, data={"asset_id": sae, "confirmar": "true"},
                    files={"fichero": ("reservas.xlsx", fichero, XLSX)}).json()
    assert r["importadas"] == 2
    res = client.get("/api/turistico/reservas?q=IMP-&desde=2025-05-01&hasta=2025-05-31", headers=admin).json()
    assert {(x["localizador"], x["canal"], x["importe_pagado"]) for x in res} == {("IMP-1", "booking", 0),
                                                                                 ("IMP-2", "airbnb", 0)}
    # reimportar el mismo fichero no duplica
    r = client.post("/api/turistico/importar", headers=admin, data={"asset_id": sae, "confirmar": "true"},
                    files={"fichero": ("reservas.xlsx", fichero, XLSX)}).json()
    assert r["importadas"] == 0 and r["omitidas"] == 2

    # CSV con cabeceras de Booking; las canceladas se ignoran
    csv = ("Número de reserva;Nombre del cliente;Check-in;Check-out;Personas;Precio;Estado\n"
           "BK-77;Ana Pérez;2025-06-01;2025-06-04;2;450,00 €;ok\n"
           "BK-78;Juan Gil;2025-06-02;2025-06-03;1;90;cancelled_by_guest\n").encode("cp1252")
    r = client.post("/api/turistico/importar", headers=admin, data={"asset_id": sae, "confirmar": "1"},
                    files={"fichero": ("booking.csv", csv, "text/csv")}).json()
    assert (r["importadas"], r["omitidas"], r["errores"]) == (1, 1, 0)
    assert r["filas"][0]["importe_total"] == 450 and r["filas"][0]["huesped"] == "Ana Pérez"
    # fichero sin columnas de fechas
    r = client.post("/api/turistico/importar", headers=admin, data={"asset_id": sae},
                    files={"fichero": ("x.csv", b"a;b\n1;2\n", "text/csv")})
    assert r.status_code == 400 and "fecha de entrada" in r.json()["detail"]
    # recepción de otro activo no puede importar aquí
    email = "recepcion.importa@inversiete.com"
    client.post("/api/admin/usuarios", headers=admin, json={
        "email": email, "nombre": "x", "password": "Provisional1",
        "asignaciones": [{"role_id": ids["roles"]["Recepción"], "asset_id": ids["assets"]["SFL"]["id"]}]})
    h = login(client, email, "Provisional1")
    client.post("/api/auth/password", headers=h, json={"actual": "Provisional1", "nueva": "ClaveDefinitiva2026"})
    assert client.post("/api/turistico/importar", headers=h, data={"asset_id": sae},
                       files={"fichero": ("reservas.xlsx", fichero, XLSX)}).status_code == 403


def test_informes(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    domicilio_fiscal(client, admin)
    # reserva de 3 noches 30/04-03/05/2025: 1 noche en abril, 2 en mayo. Se cobra (y factura) hoy: la producción
    # cuenta en el mes de la factura, no en el de la estancia
    unit = client.get(f"/api/unidades?asset_id={sae}&q=B-529", headers=admin).json()[0]
    assert client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": unit["id"], "guest": {"nombre": "Informe"}, "fecha_entrada": "2025-04-30",
        "fecha_salida": "2025-05-03", "importe_total": 330, "importe_pagado": 330}).status_code == 201
    wb = _libro(client.get(f"/api/informes/ocupacion?desde=2025-04-01&hasta=2025-05-31&asset_id={sae}", headers=admin))
    # la ocupación del edificio es solo de los apartamentos; los garajes van aparte y no computan
    filas = {f["Mes"]: f for f in _filas(wb["Turísticos"])}
    garajes = {f["Mes"]: f for f in _filas(wb["Garajes"])}
    assert garajes["abr-2025"]["Plazas"] == 242  # 64 plazas exteriores + 178 del sótano -1
    assert filas["abr-2025"]["Apartamentos"] == 300 and filas["abr-2025"]["Noches disponibles"] == 9000
    assert filas["abr-2025"]["Noches ocupadas"] == 1 and filas["abr-2025"]["Alojamiento facturado (base)"] == 0
    assert filas["may-2025"]["Noches ocupadas"] >= 2
    assert "Residencial" not in wb.sheetnames  # filtrado por activo turístico

    hoy = date.today()
    mes = f"{['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic'][hoy.month - 1]}-{hoy.year}"
    assert _filas(_libro(client.get("/api/informes/produccion?desde=2025-04-01&hasta=2025-04-30",
                                    headers=admin))["Producción"])[1]["Producción del edificio (base, sin garajes)"] == 0
    wb = _libro(client.get(f"/api/informes/produccion?desde={hoy.replace(day=1)}&hasta={hoy}", headers=admin))
    prod = {f["Activo"]: f for f in _filas(wb["Producción"])}
    assert set(prod) == {"C/ Babilonia 35", "Suite Aeropuerto", "Suite Florida"} and prod["Suite Aeropuerto"]["Mes"] == mes
    # coincide con las facturas emitidas este mes en Suite Aeropuerto
    facturas = [f for f in client.get(f"/api/facturas?asset_id={sae}&anio={hoy.year}", headers=admin).json()
                if f["fecha_expedicion"][:7] == f"{hoy:%Y-%m}"]
    base = round(sum(f["base_imponible"] for f in facturas), 2)
    p = prod["Suite Aeropuerto"]
    garaje = p["Plazas de garaje (base, aparte)"]
    assert round(p["Producción del edificio (base, sin garajes)"] + garaje, 2) == base
    assert p["Nº facturas"] == len(facturas)
    assert round(p["Alojamiento (base)"] + p["Rentas (base)"] + p["Servicios (base)"], 2) == round(base - garaje, 2)
    assert p["Alojamiento (base)"] >= 300  # 330 € con IVA del 10 %
    panel = {a["codigo"]: a for a in client.get("/api/panel", headers=admin).json()["activos"]}
    assert panel["SAE"]["produccion_mes"] == round(base - garaje, 2)
    assert panel["SAE"]["garajes_facturado_mes"] == garaje

    # morosidad: recibo de enero de 2025 sin cobrar
    bab = ids["assets"]["BAB35"]["id"]
    garaje = client.get(f"/api/unidades?asset_id={bab}&uso=garaje", headers=admin).json()[-5]
    lease = client.post("/api/alquiler/contratos", headers=admin, json={
        "unit_id": garaje["id"], "tenant": {"nombre": "Moroso", "telefono": "600111222"},
        "fecha_inicio": "2025-01-01", "fecha_fin": "2025-01-31", "renta_mensual": 100}).json()
    assert client.post("/api/alquiler/recibos/generar", headers=admin, json={"periodo": "2025-01"}).json()["creados"] == 1
    client.put(f"/api/alquiler/contratos/{lease['id']}", headers=admin, json={"estado": "finalizado"})
    wb = _libro(client.get("/api/informes/morosidad?hasta=2025-03-06", headers=admin))
    assert wb.sheetnames[0] == "Resumen"
    imp = [f for f in _filas(wb["Recibos impagados"]) if f["Inquilino"] == "Moroso"][0]
    assert (imp["Pendiente"], imp["Días de retraso"], imp["Tramo"], imp["Teléfono"]) == (121, 60, "31-60 días", "600111222")
    resumen = {f["Activo"]: f for f in _filas(wb["Resumen"])}
    assert resumen["C/ Babilonia 35"]["31-60 días"] >= 121

    # mantenimiento: costes por instalación
    w = client.post("/api/mantenimiento/ordenes", headers=admin, json={
        "asset_id": sae, "titulo": "Cambio bomba ACS", "categoria": "acs", "proveedor": "Bombas Ejemplo"}).json()
    client.post(f"/api/mantenimiento/ordenes/{w['id']}/confirmar-mantenimiento", headers=admin,
                json={"solucion": "Bomba sustituida", "coste_real": 850})
    hoy = date.today()
    wb = _libro(client.get(f"/api/informes/mantenimiento?desde={hoy}&hasta={hoy}&asset_id={sae}", headers=admin))
    assert {"Órdenes de trabajo", "Por instalación", "Por tipo", "Por proveedor"} <= set(wb.sheetnames)
    acs = [f for f in _filas(wb["Por instalación"]) if f["Instalación"] == "acs"][0]
    assert acs["Coste real"] >= 850
    prov = [f for f in _filas(wb["Por proveedor"]) if f["Proveedor"] == "Bombas Ejemplo"][0]
    assert prov["Nº OT"] == 1 and prov["Coste real"] == 850
    ot = [f for f in _filas(wb["Órdenes de trabajo"]) if f["OT"] == f"OT-{w['id']:05d}"][0]
    assert ot["Instalación"] == "acs" and ot["Estado"] == "trabajo_realizado"

    # permisos y validaciones
    assert client.get("/api/informes/inventado", headers=admin).status_code == 404
    assert client.get("/api/informes/ocupacion?desde=2025-05-01&hasta=2025-04-01", headers=admin).status_code == 400
    email = "recepcion.informes@inversiete.com"
    client.post("/api/admin/usuarios", headers=admin, json={
        "email": email, "nombre": "x", "password": "Provisional1",
        "asignaciones": [{"role_id": ids["roles"]["Recepción"], "asset_id": ids["assets"]["SFL"]["id"]}]})
    h = login(client, email, "Provisional1")
    client.post("/api/auth/password", headers=h, json={"actual": "Provisional1", "nueva": "ClaveDefinitiva2026"})
    assert client.get("/api/informes/produccion", headers=h).status_code == 403  # sin finanzas
    wb = _libro(client.get("/api/informes/ocupacion?desde=2025-04-01&hasta=2025-04-30", headers=h))
    assert {f["Activo"] for f in _filas(wb["Turísticos"])} == {"Suite Florida"}  # solo su activo


def test_resumen_diario(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    hoy = avisos.hoy()
    assert client.post("/api/mantenimiento/planes", headers=admin, json={
        "asset_id": sae, "titulo": "Inspección periódica ascensores (OCA)", "categoria": "ascensores",
        "normativa": "RD 355/2024", "periodicidad_dias": 730,
        "proxima_fecha": (hoy + timedelta(days=10)).isoformat()}).status_code == 201
    bab = ids["assets"]["BAB35"]["id"]
    garaje = client.get(f"/api/unidades?asset_id={bab}&uso=garaje", headers=admin).json()[-6]
    lease = client.post("/api/alquiler/contratos", headers=admin, json={
        "unit_id": garaje["id"], "tenant": {"nombre": "Vence Pronto"}, "fecha_inicio": "2025-01-01",
        "fecha_fin": (hoy + timedelta(days=45)).isoformat(), "renta_mensual": 80}).json()
    avisos.BANDEJA.clear()
    r = client.post("/api/admin/avisos/resumen", headers=admin).json()
    assert r["enviados"] >= 3 and r["errores"] == 0
    alfonso = [m for m in avisos.BANDEJA if m["para"] == "alfonso@inversiete.es"][0]
    assert alfonso["asunto"].startswith(f"PMS · Resumen del {hoy:%d/%m/%Y}:")
    assert "Inspección periódica ascensores (OCA)" in alfonso["texto"] and "en 10 días" in alfonso["texto"]
    assert "Vence Pronto" in alfonso["texto"] and "45 días" in alfonso["texto"]
    # la recepción no recibe los alquileres de viviendas ni el preventivo; solo, si los hay, los recibos de las
    # plazas de garaje alquiladas a clientes externos de su activo
    for m in avisos.BANDEJA:
        if m["para"] == "jaime@apartamentossuitesaeropuerto.es":
            assert "Plazas de garaje" in m["html"] and "Recibos vencidos sin cobrar (" not in m["html"]
            assert "Revisiones preventivas" not in m["html"]
    # el automático no repite a quien ya lo recibió hoy
    from app.database import SessionLocal
    with SessionLocal() as db:
        assert avisos.resumen_diario(db)["enviados"] == 0
    # correo de prueba y registro
    assert client.post("/api/admin/avisos/probar", headers=admin, json={}).json()["ok"]
    assert avisos.BANDEJA[-1]["asunto"] == "PMS · Correo de prueba"
    estado = client.get("/api/admin/avisos", headers=admin).json()
    assert estado["configurado"] and estado["remitente"] == "avisos@inversiete.es"
    assert any(e["tipo"] == "resumen" and e["ok"] for e in estado["registro"])
    client.put(f"/api/alquiler/contratos/{lease['id']}", headers=admin, json={"estado": "finalizado"})
