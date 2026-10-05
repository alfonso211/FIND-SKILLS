"""Peticiones IA: el asistente consulta con los permisos del usuario, crea Excel y archiva documentos.
Se usa un cliente simulado de la API (sin llamadas reales)."""
import json
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image

from app import ia
from test_adjuntos_ot import _usuario

HOY = date.today()


class Bloque(SimpleNamespace):
    def to_dict(self):
        return {k: v for k, v in vars(self).items()}


def texto(t):
    return Bloque(type="text", text=t)


def herramienta(i, nombre, entrada):
    return Bloque(type="tool_use", id=f"toolu_{i}", name=nombre, input=entrada)


class ClienteFalso:
    """Devuelve las respuestas guionizadas y guarda las peticiones recibidas."""
    def __init__(self, guion):
        self.guion, self.peticiones = list(guion), []
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kw):
        self.peticiones.append(json.loads(json.dumps(kw, default=str)))
        contenido = self.guion.pop(0)
        parada = "tool_use" if any(b.type == "tool_use" for b in contenido) else "end_turn"
        return SimpleNamespace(content=contenido, stop_reason=parada, usage=SimpleNamespace(output_tokens=10))


@pytest.fixture
def falso():
    def poner(guion):
        ia.CLIENTE_PRUEBAS = ClienteFalso(guion)
        return ia.CLIENTE_PRUEBAS
    yield poner
    ia.CLIENTE_PRUEBAS = None


def _jpeg() -> bytes:
    out = BytesIO()
    Image.new("RGB", (600, 800), (240, 240, 240)).save(out, "JPEG")
    return out.getvalue()


def test_sin_configurar(client, admin):
    assert client.get("/api/ia/estado", headers=admin).json() == {"puede": True, "configurado": False}
    r = client.post("/api/ia/peticion", headers=admin, json={"texto": "Hola"})
    assert r.status_code == 503


def test_informe_en_excel(client, admin, ids, falso):
    sae = ids["assets"]["SAE"]["id"]
    rec = _usuario(client, admin, ids, "recepcion.ia@inversiete.com", "Recepción", "SAE")
    client.post("/api/gastos", headers=rec, json={"asset_id": sae, "fecha": HOY.isoformat(), "categoria": "limpieza",
                                                  "concepto": "Lavandería IA", "total": 121})
    f = falso([
        [texto("Consulto los gastos."), herramienta(1, "consultar", {"tipo": "gastos", "asset_id": sae})],
        [herramienta(2, "crear_excel", {"titulo": "Gastos de limpieza", "hojas": [
            {"nombre": "Gastos", "columnas": ["Fecha", "Concepto", "Total"],
             "filas": [[HOY.isoformat(), "Lavandería IA", 121]]}]})],
        [texto("Aquí tiene el Excel con los gastos de limpieza.")],
    ])
    r = client.post("/api/ia/peticion", headers=rec, json={"texto": "Excel con los gastos de limpieza del mes"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["respuesta"] == "Aquí tiene el Excel con los gastos de limpieza."
    assert d["ficheros"][0]["nombre"] == "Gastos de limpieza.xlsx"
    # la consulta se hizo con los permisos de la recepción y devolvió sus gastos
    resultado = d["historial"][2]["content"][0]
    assert resultado["type"] == "tool_result" and "Lavandería IA" in resultado["content"]
    # petición a la API: modelo, herramientas, sistema sin mencionar al proveedor, respaldo por defecto
    p = f.peticiones[0]
    assert p["model"] == "claude-opus-5-5" and p["fallbacks"] == "default"
    assert {t["name"] for t in p["tools"]} == {"consultar", "crear_excel", "archivar_documento"}
    assert "Asistente de INVERPMS" in p["system"] and "Suite Aeropuerto" in p["system"]
    assert len(f.peticiones) == 3 and f.peticiones[2]["messages"][-1]["content"][0]["type"] == "tool_result"
    # el Excel lo descarga quien lo pidió, nadie más
    x = client.get(d["ficheros"][0]["url"], headers=rec)
    assert x.status_code == 200
    celdas = [c for fila in load_workbook(BytesIO(x.content))["Gastos"].iter_rows(values_only=True) for c in fila]
    assert "Lavandería IA" in celdas and 121 in celdas
    otra = _usuario(client, admin, ids, "recepcion.ia2@inversiete.com", "Recepción", "SAE")
    assert client.get(d["ficheros"][0]["url"], headers=otra).status_code == 404

    # siguiente turno: se reenvía el historial tal cual y se añade la nueva pregunta
    f = falso([[texto("Hay 1 gasto.")]])
    r2 = client.post("/api/ia/peticion", headers=rec, json={"texto": "¿Cuántos son?", "historial": d["historial"]})
    assert r2.json()["respuesta"] == "Hay 1 gasto."
    assert f.peticiones[0]["messages"][:len(d["historial"])] == json.loads(json.dumps(d["historial"]))


def test_leer_y_archivar_documento(client, admin, ids, falso):
    sae = ids["assets"]["SAE"]["id"]
    rec = _usuario(client, admin, ids, "recepcion.ia3@inversiete.com", "Recepción", "SAE")
    rec_sfl = _usuario(client, admin, ids, "recepcion.ia.sfl@inversiete.com", "Recepción", "SFL")
    adj = client.post("/api/ia/adjuntos", headers=rec, files={"fichero": ("factura.jpg", _jpeg(), "image/jpeg")}).json()
    entrada = {"adjunto": adj["ref"], "asset_id": sae, "tipo": "factura", "fecha": HOY.isoformat(),
               "emisor": "Electricidad IA SL", "referencia": "F-77", "apartamento": "A-151",
               "gasto": {"categoria": "mantenimiento", "concepto": "Reparación enchufe", "total": 60.5, "tipo_iva": 21}}
    f = falso([[herramienta(1, "archivar_documento", entrada)], [texto("Factura archivada.")]])
    d = client.post("/api/ia/peticion", headers=rec, json={"texto": "Archívala", "adjuntos": [adj["ref"]]}).json()
    assert d["acciones"][0]["texto"] == "Factura archivada en Suite Aeropuerto · gasto de 60.50 € anotado"
    contenido = f.peticiones[0]["messages"][0]["content"]
    assert contenido[0]["type"] == "image" and adj["ref"] in contenido[1]["text"]
    doc = next(x for x in client.get(f"/api/documentos-recibidos?asset_id={sae}", headers=rec).json()
               if x["referencia"] == "F-77")
    assert doc["unidad"] == "A-151" and doc["gasto"]["base"] == 50 and doc["gasto"]["lugar"] == "Apartamento A-151"

    # otra recepción no puede usar el adjunto ni archivar en un activo ajeno
    f = falso([[herramienta(1, "archivar_documento", entrada)], [texto("No ha sido posible.")]])
    r = client.post("/api/ia/peticion", headers=rec_sfl, json={"texto": "Archívala", "adjuntos": [adj["ref"]]})
    assert r.status_code == 400
    f = falso([[herramienta(1, "archivar_documento", entrada)], [texto("No ha sido posible.")]])
    d = client.post("/api/ia/peticion", headers=rec_sfl, json={"texto": "Archívala"}).json()
    res = d["historial"][2]["content"][0]
    assert res["is_error"] and not d["acciones"]
    # consultar reservas de un activo ajeno no devuelve nada
    f = falso([[herramienta(1, "consultar", {"tipo": "reservas", "asset_id": sae})], [texto("Sin datos.")]])
    d = client.post("/api/ia/peticion", headers=rec_sfl, json={"texto": "Reservas de Suite Aeropuerto"}).json()
    assert json.loads(d["historial"][2]["content"][0]["content"])["total_filas"] == 0
    assert "Suite Aeropuerto" not in f.peticiones[0]["system"]


def test_solo_direccion_y_recepcion(client, admin, ids, falso):
    falso([[texto("hola")]])
    gob = _usuario(client, admin, ids, "gobernanta.ia@inversiete.com", "Gobernanta / Limpieza", "SAE")
    assert client.get("/api/ia/estado", headers=gob).json()["puede"] is False
    assert client.post("/api/ia/peticion", headers=gob, json={"texto": "Hola"}).status_code == 403
    assert client.post("/api/ia/adjuntos", headers=gob, files={"fichero": ("x.jpg", _jpeg(), "image/jpeg")}
                       ).status_code == 403
