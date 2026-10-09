"""Documentos recibidos (carpeta por activo) y cuenta de gastos."""
import json
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from PIL import Image

from test_adjuntos_ot import _usuario

HOY = date.today()


def _jpeg(color=(30, 120, 200)) -> bytes:
    out = BytesIO()
    Image.new("RGB", (800, 1100), color).save(out, "JPEG")
    return out.getvalue()


def _subir(client, h, asset_id, ficheros, gasto=None, **campos):
    data = {"asset_id": str(asset_id), "tipo": "factura", "fecha": HOY.isoformat(), **campos}
    if gasto:
        data["gasto"] = json.dumps(gasto)
    return client.post("/api/documentos-recibidos", headers=h, data=data,
                       files=[("ficheros", (n, d, "image/jpeg")) for n, d in ficheros])


def test_documentos_y_gastos(client, admin, ids):
    sae, sfl = ids["assets"]["SAE"]["id"], ids["assets"]["SFL"]["id"]
    apto = next(u for u in client.get(f"/api/unidades?asset_id={sae}&q=A-150", headers=admin).json()
                if u["codigo"] == "A-150")
    rec_sae = _usuario(client, admin, ids, "recepcion.gastos.sae@inversiete.com", "Recepción", "SAE")
    rec_sfl = _usuario(client, admin, ids, "recepcion.gastos.sfl@inversiete.com", "Recepción", "SFL")
    gob = _usuario(client, admin, ids, "gobernanta.gastos@inversiete.com", "Gobernanta / Limpieza", "SAE")

    # factura de dos páginas (dos fotos: se unen en un PDF) imputada a un apartamento
    r = _subir(client, rec_sae, sae, [("p1.jpg", _jpeg()), ("p2.jpg", _jpeg((200, 30, 30)))],
               gasto={"fecha": HOY.isoformat(), "categoria": "mantenimiento", "concepto": "Cambio de grifo",
                      "ambito": "apartamento", "unit_id": apto["id"], "total": 121, "tipo_iva": 21,
                      "proveedor": "Fontanería Gasto SL", "numero_factura": "F-2026-15", "forma_pago": "transferencia", "naturaleza": "OPEX"},
               emisor="Fontanería Gasto SL", referencia="F-2026-15", unit_id=str(apto["id"]))
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["mime"] == "application/pdf" and doc["tipo_nombre"] == "Factura" and doc["unidad"] == "A-150"
    g = doc["gasto"]
    assert (g["base"], g["cuota"], g["total"]) == (100, 21, 121) and g["lugar"] == "Apartamento A-150"
    pdf = client.get(f"/api/documentos-recibidos/{doc['id']}/fichero", headers=rec_sae)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    # una carta: no es gasto
    carta = _subir(client, rec_sae, sae, [("carta.jpg", _jpeg())], tipo="carta", emisor="Ayuntamiento")
    assert carta.status_code == 201 and carta.json()["gasto"] is None
    # gasto sin documento (recibo domiciliado), con varios tipos de IVA (base indicada)
    s = client.post("/api/gastos", headers=rec_sae, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "suministros", "concepto": "Luz zonas comunes",
        "total": 250, "base": 210, "tipo_iva": 21, "pagado": True, "forma_pago": "domiciliacion",
        "naturaleza": "OPEX"})
    assert s.status_code == 400 and "justificante" in s.json()["detail"]  # pagada sin justificante: no
    s = client.post("/api/gastos", headers=rec_sae, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "suministros", "concepto": "Luz zonas comunes",
        "total": 250, "base": 210, "tipo_iva": 21, "forma_pago": "domiciliacion", "naturaleza": "OPEX"})
    assert s.status_code == 201, s.text
    s = client.post(f"/api/gastos/{s.json()['id']}/pagar", headers=rec_sae, files=[("ficheros", ("cargo.jpg", _jpeg(), "image/jpeg"))])
    assert s.status_code == 200, s.text
    assert s.json()["cuota"] == 40 and s.json()["fecha_pago"] == HOY.isoformat() and s.json()["lugar"] == "General"
    assert s.json()["pagado"] and s.json()["justificante_id"]
    assert client.get(f"/api/documentos-recibidos/{s.json()['justificante_id']}/fichero", headers=rec_sae).status_code == 200
    assert client.post("/api/gastos", headers=rec_sae, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "inventada", "concepto": "X x", "total": 1}
    ).status_code == 400
    # obligatorio: si la factura se paga por transferencia o está cargada en cuenta
    sin_forma = client.post("/api/gastos", headers=rec_sae, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "suministros", "concepto": "Agua", "total": 50,
        "naturaleza": "OPEX"})
    assert sin_forma.status_code == 400 and "cargo en cuenta" in sin_forma.json()["detail"]
    # obligatorio: OPEX (gasto corriente) o CAPEX (inversión)
    sin_nat = client.post("/api/gastos", headers=rec_sae, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "suministros", "concepto": "Agua", "total": 50,
        "forma_pago": "transferencia"})
    assert sin_nat.status_code == 400 and "CAPEX" in sin_nat.json()["detail"]
    assert set(client.get("/api/gastos/catalogos", headers=rec_sae).json()["naturalezas"]) == {"OPEX", "CAPEX"}
    formas = client.get("/api/gastos/catalogos", headers=rec_sae).json()["formas_pago"]
    assert list(formas)[:2] == ["transferencia", "domiciliacion"] and formas["domiciliacion"].startswith("Cargo en cuenta")

    lista = client.get(f"/api/gastos?asset_id={sae}", headers=rec_sae).json()
    mios = [x for x in lista["gastos"] if x["concepto"] in ("Cambio de grifo", "Luz zonas comunes")]
    assert len(mios) == 2 and lista["totales"]["sin_documento"] >= 1
    assert client.get(f"/api/gastos?asset_id={sae}&unit_id={apto['id']}", headers=rec_sae).json()["totales"]["base"] >= 100

    # cada activo ve solo lo suyo; limpieza no tiene acceso
    assert not [x for x in client.get("/api/documentos-recibidos", headers=rec_sfl).json() if x["asset_id"] == sae]
    assert not [x for x in client.get("/api/gastos", headers=rec_sfl).json()["gastos"] if x["asset_id"] == sae]
    assert client.get(f"/api/documentos-recibidos/{doc['id']}/fichero", headers=rec_sfl).status_code == 403
    assert _subir(client, rec_sfl, sae, [("x.jpg", _jpeg())]).status_code == 403
    assert client.get("/api/documentos-recibidos", headers=gob).status_code == 403
    assert client.get("/api/gastos", headers=gob).status_code == 403
    assert _subir(client, rec_sfl, sfl, [("x.jpg", _jpeg())], tipo="ticket").status_code == 201
    assert {x["asset_id"] for x in client.get("/api/documentos-recibidos", headers=admin).json()} >= {sae, sfl}

    # marcarlo pagado: solo con el justificante de pago
    e = client.put(f"/api/gastos/{g['id']}", headers=rec_sae, json={**g, "pagado": True, "base": None})
    assert e.status_code == 400 and "justificante" in e.json()["detail"]
    assert client.post(f"/api/gastos/{g['id']}/pagar", headers=rec_sae, data={"fecha_pago": "2099-01-01"},
                       files=[("ficheros", ("t.jpg", _jpeg(), "image/jpeg"))]).status_code == 400  # fecha futura
    e = client.post(f"/api/gastos/{g['id']}/pagar", headers=rec_sae, files=[("ficheros", ("t.jpg", _jpeg(), "image/jpeg"))])
    assert e.status_code == 200 and e.json()["pagado"] and e.json()["base"] == 100
    assert client.post(f"/api/gastos/{g['id']}/pagar", headers=rec_sae,
                       files=[("ficheros", ("t.jpg", _jpeg(), "image/jpeg"))]).status_code == 400  # ya pagada
    # editar otros datos de una pagada la mantiene pagada
    e2 = client.put(f"/api/gastos/{g['id']}", headers=rec_sae, json={**e.json(), "total": e.json()["liquido"],
                                                                      "base": None, "notas": "Revisada"})
    assert e2.status_code == 200 and e2.json()["pagado"] and e2.json()["justificante_id"] == e.json()["justificante_id"]

    # Excel de la cuenta de gastos
    x = client.get(f"/api/gastos/excel?asset_id={sae}&desde={HOY.replace(day=1)}&hasta={HOY}", headers=admin)
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Gastos", "Por categoría", "Por apartamento"]
    textos = [c for fila in wb["Por apartamento"].iter_rows(values_only=True) for c in fila]
    assert "A-150" in textos

    # borrar: quien lo subió (el mismo día) o la dirección; otra recepción del mismo activo, no
    otra = _usuario(client, admin, ids, "recepcion.gastos.sae2@inversiete.com", "Recepción", "SAE")
    assert client.delete(f"/api/documentos-recibidos/{doc['id']}", headers=otra).status_code == 403
    assert client.delete(f"/api/documentos-recibidos/{doc['id']}", headers=rec_sae).json()["ok"]
    queda = [x for x in client.get(f"/api/gastos?asset_id={sae}", headers=admin).json()["gastos"] if x["id"] == g["id"]]
    assert queda and queda[0]["documento_id"] is None  # el apunte del gasto se conserva
    assert client.delete(f"/api/gastos/{g['id']}", headers=admin).json()["ok"]


def test_proveedor_nuevo_desde_un_gasto(client, admin, ids):
    """Recepción da de alta en el fichero del grupo el proveedor de una factura; el gasto queda enlazado."""
    sae = ids["assets"]["SAE"]["id"]
    rec = _usuario(client, admin, ids, "recepcion.proveedor@inversiete.com", "Recepción", "SAE")
    g = client.post("/api/gastos", headers=rec, json={
        "asset_id": sae, "fecha": HOY.isoformat(), "categoria": "material", "concepto": "Sábanas",
        "proveedor": "Textiles  Nuevo Proveedor SL", "total": 60.5, "forma_pago": "tarjeta", "naturaleza": "OPEX"}).json()
    assert g["supplier_id"] is None
    assert client.get("/api/proveedores?q=Textiles", headers=rec).json() == []
    p = client.post("/api/proveedores", headers=rec, json={
        "nombre": "TEXTILES NUEVO PROVEEDOR SL", "nif": "A58818501", "pais": "España", "telefono": "910000001",
        "direccion": "C/ Lino 4", "municipio": "Madrid", "notas": "Entrega en 48 h"})
    assert p.status_code == 201, p.text
    lista = client.get(f"/api/gastos?asset_id={sae}", headers=rec).json()["gastos"]
    assert next(x for x in lista if x["id"] == g["id"])["supplier_id"] == p.json()["id"]
    # recepción no borra proveedores (son de todo el grupo); la dirección sí, y el gasto conserva el nombre
    assert client.delete(f"/api/proveedores/{p.json()['id']}", headers=rec).status_code == 403
    assert client.delete(f"/api/proveedores/{p.json()['id']}", headers=admin).json()["ok"]
    lista = client.get(f"/api/gastos?asset_id={sae}", headers=admin).json()["gastos"]
    x = next(x for x in lista if x["id"] == g["id"])
    assert x["supplier_id"] is None and x["proveedor"] == "Textiles Nuevo Proveedor SL"
    # limpieza no ve el fichero de proveedores de la cuenta de gastos ni da de alta
    gob = _usuario(client, admin, ids, "gobernanta.proveedor@inversiete.com", "Gobernanta / Limpieza", "SAE")
    assert client.post("/api/proveedores", headers=gob, json={"nombre": "No Permitido SL"}).status_code == 403
