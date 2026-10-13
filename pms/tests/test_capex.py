"""CAPEX fase 1: proyecto con presupuesto importado de un BC3, base de precios, documentos, plantillas «CAPEX tipo»,
licitación ciega (enlace y Excel protegido), comparativo con mapa de calor, adjudicación por capítulos y aprobación de
la Dirección General por correo cuando se supera el umbral de 30.000 €."""
import io
import re
from pathlib import Path

from openpyxl import load_workbook

from app import avisos

BC3 = (Path(__file__).parent / "datos" / "presupuesto_prueba.bc3").read_bytes()
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


def _proveedor(client, admin, nombre, email):
    r = client.post("/api/proveedores", headers=admin, json={"nombre": nombre, "email": email, "telefono": "600111333",
                                                               "actividad": "Reformas"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _publico(url: str) -> tuple[str, str]:
    iid, k = re.search(r"/oferta/(\d+)\?k=([\w-]+)", url).groups()
    return f"/api/publico/oferta/{iid}", k


def test_capex_licitacion_comparativo_y_aprobacion(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    provs = [_proveedor(client, admin, f"Reformas CAPEX {n} SL", f"capex{n.lower()}@inversiete.com") for n in "ABC"]

    # proyecto + presupuesto desde un BC3 con mediciones
    p = client.post("/api/capex", headers=admin, json={"asset_id": sfl, "titulo": "Reforma integral baños planta 3",
                                                       "ubicacion": "Planta 3", "presupuesto_objetivo": 45000}).json()
    assert p["codigo"].startswith("CPX-") and p["estado"] == "definicion"
    r = client.post(f"/api/capex/{p['id']}/importar-bc3", headers=admin,
                    files={"fichero": ("baño.bc3", BC3, "application/octet-stream")}, data={"reemplazar": "true"})
    assert r.status_code == 200 and r.json() == {"partidas": 3, "capitulos": 2}, r.text
    d = client.get(f"/api/capex/{p['id']}", headers=admin).json()
    assert d["totales"]["pem"] == 1291.39 and d["totales"]["pbl"] == round(1291.39 * 1.19, 2)
    alicatado = next(x for x in d["partidas"] if x["codigo"] == "ALB030")
    assert client.put(f"/api/capex/partidas/{alicatado['id']}", headers=admin, json={"cantidad": 1000}).status_code == 200

    # base de precios FIEBDC-3 y partida añadida desde la base (capítulo nuevo)
    b = client.post("/api/capex/bases/lista", headers=admin, data={"nombre": "Base de prueba"},
                    files={"fichero": ("base.bc3", BC3, "application/octet-stream")})
    assert b.status_code == 201 and b.json()["partidas"] == 3, b.text
    encontrados = client.get("/api/capex/precios/buscar", headers=admin, params={"q": "desmontaje sanitario"}).json()
    assert [x["codigo"] for x in encontrados] == ["DEM020"] and encontrados[0]["precio"] == 13.5
    nuevas = client.post(f"/api/capex/{p['id']}/partidas/desde-base", headers=admin, json={
        "precio_ids": [encontrados[0]["id"]], "capitulo": "03", "capitulo_nombre": "FONTANERÍA"}).json()
    assert nuevas[0]["cantidad"] == 0

    # documentos (el plano va a la licitación) y plantilla a partir del proyecto
    r = client.post(f"/api/capex/{p['id']}/documentos", headers=admin, data={"etiqueta": "plano"},
                    files=[("ficheros", ("planta3.pdf", PDF, "application/pdf")),
                           ("ficheros", ("planta3.dxf", b"0\nSECTION\n2\nENTITIES\n0\nENDSEC\n0\nEOF", "application/dxf"))])
    assert r.status_code == 201 and len(r.json()) == 2, r.text
    t = client.post("/api/capex/plantillas/lista", headers=admin, json={"nombre": "Reforma integral baño",
                                                                        "project_id": p["id"]}).json()
    assert t["n_partidas"] == 4 and len(t["documentos"]) == 2
    p2 = client.post("/api/capex", headers=admin, json={"asset_id": sfl, "titulo": "Baño 412",
                                                        "plantilla_id": t["id"]}).json()
    d2 = client.get(f"/api/capex/{p2['id']}", headers=admin).json()
    assert len(d2["partidas"]) == 4 and len(d2["documentos"]) == 2

    # licitación: no con partidas sin medir
    plazo = "2099-12-31"
    r = client.post(f"/api/capex/{p['id']}/licitacion", headers=admin, json={"proveedor_ids": provs, "fecha_limite": plazo})
    assert r.status_code == 400 and "sin medición" in r.json()["detail"]
    client.put(f"/api/capex/partidas/{nuevas[0]['id']}", headers=admin, json={"cantidad": 4})
    avisos.BANDEJA.clear()
    r = client.post(f"/api/capex/{p['id']}/licitacion", headers=admin, json={
        "proveedor_ids": provs, "fecha_limite": plazo, "mensaje": "Visita de obra el lunes a las 10:00."})
    assert r.status_code == 201 and all(x["ok"] and x["whatsapp"] for x in r.json()["enviados"]), r.text
    correos = {m["para"]: m for m in avisos.BANDEJA}
    assert set(correos) == {"capexa@inversiete.com", "capexb@inversiete.com", "capexc@inversiete.com"}
    adj = [n for n, _, _ in correos["capexa@inversiete.com"]["adjuntos"]]
    assert adj[0].endswith(".xlsx") and "planta3.pdf" in adj and "planta3.dxf" in adj
    enlaces = {e: _publico(m["texto"]) for e, m in correos.items()}
    # las mediciones ya no se cambian
    assert client.put(f"/api/capex/partidas/{nuevas[0]['id']}", headers=admin, json={"cantidad": 5}).status_code == 400

    # portal ciego: sin precios base; se registra la apertura
    ruta, k = enlaces["capexa@inversiete.com"]
    datos = client.get(ruta, params={"k": k}).json()
    assert datos["abierta"] and len(datos["partidas"]) == 4 and all("precio" not in x for x in datos["partidas"])
    assert client.get(ruta, params={"k": "otra"}).status_code == 404
    ids_p = {x["codigo"] + ("b" if x["capitulo"] == "03" else ""): str(x["id"]) for x in datos["partidas"]}
    base = {"DEM010": 12.5, "DEM020": 13.5, "ALB030": 38.4}
    # A: completa y barata en demoliciones; borrador y después envío
    pa = {ids_p["DEM010"]: 8.0, ids_p["DEM020"]: 12.0, ids_p["ALB030"]: 41.0, ids_p["DEM020b"]: 14.0}
    assert client.post(ruta, params={"k": k}, json={"precios": pa}).json()["estado"] == "borrador"
    assert client.post(ruta, params={"k": k}, json={"precios": pa, "plazo_dias": 20, "enviar": True}).json()["estado"] == "enviada"
    # B: por Excel protegido (no vale el de otro contratista)
    ruta_b, kb = enlaces["capexb@inversiete.com"]
    xl = client.get(f"{ruta_b}/excel", params={"k": kb})
    wb = load_workbook(io.BytesIO(xl.content))
    ws = wb["Oferta"]
    assert ws.protection.sheet and ws.column_dimensions["A"].hidden
    precios_b = {"DEM010": 13.0, "DEM020": 13.0, "ALB030": 36.0, "DEM020b": 15.0}
    vistos = set()
    for fila in ws.iter_rows(min_row=7):
        if isinstance(fila[0].value, int):
            cod = fila[2].value
            clave = cod + ("b" if cod in vistos else "")
            vistos.add(cod)
            fila[6].value = precios_b[clave]
    out = io.BytesIO()
    wb.save(out)
    ruta_c, kc = enlaces["capexc@inversiete.com"]
    assert client.post(f"{ruta_c}/excel", params={"k": kc}, files={"fichero": ("o.xlsx", out.getvalue())}).status_code == 400
    r = client.post(f"{ruta_b}/excel", params={"k": kb}, files={"fichero": ("o.xlsx", out.getvalue())})
    assert r.status_code == 200 and len(r.json()["precios"]) == 4, r.text
    client.post(ruta_b, params={"k": kb}, json={"precios": r.json()["precios"], "enviar": True})
    # C: incompleta (sin el alicatado) y muy cara en demoliciones
    client.post(ruta_c, params={"k": kc}, json={"precios": {ids_p["DEM010"]: 25.0, ids_p["DEM020"]: 13.5}, "enviar": True})

    # comparativo con mapa de calor y PDF
    cmp = client.get(f"/api/capex/{p['id']}/comparativo", headers=admin).json()
    assert len(cmp["ofertas"]) == 3
    por = {o["proveedor"][-4:-3]: o for o in cmp["ofertas"]}
    assert por["A"]["completa"] and por["B"]["completa"] and por["C"]["faltan"] == 2
    dem = next(f for f in cmp["partidas"] if f["codigo"] == "DEM010")
    niveles = {o["proveedor"][-4:-3]: dem["ofertas"][str(o["bid_id"])]["nivel"] for o in cmp["ofertas"]}
    assert niveles == {"A": "anormal", "B": "ok", "C": "exceso"}
    assert dem["precio_base"] == base["DEM010"]
    pdf = client.get(f"/api/capex/{p['id']}/comparativo.pdf", headers=admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    lic = client.get(f"/api/capex/{p['id']}", headers=admin).json()["licitaciones"][0]
    assert lic["invitaciones"][0]["aperturas"] >= 1 and lic["invitaciones"][1]["descargas"] >= 1
    client.post(f"/api/capex/licitacion/{lic['id']}/cerrar", headers=admin)

    # adjudicación por capítulos: no a una oferta incompleta; A demoliciones, B el resto (> 30.000 €)
    r = client.post(f"/api/capex/{p['id']}/adjudicar", headers=admin, json={"lotes": [
        {"bid_id": por["C"]["bid_id"], "capitulos": ["02"]}]})
    assert r.status_code == 400 and "no tiene precio" in r.json()["detail"]
    r = client.post(f"/api/capex/{p['id']}/adjudicar", headers=admin, json={"lotes": [
        {"bid_id": por["A"]["bid_id"], "capitulos": ["01"]}, {"bid_id": por["B"]["bid_id"], "capitulos": ["01", "02"]}]})
    assert r.status_code == 400 and "dos veces" in r.json()["detail"]
    avisos.BANDEJA.clear()
    r = client.post(f"/api/capex/{p['id']}/adjudicar", headers=admin, json={"lotes": [
        {"bid_id": por["A"]["bid_id"], "capitulos": ["01"]},
        {"bid_id": por["B"]["bid_id"], "capitulos": ["02", "03"]}], "nota": "B garantiza el plazo de agosto"})
    assert r.status_code == 200, r.text
    r = r.json()
    assert r["requiere_aprobacion"] and r["proyecto"]["estado"] == "pendiente_aprobacion" and r["aprobacion"][0]["ok"]
    m = next(m for m in avisos.BANDEJA if m["para"] == "barbara@inversiete.es")
    assert m["urgente"] and m["adjuntos"][0][0].endswith(".pdf") and ">Aprobar</a>" in m["html"]
    aid, ka = re.search(r"/capex/aprobacion/(\d+)\?k=([\w-]+)", m["texto"]).groups()
    # el administrador no es la Dirección General
    assert client.post(f"/api/capex/{p['id']}/aprobacion", headers=admin, json={"decision": "aprobar"}).status_code == 403

    # decisión desde el correo, sin entrar en el PMS
    assert client.get(f"/capex/aprobacion/{aid}").status_code == 200
    ap = client.get(f"/api/publico/capex-aprobacion/{aid}", params={"k": ka}).json()
    assert ap["pendiente"] and ap["resumen"]["total"] == r["total"] and len(ap["resumen"]["lotes"]) == 2
    assert client.post(f"/api/publico/capex-aprobacion/{aid}", params={"k": ka},
                       json={"decision": "rechazar"}).status_code == 400  # sin motivo
    x = client.post(f"/api/publico/capex-aprobacion/{aid}", params={"k": ka}, json={"decision": "aprobar",
                                                                                    "comentario": "Adelante"}).json()
    assert x["estado"] == "aprobada"
    assert client.post(f"/api/publico/capex-aprobacion/{aid}", params={"k": ka},
                       json={"decision": "aprobar"}).status_code == 400  # ya resuelta
    fin = client.get(f"/api/capex/{p['id']}", headers=admin).json()
    assert fin["proyecto"]["estado"] == "adjudicado" and {a["estado"] for a in fin["adjudicaciones"]} == {"aprobada"}
    assert fin["aprobaciones"][0]["canal"] == "correo" and fin["aprobaciones"][0]["decidido_en"]
    # el contratista ya no puede tocar su oferta
    assert client.post(ruta, params={"k": k}, json={"precios": pa, "enviar": True}).status_code == 400


def test_capex_bajo_umbral_y_permisos(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    s = _proveedor(client, admin, "Carpintería Umbral SL", "umbral@inversiete.com")
    p = client.post("/api/capex", headers=admin, json={"asset_id": sae, "titulo": "Panelado de puertas"}).json()
    client.post(f"/api/capex/{p['id']}/partidas", headers=admin, json={
        "capitulo": "01", "capitulo_nombre": "CARPINTERÍA", "descripcion": "Panelado de puerta de entrada",
        "unidad": "ud", "cantidad": 10, "precio": 300})
    avisos.BANDEJA.clear()
    client.post(f"/api/capex/{p['id']}/licitacion", headers=admin, json={"proveedor_ids": [s], "fecha_limite": "2099-01-31"})
    ruta, k = _publico(avisos.BANDEJA[0]["texto"])
    iid = client.get(ruta, params={"k": k}).json()["partidas"][0]["id"]
    client.post(ruta, params={"k": k}, json={"precios": {str(iid): 280}, "enviar": True})
    cmp = client.get(f"/api/capex/{p['id']}/comparativo", headers=admin).json()
    r = client.post(f"/api/capex/{p['id']}/adjudicar", headers=admin,
                    json={"lotes": [{"bid_id": cmp["ofertas"][0]["bid_id"]}]}).json()
    assert not r["requiere_aprobacion"] and r["proyecto"]["estado"] == "adjudicado" and r["total"] == round(2800 * 1.19, 2)
    # recepción no ve CAPEX
    from test_pms import _new_user
    rec = _new_user(client, admin, "rec.capex@inversiete.com", [{"role_id": ids["roles"]["Recepción"], "asset_id": sae}])
    assert client.get(f"/api/capex/{p['id']}", headers=rec).status_code == 403
    assert client.get("/api/capex", headers=rec).json() == []
