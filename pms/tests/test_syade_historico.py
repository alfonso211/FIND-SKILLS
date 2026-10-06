"""Facturación del programa anterior (SYADE): lectura de sus listados, cuadre con totales, importación sin
duplicar y suma en la producción. Datos ficticios con las rarezas de la conversión PDF -> Excel."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from app import importacion_syade

CAB_AE = ["N.Factura", None, "Fecha", "Localizador", "Nif", "Cliente", None, "Base", "Tipo %", "Iva", "Total Fra.",
          "Fianza", None]


def _xlsx(hojas: list[list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for i, filas in enumerate(hojas):
        ws = wb.create_sheet(f"Table {i + 1}")
        for f in filas:
            ws.append(f)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def facturas() -> bytes:
    h1 = [CAB_AE,
          ["AE/", None, datetime(2030, 1, 31), 900001, "X0000001A", "CLIENTE UNO", None, 100.0, 10.0, 10.0, 110.0, 0.0],
          ["AE/", 1, datetime(2030, 1, 31), 900002, 12345678, "EMPRESA DOS", 200.0, None, 10.0, 20.0, 220.0, 300.0]]
    h2 = [CAB_AE[:7] + [None] + CAB_AE[7:],  # otra página: columnas desplazadas y cabecera repetida
          ["AE/", 2, datetime(2030, 2, 28), 900003, "Y0000003B", "CLIENTE TRES", None, None, 1000.5, 10.0, 100.05,
           1100.55, 0.0],
          CAB_AE,
          ["Reg.", None, 3, None, "Totales ......", "1.300,50      130,05      1.430,55      300,00"]]
    return _xlsx([h1, h2])


def test_lectura_facturas_con_rarezas():
    r = importacion_syade.leer(facturas(), 2030)
    assert r["tipo"] == "alojamiento" and r["cuadra"], r
    assert [f["numero"] for f in r["filas"]] == ["s/n 900001 31/01/2030", "1", "2"]
    assert r["filas"][1]["nif"] == "12345678" and str(r["filas"][2]["base"]) == "1000.50"
    assert any("sin número" in a for a in r["avisos"])


def test_lectura_abonos_servicios_y_fianzas():
    ab = _xlsx([[["N.Abono", None, "Fecha", "Localizador", "Nif", "Cliente", "Base", "Tipo %", "Iva", "Total Abono",
                  "Rectif. Fra. Núm"],
                 ["AB_AE", 7, datetime(2030, 2, 28), 900003, "Y0000003B", "CLIENTE TRES", -100.0, 10.0, -10.0, -110.0, 2]],
                [["AB_AE", "8  28/02/2030     900002", "12345678", "EMPRESA DOS", -50.0, 10.0, -5.0, -55.0, 1],
                 [None, "Reg.", 2, "Totales ......", -150.0, None, -15.0, -165.0]]])
    r = importacion_syade.leer(ab, 2030)
    assert r["tipo"] == "abono" and r["cuadra"] and len(r["filas"]) == 2
    assert r["filas"][1]["numero"] == "8" and r["filas"][1]["detalle"] == {"rectifica": "1"}

    se = _xlsx([[["N.Factura", None, "Fecha", "Loc.", "Nif", "Cliente", "Base", "Tipo %", "Iva", "Total"],
                 ["SE /", 5, "31/01/203", 900001, "X0000001A", "CLIENTE UNO", 10.0, 21.0, 2.1, 12.1],
                 ["Reg.          1", "Totales ......", 10.0, 2.1, 12.1]]])
    r = importacion_syade.leer(se, 2030)
    assert r["tipo"] == "servicio" and r["cuadra"] and r["filas"][0]["fecha"] == date(2030, 1, 31)
    assert any("último dígito del año" in a for a in r["avisos"])

    fz = _xlsx([[["Fecha Rec.", "Localizador", "Nombre", "Fianza Rec.", "Fianza Dev.", "Retenciones      Fecha Dev.",
                  None, "ObservFianza"],
                 [datetime(2029, 5, 1), 800001, "CLIENTE CUATRO", 300.0, 300.0, 50.0, datetime(2030, 2, 10), "nota"],
                 [None, None, "Totales ......", None, 300.0, 50.0]]])
    r = importacion_syade.leer(fz, 2030)
    assert r["tipo"] == "fianza_devuelta" and r["cuadra"] and r["filas"][0]["detalle"]["retenida"] == 50.0

    try:
        importacion_syade.leer(_xlsx([[["Otra", "cosa"]]]))
        raise AssertionError("debía rechazarlo")
    except ValueError:
        pass


def test_importacion_y_produccion(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]

    def subir(datos, confirmar):
        return client.post("/api/historico/importar", headers=admin,
                           files={"fichero": ("datos_syade.xlsx", datos, "application/octet-stream")},
                           data={"asset_id": str(sae), "confirmar": str(confirmar).lower(), "anio": "2030"})
    prev = subir(facturas(), False).json()
    assert prev["cuadra"] and prev["nuevas"] == 3 and client.get("/api/historico", headers=admin).json() == []
    assert subir(facturas(), True).json()["importado"]
    again = subir(facturas(), True).json()  # volver a importar no duplica
    assert again["nuevas"] == 0 and again["actualizadas"] == 3
    assert subir(b"no es un excel", False).status_code == 400

    res = client.get("/api/historico", headers=admin).json()
    meses = {m["mes"]: m for m in res[0]["meses"]}
    assert meses["2030-01"]["produccion"] == 300.0 and meses["2030-01"]["fianzas_cobradas"] == 300.0
    assert meses["2030-02"]["produccion"] == 1000.5

    # el informe de producción lo suma y lo indica aparte
    r = client.get("/api/informes/produccion", headers=admin,
                   params={"desde": "2030-01-01", "hasta": "2030-02-28", "asset_id": sae})
    ws = load_workbook(io.BytesIO(r.content))["Producción"]
    filas = [[c.value for c in row] for row in ws.iter_rows()]
    cab = next(f for f in filas if f and f[0] == "Activo")
    ene = next(f for f in filas if f and f[1] == "ene-2030")
    assert ene[cab.index("Alojamiento (base)")] == 300.0
    assert ene[cab.index("De ello, del programa anterior (base)")] == 300.0

    assert client.delete(f"/api/historico?asset_id={sae}", headers=admin).json()["borrados"] == 3
    assert client.get("/api/historico", headers=admin).json() == []
