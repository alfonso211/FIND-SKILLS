"""Fechas: las facturas emitidas llevan el último día del mes (criterio de la gestoría); las recibidas, su fecha y
su vencimiento tal como figuran."""
import io
from datetime import date

from openpyxl import load_workbook

from app.facturacion import fin_de_mes


def test_fin_de_mes():
    assert fin_de_mes(date(2026, 10, 6)) == date(2026, 10, 31)
    assert fin_de_mes(date(2028, 2, 3)) == date(2028, 2, 29) and fin_de_mes(date(2027, 2, 28)) == date(2027, 2, 28)
    assert fin_de_mes(date(2026, 12, 31)) == date(2026, 12, 31) and fin_de_mes(date(2026, 4, 1)) == date(2026, 4, 30)


def test_vencimiento_de_facturas_recibidas(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    base = {"asset_id": bab, "fecha": "2025-03-10", "categoria": next(iter(client.get(
        "/api/gastos/catalogos", headers=admin).json()["categorias"])), "concepto": "Factura de prueba",
        "total": 121, "tipo_iva": 21, "numero_factura": "PRV-2030-77", "forma_pago": "transferencia", "naturaleza": "OPEX"}
    assert client.post("/api/gastos", headers=admin, json={**base, "vencimiento": "2025-03-01"}).status_code == 400
    g = client.post("/api/gastos", headers=admin, json={**base, "vencimiento": "2025-04-09"})
    assert g.status_code == 201, g.text
    assert g.json()["fecha"] == "2025-03-10" and g.json()["vencimiento"] == "2025-04-09"
    xl = client.get("/api/gastos/excel", headers=admin, params={"asset_id": bab, "desde": "2025-03-01",
                                                              "hasta": "2025-03-31"})
    ws = load_workbook(io.BytesIO(xl.content))["Gastos"]
    cab = [c.value for c in ws[1]]
    fila = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[cab.index("Nº factura")] == "PRV-2030-77")
    assert fila[cab.index("Fecha factura")].date() == date(2025, 3, 10)
    assert fila[cab.index("Vencimiento")].date() == date(2025, 4, 9)
