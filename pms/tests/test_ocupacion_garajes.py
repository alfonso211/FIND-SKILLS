"""Las plazas de garaje no computan en la ocupación ni en la producción del edificio (se informan aparte)."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from conftest import domicilio_fiscal

HOY = date.today()


def _plaza(client, h, asset_id, codigo):
    return next(u for u in client.get(f"/api/unidades?asset_id={asset_id}&q={codigo}", headers=h).json()
                if u["codigo"] == codigo)


def _panel(client, h):
    return {a["codigo"]: a for a in client.get("/api/panel", headers=h).json()["activos"]}["SAE"]


def _hoja(client, h, informe, hoja, asset_id):
    wb = load_workbook(BytesIO(client.get(f"/api/informes/{informe}?desde={HOY.replace(day=1)}&hasta={HOY}"
                                          f"&asset_id={asset_id}", headers=h).content))
    filas = list(wb[hoja].iter_rows(values_only=True))
    cab = filas[0] if filas[0][0] == "Activo" else filas[1]
    return [dict(zip(cab, f)) for f in filas[filas.index(cab) + 1:] if f[0] == "Suite Aeropuerto"]


def test_garajes_no_computan(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    domicilio_fiscal(client, admin)
    antes, ocup_antes = _panel(client, admin), _hoja(client, admin, "ocupacion", "Turísticos", sae)[0]
    prod_antes = _hoja(client, admin, "produccion", "Producción", sae)[0]

    # una plaza reservada por días (cobrada) y otra alquilada por meses (recibo cobrado)
    p1, p2 = _plaza(client, admin, sae, "S1-120"), _plaza(client, admin, sae, "EXT-30")
    r = client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": p1["id"], "guest": {"nombre": "Solo Garaje"}, "fecha_entrada": HOY.isoformat(),
        "fecha_salida": (HOY + timedelta(days=3)).isoformat(), "importe_total": 36.3, "importe_pagado": 36.3})
    assert r.status_code == 201, r.text
    l = client.post("/api/garajes/contratos", headers=admin, json={
        "unit_id": p2["id"], "cliente": {"nombre": "Externo Ocupación", "telefono": "600100200"},
        "fecha_inicio": HOY.replace(day=1).isoformat(), "renta_mensual": 50, "dia_pago": 1,
        "fecha_fin": ((HOY.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)).isoformat()})
    assert l.status_code == 201, l.text
    rec = l.json()["proximo_recibo"]
    assert client.post(f"/api/garajes/recibos/{rec['id']}/cobro", headers=admin,
                       json={"importe": rec["pendiente"]}).status_code == 200

    despues = _panel(client, admin)
    assert despues["ocupacion_hoy"] == antes["ocupacion_hoy"]  # la ocupación es solo de apartamentos
    assert sum(despues["estados"].values()) == sum(antes["estados"].values()) == despues["unidades"]
    assert despues["garajes_ocupados"] == antes["garajes_ocupados"] + 2  # se informan aparte
    assert despues["produccion_mes"] == antes["produccion_mes"]  # ni la reserva ni el recibo de la plaza
    assert round(despues["garajes_facturado_mes"] - antes.get("garajes_facturado_mes", 0), 2) == round(30 + 50, 2)

    ocup = _hoja(client, admin, "ocupacion", "Turísticos", sae)[0]
    assert ocup["Noches ocupadas"] == ocup_antes["Noches ocupadas"] and "Uso" not in ocup
    gar = _hoja(client, admin, "ocupacion", "Garajes", sae)[0]
    assert gar["Días-plaza por reservas"] >= 1 and gar["Días-plaza alquilados por meses"] >= 1
    prod = _hoja(client, admin, "produccion", "Producción", sae)[0]
    assert prod["Producción del edificio (base, sin garajes)"] == prod_antes["Producción del edificio (base, sin garajes)"]
    assert round(prod["Plazas de garaje (base, aparte)"] - prod_antes["Plazas de garaje (base, aparte)"], 2) == 80
    assert round(prod["Total facturado"] - prod_antes["Total facturado"], 2) == round(36.3 + 60.5, 2)
