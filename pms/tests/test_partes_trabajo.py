"""Parte de trabajo diario de mantenimiento y limpieza: OT terminadas y limpiezas hechas entran solas, las
actuaciones sin parte se anotan, recepción valida (antes no se imprime) y queda archivado; informe quincenal."""
import io
from datetime import date

from openpyxl import load_workbook

from app import partes_trabajo as pt
from test_adjuntos_ot import _usuario

HOY = date.today()


def test_parte_diario_validacion_e_informe(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    tec = _usuario(client, admin, ids, "tecnico.partes@inversiete.com", "Técnico Mantenimiento", "SFL")
    lim = _usuario(client, admin, ids, "limpieza.partes@inversiete.com", "Gobernanta / Limpieza", "SFL")
    rec = _usuario(client, admin, ids, "recepcion.partes@inversiete.com", "Recepción", "SFL")
    hoy = HOY.isoformat()
    # OT terminada hoy: entra sola en el parte de mantenimiento
    w = client.post("/api/mantenimiento/ordenes", headers=admin, json={"asset_id": sfl, "titulo": "Grifo que gotea"}).json()
    client.post(f"/api/mantenimiento/ordenes/{w['id']}/confirmar-mantenimiento", headers=admin,
                json={"solucion": "Cambio de cartucho"})
    # tarea sin OT anotada por el técnico
    p = client.post("/api/partes-trabajo/actuaciones", headers=tec, json={
        "asset_id": sfl, "area": "mantenimiento", "fecha": hoy, "descripcion": "Revisión del grupo de presión",
        "ubicacion": "Sala de máquinas", "horas": 1.5}).json()
    refs = {x["ref"]: x for x in p["lineas"]}
    assert f"OT-{w['id']:05d}" in refs and refs[f"OT-{w['id']:05d}"]["detalle"] == "Cambio de cartucho"
    anotada = next(x for x in p["lineas"] if x["origen"] == "actuacion")
    assert anotada["persona"] == "tecnico.partes" and not p["puede_validar"]
    # limpieza: la hecha hoy y la zona común según el programa
    u = next(x for x in client.get(f"/api/unidades?asset_id={sfl}", headers=admin).json() if x["uso"] != "garaje")
    t = client.post("/api/limpieza/extra", headers=admin, json={"asset_id": sfl, "unit_id": u["id"], "fecha": hoy}).json()
    client.post(f"/api/limpieza/{t['id']}/hecha", headers=admin)
    assert client.post("/api/partes-trabajo/actuaciones", headers=lim, json={
        "asset_id": sfl, "area": "mantenimiento", "fecha": hoy, "descripcion": "no es su área"}).status_code == 403
    pl = client.post("/api/partes-trabajo/actuaciones", headers=lim, json={
        "asset_id": sfl, "area": "limpieza", "fecha": hoy, "descripcion": "Limpieza de portal y escaleras (programa)",
        "ubicacion": "Zonas comunes · Portal A"}).json()
    assert {x["origen"] for x in pl["lineas"]} >= {"limpieza", "actuacion"}

    # sin validar no se imprime; valida recepción y queda archivado
    q = {"asset_id": sfl, "area": "mantenimiento", "fecha": hoy}
    assert client.get("/api/partes-trabajo/descargar", headers=rec, params=q).status_code == 400
    assert client.post("/api/partes-trabajo/validar", headers=tec, json=q).status_code == 403
    v = client.post("/api/partes-trabajo/validar", headers=rec, json={**q, "observaciones": "Todo correcto"}).json()
    assert v["estado"] == "validado" and v["validado_por"] == "recepcion.partes"
    pdf = client.get("/api/partes-trabajo/descargar", headers=tec, params=q)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    xl = load_workbook(io.BytesIO(client.get("/api/partes-trabajo/descargar", headers=rec, params={**q, "formato": "xlsx"}).content))
    assert any("Revisión del grupo de presión" in str(c.value) for fila in xl.active.iter_rows() for c in fila)
    docs = client.get(f"/api/documentos?asset_id={sfl}&tipo=parte_trabajo", headers=admin)
    if docs.status_code == 200:
        assert any(d["nombre"].startswith(f"Parte_mantenimiento_SFL_{hoy}") for d in docs.json())
    # validado: no se añade nada hasta reabrirlo
    assert client.post("/api/partes-trabajo/actuaciones", headers=tec, json={
        "asset_id": sfl, "area": "mantenimiento", "fecha": hoy, "descripcion": "Otra tarea"}).status_code == 400
    assert client.post("/api/partes-trabajo/reabrir", headers=rec, json=q).json()["estado"] == "abierto"
    client.post("/api/partes-trabajo/validar", headers=rec, json=q)

    # informe quincenal
    qn = 1 if HOY.day <= 15 else 2
    inf = client.get("/api/partes-trabajo/quincenal", headers=rec,
                     params={"asset_id": sfl, "anio": HOY.year, "mes": HOY.month, "quincena": qn}).json()
    m, li = inf["areas"]["mantenimiento"], inf["areas"]["limpieza"]
    assert m["por_origen"]["ot"] >= 1 and m["por_origen"]["actuacion"] == 1 and m["horas_anotadas"] == 1.5
    assert li["total"] >= 2 and hoy in li["sin_validar"] and hoy not in m["sin_validar"]
    r = client.get("/api/partes-trabajo/quincenal/descargar", headers=rec, params={
        "asset_id": sfl, "anio": HOY.year, "mes": HOY.month, "quincena": qn, "archivar": True})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert client.get("/api/partes-trabajo/quincenal", headers=rec, params={
        "asset_id": sfl, "anio": HOY.year, "mes": HOY.month, "quincena": qn}).json()["archivado"]


def test_quincenas():
    assert pt.quincena(2026, 2, 2) == (date(2026, 2, 16), date(2026, 2, 28))
    assert pt.quincena(2026, 12, 2) == (date(2026, 12, 16), date(2026, 12, 31))
    assert pt.quincena_anterior(date(2026, 10, 8)) == (2026, 9, 2)
    assert pt.quincena_anterior(date(2026, 10, 20)) == (2026, 10, 1)
