"""1.70: coste de la mano de obra propia (coste/hora de la ficha × horas de OT y del parte diario) en el informe de
mantenimiento, por OT, por técnico y por apartamento. El coste/hora solo lo ven el Director Técnico y dirección."""
from io import BytesIO

from openpyxl import load_workbook

from conftest import trabajo_realizado
from test_pms import _new_user


def _hojas(r):
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    return {ws.title: [list(f) for f in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def test_coste_mano_obra_en_informe(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    p = client.post("/api/personal", headers=admin, json={
        "asset_id": sfl, "area": "mantenimiento", "nombre": "ZZ-PRUEBA Técnico Coste", "telefono": "600999111",
        "coste_hora": 25})
    assert p.status_code in (200, 201), p.text
    assert p.json()["coste_hora"] == 25
    unit = client.get(f"/api/unidades?asset_id={sfl}", headers=admin).json()
    unit = next(u for u in unit if u["uso"] == "apartamento")
    w = client.post("/api/mantenimiento/ordenes", headers=admin, json={
        "asset_id": sfl, "unit_id": unit["id"], "titulo": "ZZ-PRUEBA grifo", "asignacion": "propio",
        "personal_id": p.json()["id"]}).json()
    r = trabajo_realizado(client, admin, w["id"], "Cambio de grifo", coste_real=40, minutos=120)
    assert r.json()["estado"] == "cerrada", r.text

    h = _hojas(client.get(f"/api/informes/mantenimiento?asset_id={sfl}", headers=admin))
    cab = h["Órdenes de trabajo"][0]
    fila = next(f for f in h["Órdenes de trabajo"] if f[0] == f"OT-{w['id']:05d}")
    assert fila[cab.index("Horas")] == 2 and fila[cab.index("Mano de obra propia")] == 50
    assert fila[cab.index("Coste total")] == 90
    tec = next(f for f in h["Horas por técnico"] if f[0] == "ZZ-PRUEBA Técnico Coste")
    cab = h["Horas por técnico"][0]
    assert tec[cab.index("Horas OT")] == 2 and tec[cab.index("Coste/hora")] == 25
    assert tec[cab.index("Coste mano de obra")] == 50
    apt = next(f for f in h["Por apartamento"] if f[1] == unit["codigo"])
    cab = h["Por apartamento"][0]
    assert apt[cab.index("Mano de obra propia")] >= 50 and apt[cab.index("Horas OT")] >= 2

    # recepción ve las horas, no el coste/hora del personal
    rec = _new_user(client, admin, "rec.coste170@inversiete.com",
                    [{"role_id": ids["roles"]["Recepción"], "asset_id": sfl}])
    h = _hojas(client.get(f"/api/informes/mantenimiento?asset_id={sfl}", headers=rec))
    assert "Horas" in h["Órdenes de trabajo"][0] and "Mano de obra propia" not in h["Órdenes de trabajo"][0]
    assert "Coste/hora" not in h["Horas por técnico"][0]
    assert "coste_hora" not in next(x for x in client.get("/api/personal", headers=rec).json()
                                    if x["id"] == p.json()["id"])
