"""Facturas recibidas: retención de IRPF (se descuenta del pago y se exporta) y pago retenido (no pagar de momento),
con aviso a quien paga (dirección, administración) y fecha de revisión."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app import avisos
from app import export_invergestion as ex
from app.database import SessionLocal
from test_adjuntos_ot import _usuario

HOY = date.today()


def test_retencion_irpf(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    base = {"asset_id": bab, "fecha": HOY.isoformat(), "categoria": "mantenimiento", "concepto": "Proyecto técnico",
            "proveedor": "Ingeniero Retención", "numero_factura": "IR-1", "forma_pago": "transferencia", "naturaleza": "CAPEX"}
    # factura de profesional: base 1.000 + IVA 210 − IRPF 15 % 150 = 1.060 (lo que figura y se paga)
    g = client.post("/api/gastos", headers=admin, json={**base, "total": 1060, "tipo_iva": 21,
                                                        "retencion_tipo": "irpf_profesional", "retencion_pct": 15})
    assert g.status_code == 201, g.text
    g = g.json()
    assert (g["base"], g["cuota"], g["retencion"], g["total"], g["liquido"]) == (1000, 210, 150, 1210, 1060)
    assert g["retencion_nombre"] == "IRPF profesionales"
    # con la base indicada (varios tipos de IVA) la cuota sale de total + retención − base
    otra = {**base, "numero_factura": "IR-2", "total": 1060, "base": 1000, "retencion_tipo": "irpf_profesional",
            "retencion_pct": 15}
    parecida = client.post("/api/gastos", headers=admin, json=otra)  # mismo proveedor, importe y fecha
    assert parecida.status_code == 409 and parecida.json()["detail"].startswith("Posible duplicado")
    g2 = client.post("/api/gastos", headers=admin, json={**otra, "confirmar_duplicado": True}).json()
    assert (g2["cuota"], g2["retencion"], g2["liquido"]) == (210, 150, 1060)
    assert client.post("/api/gastos", headers=admin, json={**base, "total": 100, "retencion_tipo": "otra"}
                       ).status_code == 400  # sin %
    assert client.post("/api/gastos", headers=admin, json={**base, "total": 100, "retencion_tipo": "inventada",
                                                           "retencion_pct": 5}).status_code == 400
    # editar sin tocar el total (lo que figura) conserva los importes
    e = client.put(f"/api/gastos/{g['id']}", headers=admin, json={**g, "total": g["liquido"]}).json()
    assert (e["base"], e["retencion"], e["liquido"]) == (1000, 150, 1060)
    cat = client.get("/api/gastos/catalogos", headers=admin).json()["retenciones"]
    assert cat["irpf_arrendamiento"] == {"nombre": "IRPF arrendamiento de inmuebles", "pct": 19}

    # exportación: retención en su columna y total = base + IVA − retención
    with SessionLocal() as db:
        filas, _, avs = ex.recibidas(db, {bab}, HOY, HOY)
    fila = next(f for f in filas if f["numero"] == "IR-1")
    assert (fila["base"], fila["cuota_iva"], fila["retencion"], fila["total"]) == ("1000.00", "210.00", "150.00",
                                                                                  "1060.00")
    assert not any("IR-1" in a and "no cuadra" in a for a in avs)

    xl = client.get("/api/gastos/excel", headers=admin, params={"asset_id": bab, "desde": HOY.isoformat(),
                                                               "hasta": HOY.isoformat()})
    ws = load_workbook(io.BytesIO(xl.content))["Gastos"]
    cab = [c.value for c in ws[1]]
    r = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[cab.index("Nº factura")] == "IR-1")
    assert (r[cab.index("Importe retención")], r[cab.index("A pagar")], r[cab.index("Retención %")]) == (150, 1060, 15)


def test_pago_retenido(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    rec = _usuario(client, admin, ids, "recepcion.retenida@inversiete.com", "Recepción", "SAE")
    otra = _usuario(client, admin, ids, "recepcion.retenida2@inversiete.com", "Recepción", "SAE")
    base = {"asset_id": sae, "fecha": HOY.isoformat(), "categoria": "mantenimiento", "concepto": "Reparación bomba",
            "proveedor": "Bombas Retenidas SL", "numero_factura": "BR-9", "total": 242, "forma_pago": "transferencia", "naturaleza": "OPEX"}
    assert client.post("/api/gastos", headers=rec, json={**base, "retener_pago": True}).status_code == 400
    assert client.post("/api/gastos", headers=rec, json={
        **base, "retener_pago": True, "retener_motivo": "Trabajo sin terminar",
        "retener_revision": (HOY - timedelta(days=1)).isoformat()}).status_code == 400

    avisos.BANDEJA.clear()
    g = client.post("/api/gastos", headers=rec, json={
        **base, "retener_pago": True, "retener_motivo": "Trabajo sin terminar: falta la prueba de presión",
        "retener_revision": HOY.isoformat()})
    assert g.status_code == 201, g.text
    g = g.json()
    assert g["pago_retenido"] and g["pago_retenido_por"] == "recepcion.retenida"
    # aviso al momento a quien paga (dirección del grupo), no a recepción
    para = {m["para"] for m in avisos.BANDEJA}
    assert "alfonso@inversiete.es" in para and "recepcion.retenida2@inversiete.com" not in para
    m = next(m for m in avisos.BANDEJA if m["para"] == "alfonso@inversiete.es")
    assert m["asunto"].startswith("PAGO RETENIDO · Suite Aeropuerto · Bombas Retenidas SL BR-9")
    assert "falta la prueba de presión" in m["texto"] and "REVISAR HOY" in m["texto"]

    # retenida: no se puede marcar pagada
    pagar = client.put(f"/api/gastos/{g['id']}", headers=rec, json={**g, "total": g["liquido"], "pagado": True})
    assert pagar.status_code == 400 and "retenido" in pagar.json()["detail"]

    # resumen diario de quien paga y panel
    with SessionLocal() as db:
        datos = avisos._datos_resumen(db, HOY)
    fila = next(f for _, f in datos["pagos_retenidos"] if f[2] == "BR-9")
    assert fila[-1] == "REVISAR HOY" and fila[6] == "recepcion.retenida"
    panel = next(a for a in client.get("/api/panel", headers=admin).json()["activos"] if a["id"] == sae)
    assert panel["pagos_retenidos"] >= 1 and panel["pagos_retenidos_revisar"] >= 1
    assert "pagos_retenidos" not in next(a for a in client.get("/api/panel", headers=rec).json()["activos"]
                                         if a["id"] == sae)
    t = client.get("/api/gastos", headers=admin, params={"asset_id": sae}).json()["totales"]
    assert t["retenidas"] >= 1 and t["retenidas_importe"] >= 242

    # cambiar la fecha de revisión vuelve a avisar; liberar: quien paga o quien la retuvo, no otra recepción
    avisos.BANDEJA.clear()
    r = client.post(f"/api/gastos/{g['id']}/retener-pago", headers=rec,
                    json={"motivo": "Pendiente de la prueba de presión", "revision": (HOY + timedelta(days=7)).isoformat()})
    assert r.status_code == 200 and r.json()["pago_retenido_revision"] == (HOY + timedelta(days=7)).isoformat()
    assert any(m["para"] == "alfonso@inversiete.es" for m in avisos.BANDEJA)
    assert client.post(f"/api/gastos/{g['id']}/liberar-pago", headers=otra, json={}).status_code == 403
    lib = client.post(f"/api/gastos/{g['id']}/liberar-pago", headers=admin, json={"nota": "Prueba superada"})
    assert lib.status_code == 200 and lib.json()["pago_retenido"] is False
    assert client.post(f"/api/gastos/{g['id']}/liberar-pago", headers=admin, json={}).status_code == 400
    pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
    ok = client.post(f"/api/gastos/{g['id']}/pagar", headers=rec, files=[("ficheros", ("transferencia.pdf", pdf, "application/pdf"))])
    assert ok.status_code == 200 and ok.json()["pagado"]
    assert client.post(f"/api/gastos/{g['id']}/retener-pago", headers=rec,
                       json={"motivo": "tarde", "revision": HOY.isoformat()}).status_code == 400  # ya pagada
