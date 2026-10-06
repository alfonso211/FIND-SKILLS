"""Carga de la ocupación actual exportada del PMS anterior y aviso de estancias vencidas."""
import io
from datetime import date, timedelta

from openpyxl import Workbook

from app import avisos, importacion_ocupacion as io_ocu
from app.database import SessionLocal

HOY = date.today()
f = lambda d: d.strftime("%d/%m/%Y")  # noqa: E731


def _listado(filas) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 3"
    ws.append(["R", "Reserva"] + [None] * 10)
    ws.append(["Localizador", None, "Blo    Por", "Esc    Pla", "Núm  Tip", "Dorm", "C/Mat  C/Ind", "Sót", "Plaza",
               "FEntrada       FSalida     Noches", "Ocupante", "Teléfonos"])
    for x in filas:
        ws.append(x)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _importar(client, h, sae, datos, confirmar):
    return client.post("/api/turistico/importar-ocupacion", headers=h, data={"asset_id": sae, "confirmar": str(confirmar).lower()},
                       files={"fichero": ("ocupacion.xlsx", datos)}).json()


def test_nombres():
    assert io_ocu.nombre_y_apellidos("ICSEL VICTORIA TOLEDO SANZONETTI") == ("Icsel Victoria", "Toledo Sanzonetti")
    assert io_ocu.nombre_y_apellidos("RASHMI MOORTHI .") == ("Rashmi", "Moorthi")
    assert io_ocu.nombre_y_apellidos("VIAJES ECUADOR S.L.") == ("VIAJES ECUADOR S.L.", None)
    assert io_ocu.telefono("643808631-691445957") == "643808631"


def test_importar_ocupacion_y_vencidas(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    vencida = (HOY - timedelta(days=40), HOY - timedelta(days=10))
    vigente = (HOY - timedelta(days=5), HOY + timedelta(days=2))
    futura = (HOY + timedelta(days=9), HOY + timedelta(days=20))
    datos = _listado([
        ["A", 91001, "A", "4º", 450, 1, "1        0", 1, 177, f"{f(vencida[0])}{f(vencida[1])}     30", "PEDRO PRUEBA PEREZ", 611000111],
        ["A", 91002, "A", "4º", "451  Est", 0, "1        0", 0, 66, f"{f(vigente[0])}{f(vigente[1])}     7", "ROSA MARIA VIGENTE .", None],
        ["R", 91003, "A", "4º", 452, 1, "1        0", None, None, f"{f(futura[0])}{f(futura[1])}     11", "LUIS FUTURO", None],
        ["A", 91004, "A", "4º", 999, 1, "1        0", None, None, f"{f(vigente[0])}{f(vigente[1])}     7", "NO EXISTE", None],
    ])
    prev = _importar(client, admin, sae, datos, False)
    assert (prev["nuevas"], prev["errores"], prev["vencidas"]) == (3, 1, 1), [x.get("motivo") for x in prev["filas"]]
    assert "A-999 no existe" in [x for x in prev["filas"] if x["estado"] == "error"][0]["motivo"]
    assert not client.get(f"/api/turistico/reservas?asset_id={sae}&q=91001", headers=admin).json()  # solo vista previa
    r = _importar(client, admin, sae, datos, True)
    assert (r["nuevas"], r["garajes"]) == (3, 2)
    res = {x["localizador"]: x for x in client.get(f"/api/turistico/reservas?asset_id={sae}&q=9100", headers=admin).json()}
    assert res["91001"]["estado"] == "checkin" and res["91001"]["huesped"] == "Pedro Prueba Perez"
    assert res["91002"]["estado"] == "checkin" and res["91003"]["estado"] == "confirmada"
    assert res["91001-G"]["unidad"] == "S1-177" and res["91002-G"]["unidad"] == "EXT-66"
    # volver a importar no duplica
    r2 = _importar(client, admin, sae, datos, True)
    assert (r2["nuevas"], r2["actualizadas"]) == (0, 3)

    # estancia vencida: sigue ocupando el plano, con su marca, y aparece en el aviso
    p = client.get(f"/api/plano/{sae}", headers=admin).json()
    celda = next(c for pl in p["plantas"] for c in pl["celdas"] if c.get("codigo") == "A-450")
    assert celda["estado"] == "alquilado" and celda["vencida"]
    v = client.get(f"/api/turistico/vencidas?asset_id={sae}", headers=admin).json()
    venc = next(x for x in v["vencidas"] if x["localizador"] == "91001")
    assert venc["dias"] == 10 and venc["whatsapp"].startswith("https://wa.me/34611000111?text=")
    assert any(x["localizador"] == "91002" for x in v["proximas"])
    panel = {a["codigo"]: a for a in client.get("/api/panel", headers=admin).json()["activos"]}
    assert panel["SAE"]["estancias_vencidas"] >= 1
    with SessionLocal() as db:
        assert any(fila[1] == "A-450" and fila[5].startswith("VENCIDA")
                   for _, fila in avisos._datos_resumen(db, HOY)["estancias_vencidas"])

    # ampliar la estancia mueve también la plaza de garaje; el check-out la libera
    nueva = (HOY + timedelta(days=20)).isoformat()
    assert client.put(f"/api/turistico/reservas/{res['91001']['id']}", headers=admin,
                      json={"fecha_salida": nueva}).status_code == 200
    g = client.get(f"/api/turistico/reservas?asset_id={sae}&q=91001-G", headers=admin).json()[0]
    assert g["fecha_salida"] == nueva
    assert not any(x["localizador"] == "91001" for x in
                   client.get(f"/api/turistico/vencidas?asset_id={sae}", headers=admin).json()["vencidas"])
    client.post(f"/api/turistico/reservas/{res['91002']['id']}/checkout", headers=admin)
    g2 = client.get(f"/api/turistico/reservas?asset_id={sae}&q=91002-G", headers=admin).json()[0]
    assert g2["estado"] == "checkout"


def test_fichero_no_valido(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    wb = Workbook()
    b = io.BytesIO()
    wb.save(b)
    r = client.post("/api/turistico/importar-ocupacion", headers=admin, data={"asset_id": sae},
                    files={"fichero": ("x.xlsx", b.getvalue())})
    assert r.status_code == 400


def test_formato_suite_florida(client, admin, ids):
    """Listado de Suite Florida: portal y «planta letra»; las columnas cambian de sitio según la página."""
    sfl = ids["assets"]["SFL"]["id"]
    futura = (HOY + timedelta(days=700), HOY + timedelta(days=730))  # lejos: no choca con otras pruebas
    fechas = f"{f(futura[0])}{f(futura[1])}     30"
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws.append(["Localizador", None, "Blo   Por", "Esc   Pla   Let  Tip", "Dorm  C/Mat C/Ind", "Sót", "Plaza",
               "FEntrada     FSalida    Noches", "Ocupante", "Teléfonos"])
    ws.append(["A", 92001, 1, "1º     J", "2        1        2", 2, 14, fechas, "ANA PRUEBA FLORIDA", 600000001])
    ws.append([None, None, 1, "1º     K", "1        1        0"] + [None] * 5)  # apartamento libre: se ignora
    ws2 = wb.create_sheet("Table 4")
    ws2.append(["Localizador", None, "Blo   Por", "Esc   Pla   Let  Tip", None, "Dorm  C/Mat C/Ind", None, "Sót",
                "Plaza", "FEntrada     FSalida    Noches", None, "Ocupante", None, "Teléfonos"])
    ws2.append(["R", 92002, 2, "1º     G", None, "1        1        0", None, 1, 189, fechas, None, "LUIS PRUEBA .",
                None, "600000002"])
    b = io.BytesIO()
    wb.save(b)
    filas = io_ocu.leer(b.getvalue())
    assert [(x["unidad"], x.get("garaje"), x["ocupante"], x["telefono"]) for x in filas] == [
        ("P1-1J", "S2-14", "ANA PRUEBA FLORIDA", "600000001"), ("P2-1G", "S1-189", "LUIS PRUEBA .", "600000002")]
    r = _importar(client, admin, sfl, b.getvalue(), True)
    assert (r["nuevas"], r["errores"], r["garajes"]) == (2, 0, 2), [x.get("motivo") for x in r["filas"]]
