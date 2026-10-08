"""Cobros de alquiler de otro programa: lectura de las facturas (PDF o ZIP), fichas de inquilinos, un ingreso
cobrado por mes, documento por inquilino, Excel de ingresos y producción del activo. Datos ficticios."""
import io
import zipfile
from datetime import date

from openpyxl import load_workbook

from app import cobros_alquiler as ca


def _factura(numero, fecha, nombre, mes, total, nif=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    lineas = ["Factura", "Simplificada", f"Número # {numero}", f"Fecha {fecha}", "EMPRESA FICTICIA S.A.",
              "Calle Falsa, 1", "Madrid (28000), Madrid, España", "A00000000", "info@ejemplo.es", nombre,
              *([nif] if nif else []), "Calle Inventada 35", "madrid (28000), Madrid, España",
              "CONCEPTO PRECIO UNIDADES SUBTOTAL TOTAL", "Alquilervivienda", mes,
              f"{total} 1 {total} {total}", f"BASE IMPONIBLE {total}", f"TOTAL {total}", "Condiciones de pago"]
    for x in lineas:
        c.drawString(50, y, x)
        y -= 18
    c.save()
    return buf.getvalue()


def test_leer_factura():
    f = ca.leer_factura(_factura("T269001", "05/06/2026", "Persona Ficticia Uno", "JUNIO 2026", "675,00€"))
    assert f == {"numero": "T269001", "fecha": "2026-06-05", "cliente": "Persona Ficticia Uno", "nif": None,
                 "direccion": "Calle Inventada 35", "total": 675.0, "mes": "2026-06"}
    g = ca.leer_factura(_factura("T269002", "31/08/2026", "Otra Persona Ficticia", "ALQUILER SEPTIEMBRE", "1.810,50€",
                                 nif="12345678Z"))
    assert g["mes"] == "2026-09" and g["total"] == 1810.5 and g["nif"] == "12345678Z"
    assert ca.meses("2025-11", "2026-02") == ["2025-11", "2025-12", "2026-01", "2026-02"]


def test_importar_cobros(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("a.pdf", _factura("T269011", "05/06/2026", "Inquilina Prueba Cobros", "JUNIO 2026", "700,00€"))
        zf.writestr("b.pdf", _factura("T269012", "10/06/2026", "Inquilino Segundo Cobros", "JUNIO 2026", "650,00€"))
        zf.writestr("leeme.txt", "no es un pdf")
    r = client.post("/api/cobros-alquiler/leer", headers=admin, data={"asset_id": bab},
                    files=[("ficheros", ("pagos.zip", z.getvalue(), "application/zip"))])
    assert r.status_code == 200, r.text
    fs = r.json()["facturas"]
    assert len(fs) == 2 and not any(f["ficha"] for f in fs)
    sfl = ids["assets"]["SFL"]["id"]
    assert client.post("/api/cobros-alquiler/leer", headers=admin, data={"asset_id": sfl},
                       files=[("ficheros", ("p.zip", z.getvalue(), "application/zip"))]).status_code == 400
    filas = [{**f, "desde": "2026-01", "hasta": "2026-10"} for f in fs]
    r = client.post("/api/cobros-alquiler/importar", headers=admin, json={"asset_id": bab, "filas": filas}).json()
    assert r == {"inquilinos": 2, "cobros": 20, "fichas_nuevas": 2}
    # repetir no duplica
    r = client.post("/api/cobros-alquiler/importar", headers=admin, json={"asset_id": bab, "filas": filas}).json()
    assert r["cobros"] == 20 and r["fichas_nuevas"] == 0
    s = client.get("/api/cobros-alquiler", headers=admin, params={"asset_id": bab, "anio": 2026}).json()
    mias = [x for x in s["inquilinos"] if x["cliente"] in ("Inquilina Prueba Cobros", "Inquilino Segundo Cobros")]
    assert len(mias) == 2 and all(x["meses"] == 10 and x["contact_id"] for x in mias)
    assert {x["total"] for x in mias} == {7000.0, 6500.0}
    pdf = client.get(f"/api/cobros-alquiler/{mias[0]['clave']}/pdf", headers=admin, params={"asset_id": bab})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(client.get("/api/cobros-alquiler/excel/ingresos", headers=admin,
                                             params={"asset_id": bab, "anio": 2026}).content))
    assert wb.active.max_row >= 22 and wb.active["A2"].value.date() == date(2026, 1, 5)
    # la ficha del inquilino ve sus cobros y cuentan en la producción del activo
    contacto = client.get(f"/api/clientes/{mias[0]['contact_id']}", headers=admin)
    if contacto.status_code == 200:
        assert contacto.json()["tipo"] == "inquilino"
    prod = client.get("/api/historico", headers=admin, params={"asset_id": bab}).json()
    marzo = next(m for a in prod for m in a["meses"] if m["mes"] == "2026-03")
    assert marzo["produccion"] >= 1350
    # borrar los cobros de un inquilino para corregir
    for x in mias:  # deja el activo sin cobros importados (otras pruebas cuentan el histórico)
        assert client.delete(f"/api/cobros-alquiler/{x['clave']}", headers=admin,
                             params={"asset_id": bab}).json()["borrados"] == 10
