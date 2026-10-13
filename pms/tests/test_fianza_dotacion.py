"""Devolución de fianza al hacer la salida, lista negra, expediente del cliente y dotación de los apartamentos."""
import io
import zipfile
from datetime import date, timedelta

import docx

from app import avisos

HOY = date.today()
COMPLETO = {"nombre": "Fausto", "apellidos": "Fianza Prueba", "documento_tipo": "DNI", "documento_num": "52345678Q",
            "num_soporte": "BAA123123", "nacionalidad": "España", "fecha_nacimiento": "1975-03-03", "sexo": "M",
            "telefono": "611222333", "email": "fausto@example.com", "direccion": "C/ Luna 1", "cp": "28013",
            "municipio": "Madrid", "pais": "España"}
IBAN = "ES9121000418450200051332"


def _reserva(client, h, asset_id, guest, loc, dias=(-3, 2)):
    """Reserva alojada (check-in) en el primer apartamento libre del activo."""
    unidades = [u for u in client.get(f"/api/unidades?asset_id={asset_id}", headers=h).json() if u["uso"] == "apartamento"]
    for u in reversed(unidades):
        r = client.post("/api/turistico/reservas", headers=h, json={
            "unit_id": u["id"], "localizador": loc, "fecha_entrada": (HOY + timedelta(days=dias[0])).isoformat(),
            "fecha_salida": (HOY + timedelta(days=dias[1])).isoformat(), "importe_total": 500, **guest})
        if r.status_code == 201:
            return r.json()
        assert "ya está reservada" in r.text or "no admite" in r.text, r.text
    raise AssertionError("sin apartamento libre")


def test_devolucion_fianza_transferencia(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    r = _reserva(client, admin, sfl, {"guest": COMPLETO}, "FZ-1")
    client.put(f"/api/turistico/reservas/{r['id']}", headers=admin, json={"estado": "checkin"})
    client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json={
        "fecha_firma": HOY.isoformat(), "fianza": 400, "motivo": ["laboral"], "acreditacion": ["dni"],
        "solo_guardar": True})
    info = client.get(f"/api/turistico/reservas/{r['id']}/fianza", headers=admin).json()
    assert info["fianza"] == 400 and info["procede"] and info["titular"] == "Fausto Fianza Prueba"
    # sin indicar la devolución no se puede hacer la salida
    assert client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin).status_code == 409
    mal = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin,
                      json={"fianza": {"forma": "transferencia", "iban": "ES0000000000000000000000"}})
    assert mal.status_code == 400 and "IBAN" in mal.json()["detail"]
    parcial = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin,
                          json={"fianza": {"forma": "transferencia", "iban": IBAN, "importe": 350}})
    assert parcial.status_code == 400 and "motivo" in parcial.json()["detail"]
    avisos.BANDEJA.clear()
    ok = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin, json={"fianza": {
        "forma": "transferencia", "iban": IBAN, "importe": 350, "retenido_motivo": "Vaso roto y mando perdido",
        "fecha": (HOY + timedelta(days=2)).isoformat()}})
    assert ok.status_code == 200, ok.text
    x = ok.json()
    assert x["estado"] == "checkout" and x["devolucion"]["estado"] == "pendiente" and x["devolucion"]["retenido"] == 50
    # correo a la Dirección (DG y DT) y WhatsApp preparado con los mismos datos
    para = sorted(m["para"] for m in avisos.BANDEJA)
    assert para == ["alfonso@inversiete.es", "barbara@inversiete.es"]
    assert "ES91 2100 0418 4502 0005 1332" in avisos.BANDEJA[0]["texto"]
    assert len(x["whatsapp"]) == 2 and all("wa.me" in w["url"] for w in x["whatsapp"])
    # queda en el registro de envíos
    reg = client.get("/api/admin/avisos", headers=admin).json()["registro"]
    assert sum(1 for e in reg if e["tipo"] == "fianza" and e["clave"] == f"fianza:{x['devolucion']['id']}") == 2
    # el IBAN queda en la ficha del cliente
    assert client.get(f"/api/terceros/{r['guest_id']}", headers=admin).json()["iban"] == IBAN
    # lista de devoluciones y transferencia hecha
    lst = client.get("/api/turistico/fianzas?estado=pendiente", headers=admin).json()
    assert any(d["id"] == x["devolucion"]["id"] for d in lst)
    d = client.post(f"/api/turistico/fianzas/{x['devolucion']['id']}/devuelta", headers=admin, json={}).json()
    assert d["estado"] == "devuelta" and d["devuelta_fecha"] == HOY.isoformat()
    # envío a gestión (INVERGESTION): FIANZAS_*.csv en el paquete
    q = {"desde": HOY.isoformat(), "hasta": HOY.isoformat(), "activo": "SFLORIDA", "pdf": "false"}
    comp = client.get("/api/exportacion/invergestion/comprobar", headers=admin, params=q).json()
    assert comp["fianzas"]["devoluciones"] >= 1
    z = client.get("/api/exportacion/invergestion", headers=admin, params=q)
    assert z.status_code == 200, z.text
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        nombre = next(n for n in zf.namelist() if n.startswith("FIANZAS_"))
        csv = zf.read(nombre).decode()
    assert "FZ-1" in csv and "TRANSFERENCIA" in csv and "350.00" in csv and "50.00" in csv and "DEVUELTA" in csv


def test_salida_efectivo_y_sin_fianza(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    a = _reserva(client, admin, sae, {"guest": {**COMPLETO, "documento_num": "11111111H", "nombre": "Efe"}}, "FZ-2")
    client.put(f"/api/turistico/reservas/{a['id']}", headers=admin, json={"estado": "checkin"})
    client.post(f"/api/turistico/reservas/{a['id']}/contrato", headers=admin, json={
        "fecha_firma": HOY.isoformat(), "fianza": 200, "motivo": ["turismo"], "acreditacion": ["dni"],
        "solo_guardar": True})
    avisos.BANDEJA.clear()
    x = client.post(f"/api/turistico/reservas/{a['id']}/checkout", headers=admin,
                    json={"fianza": {"forma": "efectivo"}}).json()
    assert x["devolucion"]["estado"] == "devuelta" and x["devolucion"]["importe"] == 200 and not avisos.BANDEJA
    # sin fianza, la salida es como siempre
    b = _reserva(client, admin, sae, {"guest": {"nombre": "Sin Fianza"}}, "FZ-3")
    client.put(f"/api/turistico/reservas/{b['id']}", headers=admin, json={"estado": "checkin"})
    assert client.post(f"/api/turistico/reservas/{b['id']}/checkout", headers=admin).status_code == 200


def test_lista_negra_y_expediente(client, admin, ids):
    sfl, sae = ids["assets"]["SFL"]["id"], ids["assets"]["SAE"]["id"]
    datos = {**COMPLETO, "documento_num": "33333333P", "nombre": "Negro"}
    r = _reserva(client, admin, sfl, {"guest": datos}, "LN-1", dias=(20, 25))
    cid = r["guest_id"]
    assert client.put(f"/api/terceros/{cid}/lista-negra", headers=admin, json={"activa": True}).status_code == 400
    f = client.put(f"/api/terceros/{cid}/lista-negra", headers=admin,
                   json={"activa": True, "motivo": "Impago y desperfectos"}).json()
    assert f["lista_negra"] and f["lista_negra_avisos"][0]["motivo"] == "Impago y desperfectos"
    # vuelve a reservar (aunque sea en otro activo, con el mismo documento): sale el motivo
    chk = client.get("/api/lista-negra/comprobar", headers=admin, params={"documento": "33333333-P"}).json()
    assert chk and chk[0]["motivo"] == "Impago y desperfectos"
    unidades = [u for u in client.get(f"/api/unidades?asset_id={sae}", headers=admin).json() if u["uso"] == "apartamento"]
    body = {"unit_id": unidades[0]["id"], "fecha_entrada": (HOY + timedelta(days=200)).isoformat(),
            "fecha_salida": (HOY + timedelta(days=203)).isoformat(), "guest": datos}
    no = client.post("/api/turistico/reservas", headers=admin, json=body)
    assert no.status_code == 409 and "LISTA NEGRA" in no.json()["detail"] and "Impago" in no.json()["detail"]
    si = client.post("/api/turistico/reservas", headers=admin, json={**body, "lista_negra_aceptada": True})
    assert si.status_code == 201, si.text
    assert any(c["id"] == cid for c in client.get("/api/lista-negra", headers=admin).json())
    # expediente: estancias, contratos, facturas y fianzas del cliente en un sitio
    client.post(f"/api/turistico/reservas/{r['id']}/cobros", headers=admin, json={"importe": 100})
    e = client.get(f"/api/terceros/{cid}/expediente", headers=admin).json()
    assert [x["localizador"] for x in e["reservas"]] == ["LN-1"] and e["lista_negra"]
    assert set(e) >= {"reservas", "contratos", "facturas", "alquileres", "fianzas"}
    client.put(f"/api/terceros/{cid}/lista-negra", headers=admin, json={"activa": False})
    assert not client.get("/api/lista-negra/comprobar", headers=admin, params={"documento": "33333333P"}).json()


def test_dotacion_habitual_especial_y_anexo(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    cat = client.get(f"/api/dotacion/catalogo?asset_id={sfl}", headers=admin).json()
    assert len(cat["articulos"]) >= 80 and "AMENITY_2BED_3K" in cat["codigos_perfil"].values()
    tips = client.get(f"/api/dotacion/habitual?asset_id={sfl}", headers=admin).json()
    assert tips and not tips[0]["fijada"] and tips[0]["items"]
    t = tips[0]
    # «Otra dotación»: artículo nuevo que queda en el catálogo
    nuevo = client.post("/api/dotacion/articulos", headers=admin,
                        json={"asset_id": sfl, "zona": "cocina", "articulo": "Cafetera de cápsulas Dolce"}).json()
    assert nuevo["codigo"].startswith("OTR")
    items = [i for i in t["items"] if i["codigo"] != "SAL08"] + [{**nuevo, "cantidad": "1"}]
    g = client.put("/api/dotacion/habitual", headers=admin, json={
        "asset_id": sfl, "tipologia": t["tipologia"], "perfil": t["perfil"], "items": items}).json()
    assert g["fijada"] and any(i["codigo"] == nuevo["codigo"] for i in g["items"])
    # reserva en un apartamento de esa tipología: dotación especial (quitar, cambiar y añadir)
    u = next(u for u in client.get(f"/api/unidades?asset_id={sfl}", headers=admin).json()
             if u["uso"] == "apartamento" and (u.get("tipologia") or "") == t["tipologia"])
    r = client.post("/api/turistico/reservas", headers=admin, json={
        "unit_id": u["id"], "localizador": "DOT-1", "fecha_entrada": (HOY + timedelta(days=300)).isoformat(),
        "fecha_salida": (HOY + timedelta(days=305)).isoformat(), "guest": {"nombre": "Dotación Prueba"}})
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    d = client.get(f"/api/turistico/reservas/{rid}/dotacion", headers=admin).json()
    assert d["habitual_fijada"] and not d["especial"]
    primero = d["items"][0]["codigo"]
    d = client.put(f"/api/turistico/reservas/{rid}/dotacion", headers=admin, json={
        "cambios": {primero: "—", "SAL02": "2"},
        "extra": [{"codigo": "CUNA", "zona": "dormitorio", "articulo": "Cuna de viaje", "cantidad": "1"}]}).json()
    cods = {i["codigo"]: i for i in d["items"]}
    assert d["especial"] and primero not in cods and cods["CUNA"]["especial"]
    # contrato con «Dotación habitual»: Anexo I con la lista
    c = client.post(f"/api/turistico/reservas/{rid}/contrato", headers=admin, json={
        "fecha_firma": HOY.isoformat(), "motivo": ["turismo"], "acreditacion": ["dni"], "permitir_huecos": True,
        "dotacion": True})
    assert c.status_code == 200, c.text
    texto = "\n".join(p.text for p in docx.Document(io.BytesIO(c.content)).paragraphs)
    celdas = [cell.text for tb in docx.Document(io.BytesIO(c.content)).tables for row in tb.rows for cell in row.cells]
    assert "ANEXO I" in texto and "Cuna de viaje (dotación especial)" in celdas and "Cafetera de cápsulas Dolce" in celdas
    # la reimpresión sale igual
    h = client.get(f"/api/turistico/reservas/{rid}/contrato", headers=admin).json()["historial"][0]
    re = client.get(f"/api/turistico/reservas/{rid}/contrato/{h['id']}", headers=admin)
    assert "Cuna de viaje (dotación especial)" in [cell.text for tb in docx.Document(io.BytesIO(re.content)).tables
                                                   for row in tb.rows for cell in row.cells]


def test_precios_reposicion_cargos_y_hojas(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    # precios: orientativos y propios del activo
    pre = {p["codigo"]: p for p in client.get(f"/api/dotacion/precios?asset_id={sae}", headers=admin).json()}
    assert pre["MEN04"]["precio"] == 2.5 and not pre["MEN04"]["fijado"]
    pre = {p["codigo"]: p for p in client.put("/api/dotacion/precios", headers=admin, json={
        "asset_id": sae, "precios": {"MEN04": 3, "SAL02": 30}}).json()}
    assert pre["MEN04"]["precio"] == 3 and pre["MEN04"]["fijado"] and pre["MEN04"]["orientativo"] == 2.5
    # dotación propia de un apartamento concreto: manda sobre la de su tipo; al quitarla vuelve la habitual
    r = _reserva(client, admin, sae, {"guest": {**COMPLETO, "documento_num": "44444444A", "nombre": "Rompe"}}, "RP-1")
    u = r["unit_id"]
    assert client.get(f"/api/dotacion/unidad/{u}", headers=admin).json()["origen"] in ("tipologia", "estandar")
    propia = client.put(f"/api/dotacion/unidad/{u}", headers=admin, json={"items": [
        {"codigo": "MEN04", "zona": "menaje", "articulo": "Vasos de agua / refresco", "cantidad": "8"},
        {"codigo": "SAL02", "zona": "salon", "articulo": "Mando a distancia TV con funda", "cantidad": "1"},
        {"codigo": "TERR1", "zona": "auxiliar", "articulo": "Mesa de terraza", "cantidad": "1"}]}).json()
    assert propia["origen"] == "apartamento" and len(propia["items"]) == 3
    assert any(x["propia"] for x in client.get(f"/api/dotacion/unidades?asset_id={sae}", headers=admin).json())
    d = client.get(f"/api/turistico/reservas/{r['id']}/dotacion", headers=admin).json()
    assert d["origen"] == "apartamento" and {x["codigo"]: x["precio"] for x in d["items"]}["MEN04"] == 3
    # hoja de revisión de salida (de la reserva y las del día en limpieza)
    h = client.get(f"/api/turistico/reservas/{r['id']}/hoja-revision.pdf", headers=admin)
    assert h.status_code == 200 and h.content.startswith(b"%PDF")
    client.put(f"/api/turistico/reservas/{r['id']}", headers=admin, json={"estado": "checkin"})
    client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json={
        "fecha_firma": HOY.isoformat(), "fianza": 100, "motivo": ["turismo"], "acreditacion": ["dni"],
        "solo_guardar": True})
    # salida con roturas: se descuentan de la fianza a su precio y el detalle va en el motivo
    x = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin, json={"fianza": {
        "forma": "efectivo", "cargos": [{"codigo": "MEN04", "articulo": "Vasos de agua / refresco", "cantidad": 2,
                                         "precio": 3}, {"codigo": "SAL02", "articulo": "Mando TV", "precio": 30}]}})
    assert x.status_code == 200, x.text
    dv = x.json()["devolucion"]
    assert dv["importe"] == 64 and dv["retenido"] == 36 and "2 × Vasos de agua" in dv["retenido_motivo"]
    assert len(dv["cargos"]) == 2
    hs = client.get(f"/api/limpieza/hojas-revision.pdf?asset_id={sae}", headers=admin)
    assert hs.status_code == 200 and hs.content.startswith(b"%PDF")
    assert client.delete(f"/api/dotacion/unidad/{u}", headers=admin).json()["origen"] != "apartamento"


def test_revision_movil_ot_factura_exceso_e_informe(client, admin, ids):
    from urllib.parse import parse_qs, urlparse
    sfl = ids["assets"]["SFL"]["id"]
    r = _reserva(client, admin, sfl, {"guest": {**COMPLETO, "documento_num": "55555555K", "nombre": "Exceso"}},
                 "EX-1", dias=(-4, 0))
    client.put(f"/api/turistico/reservas/{r['id']}", headers=admin, json={"estado": "checkin"})
    client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json={
        "fecha_firma": HOY.isoformat(), "fianza": 50, "motivo": ["turismo"], "acreditacion": ["dni"],
        "solo_guardar": True})
    # enlace al móvil de limpieza (firmado, sin usuario): marca lo roto y lo que falta
    e = client.get(f"/api/revision/reserva/{r['id']}/enlace", headers=admin).json()
    assert "/revision/" in e["url"] and "wa.me" in e["whatsapp"]
    dia = client.get(f"/api/revision/dia?asset_id={sfl}", headers=admin).json()
    assert any(x["reserva_id"] == r["id"] for x in dia["salidas"])
    u = urlparse(e["url"])
    k = {a: b[0] for a, b in parse_qs(u.query).items()}
    assert client.get(u.path).status_code == 200
    assert client.get(f"/api/publico/revision/{r['id']}", params={**k, "f": "0" * 32}).status_code == 404
    datos = client.get(f"/api/publico/revision/{r['id']}", params=k).json()
    cods = [x["codigo"] for x in datos["items"]]
    tv = next(c for c in cods if c == "SAL01")
    avisos.BANDEJA.clear()
    g = client.post(f"/api/publico/revision/{r['id']}", params=k, json={
        "nombre": "Limpieza Prueba", "nota": "Mancha en el sofá",
        "marcas": [{"codigo": tv, "estado": "roto"}, {"codigo": "MEN01", "estado": "falta", "cantidad": 2}]})
    assert g.status_code == 200 and g.json()["incidencias"] == 2
    assert any("Revisión de salida" in m["asunto"] for m in avisos.BANDEJA)  # aviso a Recepción 1
    rev = client.get(f"/api/turistico/reservas/{r['id']}/dotacion", headers=admin)
    assert client.get(f"/api/revision/reserva/{r['id']}/enlace", headers=admin).json()["revision"]["nombre"] == \
        "Limpieza Prueba" and rev.status_code == 200
    # salida: la TV rota abre OT a mantenimiento y lo que supera la fianza se factura (sin IVA, pendiente)
    x = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin, json={"fianza": {
        "forma": "efectivo", "facturar_exceso": True, "cargos": [
            {"codigo": tv, "articulo": "Smart TV", "precio": 450, "ot": True},
            {"codigo": "MEN01", "articulo": "Platos llanos", "cantidad": 2, "precio": 4}]}})
    assert x.status_code == 200, x.text
    x = x.json()
    assert x["devolucion"]["importe"] == 0 and len(x["ots"]) == 1 and x["factura_exceso"]["total"] == 408
    ot = client.get(f"/api/mantenimiento/ordenes/{x['ots'][0]}", headers=admin)
    assert ot.status_code != 200 or "Smart TV" in ot.text
    f = client.get(f"/api/facturas/{x['factura_exceso']['id']}", headers=admin).json()
    assert f["cobro"] == "pendiente" and "78.Tres" in (f["exencion"] or "") and float(f["cuota_iva"]) == 0
    # informe de roturas del periodo (pantalla y Excel)
    q = {"desde": (HOY - timedelta(days=1)).isoformat(), "hasta": (HOY + timedelta(days=3)).isoformat()}
    inf = client.get("/api/turistico/informe-roturas", headers=admin, params=q).json()
    tv_inf = next(a for a in inf["por_articulo"] if a["codigo"] == tv)
    assert tv_inf["importe"] >= 450 and tv_inf["ots"] >= 1 and inf["facturado_exceso"] >= 408
    xl = client.get("/api/turistico/informe-roturas", headers=admin, params={**q, "formato": "xlsx"})
    assert xl.status_code == 200 and xl.content[:2] == b"PK"


def test_pedido_reposicion_y_recordatorio_cobro(client, admin, ids):
    from app import fianzas
    from app.database import SessionLocal
    sae = ids["assets"]["SAE"]["id"]

    def salida(loc, doc, cargos, fianza=20):
        r = _reserva(client, admin, sae, {"guest": {**COMPLETO, "documento_num": doc, "nombre": loc,
                                                     "email": f"{loc.lower()}@example.com"}}, loc)
        client.put(f"/api/turistico/reservas/{r['id']}", headers=admin, json={"estado": "checkin"})
        client.post(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin, json={
            "fecha_firma": HOY.isoformat(), "fianza": fianza, "motivo": ["turismo"], "acreditacion": ["dni"],
            "solo_guardar": True})
        x = client.post(f"/api/turistico/reservas/{r['id']}/checkout", headers=admin, json={"fianza": {
            "forma": "efectivo", "facturar_exceso": True, "cargos": cargos}})
        assert x.status_code == 200, x.text
        return x.json()

    vasos = {"codigo": "MEN04", "articulo": "Vasos de agua / refresco", "cantidad": 3, "precio": 2.5, "pedido": True}
    a = salida("PEDA", "66666666Q", [vasos, {"codigo": "BAN09", "articulo": "Toallas de baño (100x150)",
                                             "precio": 15, "pedido": True}])
    b = salida("PEDB", "77777777B", [{**vasos, "cantidad": 2}])
    # un solo pedido de reposición pendiente de autorizar, que se va sumando
    assert a["pedido"]["nuevo"] and not b["pedido"]["nuevo"] and a["pedido"]["id"] == b["pedido"]["id"]
    p = client.get(f"/api/pedidos/{a['pedido']['id']}", headers=admin).json()
    assert p["estado"] == "pendiente" and len(p["lineas"]) == 2
    assert any(float(x["cantidad"]) == 5 for x in p["lineas"])
    # recordatorio de cobro del exceso (fianza 20 €, cargos 22,50 €)
    dv = a["devolucion"]
    assert a["factura_exceso"]["total"] == 2.5
    avisos.BANDEJA.clear()
    m = client.post(f"/api/turistico/fianzas/{dv['id']}/recordatorio", headers=admin).json()
    assert m["email"]["ok"] and m["whatsapp"].startswith("https://wa.me/34611222333")
    assert avisos.BANDEJA[-1]["para"] == "peda@example.com" and avisos.BANDEJA[-1]["adjuntos"]
    lst = {x["id"]: x for x in client.get("/api/turistico/fianzas", headers=admin).json()}
    assert lst[dv["id"]]["exceso_cobro"] == "pendiente"
    with SessionLocal() as db:
        assert fianzas.recordatorios_exceso(db, HOY) == 0  # aún no toca
        n3 = fianzas.recordatorios_exceso(db, HOY + timedelta(days=3))
        assert n3 >= 1 and fianzas.recordatorios_exceso(db, HOY + timedelta(days=3)) == 0  # una vez
        assert fianzas.recordatorios_exceso(db, HOY + timedelta(days=10)) >= 1
