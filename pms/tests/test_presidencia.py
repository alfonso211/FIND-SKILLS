"""Informe mensual a la presidencia: alquileres por tipo de estancia, servicios, gastos por proveedor (OPEX/CAPEX),
resultado con semáforo, resúmenes anuales, PDF/Excel, envío tras revisarlo y aviso el primer día laborable."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app import avisos, calendario
from app import informe_presidencia as ip
from app.database import SessionLocal
from app.models import Reservation
from conftest import domicilio_fiscal

HOY = date.today()
INI = HOY.replace(day=1)


def test_calendario_laboral():
    assert calendario.pascua(2026) == date(2026, 4, 5) and calendario.pascua(2027) == date(2027, 3, 28)
    assert calendario.primer_laborable(2026, 11) == date(2026, 11, 2)  # 1 de noviembre, domingo y festivo
    assert calendario.primer_laborable(2027, 1) == date(2027, 1, 4)  # 1 viernes festivo, 2-3 fin de semana
    assert calendario.primer_laborable(2026, 10) == date(2026, 10, 1)
    assert not calendario.laborable(date(2026, 4, 3))  # Viernes Santo


def test_clasificacion_y_semaforo():
    r = lambda n, renueva=None: Reservation(fecha_entrada=INI, fecha_salida=INI + timedelta(days=n),  # noqa: E731
                                            renueva_id=renueva)
    assert [ip.tipo_estancia(r(n)) for n in (1, 6, 7, 13, 14, 27, 28, 31)] == [
        "corta", "corta", "semana", "semana", "dos_semanas", "dos_semanas", "mes", "mes"]
    assert ip.tipo_estancia(r(30, renueva=5)) == "renovacion"
    assert [ip.semaforo(x) for x in (None, -5, 0, 12, 13, 25, 26, 40)] == [
        "rojo", "rojo", "rojo", "rojo", "naranja", "naranja", "verde", "verde"]


def test_informe_mensual(client, admin, ids):
    domicilio_fiscal(client, admin)
    sae = ids["assets"]["SAE"]["id"]
    aptos = [u for u in client.get(f"/api/unidades?asset_id={sae}", headers=admin).json() if u["uso"] != "garaje"]
    # estancias facturadas este mes: 3 noches, 7 noches y 1 mes
    for i, (noches, importe) in enumerate(((3, 300), (7, 560), (28, 1400))):
        r = client.post("/api/turistico/reservas", headers=admin, json={
            "unit_id": aptos[-30 - i]["id"], "localizador": f"PRES-{i}", "fecha_entrada": INI.isoformat(),
            "fecha_salida": (INI + timedelta(days=noches)).isoformat(), "importe_total": importe,
            "facturar_pendiente": True, "guest": {"nombre": f"Cliente Pres {i}"}})
        assert r.status_code == 201, r.text
    # gastos del mes: un proveedor con dos facturas (OPEX y CAPEX) y otro con una
    base = {"asset_id": sae, "fecha": HOY.isoformat(), "categoria": "mantenimiento", "concepto": "Trabajos",
            "tipo_iva": 21, "forma_pago": "transferencia"}
    for prov, num, total, nat in (("Proveedor Pres Uno SL", "PU-1", 121, "OPEX"), ("Proveedor Pres Uno SL", "PU-2", 242,
                                                                                 "CAPEX"),
                                  ("Proveedor Pres Dos SL", "PD-9", 60.5, "OPEX")):
        assert client.post("/api/gastos", headers=admin, json={**base, "proveedor": prov, "numero_factura": num,
                                                               "total": total, "naturaleza": nat}).status_code == 201
    r = client.get("/api/presidencia", headers=admin, params={"asset_id": sae, "anio": HOY.year, "mes": HOY.month})
    assert r.status_code == 200, r.text
    j = r.json()
    d = j["datos"]
    alq = {f["clave"]: f for f in d["alquileres"]}
    assert [f["concepto"] for f in d["alquileres"][:5]] == [t for _, t in ip.TIPOS_ESTANCIA]
    assert alq["corta"]["contratos"] >= 1 and alq["corta"]["importe"] >= round(300 / 1.1, 2) - 0.01
    assert alq["semana"]["pernoctaciones"] >= 7 and alq["mes"]["contratos"] >= 1
    assert d["total_pernoctaciones"] == sum(f["pernoctaciones"] or 0 for f in d["alquileres"])
    assert abs(d["total_ingresos"] - d["total_alquileres"] - d["total_servicios"]) < 0.01
    uno = next(g for g in d["gastos"] if g["proveedor"] == "Proveedor Pres Uno SL")
    assert uno["facturas"] == ["PU-1", "PU-2"] and uno["importe"] == 300 and uno["naturaleza"] == "Varios"
    assert next(g for g in d["gastos"] if g["proveedor"] == "Proveedor Pres Dos SL")["naturaleza"] == "OPEX"
    assert d["resultado"] == round(d["total_ingresos"] - d["total_gastos"], 2)
    assert d["semaforo"] == ip.semaforo(d["porcentaje"])
    assert d["anterior"]["anio"] == HOY.year - 1 and len(d["anterior"]["meses"]) == 12
    assert len(d["actual"]["meses"]) == HOY.month
    actual = d["actual"]["meses"][-1]
    assert actual["produccion"] == d["total_ingresos"] and actual["pernoctaciones"] >= 38
    assert j["responsable"] == "Recepción 1 · Suite Aeropuerto" and "jr@inversiete.es" in j["destinatarios"]
    assert client.get("/api/presidencia", headers=admin, params={"asset_id": sae, "anio": HOY.year + 1,
                                                                 "mes": 1}).status_code == 400

    pdf = client.get("/api/presidencia/pdf", headers=admin, params={"asset_id": sae, "anio": HOY.year, "mes": HOY.month})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf.content)
    assert len(doc) == 3
    textos = [doc[i].get_textpage().get_text_range() for i in range(3)]
    assert "Contratos inferiores a 1 semana" in textos[0] and "RESULTADO DE PRODUCCIÓN" in textos[0]
    assert f"Resumen mensual {HOY.year - 1}" in textos[1] and f"Resumen mensual {HOY.year}" in textos[2]
    xl = client.get("/api/presidencia/excel", headers=admin, params={"asset_id": sae, "anio": HOY.year,
                                                                     "mes": HOY.month})
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Producción", f"Resumen {HOY.year - 1}", f"Resumen {HOY.year}"]

    # envío: solo tras marcar que se ha revisado; por correo con el PDF; el Excel queda en documentos
    envio = {"asset_id": sae, "anio": HOY.year, "mes": HOY.month, "canal": "email"}
    assert client.post("/api/presidencia/enviar", headers=admin, json=envio).status_code == 400
    avisos.BANDEJA.clear()
    ok = client.post("/api/presidencia/enviar", headers=admin, json={**envio, "revisado": True})
    assert ok.status_code == 200, ok.text
    m = next(x for x in avisos.BANDEJA if x["para"] == "jr@inversiete.es")
    assert m["adjuntos"][0][0] == f"Informe_presidencia_SAE_{HOY.year}-{HOY.month:02d}.pdf"
    reg = ok.json()["registro"]
    assert reg["canal"] == "email" and reg["enviado_por_nombre"] == "Administrador" and reg["documento_id"]
    docs = client.get("/api/documentos-recibidos", headers=admin, params={"asset_id": sae, "tipo": "informe",
                                                                          "desde": INI.isoformat()}).json()
    assert [x["nombre"] for x in docs] == [f"Informe_presidencia_SAE_{HOY.year}-{HOY.month:02d}.xlsx"]
    # por WhatsApp: enlace con el teléfono; el reenvío sustituye el Excel guardado
    wa = client.post("/api/presidencia/enviar", headers=admin, json={**envio, "revisado": True, "canal": "whatsapp",
                                                                     "telefono": "600 111 222"}).json()
    assert wa["whatsapp_url"].startswith("https://wa.me/34600111222?text=") and wa["registro"]["canal"] == "whatsapp"
    assert len(client.get("/api/documentos-recibidos", headers=admin, params={
        "asset_id": sae, "tipo": "informe", "desde": INI.isoformat()}).json()) == 1


def test_aviso_primer_dia_laborable(client, admin, ids):
    nov = date(2030, 11, 4)  # primer laborable de noviembre de 2030 (1 viernes festivo)
    assert calendario.primer_laborable(2030, 11) == nov
    avisos.BANDEJA.clear()
    with SessionLocal() as db:
        assert avisos.informe_mensual(db, nov + timedelta(days=1)) == 0  # otro día: nada
        n = avisos.informe_mensual(db, nov)
        assert avisos.informe_mensual(db, nov) == 0  # no se repite
    para = {m["para"] for m in avisos.BANDEJA}
    assert n >= 2 and {"jaime@apartamentossuitesaeropuerto.es", "juancarlos@apartamentossuitesflorida.es"} <= para
    assert "info@apartamentossuitesaeropuerto.es" not in para  # Recepción 2: no
    m = next(x for x in avisos.BANDEJA if x["para"] == "jaime@apartamentossuitesaeropuerto.es")
    assert "octubre 2030" in m["asunto"]
