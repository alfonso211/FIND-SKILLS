"""Registro de todos los ocupantes, parte de viajeros a SES.HOSPEDAJE, encuesta del INE y firma en tablet."""
import base64
import io
from datetime import date, timedelta
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from app import avisos, firma_contrato, registro_viajeros as rv

HOY = date.today()
TITULAR = {"nombre": "Lucía", "apellidos": "de la Fuente García", "documento_tipo": "DNI", "documento_num": "12345678Z",
           "num_soporte": "BAA123456", "nacionalidad": "España", "fecha_nacimiento": "1985-03-02", "sexo": "F",
           "telefono": "600123123", "email": "lucia@example.com", "direccion": "C/ Mayor 5, 2º A", "cp": "28230",
           "municipio": "Las Rozas de Madrid", "pais": "España"}


def _reserva(client, h, asset_id, codigo, entrada, noches, adultos=2, ninos=1, canal="booking", total=330):
    u = client.get(f"/api/unidades?asset_id={asset_id}&q={codigo}", headers=h).json()[0]
    r = client.post("/api/turistico/reservas", headers=h, json={
        "unit_id": u["id"], "localizador": f"RV-{codigo}", "canal": canal, "fecha_entrada": entrada.isoformat(),
        "fecha_salida": (entrada + timedelta(days=noches)).isoformat(), "adultos": adultos, "ninos": ninos,
        "importe_total": total, "guest": TITULAR})
    assert r.status_code == 201, r.text
    return r.json()


def _ocupantes(client, h, rid):
    adulto = {"contact": {"nombre": "John", "apellidos": "Smith", "documento_tipo": "PAS", "documento_num": "X1234567",
                          "nacionalidad": "Reino Unido", "fecha_nacimiento": "1984-07-01", "sexo": "M",
                          "direccion": "10 Baker Street", "cp": "NW1 6XE", "municipio": "London", "pais": "Reino Unido"}}
    menor = {"contact": {"nombre": "Martina", "apellidos": "Smith de la Fuente", "fecha_nacimiento":
                         (HOY - timedelta(days=365 * 4)).isoformat(), "sexo": "F", "nacionalidad": "España",
                         "direccion": "C/ Mayor 5, 2º A", "cp": "28230", "municipio": "Las Rozas de Madrid",
                         "pais": "España"}, "parentesco": "HJ"}
    for o in (adulto, menor):
        r = client.post(f"/api/turistico/reservas/{rid}/ocupantes", headers=h, json=o)
        assert r.status_code == 201, r.text


def test_utilidades():
    assert rv.municipio_ine("Rozas de Madrid, Las") == "28127" and rv.municipio_ine("Madrid", "28022") == "28079"
    assert rv.municipio_ine("Villanueva de la Torre", "08000") is None  # existe, pero no en Barcelona
    assert rv.pais_iso3("Reino Unido") == "GBR" and rv.pais_iso3("francia") == "FRA" and rv.pais_iso3("ESP") == "ESP"
    assert rv.apellidos_ses("de la Fuente García") == ("de la Fuente", "García")
    assert rv.comunidad("28230") == "Madrid, Comunidad de" and rv.comunidad("08001") == "Cataluña"
    assert firma_contrato.movil_whatsapp("600 12 31 23") == "34600123123"
    assert firma_contrato.movil_whatsapp("+44 7700 900123") == "447700900123"


def test_ocupantes_ses_e_ine(client, admin, ids):
    sae = ids["assets"]["SAE"]["id"]
    entrada = date(HOY.year + 3, 3, 30)  # mes futuro sin otras reservas: la encuesta se puede comprobar exacta
    r = _reserva(client, admin, sae, "A-127", entrada, 4)  # 2 dormitorios: 3 plazas
    lista = client.get(f"/api/turistico/reservas/{r['id']}/ocupantes", headers=admin).json()
    assert lista["requeridos"] == 3 and lista["ocupantes"][0]["contact"]["municipio_ine"] == "28127"
    assert "faltan 2 ocupante(s) por registrar (1 de 3)" in lista["pendiente"]
    _ocupantes(client, admin, r["id"])
    assert client.get(f"/api/turistico/reservas/{r['id']}/ocupantes", headers=admin).json()["pendiente"] == []

    # SES.HOSPEDAJE: sin código de establecimiento no se genera
    url = f"/api/turistico/ses/partes.xml?asset_id={sae}"
    assert client.post(url, headers=admin, json=[r["id"]]).status_code == 400
    client.put(f"/api/activos/{sae}", headers=admin, json={"ses_codigo_establecimiento": "0000012345"})
    estado = client.get(f"/api/turistico/ses?asset_id={sae}&desde={entrada}&hasta={entrada}", headers=admin).json()
    fila = next(x for x in estado["reservas"] if x["id"] == r["id"])
    assert fila["completo"] and fila["ocupantes_registrados"] == 3 and fila["ses_comunicado"] is None
    res = client.post(url, headers=admin, json=[r["id"]])
    assert res.status_code == 200, res.text
    raiz = ET.fromstring(res.content)
    assert raiz.tag == f"{{{rv.NS_ALTA}}}peticion"
    sol = raiz.find("solicitud")
    assert sol.findtext("codigoEstablecimiento") == "0000012345"
    com = sol.find("comunicacion")
    ct = com.find("contrato")
    assert ct.findtext("referencia") == "RV-A-127" and ct.findtext("numPersonas") == "3"
    assert ct.findtext("fechaEntrada") == f"{entrada.isoformat()}T15:00:00" and ct.findtext("internet") == "true"
    assert ct.find("pago").findtext("tipoPago") == "PLATF"
    personas = com.findall("persona")
    assert [p.findtext("nombre") for p in personas] == ["Lucía", "John", "Martina"]
    lucia, john, martina = personas
    assert (lucia.findtext("apellido1"), lucia.findtext("apellido2")) == ("de la Fuente", "García")
    assert lucia.findtext("tipoDocumento") == "NIF" and lucia.findtext("soporteDocumento") == "BAA123456"
    assert lucia.findtext("sexo") == "M" and lucia.find("direccion").findtext("codigoMunicipio") == "28127"
    assert john.findtext("tipoDocumento") == "PAS" and john.findtext("nacionalidad") == "GBR"
    assert john.find("direccion").findtext("nombreMunicipio") == "London" and john.findtext("sexo") == "H"
    assert martina.find("numeroDocumento") is None and martina.findtext("parentesco") == "HJ"
    assert martina.findtext("telefono") == "600123123"  # contacto del titular
    assert lucia.find("parentesco") is None
    estado = client.get(f"/api/turistico/ses?asset_id={sae}&desde={entrada}&hasta={entrada}", headers=admin).json()
    assert next(x for x in estado["reservas"] if x["id"] == r["id"])["ses_comunicado"]

    # encuesta del INE del mes de entrada (30 y 31 de marzo; salida el 3 de abril)
    ine = client.get(f"/api/turistico/ine?asset_id={sae}&anio={entrada.year}&mes=3", headers=admin).json()
    res_ = {x["residencia"]: x for x in ine["residencias"]}
    assert res_["Madrid, Comunidad de"]["entradas"][29] == 2 and res_["Reino Unido"]["entradas"][29] == 1
    assert res_["Madrid, Comunidad de"]["pernoctaciones"][30] == 2
    assert ine["totales"]["viajeros_entrados"] == 3 and ine["totales"]["pernoctaciones"] == 6
    assert ine["apartamentos_ocupados"][29:] == [1, 1] and ine["totales"]["apartamentos_noche"] == 2
    assert ine["precios"] == [{"tipo": "Agencia de viajes online", "apartamentos_noche": 2, "porcentaje": 1.0,
                               "tarifa_media": 75.0}]  # 330 € IVA incl. / 1,10 / 4 noches
    xl = client.get(f"/api/turistico/ine.xlsx?asset_id={sae}&anio={entrada.year}&mes=3", headers=admin)
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Resumen", "Viajeros entrados", "Pernoctaciones", "Viajeros salidos",
                             "Ocupación diaria", "Precios", "Avisos"]


def _firma_png() -> str:
    img = Image.new("RGBA", (600, 200), (0, 0, 0, 0))
    ImageDraw.Draw(img).line([(30, 150), (200, 40), (380, 160), (560, 50)], fill=(0, 0, 0, 255), width=6)
    b = io.BytesIO()
    img.save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


@pytest.mark.skipif(not firma_contrato.disponible(), reason="sin LibreOffice")
def test_firma_en_tablet_y_envio(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    client.put(f"/api/activos/{sfl}", headers=admin, json={
        "contrato_representante": "Recepción Suite Florida", "contrato_representante_dni": "00000000T",
        "contrato_email": "info@apartamentossuitesflorida.es"})
    r = _reserva(client, admin, sfl, "P2-2B", HOY + timedelta(days=70), 3, adultos=1, ninos=0, canal="directo")
    datos = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()["datos"]
    assert datos["cliente_domicilio"] == "C/ Mayor 5, 2º A" and datos["cliente_municipio"] == "Las Rozas de Madrid"
    datos.update(capacidad=2, dormitorios=1, sin_garaje=True, fianza=200, tarjeta_titular="Lucía de la Fuente",
                 tarjeta_terminacion="4242", tarjeta_caducidad="08/29", motivo=["turismo"], acreditacion=["dni"])
    prep = client.post(f"/api/turistico/reservas/{r['id']}/contrato/firma", headers=admin, json=datos)
    assert prep.status_code == 200, prep.text
    p = prep.json()
    assert len(p["paginas"]) >= 2 and p["paginas"][0].startswith("data:image/png;base64,")
    url = f"/api/turistico/reservas/{r['id']}/contrato/{p['contrato_id']}"
    sin_aceptar = {"firma": _firma_png(), "acepta": True, "acepta_privacidad": False}
    assert client.post(url + "/firmar", headers=admin, json=sin_aceptar).status_code == 400
    vacia = Image.new("RGBA", (600, 200), (0, 0, 0, 0))
    b = io.BytesIO()
    vacia.save(b, "PNG")
    blanca = "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()
    assert client.post(url + "/firmar", headers=admin, json={**sin_aceptar, "firma": blanca,
                                                            "acepta_privacidad": True}).status_code == 400
    avisos.BANDEJA.clear()
    res = client.post(url + "/firmar", headers={**admin, "User-Agent": "Tablet de recepcion"}, json={
        "firma": _firma_png(), "acepta": True, "acepta_privacidad": True, "email": "lucia@example.com",
        "movil": "600123123"})
    assert res.status_code == 200, res.text
    out = res.json()
    envios = {e["canal"]: e for e in out["envios"]}
    assert envios["email"]["enviado"] and envios["whatsapp"]["whatsapp"].startswith("https://wa.me/34600123123?text=")
    correo = avisos.BANDEJA[-1]
    assert correo["para"] == "lucia@example.com" and correo["adjuntos"][0][2] == "application/pdf"
    assert client.post(url + "/firmar", headers=admin, json={**sin_aceptar, "acepta_privacidad": True}).status_code == 400

    pdf = client.get(url + "/pdf", headers=admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    import hashlib

    import pypdfium2 as pdfium
    assert hashlib.sha256(pdf.content).hexdigest() == out["sha256"]
    doc = pdfium.PdfDocument(pdf.content)
    texto = "\n".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))
    assert "Evidencias de la firma electrónica" in texto and "Tablet de recepcion" in texto
    assert "Lucía de la Fuente García" in texto

    # enlace del cliente (WhatsApp/correo): descarga sin usuario; un enlace inventado no sirve
    from urllib.parse import unquote, urlparse
    enlace = urlparse(unquote(envios["whatsapp"]["whatsapp"]).rsplit(" ", 1)[-1])
    assert enlace.path.startswith("/api/publico/contrato/")
    publico = client.get(enlace.path)
    assert publico.status_code == 200 and publico.content == pdf.content
    assert client.get("/api/publico/contrato/inventado").status_code == 404

    hist = client.get(f"/api/turistico/reservas/{r['id']}/contrato", headers=admin).json()["historial"]
    assert hist[0]["firmado"] and {e["canal"] for e in hist[0]["envios"]} == {"email", "whatsapp"}
