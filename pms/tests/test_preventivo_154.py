"""Versión 1.54: control del preventivo e inspecciones OCA (revisiones con documentos, deficiencias con OT)."""
import json
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app import avisos, preventivo
from app.config import settings
from app.database import SessionLocal
from app.models import PreventivePlan, WorkOrder
from test_adjuntos_ot import _usuario

HOY = date.today()
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


def _plan(client, admin, aid, **k):
    d = {"asset_id": aid, "titulo": "Plan 154", "categoria": "pci", "periodicidad_dias": 90,
         "proxima_fecha": (HOY - timedelta(days=5)).isoformat(), "proveedor": "Mantenedora PCI 154", **k}
    r = client.post("/api/mantenimiento/planes", headers=admin, json=d)
    assert r.status_code == 201, r.text
    return r.json()


def _registrar(client, h, pid, datos, ficheros=()):
    return client.post(f"/api/preventivo/planes/{pid}/revisiones", headers=h, data={"datos": json.dumps(datos)},
                       files=[("ficheros", (n, c, "application/pdf")) for n, c in ficheros])


def _control(client, admin, aid):
    return {p["id"]: p for p in client.get("/api/preventivo/control", headers=admin,
                                           params={"asset_id": aid}).json()["planes"]}


def test_revision_de_mantenimiento_con_deficiencias(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    p = _plan(client, admin, bab, titulo="Revisión trimestral PCI 154")
    c = _control(client, admin, bab)[p["id"]]
    assert c["estado"] == "vencido" and "Sin ninguna revisión registrada" in c["avisos"]
    assert any("Extintores" in x for x in c["puntos"])  # lista de comprobación de su instalación
    # la OT preventiva generada sigue marcando el vencimiento aunque el plan avance su próxima fecha
    client.post("/api/mantenimiento/planes/generar", headers=admin, json={"asset_id": bab, "dias_antelacion": 0})
    c = _control(client, admin, bab)[p["id"]]
    assert c["estado"] == "vencido" and c["ot_preventiva"]
    ot_prev = c["ot_preventiva"]
    # resultado incoherente con lo anotado
    malo = {"fecha": HOY.isoformat(), "resultado": "correcto",
            "checklist": [{"punto": "Extintores", "estado": "deficiencia", "nota": "2 caducados"}]}
    assert _registrar(client, admin, p["id"], malo).status_code == 400
    assert _registrar(client, admin, p["id"], {**malo, "resultado": "favorable"}).status_code == 400  # es de OCA
    datos = {"fecha": (HOY - timedelta(days=1)).isoformat(), "resultado": "deficiencias", "numero": "INF-154",
             "tecnico": "Técnico PCI", "proxima": (HOY + timedelta(days=365)).isoformat(),
             "checklist": [{"punto": "Extintores: presión, precinto, etiqueta y accesibilidad", "estado": "deficiencia",
                            "nota": "2 extintores caducados"},
                           {"punto": "Alumbrado de emergencia: encendido y autonomía", "estado": "ok"}],
             "deficiencias": [{"descripcion": "Puerta cortafuegos planta 2 no cierra", "gravedad": "grave"}]}
    r = _registrar(client, admin, p["id"], datos, [("informe_pci.pdf", PDF)])
    assert r.status_code == 201, r.text
    rev = r.json()
    assert len(rev["ots"]) == 2 and {o["prioridad"] for o in rev["ots"]} == {"alta", "media"}
    assert rev["documentos"][0]["nombre"] == "informe_pci.pdf"
    assert rev["aviso"] and "periodicidad" in rev["aviso"]  # el documento dice 1 año y el plan 90 días
    assert rev["work_order_id"] == ot_prev
    assert rev["plazo_subsanacion"]  # sin plazo en el documento: según la gravedad
    with SessionLocal() as db:
        w = db.get(WorkOrder, ot_prev)
        assert w.estado == "cerrada" and "INF-154" in w.solucion  # la OT preventiva queda hecha
        assert db.get(PreventivePlan, p["id"]).proxima_fecha == HOY + timedelta(days=365)
        ot = db.get(WorkOrder, rev["ots"][0]["id"])
        assert ot.plan_id == p["id"] and ot.tipo == "correctivo" and "INF-154" in ot.descripcion
    c = _control(client, admin, bab)[p["id"]]
    assert c["estado"] == "deficiencias" and len(c["deficiencias_abiertas"]) == 2
    assert any("periodicidad" in a for a in c["avisos"])
    # las OT de las deficiencias no frenan la siguiente OT preventiva del plan
    with SessionLocal() as db:
        db.get(PreventivePlan, p["id"]).proxima_fecha = HOY
        db.commit()
    assert client.post("/api/mantenimiento/planes/generar", headers=admin,
                       json={"asset_id": bab, "dias_antelacion": 0}).json()["creadas"] >= 1
    # historial con sus documentos y el estado de las OT
    h = client.get(f"/api/preventivo/planes/{p['id']}/revisiones", headers=admin).json()
    assert h["revisiones"][0]["deficiencias"][0]["ot_estado"] == "abierta"
    doc = client.get(f"/api/preventivo/documentos/{rev['documentos'][0]['id']}", headers=admin)
    assert doc.status_code == 200 and doc.content.startswith(b"%PDF")


def test_inspeccion_oca_con_plazo_de_subsanacion(client, admin, ids):
    bab = ids["assets"]["BAB35"]["id"]
    p = _plan(client, admin, bab, titulo="Inspección periódica ascensores OCA 154", categoria="ascensores", tipo="oca",
              periodicidad_dias=730, proveedor="OCA Ejemplo 154")
    assert p["tipo"] == "oca"
    c = _control(client, admin, bab)[p["id"]]
    assert "Acta o certificado" in c["puntos"][0] and [r[0] for r in c["resultados"]] == [
        "favorable", "favorable_leves", "desfavorable", "negativa"]
    assert _registrar(client, admin, p["id"], {"fecha": HOY.isoformat(), "resultado": "deficiencias",
                                               "deficiencias": [{"descripcion": "x"}]}).status_code == 400
    antes = len(avisos.BANDEJA)
    r = _registrar(client, admin, p["id"], {
        "fecha": (HOY - timedelta(days=20)).isoformat(), "resultado": "desfavorable", "numero": "ACTA-2026-154",
        "tecnico": "Inspector OCA", "proxima": (HOY - timedelta(days=20) + timedelta(days=730)).isoformat(),
        "plazo_subsanacion": (HOY - timedelta(days=1)).isoformat(),
        "deficiencias": [{"descripcion": "Falta iluminación de emergencia en cabina", "gravedad": "muy_grave"}]})
    assert r.status_code == 201, r.text
    rev = r.json()
    assert rev["aviso"] is None and rev["ots"][0]["prioridad"] == "urgente"
    with SessionLocal() as db:
        assert db.get(WorkOrder, rev["ots"][0]["id"]).tipo == "normativo"
    c = _control(client, admin, bab)[p["id"]]
    assert c["estado"] == "subsanacion_vencida" and c["color"] == "rojo"
    assert any("documento" in a for a in c["avisos"])  # el acta todavía no se ha subido
    # el acta llega después
    r = client.post(f"/api/preventivo/revisiones/{rev['id']}/documentos", headers=admin,
                    files=[("ficheros", ("acta.pdf", PDF, "application/pdf"))])
    assert r.status_code == 201 and len(r.json()["documentos"]) == 1
    assert not any("documento" in a for a in _control(client, admin, bab)[p["id"]]["avisos"])
    assert len(avisos.BANDEJA) > antes  # OT urgente: aviso inmediato
    # aviso semanal al director técnico con lo que está en rojo
    antes = len(avisos.BANDEJA)
    assert preventivo.aviso_semanal(SessionLocal(), HOY, forzar=True) >= 1
    m = [x for x in avisos.BANDEJA[antes:] if "Control de preventivo" in x["asunto"]]
    assert m and m[0]["para"] == settings.mantenimiento_email and "ACTA" not in m[0]["asunto"]
    assert "Inspección periódica ascensores OCA 154" in m[0]["texto"]
    # registro en Excel
    x = client.get("/api/preventivo/excel", headers=admin, params={"asset_id": bab})
    wb = load_workbook(BytesIO(x.content))
    assert {"Control", "Historial"} <= set(wb.sheetnames)
    assert any(row[7] == "ACTA-2026-154" for row in wb["Historial"].iter_rows(min_row=2, values_only=True))
    tot = client.get("/api/preventivo/control", headers=admin, params={"asset_id": bab}).json()["totales"]
    assert tot["oca"] >= 1 and tot["subsanacion_vencida"] >= 1


def test_plantilla_con_oca_y_permisos(client, admin, ids):
    sfl = ids["assets"]["SFL"]["id"]
    rec = _usuario(client, admin, ids, "recepcion.preventivo154@inversiete.com", "Recepción", "SFL")
    previos = {p["id"] for p in client.get("/api/mantenimiento/planes", headers=admin, params={"asset_id": sfl}).json()}
    assert client.post("/api/mantenimiento/planes/plantilla", headers=admin, json={"asset_id": sfl}).status_code == 201
    planes = [p for p in client.get("/api/mantenimiento/planes", headers=admin, params={"asset_id": sfl}).json()
              if p["id"] not in previos]
    oca = [p for p in planes if p["tipo"] == "oca"]
    assert any("ascensores" in p["titulo"] for p in oca) and all(p["tipo"] == "mantenimiento"
                                                                for p in planes if "OCA" not in p["titulo"]
                                                                and not p["titulo"].startswith("Inspección periódica"))
    r = _registrar(client, rec, oca[0]["id"], {"fecha": HOY.isoformat(), "resultado": "favorable"})
    assert r.status_code == 403
    assert _registrar(client, admin, oca[0]["id"], {"fecha": (HOY + timedelta(days=3)).isoformat(),
                                                    "resultado": "favorable"}).status_code == 400  # fecha futura
    with SessionLocal() as db:  # se deja como estaba: otras pruebas cuentan planes y OT
        for p in planes:
            db.query(WorkOrder).filter(WorkOrder.plan_id == p["id"]).delete()
            db.query(PreventivePlan).filter(PreventivePlan.id == p["id"]).delete()
        db.commit()
