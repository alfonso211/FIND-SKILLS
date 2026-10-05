"""Importación de la ocupación actual exportada del PMS anterior (listado por plantas).

Cada hoja es una planta con filas: estado (A = alojado, R = reserva), localizador, bloque, planta, número y
tipología, dormitorios, camas, sótano (0 = garaje exterior, 1 = sótano -1) y plaza de garaje, fechas de entrada y
salida con las noches (en una sola celda), ocupante y teléfonos.
"""
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

FECHAS = re.compile(r"(\d{2}/\d{2}/\d{4})\s*(\d{2}/\d{2}/\d{4})\s*(\d+)?")
SOCIEDADES = re.compile(r"\b(S\.?L\.?U?|S\.?A\.?U?|UK|LTD|LLC|GMBH|SAS|SL|SA|S\.L\.)\b\.?$", re.I)


def es_ocupacion(datos: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return False
    for ws in wb.worksheets[:2]:
        for fila in ws.iter_rows(max_row=4, values_only=True):
            textos = " ".join(str(x) for x in fila if x)
            if "Localizador" in textos and "FEntrada" in textos:
                return True
    return False


def _fecha(s: str) -> date:
    return datetime.strptime(s, "%d/%m/%Y").date()


def nombre_y_apellidos(texto: str) -> tuple[str, str | None]:
    """«ICSEL VICTORIA TOLEDO SANZONETTI» -> («Icsel Victoria», «Toledo Sanzonetti»). Las empresas, enteras."""
    t = " ".join(p for p in str(texto or "").split() if p != ".").strip()
    if not t:
        return "Sin nombre", None
    if SOCIEDADES.search(t) or " / " in t:
        return t, None
    pal = t.title().split()
    if len(pal) == 1:
        return pal[0], None
    if len(pal) == 2:
        return pal[0], pal[1]
    if len(pal) == 3:
        return pal[0], " ".join(pal[1:])
    return " ".join(pal[:-2]), " ".join(pal[-2:])


def telefono(v) -> str | None:
    if v in (None, ""):
        return None
    t = str(v).split("-")[0].split("/")[0].strip()
    return t[:-2] if t.endswith(".0") else t


def leer(datos: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    filas = []
    for ws in wb.worksheets:
        for n, f in enumerate(ws.iter_rows(values_only=True), 1):
            f = list(f) + [None] * 12
            if f[0] not in ("A", "R") or not isinstance(f[1], (int, float)):
                continue
            r = {"hoja": ws.title, "fila": n, "situacion": "alojado" if f[0] == "A" else "reserva",
                 "localizador": str(int(f[1])), "bloque": str(f[2] or "").strip(),
                 "planta": str(f[3] or "").replace("º", "").strip(), "ocupante": str(f[10] or "").strip(),
                 "telefono": telefono(f[11])}
            try:
                r["numero"] = str(int(str(f[4]).split()[0]))
                m = FECHAS.search(str(f[9] or ""))
                if not m:
                    raise ValueError(f"fechas no reconocidas: «{f[9]}»")
                r["entrada"], r["salida"] = _fecha(m.group(1)), _fecha(m.group(2))
                if r["salida"] <= r["entrada"]:
                    raise ValueError("la salida es anterior a la entrada")
                if f[8] not in (None, ""):
                    r["garaje"] = f"{'EXT' if int(f[7] or 0) == 0 else 'S1'}-{int(f[8])}"
            except (ValueError, TypeError) as e:
                r["error"] = str(e)
            filas.append(r)
    return filas
