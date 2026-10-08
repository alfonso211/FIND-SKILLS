"""Parte de trabajo diario de mantenimiento y de limpieza, y su informe quincenal.

El parte de un día lo forman, solas, las órdenes de trabajo terminadas ese día (mantenimiento) y las limpiezas
hechas ese día (limpieza), más las actuaciones que no estaban en ningún parte y se anotan a mano (zonas comunes
según el programa de limpieza, tareas de mantenimiento sin OT…). Recepción lo valida: entonces se congela, se
puede imprimir y su PDF y su Excel quedan archivados en los documentos del activo.
"""
import io
from collections import Counter
from datetime import date, datetime, time, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import AREAS_PARTE, CleaningTask, Unit, User, WorkEntry, WorkOrder, WorkReport

TIPOS_LIMPIEZA = {"salida": "Limpieza de salida", "contratada": "Limpieza contratada", "extra": "Limpieza extra"}
ORIGENES = {"ot": "Orden de trabajo", "limpieza": "Limpieza del parte", "actuacion": "Actuación anotada"}


def _dia(f: date) -> tuple[datetime, datetime]:
    return datetime.combine(f, time.min), datetime.combine(f + timedelta(days=1), time.min)


def _nombres(db: Session, ids) -> dict[int, str]:
    ids = {i for i in ids if i}
    return dict(db.execute(select(User.id, User.nombre).where(User.id.in_(ids or {-1}))).all())


def lineas(db: Session, asset_id: int, area: str, f: date) -> list[dict]:
    """Trabajos del día de un área, en orden de hora: OT terminadas o limpiezas hechas, y actuaciones anotadas."""
    from .routers.personal import _lugar
    ini, fin = _dia(f)
    out: list[dict] = []
    if area == "mantenimiento":
        ots = list(db.scalars(select(WorkOrder).where(WorkOrder.asset_id == asset_id, WorkOrder.conf_mto_fecha >= ini,
                                                      WorkOrder.conf_mto_fecha < fin)))
        nombres = _nombres(db, [w.conf_mto_por for w in ots])
        for w in ots:
            out.append({"origen": "ot", "id": w.id, "ref": f"OT-{w.id:05d}", "hora": f"{w.conf_mto_fecha:%H:%M}",
                        "ubicacion": _lugar(db, asset_id, w.unit_id, w.zona), "trabajo": w.titulo,
                        "detalle": w.solucion or "", "categoria": w.categoria,
                        "persona": w.asignado_a or nombres.get(w.conf_mto_por) or "", "horas": None})
    else:
        tareas = list(db.scalars(select(CleaningTask).where(CleaningTask.asset_id == asset_id,
                                                            CleaningTask.estado == "hecha", CleaningTask.hecha >= ini,
                                                            CleaningTask.hecha < fin)))
        nombres = _nombres(db, [t.validada_por for t in tareas])
        for t in tareas:
            out.append({"origen": "limpieza", "id": t.id, "ref": f"L-{t.id:05d}", "hora": f"{t.hecha:%H:%M}",
                        "ubicacion": f"{t.unit.uso} {t.unit.codigo}" if t.unit else "",
                        "trabajo": TIPOS_LIMPIEZA.get(t.tipo, t.tipo), "detalle": t.nota or "", "categoria": t.tipo,
                        "persona": "", "validada_por": nombres.get(t.validada_por), "horas": None})
    for e in db.scalars(select(WorkEntry).where(WorkEntry.asset_id == asset_id, WorkEntry.area == area,
                                                WorkEntry.fecha == f)):
        u = db.get(Unit, e.unit_id) if e.unit_id else None
        out.append({"origen": "actuacion", "id": e.id, "ref": f"A-{e.id:05d}", "hora": f"{e.creada:%H:%M}",
                    "ubicacion": (f"{u.uso} {u.codigo}" if u else "") + (" · " if u and e.ubicacion else "")
                    + (e.ubicacion or ""), "trabajo": e.descripcion, "detalle": "", "categoria": "actuacion",
                    "persona": e.persona or "", "horas": float(e.horas) if e.horas is not None else None})
    return sorted(out, key=lambda x: (x["hora"], x["ref"]))


def parte(db: Session, asset_id: int, area: str, f: date) -> tuple[WorkReport | None, list[dict]]:
    """(parte guardado o None, líneas). Validado, las líneas son las congeladas al validar."""
    r = db.scalar(select(WorkReport).where(WorkReport.asset_id == asset_id, WorkReport.area == area,
                                           WorkReport.fecha == f))
    if r is not None and r.estado == "validado" and r.lineas is not None:
        return r, r.lineas
    return r, lineas(db, asset_id, area, f)


# --------------------------------------------------------------------------- documentos del parte
def _esc(s) -> str:
    return str(s or "").replace("&", "&amp;").replace("<", "&lt;")


def _doc_pdf(titulo: str, asset, company, ancho_mm: float = 186):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Spacer

    from . import marca
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm,
                            bottomMargin=10 * mm, title=titulo)
    ancho = landscape(A4)[0] - 24 * mm
    el = []
    cab = marca.cabecera_pdf(marca.clave_sociedad(company.cif if company else None), marca.clave_activo(asset.codigo),
                             ancho, 14 * mm)
    if cab is not None:
        el += [cab, Spacer(1, 3 * mm)]
    return out, doc, ancho, el


def _estilos():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    st = getSampleStyleSheet()
    return (ParagraphStyle("p", parent=st["BodyText"], fontSize=8.5, leading=10.5),
            ParagraphStyle("h", parent=st["Title"], fontSize=15, alignment=0, spaceAfter=2))


def _tabla(datos, anchos, repetir=1):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    from . import marca
    t = Table(datos, colWidths=anchos, repeatRows=repetir)
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(marca.NEGRO)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cfc6b4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f3ea")])]))
    return t


def pdf_parte(asset, company, r: WorkReport, filas: list[dict], validador: str | None) -> bytes:
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer
    titulo = f"Parte de trabajo · {AREAS_PARTE[r.area]} · {asset.nombre}"
    out, doc, ancho, el = _doc_pdf(f"{titulo} {r.fecha:%d-%m-%Y}", asset, company)
    p, h = _estilos()
    el += [Paragraph(titulo, h),
           Paragraph(f"{r.fecha:%d/%m/%Y} · {len(filas)} trabajo(s) · validado por <b>{_esc(validador)}</b> el "
                     f"{r.validado_en:%d/%m/%Y %H:%M}", p), Spacer(1, 3 * mm)]
    datos = [["Hora", "Ref.", "Origen", "Ubicación", "Trabajo realizado", "Persona", "Horas"]]
    for x in filas:
        datos.append([x["hora"], x["ref"], Paragraph(_esc(ORIGENES.get(x["origen"])), p), Paragraph(_esc(x["ubicacion"]), p),
                      Paragraph(f"<b>{_esc(x['trabajo'])}</b>" + (f"<br/>{_esc(x['detalle'])}" if x["detalle"] else ""), p),
                      Paragraph(_esc(x["persona"]), p), f"{x['horas']:.2f}".replace(".", ",") if x.get("horas") else ""])
    if len(datos) == 1:
        datos.append(["", "", "", "", Paragraph("Sin trabajos registrados este día.", p), "", ""])
    el.append(_tabla(datos, [14 * mm, 20 * mm, 30 * mm, 48 * mm, ancho - 160 * mm, 34 * mm, 14 * mm]))
    if r.observaciones:
        el += [Spacer(1, 3 * mm), Paragraph(f"<b>Observaciones:</b> {_esc(r.observaciones)}", p)]
    doc.build(el)
    return out.getvalue()


def excel_parte(asset, r: WorkReport, filas: list[dict], validador: str | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Parte de trabajo"
    ws.append([f"Parte de trabajo · {AREAS_PARTE[r.area]} · {asset.nombre} · {r.fecha:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Validado por {validador or ''} el {r.validado_en:%d/%m/%Y %H:%M}" if r.validado_en else "Sin validar"])
    cab = ["Hora", "Referencia", "Origen", "Ubicación", "Trabajo", "Detalle / solución", "Persona", "Horas"]
    ws.append(cab)
    for c in ws[3]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
    for x in filas:
        ws.append([x["hora"], x["ref"], ORIGENES.get(x["origen"]), x["ubicacion"], x["trabajo"], x["detalle"],
                   x["persona"], x.get("horas")])
    if r.observaciones:
        ws.append([])
        ws.append(["Observaciones", r.observaciones])
    for col, w in zip("ABCDEFGH", (8, 12, 20, 30, 40, 50, 22, 8), strict=True):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- informe quincenal
def quincena(anio: int, mes: int, q: int) -> tuple[date, date]:
    ini = date(anio, mes, 1 if q == 1 else 16)
    fin = date(anio, mes, 15) if q == 1 else (date(anio + (mes == 12), mes % 12 + 1, 1) - timedelta(days=1))
    return ini, fin


def quincena_anterior(hoy: date) -> tuple[int, int, int]:
    """La última quincena terminada: (año, mes, 1|2)."""
    if hoy.day > 15:
        return hoy.year, hoy.month, 1
    prev = hoy.replace(day=1) - timedelta(days=1)
    return prev.year, prev.month, 2


def informe(db: Session, asset_id: int, desde: date, hasta: date) -> dict:
    """Totales de la quincena por área: trabajos, por origen, por persona y por día, y días sin validar."""
    out = {"desde": desde.isoformat(), "hasta": hasta.isoformat(), "areas": {}}
    validados = {(r.area, r.fecha): r for r in db.scalars(select(WorkReport).where(
        WorkReport.asset_id == asset_id, WorkReport.fecha >= desde, WorkReport.fecha <= hasta))}
    for area in AREAS_PARTE:
        dias, origen, personas, categorias, todas = [], Counter(), Counter(), Counter(), []
        horas, sin_validar = 0.0, []
        d = desde
        while d <= hasta:
            r = validados.get((area, d))
            filas = r.lineas if r is not None and r.estado == "validado" and r.lineas is not None else lineas(db, asset_id, area, d)
            if filas and not (r is not None and r.estado == "validado"):
                sin_validar.append(d.isoformat())
            dias.append({"fecha": d.isoformat(), "trabajos": len(filas),
                         "validado": bool(r is not None and r.estado == "validado")})
            for x in filas:
                origen[x["origen"]] += 1
                categorias[x.get("categoria") or "otro"] += 1
                if x.get("persona"):
                    for persona in [s.strip() for s in str(x["persona"]).split(",") if s.strip()]:
                        personas[persona] += 1
                horas += x.get("horas") or 0
                todas.append({**x, "fecha": d.isoformat()})
            d += timedelta(days=1)
        out["areas"][area] = {"nombre": AREAS_PARTE[area], "total": len(todas), "por_origen": dict(origen),
                              "por_tipo": dict(categorias), "por_persona": dict(personas.most_common()),
                              "horas_anotadas": round(horas, 2), "dias": dias, "sin_validar": sin_validar,
                              "trabajos": todas}
    return out


def pdf_informe(asset, company, inf: dict) -> bytes:
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, Spacer
    d0, d1 = date.fromisoformat(inf["desde"]), date.fromisoformat(inf["hasta"])
    titulo = f"Informe quincenal de trabajos · {asset.nombre}"
    out, doc, ancho, el = _doc_pdf(f"{titulo} {d0:%d-%m-%Y}", asset, company)
    p, h = _estilos()
    el += [Paragraph(titulo, h), Paragraph(f"Del {d0:%d/%m/%Y} al {d1:%d/%m/%Y}", p), Spacer(1, 3 * mm)]
    resumen = [["Área", "Trabajos", "Órdenes de trabajo", "Limpiezas del parte", "Actuaciones anotadas",
                "Horas anotadas", "Días sin validar"]]
    for a in inf["areas"].values():
        o = a["por_origen"]
        resumen.append([a["nombre"], a["total"], o.get("ot", 0), o.get("limpieza", 0), o.get("actuacion", 0),
                        f"{a['horas_anotadas']:.2f}".replace(".", ","), len(a["sin_validar"])])
    el += [_tabla(resumen, [40 * mm, 22 * mm, 36 * mm, 36 * mm, 38 * mm, 30 * mm, ancho - 202 * mm]), Spacer(1, 4 * mm)]
    for a in inf["areas"].values():
        if not a["total"]:
            continue
        el.append(Paragraph(f"<b>{_esc(a['nombre'])}</b> · por persona: " + (", ".join(
            f"{_esc(k)} ({v})" for k, v in a["por_persona"].items()) or "sin indicar"), p))
        if a["sin_validar"]:
            el.append(Paragraph("Días con trabajos sin validar: " + ", ".join(
                f"{date.fromisoformat(x):%d/%m}" for x in a["sin_validar"]), p))
        el.append(Spacer(1, 2 * mm))
    for a in inf["areas"].values():
        if not a["trabajos"]:
            continue
        el += [PageBreak(), Paragraph(f"{_esc(a['nombre'])} · detalle de trabajos", h)]
        datos = [["Fecha", "Ref.", "Ubicación", "Trabajo", "Persona"]]
        for x in a["trabajos"]:
            datos.append([f"{date.fromisoformat(x['fecha']):%d/%m}", x["ref"], Paragraph(_esc(x["ubicacion"]), p),
                          Paragraph(f"<b>{_esc(x['trabajo'])}</b>" + (f"<br/>{_esc(x['detalle'])}" if x["detalle"] else ""), p),
                          Paragraph(_esc(x["persona"]), p)])
        el.append(_tabla(datos, [16 * mm, 20 * mm, 50 * mm, ancho - 126 * mm, 40 * mm]))
    doc.build(el)
    return out.getvalue()


def excel_informe(asset, inf: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    negrita, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Informe quincenal de trabajos · {asset.nombre} · {inf['desde']} a {inf['hasta']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Área", "Trabajos", "Órdenes de trabajo", "Limpiezas del parte", "Actuaciones anotadas",
               "Horas anotadas", "Días sin validar"])
    for c in ws[2]:
        c.font, c.fill = negrita, fondo
    for a in inf["areas"].values():
        o = a["por_origen"]
        ws.append([a["nombre"], a["total"], o.get("ot", 0), o.get("limpieza", 0), o.get("actuacion", 0),
                   a["horas_anotadas"], len(a["sin_validar"])])
    ws.append([])
    ws.append(["Área", "Persona", "Trabajos"])
    for c in ws[ws.max_row]:
        c.font, c.fill = negrita, fondo
    for a in inf["areas"].values():
        for k, v in a["por_persona"].items():
            ws.append([a["nombre"], k, v])
    for col, w in zip("ABCDEFG", (18, 28, 18, 20, 22, 15, 16), strict=True):
        ws.column_dimensions[col].width = w
    dt = wb.create_sheet("Trabajos")
    dt.append(["Fecha", "Área", "Referencia", "Origen", "Ubicación", "Trabajo", "Detalle", "Persona", "Horas"])
    for c in dt[1]:
        c.font, c.fill = negrita, fondo
    for a in inf["areas"].values():
        for x in a["trabajos"]:
            dt.append([date.fromisoformat(x["fecha"]), a["nombre"], x["ref"], ORIGENES.get(x["origen"]), x["ubicacion"],
                       x["trabajo"], x["detalle"], x["persona"], x.get("horas")])
    for fila in dt.iter_rows(min_row=2):
        fila[0].number_format = "DD/MM/YYYY"
    for col, w in zip("ABCDEFGHI", (11, 14, 12, 18, 28, 40, 45, 22, 8), strict=True):
        dt.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

