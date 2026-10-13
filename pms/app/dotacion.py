"""Dotación estándar de los apartamentos turísticos (categoría 3 llaves): catálogo maestro por zonas con la cantidad
de cada perfil. Los perfiles siguen la capacidad base de aprovisionamiento:

    estudio   Estudio, 2 plazas                       (AMENITY_STUDIO_3K)
    1d        1 dormitorio, 2 a 4 plazas              (AMENITY_1BED_3K)
    2d        2 dormitorios, 4 a 6 plazas             (AMENITY_2BED_3K)
    vip       VIP / Premium: misma base con calidades superiores (AMENITY_VIP_3K)

En vajilla y cubertería la regla habitual es «nº de plazas + 2» unidades de repuesto. Un «—» o «No» significa que ese
perfil no lo lleva. Recepción parte de esta lista para fijar la dotación habitual de cada apartamento."""

PERFILES = {"estudio": "Estudio (2 plazas)", "1d": "1 dormitorio (2-4 plazas)", "2d": "2 dormitorios (4-6 plazas)",
            "vip": "VIP / Premium"}
CODIGO_PERFIL = {"estudio": "AMENITY_STUDIO_3K", "1d": "AMENITY_1BED_3K", "2d": "AMENITY_2BED_3K", "vip": "AMENITY_VIP_3K"}

ZONAS = {
    "salon": "Zona de estar / Salón - comedor",
    "cocina": "Cocina: electrodomésticos y gran menaje",
    "menaje": "Cocina: vajilla, cubertería y pequeño menaje",
    "dormitorio": "Dormitorios",
    "bano": "Baño(s)",
    "auxiliar": "Limpieza, climatización y elementos auxiliares",
}

# (código, zona, artículo, estudio, 1 dormitorio, 2 dormitorios, VIP)
_TABLA = [
    # 1. Salón - comedor
    ("SAL01", "salon", "Smart TV", '43"', '50"', '50"-55"', '55"-65" OLED/QLED + barra de sonido'),
    ("SAL02", "salon", "Mando a distancia TV con funda", "1", "1", "1", "1"),
    ("SAL03", "salon", "Mesa de comedor", "Integrada / barra", "1 (4 plazas)", "1 (6 plazas)", "1 de diseño (4-6 plazas)"),
    ("SAL04", "salon", "Sillas de comedor", "2", "4", "6", "4 a 6 tapizadas"),
    ("SAL05", "salon", "Sofá / sofá-cama con apertura italiana", "1 (2 plazas)", "1 sofá-cama 140", "1 sofá-cama 160",
     "Sofá de diseño, apertura premium"),
    ("SAL06", "salon", "Mesa de centro / auxiliar", "1", "1", "1-2", "1-2 de diseño"),
    ("SAL07", "salon", "Lámpara de pie / luz indirecta", "1", "1", "1", "1 regulable (dimmer)"),
    ("SAL08", "salon", "Cuadros / elementos decorativos", "1-2", "2-3", "3-4", "Set de decoración"),
    ("SAL09", "salon", "Espejo de cuerpo entero", "1", "1", "1", "1 retroiluminado"),
    ("SAL10", "salon", "Enchufes USB / tomas en pared", "2", "2-4", "4", "USB-C rápidos + cargador inalámbrico"),
    # 2. Cocina: electrodomésticos y gran menaje
    ("COC01", "cocina", "Placa de cocción (inducción / vitro)", "2 fuegos", "2-3 fuegos", "3-4 fuegos",
     "3-4 fuegos inducción flexible"),
    ("COC02", "cocina", "Campana extractora", "1", "1", "1", "1 integrada silenciosa"),
    ("COC03", "cocina", "Frigorífico", "Bajo encimera / combi", "Combi no frost", "Combi no frost",
     "Combi integrable gran formato"),
    ("COC04", "cocina", "Microondas con grill", "1", "1", "1", "1 integrable"),
    ("COC05", "cocina", "Horno convencional", "No", "Opcional", "1 integrado", "1 pirolítico / multifunción"),
    ("COC06", "cocina", "Lavavajillas", "45 cm o compacto", "45 cm", "60 cm", "60 cm integrado silencioso"),
    ("COC07", "cocina", "Lavadora / lavadora-secadora", "Central o 1", "1", "1", "1 lavadora-secadora integrada"),
    ("COC08", "cocina", "Cafetera de cápsulas", "1", "1", "1", "1 Nespresso / superautomática + espumador"),
    ("COC09", "cocina", "Hervidor de agua", "1", "1", "1", "1 de temperatura regulable"),
    ("COC10", "cocina", "Tostadora", "1", "1", "1", "1 doble ranura ancha"),
    ("COC11", "cocina", "Batidora de mano con vaso", "1", "1", "1", "1 con set de accesorios"),
    ("COC12", "cocina", "Exprimidor", "1", "1", "1", "1 eléctrico"),
    # 3. Vajilla, cubertería y pequeño menaje
    ("MEN01", "menaje", "Platos llanos", "4", "6", "8", "8-10 porcelana / gres alta gama"),
    ("MEN02", "menaje", "Platos hondos", "4", "6", "8", "8-10 porcelana / gres alta gama"),
    ("MEN03", "menaje", "Platos de postre", "4", "6", "8", "8-10 porcelana / gres alta gama"),
    ("MEN04", "menaje", "Vasos de agua / refresco", "4", "6", "8", "8-10 cristal fino"),
    ("MEN05", "menaje", "Copas de vino", "2", "4", "6", "6-8 tinto / blanco"),
    ("MEN06", "menaje", "Copas de cava / champán", "—", "—", "—", "4-6"),
    ("MEN07", "menaje", "Tazas de café con plato", "4", "6", "8", "8"),
    ("MEN08", "menaje", "Tazas de desayuno (mugs)", "4", "6", "8", "8"),
    ("MEN09", "menaje", "Tenedores", "4", "6", "8", "8-10 acero 18/10"),
    ("MEN10", "menaje", "Cuchillos de mesa", "4", "6", "8", "8-10 acero 18/10"),
    ("MEN11", "menaje", "Cucharas", "4", "6", "8", "8-10 acero 18/10"),
    ("MEN12", "menaje", "Cucharillas", "4", "6", "8", "8-10 acero 18/10"),
    ("MEN13", "menaje", "Cuchillos chuleteros", "2", "4", "6", "6"),
    ("MEN14", "menaje", "Sartenes (pequeña, mediana, grande)", "2", "3", "3", "3 aluminio forjado / antiadherente"),
    ("MEN15", "menaje", "Cazuela, cacerola y olla con tapa", "2 piezas", "3 piezas", "3 piezas",
     "3 piezas acero inoxidable inducción"),
    ("MEN16", "menaje", "Set de cuchillos (chef, panero y puntilla)", "1 set", "1 set", "1 set",
     "1 set con taco o barra imantada"),
    ("MEN17", "menaje", "Tabla de corte", "1", "1-2", "2", "2 bambú / acacia"),
    ("MEN18", "menaje", "Ensaladera / bol grande", "1", "1", "2", "2"),
    ("MEN19", "menaje", "Utensilios: cucharón, espumadera, espátula y pinzas", "1 set", "1 set", "1 set",
     "1 set silicona / madera"),
    ("MEN20", "menaje", "Abrelatas", "1", "1", "1", "1"),
    ("MEN21", "menaje", "Sacacorchos / descorchador", "1", "1", "1", "1 profesional de dos tiempos"),
    ("MEN22", "menaje", "Pelador", "1", "1", "1", "1"),
    ("MEN23", "menaje", "Colador de pasta / verduras", "1", "1", "1", "1"),
    ("MEN24", "menaje", "Escurreplatos", "1", "1", "1", "Oculto / integrado"),
    ("MEN25", "menaje", "Cubo de basura con separación de reciclaje", "1", "1", "1", "1 doble / triple con pedal"),
    ("MEN26", "menaje", "Salvamanteles térmicos", "2", "2", "3", "3-4"),
    # 4. Dormitorios
    ("DOR01", "dormitorio", "Cama principal", "150/160", "160/180", "Principal 160/180 + secundario 2x90",
     "King size 180x200 o 200x200"),
    ("DOR02", "dormitorio", "Canapé o base tapizada con cabecero", "1", "1", "2", "Tapizado de diseño"),
    ("DOR03", "dormitorio", "Mesillas de noche con luz de lectura", "1-2", "2", "4 (2 por habitación)",
     "2 a 4 con USB / cargador inductivo"),
    ("DOR04", "dormitorio", "Smart TV en dormitorio principal", "No", "Opcional", "Opcional", '1 de 43"-50" en pared'),
    ("DOR05", "dormitorio", "Armario empotrado o exento", "1 cuerpo", "2 cuerpos", "2 cuerpos por habitación",
     "Con iluminación LED interior"),
    ("DOR06", "dormitorio", "Perchas antideslizantes de madera", "8", "12", "20 (10 por habitación)", "16-24 madera maciza"),
    ("DOR07", "dormitorio", "Caja fuerte electrónica (apta para portátil)", "1", "1", "1", "1 motorizada con luz"),
    ("DOR08", "dormitorio", "Almohadas", "2", "4 (2 firmes, 2 medias)", "6-8", "Carta de almohadas"),
    ("DOR09", "dormitorio", "Protectores de colchón y almohada", "1 juego", "1 juego", "2 juegos",
     "Doble juego impermeable transpirable"),
    ("DOR10", "dormitorio", "Juego de sábanas (200+ hilos)", "1 por cama", "1 por cama", "1 por cama",
     "Algodón egipcio / satén 300+ hilos"),
    ("DOR11", "dormitorio", "Nórdico / edredón con funda y cuadrantes", "1", "1", "2", "Nórdico 4 estaciones + cojines"),
    ("DOR12", "dormitorio", "Manta auxiliar / plaid", "1", "1", "2", "2 plaids"),
    # 5. Baño(s)
    ("BAN01", "bano", "Plato de ducha con mampara", "1", "1", "1 o 2 baños", "Ducha efecto lluvia termostática"),
    ("BAN02", "bano", "Espejo de lavabo", "1", "1", "1 o 2", "Retroiluminado LED antivaho"),
    ("BAN03", "bano", "Espejo de aumento", "Opcional", "1", "1", "1 articulado iluminado"),
    ("BAN04", "bano", "Secador de pelo (mín. 1800 W)", "1", "1", "1", "1 iónico alta potencia"),
    ("BAN05", "bano", "Papelera de baño con pedal", "1", "1", "1 o 2", "Inox / negro mate cierre amortiguado"),
    ("BAN06", "bano", "Escobillero de váter cerrado", "1", "1", "1 o 2", "1 o 2 a juego con grifería"),
    ("BAN07", "bano", "Dosificadores (gel, champú, jabón de manos)", "1 set", "1 set", "1 o 2 sets",
     "Dispensadores recargables premium"),
    ("BAN08", "bano", "Toalleros y perchas", "2 perchas", "2-3 perchas", "4 perchas", "Toallero radiador eléctrico"),
    ("BAN09", "bano", "Toallas de baño (100x150)", "2", "4", "6", "6 de 550-600 g/m²"),
    ("BAN10", "bano", "Toallas de lavabo", "2", "4", "6", "6 de 550-600 g/m²"),
    ("BAN11", "bano", "Alfombrín de baño", "2", "4", "6", "6"),
    ("BAN12", "bano", "Albornoces y zapatillas de cortesía", "—", "—", "—", "2 a 4"),
    # 6. Limpieza, climatización y auxiliares
    ("AUX01", "auxiliar", "Climatización frío / calor (termostato digital)", "1 split / rejilla", "Salón + dormitorio",
     "Salón + 2 dormitorios", "Por zonas / domotizada"),
    ("AUX02", "auxiliar", "Cerradura electrónica / inteligente", "1", "1", "1", "Teclado / móvil / código PMS"),
    ("AUX03", "auxiliar", "Plancha de vapor y tabla de planchar", "1 tabla compacta", "1 juego", "1 juego",
     "Centro de planchado / vapor vertical"),
    ("AUX04", "auxiliar", "Tendedero plegable", "1", "1", "1", "1"),
    ("AUX05", "auxiliar", "Cubo de fregar con escurridor y fregona", "1", "1", "1", "1 en armario técnico"),
    ("AUX06", "auxiliar", "Escoba y recogedor", "1", "1", "1", "1"),
    ("AUX07", "auxiliar", "Aspirador / escoba eléctrica", "Opcional", "Opcional", "Opcional", "1 vertical inalámbrico"),
    ("AUX08", "auxiliar", "Extintor homologado (DB-SI)", "1 (o en pasillo)", "1 (o en pasillo)", "1 (o en pasillo)",
     "1 (o en pasillo común)"),
    ("AUX09", "auxiliar", "Manta apagafuegos en cocina", "1", "1", "1", "1"),
    ("AUX10", "auxiliar", "Botiquín de primeros auxilios", "1", "1", "1", "1"),
]

# Precio de reposición orientativo por unidad (€, IVA incluido, gama media de hostelería). Cada activo los ajusta en
# «Precios de reposición»; las instalaciones fijas (enchufes, plato de ducha, climatización, cerradura) no llevan.
PRECIOS_REPOSICION = {
    "SAL01": 450, "SAL02": 25, "SAL03": 180, "SAL04": 45, "SAL05": 650, "SAL06": 80, "SAL07": 45, "SAL08": 30,
    "SAL09": 60,
    "COC01": 250, "COC02": 150, "COC03": 450, "COC04": 120, "COC05": 350, "COC06": 400, "COC07": 400, "COC08": 90,
    "COC09": 25, "COC10": 25, "COC11": 30, "COC12": 20,
    "MEN01": 4, "MEN02": 4, "MEN03": 3.5, "MEN04": 2.5, "MEN05": 4, "MEN06": 4, "MEN07": 5, "MEN08": 5,
    "MEN09": 2, "MEN10": 2.5, "MEN11": 2, "MEN12": 1.5, "MEN13": 3, "MEN14": 25, "MEN15": 30, "MEN16": 35,
    "MEN17": 12, "MEN18": 10, "MEN19": 15, "MEN20": 6, "MEN21": 8, "MEN22": 4, "MEN23": 10, "MEN24": 15,
    "MEN25": 35, "MEN26": 6,
    "DOR01": 450, "DOR02": 250, "DOR03": 70, "DOR04": 350, "DOR05": 300, "DOR06": 2, "DOR07": 120, "DOR08": 20,
    "DOR09": 25, "DOR10": 45, "DOR11": 70, "DOR12": 25,
    "BAN02": 60, "BAN03": 25, "BAN04": 35, "BAN05": 15, "BAN06": 12, "BAN07": 20, "BAN08": 15, "BAN09": 15,
    "BAN10": 8, "BAN11": 10, "BAN12": 30,
    "AUX03": 60, "AUX04": 25, "AUX05": 20, "AUX06": 10, "AUX07": 120, "AUX08": 40, "AUX09": 15, "AUX10": 25,
}

CATALOGO = [{"codigo": c, "zona": z, "articulo": a, "cantidades": dict(zip(PERFILES, cant))}
            for c, z, a, *cant in _TABLA]
POR_CODIGO = {x["codigo"]: x for x in CATALOGO}
NO_INCLUIDO = {"—", "No", ""}


def perfil_de(tipologia: str | None, vip: bool = False) -> str:
    """Perfil de dotación según la tipología de la unidad (estudio, 1 dormitorio, 2 dormitorios)."""
    if vip:
        return "vip"
    t = (tipologia or "").lower()
    if "2" in t:
        return "2d"
    if "1" in t:
        return "1d"
    if "est" in t or t in ("e", "st"):
        return "estudio"
    return "1d"


def estandar(perfil: str) -> list[dict]:
    """Dotación estándar de un perfil, lista para que recepción la ajuste (sin lo que ese perfil no lleva)."""
    return [{"codigo": x["codigo"], "zona": x["zona"], "articulo": x["articulo"], "cantidad": x["cantidades"][perfil]}
            for x in CATALOGO if x["cantidades"][perfil] not in NO_INCLUIDO]


# --------------------------------------------------------------------------- dotación habitual y de la reserva
def _db():  # imports locales: este módulo también se usa sin base de datos (catálogo)
    from sqlalchemy import select
    from .models import DotacionArticulo, DotacionHabitual
    return select, DotacionArticulo, DotacionHabitual


def precios(db, asset_id: int) -> dict[str, float]:
    """Precio de reposición de cada artículo en el activo: el fijado por el activo o el orientativo."""
    from sqlalchemy import select
    from .models import DotacionPrecio
    out = {k: float(v) for k, v in PRECIOS_REPOSICION.items()}
    out.update({p.codigo: float(p.precio) for p in db.scalars(select(DotacionPrecio).where(
        DotacionPrecio.asset_id == asset_id))})
    return out


def propia(db, unit_id: int):
    from sqlalchemy import select
    from .models import DotacionUnidad
    return db.scalar(select(DotacionUnidad).where(DotacionUnidad.unit_id == unit_id))


def articulos(db, asset_id: int | None) -> list[dict]:
    """Catálogo para elegir: el maestro 3 llaves y los artículos creados con «Otra dotación» (del activo o comunes)."""
    select, Art, _ = _db()
    propios = db.scalars(select(Art).where((Art.asset_id == asset_id) | Art.asset_id.is_(None)).order_by(Art.codigo))
    return CATALOGO + [{"codigo": a.codigo, "zona": a.zona, "articulo": a.articulo, "cantidades": {}, "propio": True}
                       for a in propios]


def nuevo_codigo(db) -> str:
    select, Art, _ = _db()
    n = len(list(db.scalars(select(Art.id)))) + 1
    while db.scalar(select(Art.id).where(Art.codigo == f"OTR{n:03d}")):
        n += 1
    return f"OTR{n:03d}"


def habitual(db, asset_id: int, tipologia: str | None):
    select, _, Hab = _db()
    return db.scalar(select(Hab).where(Hab.asset_id == asset_id, Hab.tipologia == (tipologia or "")))


def base_unidad(db, unit) -> tuple[list[dict], bool]:
    """(dotación habitual de la unidad, si Recepción ya la ha fijado): la propia del apartamento, si la tiene; si no,
    la de su tipología; y si tampoco, la estándar de su perfil."""
    return base_con_origen(db, unit)[:2]


def base_con_origen(db, unit) -> tuple[list[dict], bool, str]:
    p = propia(db, unit.id)
    if p:
        return list(p.items), True, "apartamento"
    h = habitual(db, unit.asset_id, unit.tipologia)
    if h:
        return list(h.items), True, "tipologia"
    return estandar(perfil_de(unit.tipologia)), False, "estandar"


def limpiar(items: list[dict]) -> list[dict]:
    """Valida y normaliza una lista [{codigo, zona, articulo, cantidad}] (sin repetidos ni vacíos)."""
    out, vistos = [], set()
    for x in items or []:
        cod = str(x.get("codigo") or "").strip().upper()[:12]
        art = str(x.get("articulo") or "").strip()[:160]
        cant = str(x.get("cantidad") if x.get("cantidad") is not None else "").strip()[:80]
        zona = x.get("zona") if x.get("zona") in ZONAS else "auxiliar"
        if not cod or not art or cant in NO_INCLUIDO or cod in vistos:
            continue
        vistos.add(cod)
        out.append({"codigo": cod, "zona": zona, "articulo": art, "cantidad": cant})
    orden = list(ZONAS)
    return sorted(out, key=lambda x: (orden.index(x["zona"]), x["codigo"]))


def de_reserva(db, r) -> dict:
    """Dotación que se entrega en una reserva: la habitual con los cambios (añadir / quitar / otra cantidad) y la
    dotación especial (artículos añadidos)."""
    base, fijada, origen = base_con_origen(db, r.unit)
    pre = precios(db, r.unit.asset_id)
    d = r.dotacion or {}
    cambios = d.get("cambios") or {}
    final = []
    for x in base:
        if x["codigo"] in cambios:
            c = str(cambios[x["codigo"]]).strip()
            if c in NO_INCLUIDO:
                continue
            x = {**x, "cantidad": c, "cambiado": True}
        final.append(x)
    extra = limpiar(d.get("extra") or [])
    codigos = {x["codigo"] for x in final}
    final += [{**x, "especial": True} for x in extra if x["codigo"] not in codigos]
    orden = list(ZONAS)
    final.sort(key=lambda x: (orden.index(x["zona"]) if x["zona"] in orden else 99, x["codigo"]))
    for x in final:
        x["precio"] = pre.get(x["codigo"])
    return {"habitual": base, "habitual_fijada": fijada, "origen": origen, "perfil": perfil_de(r.unit.tipologia),
            "tipologia": r.unit.tipologia, "cambios": cambios, "extra": extra, "items": final,
            "especial": bool(cambios or extra)}


def _bordes(tabla) -> None:
    """Bordes finos en la tabla (la plantilla del contrato no trae el estilo «Table Grid»)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "999999")):
            e.set(qn(k), v)
        b.append(e)
    pr.append(b)


def anexo(contenido: bytes, items: list[dict], titulo: str, subtitulo: str) -> bytes:
    """Añade al contrato (.docx) el Anexo I · Inventario y dotación del apartamento, con casilla de comprobación."""
    import io

    import docx
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    doc = docx.Document(io.BytesIO(contenido))
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()
    r = h.add_run(titulo)
    r.bold = True
    r.font.size = Pt(13)
    s = doc.add_paragraph(subtitulo)
    s.runs[0].font.size = Pt(9)
    t = doc.add_table(rows=1, cols=4)
    _bordes(t)
    for i, txt in enumerate(("Zona", "Artículo", "Cantidad", "Revisado")):
        c = t.rows[0].cells[i]
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
    for x in items:
        fila = t.add_row().cells
        fila[0].text = ZONAS.get(x["zona"], x["zona"]).split(":")[0].split("/")[0].strip()
        fila[1].text = x["articulo"] + (" (dotación especial)" if x.get("especial") else "")
        fila[2].text = str(x["cantidad"])
        fila[3].text = "☐"
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph()
    doc.add_paragraph("El cliente declara recibir el apartamento con la dotación indicada en buen estado de uso. "
                      "Los elementos que falten o se dañen durante la estancia podrán descontarse de la fianza.")
    doc.add_paragraph("Firma del cliente:\t\t\t\t\tPor la propiedad:")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def hoja_pdf(hojas: list[dict]) -> bytes:
    """Hoja de revisión de salida: una página por apartamento con su dotación por zonas y las casillas OK / Falta /
    Roto, el precio de reposición y las firmas. `hojas`: [{empresa, activo, apartamento, lineas: [texto], items}]."""
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from . import marca
    negro, linea, fondo = colors.HexColor(marca.NEGRO), colors.HexColor("#ddd6c8"), colors.HexColor("#f7f3ea")
    ss = getSampleStyleSheet()
    n = ParagraphStyle("n", parent=ss["Normal"], fontName="Helvetica", fontSize=8.5, leading=10.5)
    b = ParagraphStyle("b", parent=n, fontName="Helvetica-Bold")
    tit = ParagraphStyle("t", parent=n, fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=negro)
    zona = ParagraphStyle("z", parent=n, fontName="Helvetica-Bold", fontSize=9, textColor=colors.white)

    def esc(t):
        return str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title="Hoja de revisión de salida", creator=marca.NOMBRE)
    anchos = [72 * mm, 30 * mm, 13 * mm, 13 * mm, 13 * mm, 16 * mm, 25 * mm]

    def caja():  # casilla vacía para marcar a mano
        return Table([[""]], colWidths=[3.6 * mm], rowHeights=[3.6 * mm],
                     style=[("BOX", (0, 0), (-1, -1), 0.7, negro)])
    historia = []
    for i, h in enumerate(hojas):
        if i:
            historia.append(PageBreak())
        historia += [Paragraph(f"HOJA DE REVISIÓN DE SALIDA · {esc(h['apartamento'])}", tit),
                     Paragraph(esc(h["empresa"]) + " · " + esc(h["activo"]), b)]
        historia += [Paragraph(esc(x), n) for x in h["lineas"]]
        historia.append(Spacer(1, 4))
        filas = [[Paragraph(x, b) for x in ("Artículo", "Cantidad", "OK", "Falta", "Roto", "Nº", "Reposición €")]]
        estilo = [("GRID", (0, 0), (-1, -1), 0.4, linea), ("BACKGROUND", (0, 0), (-1, 0), fondo),
                  ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                  ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("ALIGN", (2, 1), (4, -1), "CENTER")]
        for z, nz in ZONAS.items():
            del_z = [x for x in h["items"] if x["zona"] == z]
            if not del_z:
                continue
            filas.append([Paragraph(esc(nz), zona), "", "", "", "", "", ""])
            k = len(filas) - 1
            estilo += [("SPAN", (0, k), (-1, k)), ("BACKGROUND", (0, k), (-1, k), negro)]
            for x in del_z:
                p = x.get("precio")
                filas.append([Paragraph(esc(x["articulo"]) + (" <i>(especial)</i>" if x.get("especial") else ""), n),
                              Paragraph(esc(x["cantidad"]), n), caja(), caja(), caja(), "",
                              Paragraph(f"{p:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if p else "—",
                                        n)])
        t = Table(filas, colWidths=anchos, repeatRows=1)
        t.setStyle(TableStyle(estilo))
        historia += [t, Spacer(1, 6),
                     Paragraph("Observaciones (manchas, desperfectos, objetos olvidados): "
                               "_______________________________________________________________", n),
                     Spacer(1, 10),
                     Table([[Paragraph("Revisado por (limpieza / gobernanta):<br/><br/>Fecha y firma:", n),
                             Paragraph("Recepción (cargos en la fianza):<br/><br/>Fecha y firma:", n)]],
                           colWidths=[91 * mm, 91 * mm],
                           style=[("BOX", (0, 0), (0, 0), 0.4, linea), ("BOX", (1, 0), (1, 0), 0.4, linea),
                                  ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 14)])]
    doc.build(historia)
    return buf.getvalue()


def hoja_de_reserva(db, r) -> dict:
    """Datos de la hoja de revisión de una reserva (su dotación, con la especial y los precios)."""
    u = r.unit
    g = r.guest
    return {"empresa": u.asset.company.nombre, "activo": u.asset.nombre,
            "apartamento": f"{u.codigo}" + (f" ({u.tipologia})" if u.tipologia else ""),
            "lineas": [f"Reserva {r.localizador or r.id} · {g.nombre} {g.apellidos or ''}".strip()
                       + f" · Estancia {r.fecha_entrada:%d/%m/%Y} – {r.fecha_salida:%d/%m/%Y}",
                       "Marque cada artículo. Lo que falte o esté roto se descuenta de la fianza a su precio de "
                       "reposición; indique el número de unidades en «Nº»."],
            "items": de_reserva(db, r)["items"]}
