"""Lectura de reservas desde Excel (.xlsx) o CSV.

Las columnas se reconocen por su nombre, sin importar mayúsculas, tildes ni el orden, y admiten los nombres
habituales de las exportaciones de Booking y similares («Número de reserva», «Check-in», «Precio»...).
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# campo -> nombres de columna admitidos (normalizados: minúsculas, sin tildes ni signos)
COLUMNAS = {
    "localizador": ["localizador", "numero de reserva", "n reserva", "reserva", "referencia", "booking number",
                    "reservation number", "id reserva", "codigo reserva"],
    "unidad": ["unidad", "apartamento", "habitacion", "numero de habitacion", "room", "unit"],
    "fecha_entrada": ["entrada", "fecha entrada", "fecha de entrada", "llegada", "fecha llegada", "check in",
                      "checkin", "arrival"],
    "fecha_salida": ["salida", "fecha salida", "fecha de salida", "check out", "checkout", "departure"],
    "adultos": ["adultos", "adults", "personas", "huespedes", "pax", "numero de personas"],
    "ninos": ["ninos", "children", "menores"],
    "canal": ["canal", "origen", "channel", "fuente", "agencia"],
    "importe_total": ["importe", "importe total", "total", "precio", "price", "precio total", "tarifa"],
    "nombre": ["nombre", "nombre del cliente", "nombre huesped", "first name", "nombre cliente", "cliente",
               "huesped", "guest name", "nombre del huesped", "reservado por", "booker name"],
    "apellidos": ["apellidos", "apellido", "last name", "surname"],
    "documento_num": ["documento", "dni", "dni nie pasaporte", "pasaporte", "nie", "numero de documento"],
    "nacionalidad": ["nacionalidad", "pais", "country", "nationality"],
    "email": ["email", "correo", "correo electronico", "e mail"],
    "telefono": ["telefono", "movil", "phone", "telefono movil"],
    "notas": ["notas", "observaciones", "comentarios", "remarks", "peticiones"],
    "estado": ["estado", "status"],
}
OBLIGATORIAS = ("fecha_entrada", "fecha_salida")
CANALES = {"booking": "booking", "booking.com": "booking", "airbnb": "airbnb", "expedia": "expedia", "web": "web",
           "directo": "directo", "agencia": "agencia"}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


_ALIAS = {_norm(alias): campo for campo, aliases in COLUMNAS.items() for alias in aliases}


def _filas_crudas(datos: bytes, nombre: str) -> list[list]:
    if datos[:2] == b"PK":  # .xlsx
        wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(f) for f in ws.iter_rows(values_only=True)]
    if datos[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("Formato .xls antiguo: ábralo en Excel y guárdelo como .xlsx")
    for cod in ("utf-8-sig", "cp1252"):
        try:
            texto = datos.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("No se reconoce el fichero: use Excel (.xlsx) o CSV")
    sep = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
    return list(csv.reader(io.StringIO(texto), delimiter=sep))


def leer(datos: bytes, nombre: str) -> tuple[list[dict], list[str]]:
    """Devuelve las filas como diccionarios (campo -> valor) con su nº de fila, y las columnas no reconocidas."""
    filas = [f for f in _filas_crudas(datos, nombre) if any(c not in (None, "") for c in f)]
    if not filas:
        raise ValueError("El fichero está vacío")
    # la cabecera es la primera fila donde se reconocen al menos las fechas (puede haber títulos encima)
    for i, f in enumerate(filas[:10]):
        campos = [_ALIAS.get(_norm(c)) for c in f]
        if all(o in campos for o in OBLIGATORIAS):
            break
    else:
        raise ValueError("No se encuentran las columnas de fecha de entrada y de salida. Use la plantilla.")
    ignoradas = [str(c) for c, campo in zip(filas[i], campos) if c not in (None, "") and not campo]
    out = []
    for n, f in enumerate(filas[i + 1:], start=i + 2):
        d = {"fila": n}
        for campo, valor in zip(campos, f):
            if campo and campo not in d and valor not in (None, ""):
                d[campo] = valor
        out.append(d)
    return out, ignoradas


def fecha(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):  # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha no válida «{v}» (use dd/mm/aaaa)")


def importe(v) -> float:
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if "," in s and "." in s:  # 1.234,56 o 1,234.56
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        raise ValueError(f"importe no válido «{v}»")


def entero(v, defecto: int) -> int:
    if v in (None, ""):
        return defecto
    try:
        return int(float(str(v).replace(",", ".")))
    except ValueError:
        raise ValueError(f"número no válido «{v}»")


def canal(v) -> str:
    n = _norm(v)
    for k, c in CANALES.items():
        if _norm(k) in n:
            return c
    return "otros" if n else "directo"


def plantilla() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reservas"
    cab = ["Localizador", "Unidad", "Entrada", "Salida", "Adultos", "Niños", "Canal", "Importe total", "Nombre",
           "Apellidos", "Documento", "Nacionalidad", "Email", "Teléfono", "Notas"]
    ws.append(cab)
    ws.append(["BK-123456", "P1-1A", date.today() + timedelta(days=7), date.today() + timedelta(days=10), 2, 0,
               "booking", 345.50, "María", "García López", "12345678Z", "ESP", "maria@ejemplo.com", "600000000", ""])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="13294B")
    for col, ancho in zip("ABCDEFGHIJKLMNO", (14, 10, 12, 12, 9, 7, 11, 13, 14, 20, 13, 12, 24, 13, 30)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in fila:
            c.number_format = "DD/MM/YYYY"
    ins = wb.create_sheet("Instrucciones")
    for linea in (
            "Una fila por reserva. Obligatorias: Entrada y Salida. Borre la fila de ejemplo.",
            "Unidad: código del apartamento tal como está en el PMS (P1-1A, A-127...). Si se deja vacía, el PMS "
            "asigna un apartamento libre con capacidad suficiente.",
            "Fechas en formato dd/mm/aaaa. Importe total con IVA incluido.",
            "Canal: directo, booking, airbnb, expedia, web, agencia u otros.",
            "Las reservas con estado «cancelada» se ignoran. Un localizador que ya existe en el activo no se duplica.",
            "Huésped: si el documento ya existe en el PMS se reutiliza su ficha; si no, se crea.",
            "Al importar no se registra ningún cobro ni se emiten facturas: los cobros se registran después en cada "
            "reserva con el botón «Cobro».",
            "También se admiten las exportaciones de Booking (columnas «Número de reserva», «Check-in», «Precio»...)."):
        ins.append([linea])
    ins.column_dimensions["A"].width = 130
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
