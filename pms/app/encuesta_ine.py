"""Encuesta de Ocupación en Apartamentos Turísticos del INE (cuestionario mensual por establecimiento).

Con los ocupantes registrados de cada reserva se calcula lo que pide el cuestionario:
- viajeros entrados y pernoctaciones por día y por lugar de residencia (comunidad autónoma para los
  residentes en España, país para los residentes en el extranjero);
- apartamentos y plazas disponibles y ocupados cada día;
- tarifa media diaria por apartamento ocupado (sin IVA) por tipo de cliente.
El resultado se entrega en Excel para trasladarlo al cuestionario del INE (sistema IRIA). Si una reserva no tiene
registrados todos sus ocupantes, los que faltan se cuentan con la residencia del titular (y se avisa).
"""
from calendar import monthrange
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal

from sqlalchemy import select

from . import registro_viajeros as rv
from .models import Asset, Reservation, Unit

ESTADOS = ("confirmada", "checkin", "checkout")
IVA_ALOJAMIENTO = Decimal("1.10")
TIPO_CLIENTE = {  # canal de la reserva -> tipo de cliente del cuestionario (precios)
    "directo": "Particulares (contratación directa)", "web": "Contratación directa online del establecimiento",
    "booking": "Agencia de viajes online", "expedia": "Agencia de viajes online", "airbnb": "Agencia de viajes online",
    "agencia": "Agencia de viajes tradicional", "otros": "Otros",
}
EXTRANJERO = "Extranjero"


def residencia(c, titular=None) -> tuple[str, str]:
    """(grupo, detalle): ("España", comunidad autónoma) o ("Extranjero", país)."""
    for persona in (c, titular):
        if persona is None:
            continue
        iso = rv.pais_iso3(persona.pais) or rv.pais_iso3(persona.nacionalidad)
        if iso == "ESP":
            return "España", rv.comunidad(persona.cp, persona.municipio_ine) or "España (comunidad sin determinar)"
        if iso:
            return EXTRANJERO, rv.nombre_pais(iso)
    return EXTRANJERO, "Residencia sin determinar"


def calcular(db, asset: Asset, anio: int, mes: int) -> dict:
    ini = date(anio, mes, 1)
    dias = monthrange(anio, mes)[1]
    fin = date(anio, mes, dias)
    unidades = list(db.scalars(select(Unit).where(Unit.asset_id == asset.id, Unit.uso != "garaje")))
    abiertas = [u for u in unidades if u.estado != "fuera_servicio"]
    reservas = list(db.scalars(select(Reservation).join(Unit).where(
        Unit.asset_id == asset.id, Unit.uso != "garaje", Reservation.estado.in_(ESTADOS),
        Reservation.fecha_entrada <= fin, Reservation.fecha_salida > ini)))

    entradas = defaultdict(lambda: [0] * dias)
    salidas = defaultdict(lambda: [0] * dias)
    pernoct = defaultdict(lambda: [0] * dias)
    aptos = [0] * dias
    ingresos, aptos_cliente = defaultdict(Decimal), defaultdict(int)
    avisos = []
    for r in reservas:
        ocupantes = [o.contact for o in r.ocupantes] or [r.guest]
        faltan = max(0, r.adultos + r.ninos - len(ocupantes))
        if faltan:
            avisos.append(f"Reserva {r.localizador or r.id} ({r.unit.codigo}): {faltan} ocupante(s) sin registrar, "
                          "contados con la residencia del titular")
        claves = [residencia(c, r.guest) for c in ocupantes] + [residencia(r.guest)] * faltan
        noches_total = (r.fecha_salida - r.fecha_entrada).days or 1
        por_noche = Decimal(str(r.importe_total or 0)) / IVA_ALOJAMIENTO / noches_total
        tipo = TIPO_CLIENTE.get(r.canal, "Otros")
        for i in range(dias):
            d = ini + timedelta(days=i)
            for k in claves:
                if r.fecha_entrada == d:
                    entradas[k][i] += 1
                if r.fecha_salida == d:
                    salidas[k][i] += 1
                if r.fecha_entrada <= d < r.fecha_salida:
                    pernoct[k][i] += 1
            if r.fecha_entrada <= d < r.fecha_salida:
                aptos[i] += 1
                ingresos[tipo] += por_noche
                aptos_cliente[tipo] += 1
    plazas = sum(u.capacidad or 0 for u in abiertas)
    plazas_ocupadas = [sum(v[i] for v in pernoct.values()) for i in range(dias)]
    total_aptos_noche = sum(aptos)
    precios = [{"tipo": t, "apartamentos_noche": aptos_cliente[t],
                "porcentaje": aptos_cliente[t] / total_aptos_noche if total_aptos_noche else 0,
                "tarifa_media": float(round(ingresos[t] / aptos_cliente[t], 2)) if aptos_cliente[t] else 0}
               for t in dict.fromkeys(TIPO_CLIENTE.values()) if aptos_cliente[t]]
    orden = lambda k: (k[0] != "España", k[1])  # noqa: E731
    residencias = sorted(set(entradas) | set(pernoct) | set(salidas), key=orden)
    return {
        "activo": asset.nombre, "registro": asset.num_registro_turistico, "anio": anio, "mes": mes, "dias": dias,
        "apartamentos_disponibles": len(abiertas), "plazas_disponibles": plazas,
        "apartamentos_ocupados": aptos, "plazas_ocupadas": plazas_ocupadas,
        "residencias": [{"grupo": g, "residencia": d, "entradas": entradas[(g, d)], "salidas": salidas[(g, d)],
                         "pernoctaciones": pernoct[(g, d)]} for g, d in residencias],
        "totales": {"viajeros_entrados": sum(sum(v) for v in entradas.values()),
                    "pernoctaciones": sum(plazas_ocupadas), "apartamentos_noche": total_aptos_noche,
                    "ocupacion_apartamentos": total_aptos_noche / (len(abiertas) * dias) if abiertas else 0,
                    "ocupacion_plazas": sum(plazas_ocupadas) / (plazas * dias) if plazas else 0,
                    "tarifa_media": float(round(sum(ingresos.values()) / total_aptos_noche, 2)) if total_aptos_noche else 0},
        "precios": precios, "avisos": avisos,
        "sin_capacidad": [u.codigo for u in abiertas if not u.capacidad],
    }


def excel(datos: dict) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    from .routers.informes import ENTERO, EUR, PCT, _hoja

    dias = datos["dias"]
    t = datos["totales"]
    wb = Workbook()
    wb.remove(wb.active)
    titulo = f"{datos['activo']} · {datos['mes']:02d}/{datos['anio']}"
    _hoja(wb, "Resumen", ["Concepto", "Valor"], [
        ["Establecimiento", datos["activo"]], ["Nº de registro turístico", datos["registro"] or ""],
        ["Mes", f"{datos['mes']:02d}/{datos['anio']}"],
        ["Apartamentos disponibles (abiertos)", datos["apartamentos_disponibles"]],
        ["Plazas disponibles", datos["plazas_disponibles"]],
        ["Viajeros entrados", t["viajeros_entrados"]], ["Pernoctaciones", t["pernoctaciones"]],
        ["Apartamentos ocupados (suma de noches)", t["apartamentos_noche"]],
        ["Grado de ocupación por apartamentos", t["ocupacion_apartamentos"]],
        ["Grado de ocupación por plazas", t["ocupacion_plazas"]],
        ["Tarifa media diaria por apartamento ocupado, sin IVA", t["tarifa_media"]],
        ["Personal empleado (completar)", None],
    ], nota=f"Encuesta de Ocupación en Apartamentos Turísticos (INE) · {titulo}. "
            "Datos calculados con los ocupantes registrados de cada reserva.")
    ws = wb["Resumen"]
    for fila in ws.iter_rows(min_row=2):
        concepto = str(fila[0].value or "")
        if concepto.startswith("Grado"):
            fila[1].number_format = PCT
        elif concepto.startswith("Tarifa"):
            fila[1].number_format = EUR
    cab_dias = [str(i + 1) for i in range(dias)]
    for hoja, campo in (("Viajeros entrados", "entradas"), ("Pernoctaciones", "pernoctaciones"),
                        ("Viajeros salidos", "salidas")):
        filas = [[x["grupo"], x["residencia"], *x[campo], sum(x[campo])] for x in datos["residencias"]]
        _hoja(wb, hoja, ["Residencia", "Comunidad autónoma / país", *cab_dias, "Total mes"], filas,
              {i: ENTERO for i in range(2, dias + 3)}, totales=list(range(2, dias + 3)),
              nota=f"{hoja} por día y lugar de residencia · {titulo}")
    _hoja(wb, "Ocupación diaria", ["Día", "Apartamentos disponibles", "Apartamentos ocupados", "Plazas disponibles",
                                   "Plazas ocupadas (pernoctaciones)"],
          [[i + 1, datos["apartamentos_disponibles"], datos["apartamentos_ocupados"][i], datos["plazas_disponibles"],
            datos["plazas_ocupadas"][i]] for i in range(dias)], {1: ENTERO, 2: ENTERO, 3: ENTERO, 4: ENTERO},
          totales=[2, 4], nota=f"Ocupación diaria · {titulo}")
    _hoja(wb, "Precios", ["Tipo de cliente", "Apartamentos ocupados (noches)", "% sobre el total",
                          "Tarifa media diaria sin IVA"],
          [[p["tipo"], p["apartamentos_noche"], p["porcentaje"], p["tarifa_media"]] for p in datos["precios"]],
          {1: ENTERO, 2: PCT, 3: EUR}, nota="Tipo de cliente según el canal de la reserva")
    avisos = datos["avisos"] + ([f"Unidades sin capacidad indicada (no suman plazas): {', '.join(datos['sin_capacidad'])}"]
                                if datos["sin_capacidad"] else [])
    ws = wb.create_sheet("Avisos")
    ws.append(["Revise antes de enviar al INE"])
    ws["A1"].font = Font(bold=True)
    for a in avisos or ["Sin avisos: todos los ocupantes están registrados."]:
        ws.append([a])
    ws.column_dimensions["A"].width = 120
    wb.properties.creator = "INVERPMS"
    wb.properties.title = f"Encuesta INE apartamentos turísticos {titulo}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
