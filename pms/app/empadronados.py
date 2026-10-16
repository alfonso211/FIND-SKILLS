"""Domicilio habitual en la misma calle que los apartamentos turísticos («empadronados»).

Un huésped cuyo domicilio habitual es la calle del propio edificio no está de paso: el contrato de alojamiento
turístico no puede continuar salvo que sea un cliente antiguo (se indica la fecha de su primera estancia).
El documento «empadronados» cuenta las personas y los apartamentos que están ahora en esa situación.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .models import MODALIDADES_RESERVA, Asset, Contact, Reservation, Unit

# tipos de vía y palabras de enlace que no distinguen una calle de otra
_VIA = {"c", "cl", "calle", "av", "avd", "avda", "avenida", "pº", "po", "paseo", "pza", "plaza", "pl", "camino",
        "cno", "ctra", "carretera", "ronda", "rda", "travesia", "trva", "glorieta", "gta", "via", "callejon",
        "urbanizacion", "urb", "pasaje", "psje", "del", "de", "la", "las", "los", "el", "y"}


def _norm(s: str | None) -> str:
    s = unicodedata.normalize("NFKD", (s or "").lower())
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9ñ]+", " ", s).strip()


def calle(direccion: str | None) -> str:
    """Nombre de la calle sin tipo de vía, número, piso ni puerta: «C/ de Alcalá, 35 3ºB» → «alcala»."""
    palabras = []
    for p in _norm(direccion).split():
        if p[0].isdigit() or p in ("n", "no", "num", "numero", "s", "sn"):
            if palabras:
                break  # a partir del número ya es portal, piso o puerta
            continue
        if p not in _VIA:
            palabras.append(p)
    return " ".join(palabras)


def misma_calle(domicilio: str | None, municipio: str | None, a: Asset) -> bool:
    """El domicilio está en la calle del activo (y, si consta, en el mismo municipio)."""
    propia, suya = calle(a.direccion), calle(domicilio)
    if len(propia) < 3 or not suya:
        return False
    if municipio and a.municipio and _norm(municipio) != _norm(a.municipio):
        return False
    return f" {propia} " in f" {suya} "


def activos_turisticos(db: Session) -> list[Asset]:
    return list(db.scalars(select(Asset).where(Asset.modalidad.in_(MODALIDADES_RESERVA)).order_by(Asset.codigo)))


def coincide(db: Session, domicilio: str | None, municipio: str | None) -> Asset | None:
    """Activo turístico en cuya calle está el domicilio, si lo hay."""
    for a in activos_turisticos(db):
        if misma_calle(domicilio, municipio, a):
            return a
    return None


# --------------------------------------------------------------------------- documento «empadronados»
def datos(db: Session, hoy: date | None = None, activos: set[int] | None = None) -> dict:
    """Personas alojadas ahora (o con llegada confirmada en curso) cuyo domicilio habitual es la calle del edificio.
    `activos` limita a los activos visibles para el usuario (None = todos)."""
    hoy = hoy or date.today()
    filas, resumen = [], []
    for a in activos_turisticos(db):
        if activos is not None and a.id not in activos:
            continue
        reservas = db.scalars(select(Reservation).join(Unit, Unit.id == Reservation.unit_id).where(
            Unit.asset_id == a.id, Unit.uso != "garaje",
            or_(Reservation.estado == "checkin",
                (Reservation.estado == "confirmada") & (Reservation.fecha_entrada <= hoy)
                & (Reservation.fecha_salida > hoy))).order_by(Unit.codigo, Reservation.id))
        personas, apartamentos = set(), set()
        for r in reservas:
            datos_c = r.datos_contrato or {}
            gente: list[Contact] = [r.guest] + [o.contact for o in r.ocupantes if o.contact_id != r.guest_id]
            for p in gente:
                dom, mun = p.direccion, p.municipio
                if p.id == r.guest_id and not dom:
                    dom, mun = datos_c.get("cliente_domicilio"), datos_c.get("cliente_municipio")
                if not misma_calle(dom, mun, a) or p.id in personas:
                    continue
                personas.add(p.id)
                apartamentos.add(r.unit_id)
                filas.append({"activo": a.nombre, "apartamento": r.unit.codigo,
                              "reserva": r.localizador or f"R-{r.id}",
                              "persona": f"{p.nombre} {p.apellidos or ''}".strip(),
                              "titular": p.id == r.guest_id, "domicilio": dom or "",
                              "entrada": r.fecha_entrada, "salida": r.fecha_salida,
                              "cliente_antiguo": bool(datos_c.get("cliente_antiguo")) and p.id == r.guest_id,
                              "primera_estancia": datos_c.get("cliente_antiguo_desde") if p.id == r.guest_id else None})
        resumen.append({"activo": a.nombre, "calle": a.direccion or "", "personas": len(personas),
                        "apartamentos": len(apartamentos)})
    return {"fecha": hoy, "resumen": resumen, "filas": filas,
            "total_personas": sum(x["personas"] for x in resumen),
            "total_apartamentos": sum(x["apartamentos"] for x in resumen)}


def excel(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    neg = Font(bold=True)
    ws.append([f"Empadronados · domicilio habitual en la calle de los apartamentos · {d['fecha']:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Activo", "Calle", "Personas", "Apartamentos"])
    for c in ws[3]:
        c.font = neg
    for x in d["resumen"]:
        ws.append([x["activo"], x["calle"], x["personas"], x["apartamentos"]])
    ws.append(["TOTAL", None, d["total_personas"], d["total_apartamentos"]])
    for c in ws[ws.max_row]:
        c.font = neg
    for col, ancho in zip("ABCD", (34, 40, 12, 14)):
        ws.column_dimensions[col].width = ancho

    det = wb.create_sheet("Detalle")
    det.append(["Activo", "Apartamento", "Reserva", "Persona", "Titular", "Domicilio habitual", "Entrada", "Salida",
                "Cliente antiguo", "Primera estancia"])
    for c in det[1]:
        c.font = neg
    for f in d["filas"]:
        det.append([f["activo"], f["apartamento"], f["reserva"], f["persona"], "Sí" if f["titular"] else "No",
                    f["domicilio"], f["entrada"], f["salida"], "Sí" if f["cliente_antiguo"] else "No",
                    f["primera_estancia"] or ""])
    for col, ancho in zip("ABCDEFGHIJ", (30, 12, 16, 32, 8, 44, 12, 12, 14, 16)):
        det.column_dimensions[col].width = ancho
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
