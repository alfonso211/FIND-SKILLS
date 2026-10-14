"""CAPEX fase 2 (obra): semáforo de la documentación PRL de cada contrata, condiciones del contrato (retención de
garantía, plazo) y borrador del contrato de obra en Word.

La obra no empieza (acta de replanteo) hasta que cada contrata adjudicada tiene el contrato firmado y toda su
documentación PRL válida y en vigor (RD 171/2004, RD 1627/1997 y Ley 32/2006)."""
import io
from datetime import date, timedelta

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import capex
from .models import TIPOS_PRL_CAPEX, Asset, CapexAward, CapexPrl, CapexProject, Company, Supplier

AVISO_DIAS = 30  # «caduca pronto»: amarillo con este margen
COLOR = {"verde": "Válido", "amarillo": "Revisar", "rojo": "Falta o no válido", "gris": "No aplica"}


def asegurar_prl(db: Session, a: CapexAward) -> list[CapexPrl]:
    """Filas de documentación PRL de la contrata (se crean la primera vez, una por tipo)."""
    filas = {x.tipo: x for x in db.scalars(select(CapexPrl).where(CapexPrl.award_id == a.id))}
    for tipo in TIPOS_PRL_CAPEX:
        if tipo not in filas:
            filas[tipo] = CapexPrl(award_id=a.id, tipo=tipo, estado="falta")
            db.add(filas[tipo])
    db.flush()
    return [filas[t] for t in TIPOS_PRL_CAPEX]


def semaforo(x: CapexPrl, hoy: date | None = None) -> tuple[str, str]:
    """(color, texto) del documento: verde válido · amarillo por revisar o caduca en 30 días · rojo falta, rechazado
    o vencido · gris no aplica."""
    hoy = hoy or date.today()
    if x.estado == "no_aplica":
        return "gris", "No aplica"
    if x.estado == "por_revisar":
        return "amarillo", "Subido por el contratista: pendiente de revisar"
    if x.estado == "rechazado":
        return "rojo", "Rechazado" + (f": {x.nota}" if x.nota else "")
    if x.estado != "valido":
        return "rojo", "Falta"
    if x.vence and x.vence < hoy:
        return "rojo", f"Vencido el {x.vence:%d/%m/%Y}"
    if x.vence and x.vence <= hoy + timedelta(days=AVISO_DIAS):
        return "amarillo", f"Válido; caduca el {x.vence:%d/%m/%Y}"
    return "verde", "Válido" + (f" hasta el {x.vence:%d/%m/%Y}" if x.vence else "")


def en_regla(x: CapexPrl, hoy: date | None = None) -> bool:
    """Cuenta para el acta de replanteo: válido y en vigor, o no aplica (por revisar todavía no)."""
    color, _ = semaforo(x, hoy)
    return color in ("verde", "gris") or (color == "amarillo" and x.estado == "valido")


def pendientes_replanteo(db: Session, p: CapexProject, awards: list[CapexAward]) -> list[str]:
    """Lo que falta para poder firmar el acta de replanteo (vacío = se puede)."""
    faltan = []
    if p.estado != "adjudicado":
        faltan.append("El proyecto debe estar adjudicado (y aprobado, si hacía falta)")
    if not awards:
        faltan.append("No hay contratas adjudicadas")
    for a in awards:
        s = db.get(Supplier, a.supplier_id)
        if not a.contrato_firmado:
            faltan.append(f"{s.nombre}: falta subir el contrato firmado")
        malos = [TIPOS_PRL_CAPEX[x.tipo][0] for x in asegurar_prl(db, a) if not en_regla(x)]
        if malos:
            faltan.append(f"{s.nombre}: documentación PRL pendiente ({'; '.join(malos)})")
    return faltan


def capitulos_adjudicados(db: Session, p: CapexProject, a: CapexAward) -> list[tuple[str, str]]:
    caps = capex.capitulos(capex.partidas(db, p.id))
    return [c for c in caps if not a.capitulos or c[0] in a.capitulos]


def _sociedad(db: Session, asset: Asset) -> Company:
    return db.get(Company, asset.propietaria_id or asset.company_id)


def _domicilio(x) -> str:
    partes = [x.direccion, " ".join(v for v in (getattr(x, "cp", None), getattr(x, "municipio", None)) if v),
              getattr(x, "provincia", None)]
    return ", ".join(v for v in partes if v) or "[domicilio]"


def contrato_docx(db: Session, p: CapexProject, a: CapexAward) -> bytes:
    """Borrador del contrato de obra con los datos de la adjudicación. Lo revisa la Dirección Técnica (y, si hace
    falta, la asesoría jurídica) antes de firmarlo."""
    asset = db.get(Asset, p.asset_id)
    soc = _sociedad(db, asset)
    s = db.get(Supplier, a.supplier_id)
    rep = (soc.contratos or {}) if soc else {}
    importe = float(a.importe)
    iva = capex.r2(importe * float(p.iva_pct) / 100)
    caps = capitulos_adjudicados(db, p, a)
    ret = float(a.retencion_pct)

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)

    def par(texto: str, negrita: bool = False, centro: bool = False, tam: int | None = None):
        q = d.add_paragraph()
        r = q.add_run(texto)
        r.bold = negrita
        if tam:
            r.font.size = Pt(tam)
        if centro:
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return q

    aviso = par("BORRADOR — revisar antes de firmar", negrita=True, centro=True)
    aviso.runs[0].font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)
    par("CONTRATO DE EJECUCIÓN DE OBRA", negrita=True, centro=True, tam=14)
    par(f"{p.codigo} · {p.titulo}", centro=True)
    par(f"En {asset.municipio or '[lugar]'}, a [fecha de firma].")

    par("REUNIDOS", negrita=True)
    rep_txt = (f", representada por D./D.ª {rep['representante']}" + (f", con DNI {rep['representante_dni']}"
               if rep.get("representante_dni") else "") + (f", en calidad de {rep['representante_cargo']}"
               if rep.get("representante_cargo") else "")) if rep.get("representante") else ", representada por [representante]"
    par(f"De una parte, {soc.nombre if soc else '[sociedad]'}, con CIF {(soc.cif if soc else None) or '[CIF]'} "
        f"y domicilio en {_domicilio(soc) if soc else '[domicilio]'}{rep_txt} (en adelante, la PROPIEDAD).")
    par(f"De otra parte, {s.nombre}, con NIF {s.nif or '[NIF]'} y domicilio en {_domicilio(s)}, representada por "
        f"{s.persona_contacto or '[representante]'} (en adelante, el CONTRATISTA).")
    par("Ambas partes se reconocen capacidad suficiente para obligarse y")

    par("ACUERDAN", negrita=True)
    clausulas = [
        ("Objeto", f"El CONTRATISTA ejecutará los trabajos «{p.titulo}» en {asset.nombre}"
         f"{f' ({p.ubicacion})' if p.ubicacion else ''}, {asset.direccion or '[dirección]'}, conforme a las "
         "mediciones, planos y pliegos de la petición de ofertas y a su oferta, que forman parte de este contrato. "
         + (f"Capítulos adjudicados: {'; '.join(f'{c} {n}' for c, n in caps)}." if caps else "")),
        ("Precio", f"El precio cerrado es de {capex.eur(importe)} (ejecución material, gastos generales y beneficio "
         f"industrial), más el IVA del {float(p.iva_pct):g} % ({capex.eur(iva)}): total {capex.eur(importe + iva)}. "
         "Las unidades de obra no previstas o los excesos de medición solo se abonan con aprobación previa y por "
         "escrito de la Dirección Técnica de la PROPIEDAD."),
        ("Plazo", (f"{a.plazo_dias} días naturales" if a.plazo_dias else "[plazo] días naturales")
         + " a contar desde la firma del acta de replanteo. La obra no podrá comenzar sin dicha acta."),
        ("Forma de pago", "Mediante certificaciones mensuales a origen, conformadas por la Dirección Técnica, con la "
         "factura correspondiente. [Plazo de pago acordado]."),
        ("Retención de garantía", f"De cada certificación se retendrá el {ret:g} % de su importe (sin IVA) como "
         f"garantía de la correcta ejecución. La cantidad retenida se devolverá a los {a.garantia_meses} meses de la "
         "firma del acta de recepción, una vez subsanados los defectos que se hubieran señalado."),
        ("Prevención de riesgos laborales", "El CONTRATISTA cumplirá la Ley 31/1995 de Prevención de Riesgos "
         "Laborales, el RD 1627/1997 (obras de construcción), el RD 171/2004 (coordinación de actividades "
         "empresariales) y la Ley 32/2006 (subcontratación en el sector de la construcción). Antes del acta de "
         "replanteo aportará: seguro de responsabilidad civil en vigor y su pago, certificados de estar al corriente "
         "con Hacienda y con la Seguridad Social, alta de sus trabajadores, su formación en PRL, inscripción en el "
         "REA, evaluación de riesgos o Plan de Seguridad y Salud y designación del recurso preventivo, y la "
         "mantendrá vigente durante toda la obra."),
        ("Subcontratación", "Requiere autorización previa y por escrito de la PROPIEDAD. Las subcontratas aportarán "
         "la misma documentación PRL y se anotarán en el Libro de Subcontratación."),
        ("Seguros y responsabilidad", "El CONTRATISTA responde de los daños que cause a la PROPIEDAD, a sus "
         "ocupantes o a terceros durante la ejecución, y mantendrá vigente su seguro de responsabilidad civil."),
        ("Recepción", "Terminada la obra se firmará el acta de recepción. Si hubiera defectos, se recogerán en ella "
         "con un plazo para subsanarlos."),
        ("Resolución", "Son causas de resolución el incumplimiento grave del contrato, de los plazos o de las "
         "obligaciones de prevención de riesgos laborales, además de las previstas en la ley."),
        ("Jurisdicción", f"Las partes se someten a los juzgados y tribunales de {asset.municipio or '[lugar]'}."),
    ]
    for n, (titulo, texto) in enumerate(clausulas, 1):
        q = d.add_paragraph()
        q.add_run(f"{n}. {titulo}. ").bold = True
        q.add_run(texto)
        q.paragraph_format.space_after = Pt(6)

    par("Y en prueba de conformidad, firman por duplicado en el lugar y fecha indicados.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "LA PROPIEDAD", "EL CONTRATISTA"
    t.cell(1, 0).text = f"\n\n\n{soc.nombre if soc else ''}"
    t.cell(1, 1).text = f"\n\n\n{s.nombre}"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
