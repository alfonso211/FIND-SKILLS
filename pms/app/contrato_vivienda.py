"""Contrato de arrendamiento de vivienda habitual (LAU) sobre la plantilla Word del grupo (Babilonia 35).

- La hoja de control interno (2 primeras páginas) no se imprime: se convierte en el checklist del expediente.
- Cada hueco «[…]» de la plantilla va en su propio fragmento resaltado: los párrafos se localizan por su inicio y
  los huecos se rellenan en orden. Lo rellenado pierde el resaltado; lo que falta se queda resaltado.
- El Anexo I solo lleva lo que se entrega con la vivienda (inventario marcado «se entrega»).
"""
import copy
import io
import re
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import docx
from docx.oxml.ns import qn

PLANTILLA = Path(__file__).parent / "plantillas" / "contrato_vivienda_LAU.docx"
HUECO = re.compile(r"\[[^\]]*\]")
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]
MODALIDADES_GARANTIA = {"deposito": "Depósito en metálico", "aval": "Aval bancario", "fiador": "Fiador solidario",
                        "ninguna": "Sin garantía adicional"}
FORMAS_PAGO = {"transferencia": "Transferencia bancaria", "cheque": "Cheque nominativo", "aval": "Aval bancario",
               "efectivo": "Efectivo"}
CONCEPTOS = {"renta_inicial": "Renta y gastos del primer periodo", "fianza": "Fianza legal (1 mensualidad)",
             "garantia": "Garantía adicional", "otro": "Otro concepto"}
LIMITE_EFECTIVO = Decimal("1000")  # Ley 7/2012: pagos con empresario o profesional

# Inventario tipo (Anexo I de la plantilla): se copia a cada vivienda y se ajusta
INVENTARIO_TIPO = [
    ("Cocina", "Frigorífico"), ("Cocina", "Placa de cocción"), ("Cocina", "Horno / microondas"),
    ("Cocina", "Campana extractora"), ("Cocina", "Lavadora / lavavajillas"), ("Cocina", "Menaje (detallar)"),
    ("Salón", "Sofá"), ("Salón", "Mesa y sillas"), ("Salón", "Mueble TV / estanterías"),
    ("Dormitorio 1", "Cama, somier y colchón"), ("Dormitorio 1", "Mesillas / armario"),
    ("Dormitorio 2", "Cama, somier y colchón"), ("Baño", "Espejo / mueble lavabo / mampara"),
    ("General", "Lámparas y luminarias"), ("General", "Cortinas / estores"),
    ("Instalaciones", "Caldera / termo (marca, modelo, nº serie)"), ("Instalaciones", "Equipos de climatización"),
]
LLAVES_TIPO = ["Llave puerta de la vivienda", "Llave portal", "Llave buzón", "Llave trastero / mando garaje", "Otros"]
CONTADORES_TIPO = ["Electricidad", "Gas", "Agua fría", "Agua caliente (si existe)", "Calefacción (repartidores)"]


def inventario_tipo() -> list[dict]:
    return [{"estancia": e, "elemento": x, "uds": 1, "estado": "B", "obs": "", "entrega": True} for e, x in INVENTARIO_TIPO]


# --------------------------------------------------------------------------- checklist (hoja de control)
# (clave, fase, texto, condición, plazo: (desde, días) o None, documento a adjuntar)
CHECKLIST = [
    ("identidad", "antes", "Identidad: copia de DNI/NIE/pasaporte de todos los arrendatarios y del fiador", None, None, True),
    ("solvencia", "antes", "Solvencia: últimas 3 nóminas o IRPF, contrato laboral / vida laboral; consulta a "
                           "fichero de solvencia con consentimiento (cláusula 23)", None, None, True),
    ("cee", "antes", "Certificado de Eficiencia Energética vigente y registrado; entregar copia y etiqueta "
                     "(RD 390/2021)", None, None, True),
    ("ibi_tasa", "antes", "Importe anual de IBI y tasa de residuos (último recibo) para la cláusula 6", None, None, True),
    ("inventario", "antes", "Inventario (Anexo I) y reportaje fotográfico con fecha, firmados por ambas partes",
     None, None, True),
    ("acta_entrega", "antes", "Acta de entrega con lecturas de contadores y estado de instalaciones (Anexo II)",
     None, None, True),
    ("cobros", "antes", "Cobro de la primera mensualidad, la fianza y la garantía adicional ANTES de entregar llaves",
     None, None, False),
    ("sepa", "antes", "Mandato SEPA firmado (la renta se domicilia)", "sepa", None, True),
    ("aval", "antes", "Aval bancario a primer requerimiento (Anexo III)", "aval", None, True),
    ("fiador", "antes", "Fiador: identidad y solvencia del fiador solidario (cláusula 22)", "fiador", None, True),
    ("contrato_firmado", "despues", "Contrato y anexos firmados por todas las partes (copia escaneada)", None,
     ("firma", 0), True),
    ("deposito_fianza", "despues", "Depositar la fianza en la Agencia de Vivienda Social de la Comunidad de Madrid "
                                   "(telemático; sanción hasta el 50 %)", None, ("firma", 30), True),
    ("suministros", "despues", "Cambio de titularidad de suministros al arrendatario y baja/lectura de los de la "
                               "sociedad (cláusula 6.4)", None, ("entrega", 15), False),
    ("seguro_hogar", "despues", "Póliza y recibo del seguro de hogar del inquilino (cláusula 11.4)", None,
     ("firma", 15), True),
    ("alta_pms", "despues", "Alta del contrato en el PMS: renta, IBI/tasa, actualización anual, vencimientos, "
                            "fianza y garantía", None, None, False),
    ("avisos", "despues", "Avisos programados en la agenda: actualización anual, revisión de caldera/gas, "
                          "vencimientos", None, None, False),
]


def _d(x) -> date | None:
    if not x:
        return None
    return x if isinstance(x, date) else date.fromisoformat(str(x)[:10])


def aplica(cond: str | None, datos: dict) -> bool:
    if cond == "sepa":
        return datos.get("forma_pago") == "sepa"
    if cond == "aval":
        return datos.get("garantia_modalidad") == "aval"
    if cond == "fiador":
        return datos.get("garantia_modalidad") == "fiador"
    return True


def checklist(lease, unit, docs_por_clave: dict[str, int], hoy: date | None = None) -> list[dict]:
    """Estado de cada punto: hecho / no_aplica / pendiente (con fecha límite y si está vencido). Algunos puntos se
    comprueban solos con los datos del PMS (CEE vigente, IBI, cobros, alta y avisos)."""
    hoy = hoy or date.today()
    exp = lease.expediente or {}
    datos, marcas = exp.get("datos") or {}, exp.get("checklist") or {}
    ficha = unit.ficha or {}
    firma, entrega = _d(datos.get("fecha_firma")), _d(datos.get("fecha_entrega")) or _d(datos.get("fecha_firma"))
    out = []
    for clave, fase, texto, cond, plazo, doc in CHECKLIST:
        m = marcas.get(clave) or {}
        auto, motivo = None, None
        if clave == "cee":
            vig = _d(ficha.get("cee_vigencia"))
            if ficha.get("cee_letra") and vig and vig >= (firma or hoy):
                auto = True
            else:
                motivo = "Falta el CEE en la ficha de la vivienda o está caducado"
        elif clave == "ibi_tasa":
            auto = ficha.get("ibi_anual") not in (None, "") and ficha.get("tasa_residuos_anual") not in (None, "")
            motivo = None if auto else "Faltan el IBI o la tasa de residuos en la ficha de la vivienda"
        elif clave == "cobros":
            pend = cobros_pendientes(lease)
            auto = not pend
            motivo = None if auto else "Falta cobrar: " + ", ".join(pend)
        elif clave == "alta_pms":
            auto = lease.estado == "vigente" and bool(lease.renta_mensual) and lease.fecha_fin is not None
        elif clave == "avisos":
            auto = bool(exp.get("agenda"))
        estado = "no_aplica" if not aplica(cond, datos) else m.get("estado") or (
            "hecho" if auto or (doc and docs_por_clave.get(clave)) else "pendiente")
        limite = None
        if plazo and estado == "pendiente":
            base = firma if plazo[0] == "firma" else entrega
            limite = base + timedelta(days=plazo[1]) if base else None
        out.append({"clave": clave, "fase": fase, "texto": texto, "estado": estado, "auto": bool(auto),
                    "motivo": motivo if estado == "pendiente" else None, "documento": doc,
                    "documentos": docs_por_clave.get(clave, 0), "limite": limite.isoformat() if limite else None,
                    "vencido": bool(limite and limite < hoy), "nota": m.get("nota"),
                    "marcado_por": m.get("usuario"), "marcado_en": m.get("fecha"),
                    "condicional": cond is not None})
    return out


# --------------------------------------------------------------------------- importes
def dinero(x) -> Decimal:
    return Decimal(str(x or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def euros(x) -> str:
    return f"{dinero(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


_UNI = ["cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve", "diez", "once", "doce",
        "trece", "catorce", "quince", "dieciséis", "diecisiete", "dieciocho", "diecinueve", "veinte", "veintiuno",
        "veintidós", "veintitrés", "veinticuatro", "veinticinco", "veintiséis", "veintisiete", "veintiocho",
        "veintinueve"]
_DEC = ["", "", "", "treinta", "cuarenta", "cincuenta", "sesenta", "setenta", "ochenta", "noventa"]
_CEN = ["", "ciento", "doscientos", "trescientos", "cuatrocientos", "quinientos", "seiscientos", "setecientos",
        "ochocientos", "novecientos"]


def _apocope(t: str) -> str:
    """«uno» delante de un sustantivo o de «mil»: veintiuno -> veintiún, treinta y uno -> treinta y un."""
    if t.endswith("veintiuno"):
        return t[:-9] + "veintiún"
    return t[:-3] + "un" if t.endswith("uno") else t


def _letras(n: int, apocope: bool = False) -> str:
    if n < 30:
        t = _UNI[n]
    elif n < 100:
        d, u = divmod(n, 10)
        t = _DEC[d] + (f" y {_UNI[u]}" if u else "")
    elif n < 1000:
        c, r = divmod(n, 100)
        t = "cien" if n == 100 else _CEN[c] + (f" {_letras(r)}" if r else "")
    elif n < 1_000_000:
        m, r = divmod(n, 1000)
        t = ("mil" if m == 1 else f"{_letras(m, True)} mil") + (f" {_letras(r)}" if r else "")
    else:
        mm, r = divmod(n, 1_000_000)
        t = ("un millón" if mm == 1 else f"{_letras(mm, True)} millones") + (f" {_letras(r)}" if r else "")
    return _apocope(t) if apocope else t


def en_letra(x) -> str:
    """1250.5 -> «mil doscientos cincuenta euros con cincuenta céntimos»."""
    v = dinero(x)
    e, c = int(v), int((v - int(v)) * 100)
    t = f"{_letras(e, True)} {'euro' if e == 1 else 'euros'}"
    if c:
        t += f" con {_letras(c, True)} {'céntimo' if c == 1 else 'céntimos'}"
    return t


# --------------------------------------------------------------------------- datos del contrato
def cobros_pendientes(lease) -> list[str]:
    exp = lease.expediente or {}
    d = exp.get("datos") or {}
    cobrado = {}
    for e in exp.get("entregas") or []:
        if e.get("anulada"):
            continue
        cobrado[e["concepto"]] = cobrado.get(e["concepto"], Decimal(0)) + dinero(e["importe"])
    falta = []
    if dinero(d.get("importe_inicial")) > cobrado.get("renta_inicial", 0):
        falta.append("primer periodo")
    if dinero(lease.fianza) > cobrado.get("fianza", 0):
        falta.append("fianza")
    if d.get("garantia_modalidad", "deposito") == "deposito" and dinero(lease.garantia_adicional) > cobrado.get("garantia", 0):
        falta.append("garantía adicional")
    return falta


def referencia(lease) -> str:
    """Ref. del contrato tal como se imprime: BAB35-2B-2026."""
    u, d = lease.unit, (lease.expediente or {}).get("datos") or {}
    anio = (_d(d.get("fecha_firma")) or lease.fecha_inicio).year
    return f"{u.asset.codigo}-{u.codigo.removeprefix(f'{u.asset.codigo}-')}-{anio}"


def vivienda_texto(lease) -> str:
    u, f = lease.unit, lease.unit.ficha or {}
    puerta = f.get("puerta") or (u.codigo.split("-")[-1] if "-" in u.codigo else u.codigo)
    return f"{u.asset.direccion or u.asset.nombre}" + (f", planta {u.planta}" if u.planta else "") + f", puerta {puerta}"


def _persona(c) -> dict:
    return {"nombre": f"{c.nombre} {c.apellidos or ''}".strip(), "doc_tipo": c.documento_tipo or "DNI",
            "doc_num": c.documento_num, "nacionalidad": c.nacionalidad,
            "domicilio": ", ".join(x for x in (c.direccion, c.cp, c.municipio, c.pais) if x) or None,
            "telefono": c.telefono, "email": c.email}


def valores(db, lease) -> dict:
    """Todos los datos para rellenar el contrato, tomados de la sociedad, el activo, la vivienda, el inquilino y el
    expediente. Las claves coinciden con los campos de CAMPOS."""
    from .models import Contact
    u, a = lease.unit, lease.unit.asset
    soc = a.propietaria or a.company
    sc, f = soc.contratos or {}, u.ficha or {}
    exp = lease.expediente or {}
    d = exp.get("datos") or {}
    t1 = _persona(lease.tenant)
    t2 = _persona(db.get(Contact, d["arrendatario2_id"])) if d.get("arrendatario2_id") else None
    fi = _persona(db.get(Contact, d["fiador_id"])) if d.get("fiador_id") and d.get("garantia_modalidad") == "fiador" else None
    firma, inicio = _d(d.get("fecha_firma")), lease.fecha_inicio
    renta = dinero(lease.renta_mensual)
    ibi, tasa = f.get("ibi_anual"), f.get("tasa_residuos_anual")
    cuota = (dinero(ibi) + dinero(tasa)) / 12 if ibi not in (None, "") and tasa not in (None, "") else None
    puerta = f.get("puerta") or (u.codigo.split("-")[-1] if "-" in u.codigo else u.codigo)
    entrega = d.get("fecha_entrega") or d.get("fecha_firma")
    v = {
        "ref_unidad": u.codigo.removeprefix(f"{a.codigo}-"), "ref_anio": str((firma or inicio).year),
        "firma_dia": str(firma.day) if firma else None, "firma_mes": MESES[firma.month - 1] if firma else None,
        "firma_anio": str(firma.year) if firma else None,
        "rm_tomo": sc.get("rm_tomo"), "rm_folio": sc.get("rm_folio"), "rm_hoja": sc.get("rm_hoja"),
        "representante": sc.get("representante"), "representante_dni": sc.get("representante_dni"),
        "representante_cargo": sc.get("representante_cargo"), "representante_poder": sc.get("representante_poder"),
        "t1_nombre": t1["nombre"], "t1_doc_tipo": t1["doc_tipo"], "t1_doc_num": t1["doc_num"],
        "t1_nacionalidad": t1["nacionalidad"], "t1_domicilio": d.get("domicilio_anterior") or t1["domicilio"],
        "t1_telefono": t1["telefono"], "t1_email": t1["email"],
        "planta": u.planta, "puerta": puerta, "cp": f.get("cp") or a.cp,
        "sup_construida": _num(u.superficie_m2), "sup_util": _num(f.get("superficie_util")),
        "distribucion": f.get("distribucion"), "ref_catastral": u.ref_catastral,
        "registro_num": re.sub(r"^n[ºo°]\.?\s*|\s+de Madrid$", "", str(f.get("registro_propiedad") or "")) or None, "finca": f.get("finca"),
        "anejos": d.get("anejos") or u.anejos or "ninguno",
        "cee_letra": f.get("cee_letra"), "cee_registro": f.get("cee_registro"),
        "cee_vigencia": _fecha(f.get("cee_vigencia")),
        "convivientes": d.get("convivientes") or "ninguna otra persona",
        "max_ocupantes": str(d.get("max_ocupantes") or u.capacidad or "") or None,
        "fecha_vencimiento": _fecha(lease.fecha_fin),
        "renta_cifra": euros(renta), "renta_letra": en_letra(renta), "renta_anual": euros(renta * 12),
        "iban_arrendatario": d.get("iban_arrendatario") if d.get("forma_pago") == "sepa" else "—",
        "iban_arrendadora": sc.get("iban"),
        "importe_inicial": euros(d["importe_inicial"]) if d.get("importe_inicial") not in (None, "") else None,
        "periodo_desde": _fecha(d.get("periodo_desde")), "periodo_hasta": _fecha(d.get("periodo_hasta")),
        "ibi_anual": euros(ibi) if ibi not in (None, "") else None,
        "tasa_anual": euros(tasa) if tasa not in (None, "") else None,
        "cuota_tributos": euros(cuota) if cuota is not None else None,
        "sin_contador": f.get("sin_contador") or "no aplica",
        "sin_contador_importe": euros(f.get("sin_contador_importe") or 0),
        "fianza": euros(lease.fianza) if lease.fianza is not None else None,
        "garantia": euros(lease.garantia_adicional) if lease.garantia_adicional is not None else None,
        "animales_detalle": d.get("animales_detalle") or "—",
        "seguro_capital": euros(d.get("seguro_capital") or 150000).replace(",00", ""),
        "email_arrendadora": sc.get("email_notificaciones"), "t1_email_notif": t1["email"],
        "email_rgpd": sc.get("email_rgpd") or sc.get("email_notificaciones"),
        "inv_fecha": _fecha(entrega), "fotos": str(d["fotos"]) if d.get("fotos") else None,
        "entrega_fecha_hora": (_fecha(entrega) + (f", {d['hora_entrega']} h" if d.get("hora_entrega") else "")) if entrega else None,
        "entrega_presentes": d.get("presentes"), "bombin": f.get("bombin"),
        "otros_docs": d.get("otros_docs") or "—",
        "telefono_averias": sc.get("telefono_averias"), "email_averias": sc.get("email_notificaciones"),
        "paginas": d.get("paginas"),
        "sociedad": soc.nombre, "cif": soc.cif,
    }
    if t2:
        v.update({"t2_nombre": t2["nombre"], "t2_doc_tipo": t2["doc_tipo"], "t2_doc_num": t2["doc_num"],
                  "t2_telefono": t2["telefono"], "t2_email": t2["email"]})
    if fi:
        v.update({"f_nombre": fi["nombre"], "f_doc_tipo": fi["doc_tipo"], "f_doc_num": fi["doc_num"],
                  "f_domicilio": fi["domicilio"]})
    return v


def _num(x) -> str | None:
    if x in (None, ""):
        return None
    return f"{float(x):.2f}".rstrip("0").rstrip(".").replace(".", ",")


def _fecha(x) -> str | None:
    x = _d(x)
    return x.strftime("%d/%m/%Y") if x else None


# (inicio del párrafo, campos de sus huecos en orden). None = hueco que se deja tal cual (se rellena a mano).
CAMPOS = [
    ("Ref.:", ["ref_unidad", "ref_anio"]),
    ("En Madrid, a", ["firma_dia", "firma_mes", "firma_anio"]),
    ("DE UNA PARTE", ["rm_tomo", "rm_folio", "rm_hoja", "representante", "representante_dni", "representante_cargo",
                      "representante_poder"]),
    ("DE OTRA PARTE", ["t1_nombre", "t1_doc_tipo", "t1_doc_num", "t1_nacionalidad", "t1_domicilio", "t1_telefono",
                       "t1_email"]),
    ("Y D./Dña.", ["t2_nombre", "t2_doc_tipo", "t2_doc_num", "t2_telefono", "t2_email"]),
    ("Y, en su caso, como FIADOR", ["f_nombre", "f_doc_tipo", "f_doc_num", "f_domicilio"]),
    ("I. Titularidad.", ["planta", "puerta", "cp", "sup_construida", "sup_util", "distribucion", "ref_catastral",
                         "registro_num", "finca", "anejos"]),
    ("II. Estado y eficiencia", ["cee_letra", "cee_registro", "cee_vigencia"]),
    ("1.2", ["convivientes", "max_ocupantes"]),
    ("3.1", ["fecha_vencimiento"]),
    ("4.1", ["renta_cifra", "renta_letra", "renta_anual"]),
    ("4.2", ["iban_arrendatario", "iban_arrendadora"]),
    ("4.3", ["importe_inicial", "periodo_desde", "periodo_hasta"]),
    ("6.2", ["ibi_anual", "tasa_anual", "cuota_tributos"]),
    ("6.4", ["sin_contador", "sin_contador_importe"]),
    ("7.1", ["fianza"]),
    ("7.2", ["garantia"]),
    ("11.2", ["animales_detalle"]),
    ("11.4", ["seguro_capital"]),
    ("21.1", ["email_arrendadora", "t1_email_notif"]),
    ("23.1", ["email_rgpd"]),
    ("Y en prueba de conformidad", ["paginas"]),
    ("Vivienda: C/", ["planta", "puerta", "inv_fecha"]),
    ("Se adjunta reportaje", ["fotos"]),
    ("Fecha y hora de entrega", ["entrega_fecha_hora", "entrega_presentes"]),
    ("☐ Otros:", ["otros_docs"]),
    ("Urgencias:", ["telefono_averias", "email_averias"]),
]
# Datos imprescindibles para imprimir (los demás pueden ir en blanco para completarlos a mano)
OBLIGATORIOS = {
    "firma_dia": "Fecha de firma", "rm_tomo": "Registro Mercantil de la sociedad (tomo, folio, hoja)",
    "representante": "Representante de la sociedad", "representante_dni": "DNI del representante",
    "representante_cargo": "Cargo del representante", "representante_poder": "Escritura de poder del representante",
    "t1_doc_num": "Documento del arrendatario", "t1_nacionalidad": "Nacionalidad del arrendatario",
    "t1_domicilio": "Domicilio anterior del arrendatario", "t1_telefono": "Teléfono del arrendatario",
    "t1_email": "Correo del arrendatario", "cp": "Código postal de la vivienda",
    "sup_construida": "Superficie construida", "sup_util": "Superficie útil", "distribucion": "Distribución",
    "ref_catastral": "Referencia catastral", "registro_num": "Registro de la Propiedad", "finca": "Finca registral",
    "cee_letra": "Certificado energético: calificación", "cee_registro": "Certificado energético: nº de registro",
    "cee_vigencia": "Certificado energético: vigencia", "max_ocupantes": "Número máximo de ocupantes",
    "fecha_vencimiento": "Fecha de vencimiento (7 años)", "iban_arrendadora": "IBAN de la sociedad",
    "importe_inicial": "Importe del primer periodo", "periodo_desde": "Primer periodo: desde",
    "periodo_hasta": "Primer periodo: hasta", "ibi_anual": "IBI anual", "tasa_anual": "Tasa de residuos anual",
    "fianza": "Fianza", "email_arrendadora": "Correo de notificaciones de la sociedad",
    "telefono_averias": "Teléfono de averías",
}


def faltan(db, lease) -> list[str]:
    v = valores(db, lease)
    d = (lease.expediente or {}).get("datos") or {}
    out = [txt for k, txt in OBLIGATORIOS.items() if not v.get(k)]
    if d.get("forma_pago") == "sepa" and not d.get("iban_arrendatario"):
        out.append("IBAN del arrendatario (domiciliación)")
    if d.get("garantia_modalidad") in (None, "", "deposito", "aval", "fiador") and lease.garantia_adicional is None \
            and d.get("garantia_modalidad") != "ninguna":
        out.append("Importe de la garantía adicional")
    if d.get("garantia_modalidad") == "fiador" and not v.get("f_nombre"):
        out.append("Datos del fiador")
    if lease.fianza is not None and dinero(lease.fianza) > dinero(lease.renta_mensual):
        out.append("La fianza legal es de una mensualidad (art. 36.1 LAU)")
    if lease.garantia_adicional is not None and dinero(lease.garantia_adicional) > 2 * dinero(lease.renta_mensual):
        out.append("La garantía adicional no puede superar dos mensualidades (art. 36.5 LAU)")
    return out


# --------------------------------------------------------------------------- relleno del Word
def _pon(run, valor: str) -> None:
    run.text = valor
    run.font.highlight_color = None


def _borra(par) -> None:
    par._element.getparent().remove(par._element)


def _marca_casilla(par, textos_marcados: list[str]) -> None:
    """Cambia «☐» por «☒» delante de cada opción elegida (busca el texto que sigue a la casilla)."""
    runs = par.runs
    for i, r in enumerate(runs):
        if "☐" not in r.text:
            continue
        partes = r.text.split("☐")
        nuevo = partes[0]
        for j, resto in enumerate(partes[1:]):
            siguiente = resto if resto.strip() else (runs[i + 1].text if i + 1 < len(runs) else "")
            marcado = any(siguiente.strip().lower().startswith(t.lower()) for t in textos_marcados)
            nuevo += ("☒" if marcado else "☐") + resto
        r.text = nuevo


def _todos_parrafos(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def generar(db, lease, paginas: int | None = None) -> bytes:
    v = valores(db, lease)
    if paginas:
        v["paginas"] = str(paginas)
    exp = lease.expediente or {}
    d = exp.get("datos") or {}
    a = lease.unit.asset
    doc = docx.Document(str(PLANTILLA))
    body = doc.element.body

    # 1) Fuera la hoja de control interno: todo hasta el primer salto de página
    for el in list(body.iterchildren()):
        es_salto = any(b.get(qn("w:type")) == "page" for b in el.iter(qn("w:br")))
        if el.tag == qn("w:sectPr"):
            break
        body.remove(el)
        if es_salto:
            break

    # 2) Partes opcionales: segundo arrendatario y fiador
    for p in list(doc.paragraphs):
        if p.text.startswith("Y D./Dña.") and not v.get("t2_nombre"):
            _borra(p)
        elif p.text.startswith("Y, en su caso, como FIADOR") and not v.get("f_nombre"):
            _borra(p)

    # 3) Huecos por párrafo
    usados = set()
    for p in _todos_parrafos(doc):
        texto = p.text.strip()
        for inicio, campos in CAMPOS:
            if texto.startswith(inicio) and (inicio, id(p._element)) not in usados:
                usados.add((inicio, id(p._element)))
                huecos = [r for r in p.runs if HUECO.fullmatch(r.text.strip() or "x")]
                for r, campo in zip(huecos, campos):
                    if campo and v.get(campo):
                        _pon(r, str(v[campo]))
                # El importe en letra ya lleva «euros»: no repetirlo con el texto fijo que sigue
                for r, sig in zip(p.runs, p.runs[1:]):
                    if r.text.endswith(("euros", "euro", "céntimos", "céntimo")) and sig.text.startswith(" euros"):
                        sig.text = sig.text[len(" euros"):]
                break

    # 4) Sociedad arrendadora y dirección del activo (la plantilla trae Comercial del Campo y Babilonia 35)
    for p in _todos_parrafos(doc):
        for r in p.runs:
            if "COMERCIAL DEL CAMPO, S.A." in r.text and v["sociedad"]:
                r.text = r.text.replace("COMERCIAL DEL CAMPO, S.A.", v["sociedad"].replace(" S.A.", ", S.A."))
            if "A28362309" in r.text and v["cif"]:
                r.text = r.text.replace("A28362309", v["cif"])
            if "BAB35-" in r.text and a.codigo != "BAB35":
                r.text = r.text.replace("BAB35-", f"{a.codigo}-")
            if "C/ Babilonia, 35" in r.text and a.codigo != "BAB35" and a.direccion:
                r.text = r.text.replace("C/ Babilonia, 35", a.direccion)

    # 5) Casillas: modalidad de garantía, animales y documentación aportada (Anexo III)
    marcas = (exp.get("checklist") or {})
    modalidad = d.get("garantia_modalidad", "deposito")
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("7.2"):
            _marca_casilla(p, {"deposito": ["depósito"], "aval": ["aval"], "fiador": ["fianza personal"]}.get(modalidad, []))
        elif t.startswith("11.2"):
            _marca_casilla(p, ["se permiten"] if d.get("animales") else ["no se permiten"])
        elif t.startswith("☐"):
            claves = {"Copia de DNI": "identidad", "Contrato de trabajo": "solvencia", "Informe de vida": "solvencia",
                      "Certificado de titularidad": "sepa", "Aval bancario": "aval", "Póliza y recibo": "seguro_hogar",
                      "Copia del Certificado": "cee", "Último recibo de IBI": "ibi_tasa",
                      "Justificante de depósito": "deposito_fianza", "Recomendaciones de uso": None,
                      "Normas de régimen": None, "Otros:": None}
            for ini, clave in claves.items():
                if t[1:].strip().startswith(ini):
                    entregado = (clave and (marcas.get(clave) or {}).get("estado") == "hecho") or \
                        (ini == "Recomendaciones de uso") or (ini == "Otros:" and d.get("otros_docs"))
                    if entregado and p.runs and p.runs[0].text.startswith("☐"):
                        p.runs[0].text = p.runs[0].text.replace("☐", "☒", 1)

    # 6) Firmas (tabla de firma): representante, arrendatarios y fiador
    firmas = doc.tables[0] if doc.tables else None
    if firmas is not None:
        nombres = [[v.get("representante")], [v.get("t1_nombre"), v.get("t2_nombre")], [v.get("f_nombre")]]
        for celda, ns in zip(firmas.rows[0].cells, nombres):
            huecos = [r for p in celda.paragraphs for r in p.runs if r.text.strip() == "[nombre]"]
            for i, r in enumerate(huecos):
                valor = ns[i] if i < len(ns) else None
                if valor:
                    _pon(r, valor)
                elif i > 0 or ns is nombres[2]:
                    _pon(r, "—")

    # 7) Anexo I: solo lo que se entrega con la vivienda
    inv = [x for x in (exp.get("inventario") if exp.get("inventario") is not None else lease.unit.inventario
                       or inventario_tipo()) if x.get("entrega", True)]
    tabla = next((t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Estancia"), None)
    if tabla is not None:
        modelo = copy.deepcopy(tabla.rows[1]._tr)
        for row in list(tabla.rows)[1:]:
            tabla._tbl.remove(row._tr)
        for x in inv:
            tr = copy.deepcopy(modelo)
            tabla._tbl.append(tr)
            fila = tabla.rows[-1]
            for celda, val in zip(fila.cells, [x.get("estancia"), x.get("elemento"), x.get("uds"), x.get("estado"),
                                               x.get("obs")]):
                par = celda.paragraphs[0]
                for r in par.runs[1:]:
                    r._element.getparent().remove(r._element)
                if par.runs:
                    par.runs[0].text = str(val if val not in (None, "") else "")
                else:
                    par.add_run(str(val if val not in (None, "") else ""))

    # 8) Anexo II: llaves y contadores de la ficha de la vivienda; lecturas del acta si se tomaron
    ficha = lease.unit.ficha or {}
    acta = d.get("acta") or {}
    for t in doc.tables:
        cab = t.rows[0].cells[0].text.strip()
        if cab == "Elemento":
            llaves = {x.get("elemento"): x for x in ficha.get("llaves") or []}
            for row in list(t.rows)[1:]:
                nombre = row.cells[0].text.split(" (")[0].strip()
                x = llaves.get(nombre) or {}
                _celda(row.cells[1], x.get("uds"))
                _celda(row.cells[2], x.get("obs"))
                if v.get("bombin"):
                    for p in row.cells[0].paragraphs:
                        for r in p.runs:
                            if r.text.strip() == "[marca/modelo]":
                                _pon(r, v["bombin"])
        elif cab == "Suministro" and t.rows[0].cells[2].text.strip() == "Lectura":
            cont = {x.get("suministro"): x for x in ficha.get("contadores") or []}
            lect = acta.get("lecturas") or {}
            for row in list(t.rows)[1:]:
                nombre = row.cells[0].text.strip()
                _celda(row.cells[1], (cont.get(nombre) or {}).get("cups"))
                _celda(row.cells[2], lect.get(nombre))
                _celda(row.cells[3], (cont.get(nombre) or {}).get("titular"))

    # 9) Huecos sin dato del Anexo V (se rellenan a mano a la salida): línea en blanco
    en_anexo_v = False
    for p in doc.paragraphs:
        if p.text.startswith("ANEXO V"):
            en_anexo_v = True
        if en_anexo_v:
            for r in p.runs:
                if HUECO.fullmatch(r.text.strip() or "x"):
                    _pon(r, "______________")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def generar_final(db, lease) -> bytes:
    """Contrato listo para imprimir: con LibreOffice se cuenta el nº de páginas y se vuelve a generar con él."""
    from . import firma_contrato
    if (lease.expediente or {}).get("datos", {}).get("paginas") or not firma_contrato.disponible():
        return generar(db, lease)
    borrador = generar(db, lease, paginas=99)  # mismo ancho aproximado que el número definitivo
    try:
        n = paginas_pdf(firma_contrato.a_pdf(borrador))
    except Exception:  # noqa: BLE001 — sin conversor se imprime sin el nº (se escribe a mano)
        return generar(db, lease)
    return generar(db, lease, paginas=n)


def _celda(celda, valor) -> None:
    if valor in (None, ""):
        return
    par = celda.paragraphs[0]
    if par.runs:
        par.runs[0].text = str(valor)
    else:
        par.add_run(str(valor))


def paginas_pdf(pdf: bytes) -> int:
    import pypdfium2 as pdfium
    return len(pdfium.PdfDocument(pdf))


# --------------------------------------------------------------------------- agenda
def _fh(dia: date, hora: str | None = None):
    from datetime import datetime, time
    try:
        h = time.fromisoformat(hora) if hora else None
    except ValueError:
        h = None
    return datetime.combine(dia, h or time(0, 0)), h is None


def agenda(lease) -> list[dict]:
    """Citas y tareas que genera el expediente para quien prepara el contrato. clave -> evento (sin guardar)."""
    exp = lease.expediente or {}
    d = exp.get("datos") or {}
    ref = f"{lease.unit.codigo} · {lease.tenant.nombre} {lease.tenant.apellidos or ''}".strip()
    firma, entrega = _d(d.get("fecha_firma")), _d(d.get("fecha_entrega")) or _d(d.get("fecha_firma"))
    out = []

    def ev(clave, tipo, titulo, dia, hora=None, descripcion=None, repeticion=None, hasta=None, aviso=1440,
           prioridad="normal"):
        if not dia:
            return
        inicio, todo = _fh(dia, hora)
        out.append({"clave": clave, "tipo": tipo, "titulo": f"{titulo} — {ref}", "inicio": inicio,
                    "todo_el_dia": todo, "descripcion": descripcion, "repeticion": repeticion,
                    "repetir_hasta": hasta, "aviso_min": aviso, "prioridad": prioridad})

    ev("firma", "reunion", "Firma del contrato de alquiler", firma, d.get("hora_firma"),
       "Llevar 2 ejemplares impresos del contrato y anexos, recibos de entrega e inventario. Escanear el firmado "
       "y subirlo a la carpeta del expediente.", prioridad="alta")
    if entrega and (entrega != firma or d.get("hora_entrega") != d.get("hora_firma")):
        ev("entrega", "reunion", "Entrega de llaves y acta (Anexo II)", entrega, d.get("hora_entrega"),
           "Tomar lecturas de contadores, revisar instalaciones y reportaje fotográfico con fecha.")
    pend = cobros_pendientes(lease)
    if d.get("garantia_modalidad") == "aval":
        pend.append("aval bancario (original)")
    if pend:
        ev("cobros", "tarea", "Cobrar antes de entregar llaves", firma,
           descripcion="Pendiente: " + ", ".join(pend) + ". Registrar cada entrega en el expediente y firmar el recibo. "
                       "Efectivo no admitido para la renta ni para importes de 1.000 € o más.", prioridad="alta")
    if firma:
        ev("deposito_fianza", "tarea", "Depositar la fianza (Agencia de Vivienda Social)", firma + timedelta(days=25),
           descripcion="Plazo: 30 días desde la firma. Subir el justificante a la carpeta y entregar copia al inquilino.",
           prioridad="alta")
        ev("seguro_hogar", "tarea", "Pedir póliza y recibo del seguro de hogar del inquilino",
           firma + timedelta(days=15), "Cláusula 11.4. Subir la póliza a la carpeta.")
    if entrega:
        ev("suministros", "tarea", "Comprobar cambio de titularidad de suministros", entrega + timedelta(days=15),
           "Cláusula 6.4: 15 días desde la entrega de llaves. Dar de baja / facturar los que sigan a nombre de la sociedad.")
    ini, fin = lease.fecha_inicio, lease.fecha_fin
    if ini:
        aniv = date(ini.year + 1, ini.month, min(ini.day, 28))
        ev("actualizacion", "recordatorio", "Actualizar la renta (IPC con tope IRAV) y notificar por escrito",
           aniv - timedelta(days=30), descripcion="Cláusula 5: la renta actualizada es exigible desde el mes siguiente "
                                                  "a la notificación escrita.", repeticion="anual", hasta=fin)
        ev("revision_caldera", "recordatorio", "Pedir al inquilino el justificante de revisión de caldera/gas",
           date(ini.year + 1, ini.month, min(ini.day, 28)), "Cláusula 6.5: revisión periódica a cargo del arrendatario.",
           repeticion="anual", hasta=fin)
    if fin:
        ev("preaviso", "recordatorio", "Decidir renovación: preaviso de la arrendadora (4 meses)",
           fin - timedelta(days=150), "Cláusula 3.2: notificar con al menos 4 meses de antelación al vencimiento.",
           prioridad="alta")
    venc_aval = _d(d.get("aval_vencimiento"))
    if d.get("garantia_modalidad") == "aval" and venc_aval:
        ev("aval_vencimiento", "recordatorio", "Vence el aval bancario: pedir renovación", venc_aval - timedelta(days=60),
           prioridad="alta")
    return out
