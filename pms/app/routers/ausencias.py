"""Personal → Vacaciones y ausencias: solicitudes de vacaciones, días libres, bajas, permisos y faltas del personal
(usuarios del PMS y personal de subcontratas por administración) con su flujo de aprobación (ver app/ausencias.py).
Todos ven el calendario de su activo para saber quién falta."""
import io
from datetime import date, datetime
from html import escape

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .. import ausencias as au
from .. import avisos
from ..database import get_db
from ..models import (COLECTIVOS_AUSENCIA, ESTADOS_AUSENCIA, TIPOS_AUSENCIA, Absence, Asset, StaffMember, User)
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404

router = APIRouter(prefix="/api/ausencias", tags=["ausencias"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VIGENTES = ("pendiente", "visto_bueno", "aprobada")


class AbsenceIn(BaseModel):
    asset_id: int
    user_id: int | None = None  # vacío y sin staff_id: la propia persona que la solicita
    staff_id: int | None = None
    tipo: str
    desde: date
    hasta: date
    motivo: str | None = Field(default=None, max_length=1000)
    aprobar: bool = False  # quien puede aprobarla la deja aprobada al registrarla


class NotaIn(BaseModel):
    nota: str | None = Field(default=None, max_length=1000)


def _activo(db: Session, scope: Scope, asset_id: int) -> Asset:
    scope.require_asset("activos.ver", asset_id)
    return get_or_404(db, Asset, asset_id)


def _out(db: Session, scope: Scope, x: Absence, nombres: dict[int, str]) -> dict:
    d = x.to_dict()
    nat, lab = au.dias(x.desde, x.hasta)
    detalle = au.ve_detalle(db, scope, x)
    d.update(dias_naturales=nat, dias_laborables=lab, colectivo_nombre=COLECTIVOS_AUSENCIA.get(x.colectivo, x.colectivo),
             tipo_nombre=TIPOS_AUSENCIA.get(x.tipo, x.tipo) if detalle else "Ausencia",
             estado_nombre=ESTADOS_AUSENCIA.get(x.estado, x.estado), aprueba=au.quien_aprueba(x.colectivo),
             visto_bueno_nombre=nombres.get(x.visto_bueno_por), resuelto_nombre=nombres.get(x.resuelto_por),
             creado_nombre=nombres.get(x.creado_por), propia=x.user_id == scope.user.id,
             puede_aprobar=au.puede_aprobar(db, scope, x), puede_visto_bueno=au.puede_visto_bueno(db, scope, x),
             puede_anular=au.puede_anular(db, scope, x))
    if not detalle:
        d.update(tipo="ausencia", motivo=None)
    return d


def _nombres(db: Session, filas: list[Absence]) -> dict[int, str]:
    ids = {i for x in filas for i in (x.visto_bueno_por, x.resuelto_por, x.creado_por) if i}
    return dict(db.execute(select(User.id, User.nombre).where(User.id.in_(ids or {-1}))).all())


def _uno(db: Session, scope: Scope, x: Absence) -> dict:
    return _out(db, scope, x, _nombres(db, [x]))


def _avisar(db: Session, usuarios: list[User], asunto: str, texto: str, excluir: int | None = None) -> None:
    """Aviso por correo (si está configurado); un fallo no impide guardar la solicitud."""
    if not avisos.configurado():
        return
    for u in {u.id: u for u in usuarios}.values():
        if u.id == excluir or not u.email:
            continue
        try:
            avisos.enviar(u.email, asunto, texto, avisos._html(asunto, f"<p>{escape(texto)}</p>"))
        except Exception:  # noqa: BLE001
            pass


def _texto(x: Absence, a: Asset, tipo_visible: bool = True) -> str:
    tipo = TIPOS_AUSENCIA.get(x.tipo, x.tipo) if tipo_visible and x.tipo != "baja_medica" else "Ausencia"
    rango = f"el {x.desde:%d/%m/%Y}" if x.desde == x.hasta else f"del {x.desde:%d/%m/%Y} al {x.hasta:%d/%m/%Y}"
    return f"{x.persona} ({COLECTIVOS_AUSENCIA.get(x.colectivo, x.colectivo)}, {a.nombre}): {tipo.lower()} {rango}."


def _avisos_nueva(db: Session, x: Absence, a: Asset, autor: int) -> None:
    pend = " Pendiente de aprobar en Personal → Vacaciones y ausencias."
    if x.colectivo == "mantenimiento":  # se pone en conocimiento de todos; autoriza el director técnico
        _avisar(db, au.usuarios_con(db, a.id), "Ausencia de mantenimiento", _texto(x, a)
                + (" Aprobada." if x.estado == "aprobada" else " La autoriza el director técnico."), autor)
        if x.estado != "aprobada":
            _avisar(db, au.usuarios_con(db, a.id, "personal.autorizar_mto"), "Ausencia por autorizar",
                    _texto(x, a) + pend, autor)
    elif x.estado == "aprobada":
        return
    elif x.colectivo == "recepcion2":
        r1 = au.recepcion_1(db, a)
        _avisar(db, [r1] if r1 else [], "Ausencia de Recepción 2: falta su visto bueno",
                _texto(x, a) + " Dé su visto bueno en Personal → Vacaciones y ausencias; después la aprueba dirección.",
                autor)
    elif x.colectivo in au.PERSONAL_R1:
        r1 = au.recepcion_1(db, a)
        _avisar(db, [r1] if r1 else [], "Ausencia por aprobar", _texto(x, a) + pend, autor)
    else:
        _avisar(db, au.usuarios_con(db, a.id, "personal.autorizar"), "Ausencia por aprobar", _texto(x, a) + pend,
                autor)


# --------------------------------------------------------------------------- consultas
@router.get("/catalogos")
def catalogs():
    return {"tipos": TIPOS_AUSENCIA, "colectivos": COLECTIVOS_AUSENCIA, "estados": ESTADOS_AUSENCIA}


@router.get("/personas")
def people(asset_id: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Para quién puede registrar ausencias: Recepción 1 y dirección, para todo el personal del activo; el resto,
    solo para sí mismo."""
    a = _activo(db, scope, asset_id)
    yo = {"clave": "u", "id": scope.user.id, "nombre": scope.user.nombre,
          "colectivo": au.colectivo_usuario(db, scope.user, a)}
    gest = au.gestiona(db, scope, a)
    tecnico = scope.can_asset("personal.autorizar_mto", a.id)
    out = [yo]
    if gest or tecnico:
        for u in sorted(au.usuarios_con(db, a.id), key=lambda u: u.nombre):
            col = au.colectivo_usuario(db, u, a)
            if u.id != scope.user.id and (gest or col == "mantenimiento"):
                out.append({"clave": "u", "id": u.id, "nombre": u.nombre, "colectivo": col})
        staff = db.scalars(select(StaffMember).where(
            StaffMember.activo, or_(StaffMember.asset_id.is_(None), StaffMember.asset_id == a.id))
            .order_by(StaffMember.area, StaffMember.nombre))
        for p in staff:
            col = au.colectivo_staff(p)
            if gest or col == "mantenimiento":
                out.append({"clave": "s", "id": p.id, "colectivo": col,
                            "nombre": p.nombre + (f" ({p.empresa})" if p.empresa else "")})
    for p in out:
        p["colectivo_nombre"] = COLECTIVOS_AUSENCIA[p["colectivo"]]
        p["aprueba"] = au.quien_aprueba(p["colectivo"])
    return {"personas": out, "gestiona": gest, "r1": getattr(au.recepcion_1(db, a), "nombre", None)}


@router.get("")
def list_absences(asset_id: int, desde: date | None = None, hasta: date | None = None, estado: str | None = None,
                  scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _activo(db, scope, asset_id)
    stmt = select(Absence).where(Absence.asset_id == asset_id)
    if desde:
        stmt = stmt.where(Absence.hasta >= desde)
    if hasta:
        stmt = stmt.where(Absence.desde <= hasta)
    if estado == "vigentes":
        stmt = stmt.where(Absence.estado.in_(VIGENTES))
    elif estado:
        stmt = stmt.where(Absence.estado == estado)
    filas = list(db.scalars(stmt.order_by(Absence.desde, Absence.persona)))
    nombres = _nombres(db, filas)
    return [_out(db, scope, x, nombres) for x in filas]


# --------------------------------------------------------------------------- alta y tramitación
@router.post("", status_code=201)
def create_absence(data: AbsenceIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    a = _activo(db, scope, data.asset_id)
    if data.tipo not in TIPOS_AUSENCIA:
        bad_request(f"Tipo no válido. Opciones: {', '.join(TIPOS_AUSENCIA)}")
    if data.hasta < data.desde:
        bad_request("La fecha final es anterior a la inicial")
    if (data.hasta - data.desde).days > 365:
        bad_request("Una ausencia no puede pasar de un año: divídala en varias")
    if data.user_id and data.staff_id:
        bad_request("Indique un usuario o una persona del personal, no los dos")
    gest = au.gestiona(db, scope, a)
    if data.staff_id:
        p = get_or_404(db, StaffMember, data.staff_id)
        if p.asset_id not in (None, a.id):
            bad_request("Esa persona no trabaja en este activo")
        colectivo, persona, uid = au.colectivo_staff(p), p.nombre, None
    else:
        u = get_or_404(db, User, data.user_id) if data.user_id else scope.user
        colectivo, persona, uid = au.colectivo_usuario(db, u, a), u.nombre, u.id
    propia = uid == scope.user.id
    tecnico = colectivo == "mantenimiento" and scope.can_asset("personal.autorizar_mto", a.id)
    if not (propia or gest or tecnico):
        raise HTTPException(403, "Solo Recepción 1 o dirección registran las ausencias de otras personas")
    choque = db.scalar(select(Absence).where(
        Absence.estado.in_(VIGENTES), Absence.desde <= data.hasta, Absence.hasta >= data.desde,
        Absence.staff_id == data.staff_id if data.staff_id else Absence.user_id == uid))
    if choque:
        bad_request(f"{persona} ya tiene una ausencia del {choque.desde:%d/%m/%Y} al {choque.hasta:%d/%m/%Y}")
    x = Absence(asset_id=a.id, user_id=uid, staff_id=data.staff_id, persona=persona, colectivo=colectivo,
                tipo=data.tipo, desde=data.desde, hasta=data.hasta, motivo=data.motivo or None,
                creado_por=scope.user.id)
    db.add(x)
    db.flush()
    if data.aprobar:
        if not au.puede_aprobar(db, scope, x):
            raise HTTPException(403, f"Esta ausencia la aprueba: {au.quien_aprueba(colectivo)}")
        x.estado, x.resuelto_por, x.resuelto_en = "aprobada", scope.user.id, datetime.now()
    audit(db, scope.user, "crear", "ausencia", x.id, {"persona": persona, "tipo": x.tipo, "desde": x.desde.isoformat(),
                                                      "hasta": x.hasta.isoformat(), "estado": x.estado})
    db.commit()
    _avisos_nueva(db, x, a, scope.user.id)
    out = _uno(db, scope, x)
    # coincidencias con otras personas del mismo colectivo (para organizar el servicio)
    out["coinciden"] = [o.persona for o in db.scalars(select(Absence).where(
        Absence.asset_id == a.id, Absence.colectivo == colectivo, Absence.id != x.id,
        Absence.estado.in_(VIGENTES), Absence.desde <= x.hasta, Absence.hasta >= x.desde))]
    return out


def _resolver_aviso(db: Session, x: Absence, a: Asset, que: str) -> None:
    u = db.get(User, x.user_id) if x.user_id else None
    if u is not None:
        _avisar(db, [u], f"Su ausencia: {que}", _texto(x, a) + f" Estado: {que}." + (f" {x.nota}" if x.nota else ""))


@router.post("/{xid}/visto-bueno")
def give_consent(xid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, Absence, xid)
    a = _activo(db, scope, x.asset_id)
    if not au.puede_visto_bueno(db, scope, x):
        raise HTTPException(403, "El visto bueno lo da Recepción 1, solo en las ausencias pendientes de Recepción 2")
    x.estado, x.visto_bueno_por, x.visto_bueno_en = "visto_bueno", scope.user.id, datetime.now()
    audit(db, scope.user, "visto_bueno", "ausencia", x.id, {"persona": x.persona})
    db.commit()
    _avisar(db, au.usuarios_con(db, a.id, "personal.autorizar"), "Ausencia de Recepción 2 por aprobar",
            _texto(x, a) + " Recepción 1 ha dado su visto bueno. Apruébela en Personal → Vacaciones y ausencias.",
            scope.user.id)
    return _uno(db, scope, x)


@router.post("/{xid}/aprobar")
def approve(xid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, Absence, xid)
    a = _activo(db, scope, x.asset_id)
    if not au.puede_aprobar(db, scope, x):
        msg = "Falta el visto bueno de Recepción 1" if x.colectivo == "recepcion2" and x.estado == "pendiente" \
            else f"Esta ausencia la aprueba: {au.quien_aprueba(x.colectivo)}"
        raise HTTPException(403, msg)
    x.estado, x.resuelto_por, x.resuelto_en = "aprobada", scope.user.id, datetime.now()
    audit(db, scope.user, "aprobar", "ausencia", x.id, {"persona": x.persona})
    db.commit()
    _resolver_aviso(db, x, a, "aprobada")
    return _uno(db, scope, x)


@router.post("/{xid}/denegar")
def deny(xid: int, data: NotaIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, Absence, xid)
    a = _activo(db, scope, x.asset_id)
    # Recepción 1 también puede no dar el visto bueno a Recepción 2
    if not (au.puede_aprobar(db, scope, x) or au.puede_visto_bueno(db, scope, x)):
        raise HTTPException(403, f"Esta ausencia la aprueba: {au.quien_aprueba(x.colectivo)}")
    if not (data.nota or "").strip():
        bad_request("Indique el motivo")
    x.estado, x.resuelto_por, x.resuelto_en, x.nota = "denegada", scope.user.id, datetime.now(), data.nota.strip()
    audit(db, scope.user, "denegar", "ausencia", x.id, {"persona": x.persona, "nota": x.nota})
    db.commit()
    _resolver_aviso(db, x, a, "denegada")
    return _uno(db, scope, x)


@router.post("/{xid}/anular")
def cancel(xid: int, data: NotaIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, Absence, xid)
    _activo(db, scope, x.asset_id)
    if not au.puede_anular(db, scope, x):
        raise HTTPException(403, "No puede anular esta ausencia")
    x.estado, x.nota = "anulada", (data.nota or "").strip() or None
    audit(db, scope.user, "anular", "ausencia", x.id, {"persona": x.persona, "nota": x.nota})
    db.commit()
    return _uno(db, scope, x)


# --------------------------------------------------------------------------- resumen anual y Excel
def _resumen(db: Session, scope: Scope, asset_id: int, anio: int) -> list[dict]:
    ini, fin = date(anio, 1, 1), date(anio, 12, 31)
    filas = db.scalars(select(Absence).where(Absence.asset_id == asset_id, Absence.estado == "aprobada",
                                             Absence.desde <= fin, Absence.hasta >= ini))
    por: dict[tuple, dict] = {}
    for x in filas:
        if not au.ve_detalle(db, scope, x) and x.tipo == "baja_medica":
            tipo = "otro"
        else:
            tipo = x.tipo
        nat, lab = au.dias(max(x.desde, ini), min(x.hasta, fin))
        r = por.setdefault((x.colectivo, x.persona), {
            "persona": x.persona, "colectivo": COLECTIVOS_AUSENCIA.get(x.colectivo, x.colectivo),
            "tipos": {}, "naturales": 0, "laborables": 0})
        t = r["tipos"].setdefault(tipo, {"naturales": 0, "laborables": 0})
        t["naturales"] += nat
        t["laborables"] += lab
        r["naturales"] += nat
        r["laborables"] += lab
    return sorted(por.values(), key=lambda r: (r["colectivo"], r["persona"]))


@router.get("/resumen")
def summary(asset_id: int, anio: int | None = None, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _activo(db, scope, asset_id)
    anio = anio or date.today().year
    return {"anio": anio, "tipos": TIPOS_AUSENCIA, "personas": _resumen(db, scope, asset_id, anio)}


@router.get("/excel")
def excel(asset_id: int, anio: int | None = None, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    a = _activo(db, scope, asset_id)
    anio = anio or date.today().year
    negrita, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
    wb = Workbook()
    ws = wb.active
    ws.title = "Ausencias"
    filas = list_absences(asset_id, date(anio, 1, 1), date(anio, 12, 31), None, scope, db)
    ws.append(["Persona", "Colectivo", "Tipo", "Desde", "Hasta", "Días naturales", "Días laborables", "Estado",
               "Aprueba", "Visto bueno R1", "Resuelta por", "Motivo / nota"])
    for x in filas:
        ws.append([x["persona"], x["colectivo_nombre"], x["tipo_nombre"], date.fromisoformat(x["desde"]),
                   date.fromisoformat(x["hasta"]), x["dias_naturales"], x["dias_laborables"], x["estado_nombre"],
                   x["aprueba"], x["visto_bueno_nombre"] or "", x["resuelto_nombre"] or "",
                   " · ".join(t for t in (x["motivo"], x["nota"]) if t)])
    ws2 = wb.create_sheet("Resumen aprobadas")
    ws2.append(["Persona", "Colectivo", *TIPOS_AUSENCIA.values(), "Total naturales", "Total laborables"])
    for r in _resumen(db, scope, asset_id, anio):
        ws2.append([r["persona"], r["colectivo"], *(r["tipos"].get(t, {}).get("naturales", 0) for t in TIPOS_AUSENCIA),
                    r["naturales"], r["laborables"]])
    for hoja in (ws, ws2):
        for c in hoja[1]:
            c.font, c.fill = negrita, fondo
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    for fila in ws.iter_rows(min_row=2):
        for c in fila[3:5]:
            c.number_format = "DD/MM/YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="Ausencias_{a.codigo}_{anio}.xlsx"'})
