"""Documentos recibidos (carpeta de cada activo) y cuenta de gastos.

Recepción y dirección suben los documentos escaneados que llegan (facturas, tickets, cartas…): se guardan cifrados
en la carpeta del activo. Si el documento es un gasto, se anota en la cuenta de gastos del activo, general o de un
apartamento concreto. Cada activo ve solo lo suyo; quien gestiona la sociedad o el grupo, todo. Limpieza y
mantenimiento no tienen acceso.
"""
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .. import adjuntos, ausencias, avisos, documentos
from ..database import get_db
from ..facturacion import FORMAS_PAGO, dinero
from ..models import (TIPOS_DOC_RESERVADOS, AMBITOS_GASTO, CATEGORIAS_GASTO, TIPOS_DOCUMENTO, Asset, Company, Expense,
                      PresidencyReport, ReceivedDocument, Supplier, Unit, User)
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404, scoped

router = APIRouter(prefix="/api", tags=["documentos y gastos"])

IVAS = (0, 4, 5, 10, 21)
# Cómo se paga la factura del proveedor (obligatorio): por transferencia nuestra o cargada en cuenta por él
# Naturaleza del gasto: lo decide quien registra la factura
NATURALEZAS = {"OPEX": "OPEX · gasto corriente (mantenimiento, suministros, limpieza…)",
               "CAPEX": "CAPEX · inversión (obra, reforma, equipamiento que dura varios años)"}
# Retenciones que se practican al proveedor: (nombre, % habitual). «otra»: el % lo indica quien registra
RETENCIONES = {
    "irpf_profesional": ("IRPF profesionales", 15),
    "irpf_inicio": ("IRPF profesionales (inicio de actividad)", 7),
    "irpf_arrendamiento": ("IRPF arrendamiento de inmuebles", 19),
    "irpf_modulos": ("IRPF actividades en módulos (obras, transporte…)", 1),
    "otra": ("Otra retención", None),
}
FORMAS_PAGO_GASTO = {"transferencia": "Transferencia", "domiciliacion": "Cargo en cuenta (domiciliación)",
                     **{k: v for k, v in FORMAS_PAGO.items() if k not in ("transferencia", "domiciliacion")}}


class ExpenseIn(BaseModel):
    asset_id: int | None = None  # al crear el gasto sin documento
    documento_id: int | None = None
    fecha: date
    categoria: str
    naturaleza: str | None = None  # OPEX o CAPEX (obligatorio)
    concepto: str = Field(min_length=2, max_length=300)
    ambito: str = "general"
    unit_id: int | None = None
    ambito_detalle: str | None = Field(default=None, max_length=200)
    proveedor: str | None = Field(default=None, max_length=200)
    numero_factura: str | None = Field(default=None, max_length=60)
    vencimiento: date | None = None  # vencimiento de la factura (tal como figura en ella)
    # total de la factura tal como figura en ella: IVA incluido y, si la hay, retención ya descontada (negativo: abono)
    total: float = Field(gt=-1_000_000, lt=1_000_000)
    tipo_iva: float = 21
    base: float | None = None  # si se indica (varios tipos de IVA), la cuota es total + retención − base
    retencion_tipo: str | None = None  # RETENCIONES
    retencion_pct: float = Field(default=0, ge=0, le=50)
    forma_pago: str | None = None
    pagado: bool = False
    fecha_pago: date | None = None
    notas: str | None = None
    # al registrarla: no pagar de momento (se avisa a quien paga) hasta revisarla en esa fecha
    retener_pago: bool = False
    retener_motivo: str | None = Field(default=None, max_length=300)
    retener_revision: date | None = None
    confirmar_duplicado: bool = False  # registrarla aunque se parezca a otra ya registrada


# --------------------------------------------------------------------------- duplicados
def _num(s: str | None) -> str:
    """Nº de factura comparable: «F26/5594», «F26-5594» y «f26 5594» son el mismo."""
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def _prov(s: str | None) -> str:
    s = unicodedata.normalize("NFKD", (s or "").lower()).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\b(s ?l ?u?|s ?a ?u?|s ?coop|c ?b|s ?l ?l)\b", " ", s)  # forma jurídica
    return re.sub(r"\s+", "", s)


def _euros(x) -> str:
    return f"{float(x):,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _describe(db: Session, g: Expense) -> str:
    a = db.get(Asset, g.asset_id)
    return (f"{g.proveedor or 'sin proveedor'} nº {g.numero_factura or 's/n'} de {g.fecha:%d/%m/%Y}, "
            f"{_euros(float(g.total) - float(g.retencion or 0))} ({a.nombre}, apunte {g.id})")


def comprobar_duplicado(db: Session, asset_id: int, proveedor: str | None, supplier_id: int | None,
                        numero: str | None, liquido: Decimal, fecha: date, excluir: int | None = None,
                        confirmado: bool = False) -> None:
    """Mismo proveedor y mismo nº de factura (en la misma sociedad): no se registra (409).
    Parecida (mismo proveedor, importe y fecha; o mismo nº e importe con otro nombre de proveedor): se avisa y
    solo se registra si quien la sube lo confirma."""
    company = db.get(Asset, asset_id).company_id
    activos = set(db.scalars(select(Asset.id).where(Asset.company_id == company)))
    q = select(Expense).where(Expense.asset_id.in_(activos), Expense.fecha >= fecha - timedelta(days=730),
                              Expense.fecha <= fecha + timedelta(days=730))
    if excluir:
        q = q.where(Expense.id != excluir)
    num, prov = _num(numero), _prov(proveedor)
    posibles = []
    for g in db.scalars(q):
        mismo_prov = bool((supplier_id and g.supplier_id == supplier_id) or (prov and _prov(g.proveedor) == prov))
        mismo_num = bool(num) and _num(g.numero_factura) == num
        mismo_importe = dinero(float(g.total) - float(g.retencion or 0)) == liquido
        if mismo_prov and mismo_num:
            raise HTTPException(409, f"Factura duplicada: {_describe(db, g)} ya está registrada. No se registra "
                                     "de nuevo; si es otra factura, revise el nº de factura o el proveedor.")
        if (mismo_prov and mismo_importe and g.fecha == fecha) or (mismo_num and mismo_importe):
            posibles.append(_describe(db, g))
    if posibles and not confirmado:
        raise HTTPException(409, "Posible duplicado: se parece a " + "; ".join(posibles[:3])
                            + ". Compruébelo antes de registrarla.")


def _comprobar(db: Session, asset_id: int, v: dict, excluir: int | None = None, confirmado: bool = False) -> None:
    comprobar_duplicado(db, asset_id, v["proveedor"], v["supplier_id"], v["numero_factura"],
                        v["total"] - v["retencion"], v["fecha"], excluir=excluir, confirmado=confirmado)


class RetenerIn(BaseModel):
    motivo: str = Field(min_length=3, max_length=300)
    revision: date


class LiberarIn(BaseModel):
    nota: str | None = Field(default=None, max_length=300)


def _ver(scope: Scope, asset_id: int) -> None:
    scope.require_asset("documentos.ver", asset_id)


def _unidad(db: Session, asset_id: int, unit_id: int | None) -> Unit | None:
    if not unit_id:
        return None
    u = get_or_404(db, Unit, unit_id)
    if u.asset_id != asset_id:
        bad_request("La unidad no pertenece al activo")
    return u


def _valores_gasto(db: Session, asset_id: int, data: ExpenseIn) -> dict:
    if data.categoria not in CATEGORIAS_GASTO:
        bad_request(f"Categoría no válida. Opciones: {', '.join(CATEGORIAS_GASTO)}")
    if data.naturaleza not in NATURALEZAS:
        bad_request("Indique si el gasto es OPEX (gasto corriente) o CAPEX (inversión)")
    if data.ambito not in AMBITOS_GASTO:
        bad_request("Ámbito no válido: general, apartamento u otro")
    if data.ambito == "apartamento" and not data.unit_id:
        bad_request("Indique el apartamento del gasto")
    if data.tipo_iva not in IVAS:
        bad_request(f"Tipo de IVA no válido: {', '.join(map(str, IVAS))} %")
    if data.fecha > date.today():
        bad_request("La fecha de la factura no puede ser posterior a hoy: ponga la que figura en la factura")
    if data.vencimiento and data.vencimiento < data.fecha:
        bad_request("El vencimiento no puede ser anterior a la fecha de la factura")
    if not data.forma_pago:
        bad_request("Indique cómo se paga la factura: transferencia o cargo en cuenta (domiciliación)")
    if data.forma_pago not in FORMAS_PAGO_GASTO:
        bad_request("Forma de pago no válida")
    _unidad(db, asset_id, data.unit_id)
    if data.retencion_tipo and data.retencion_tipo not in RETENCIONES:
        bad_request(f"Retención no válida. Opciones: {', '.join(RETENCIONES)}")
    pct = Decimal(str(data.retencion_pct)) if data.retencion_tipo else Decimal(0)
    if data.retencion_tipo and pct <= 0:
        bad_request("Indique el % de la retención")
    liquido = dinero(data.total)  # lo que figura como total de la factura (retención ya descontada)
    if data.base is not None:
        base = dinero(data.base)
        retencion = dinero(base * pct / 100)
        if abs(base) > abs(liquido + retencion):
            bad_request("La base no puede ser mayor que el total")
    else:
        base = dinero(liquido / (1 + (Decimal(str(data.tipo_iva)) - pct) / 100))
        retencion = dinero(base * pct / 100)
    total = liquido + retencion  # base + IVA
    prov = " ".join((data.proveedor or "").split()) or None
    supplier = db.scalar(select(Supplier).where(Supplier.nombre == prov)) if prov else None
    return {"fecha": data.fecha, "categoria": data.categoria, "naturaleza": data.naturaleza,
            "concepto": data.concepto.strip(),
            "ambito": data.ambito, "unit_id": data.unit_id if data.ambito == "apartamento" else None,
            "ambito_detalle": data.ambito_detalle if data.ambito == "otro" else None, "proveedor": prov,
            "supplier_id": supplier.id if supplier else None, "numero_factura": data.numero_factura,
            "vencimiento": data.vencimiento,
            "total": total, "base": base, "cuota": total - base, "tipo_iva": data.tipo_iva,
            "forma_pago": data.forma_pago, "pagado": data.pagado,
            "fecha_pago": (data.fecha_pago or date.today()) if data.pagado else None, "notas": data.notas,
            "retencion_tipo": data.retencion_tipo or None, "retencion_pct": pct, "retencion": retencion}


SIN_JUSTIFICANTE = ("Para marcar la factura como pagada hay que adjuntar antes el justificante de pago "
                    "(botón «Pagado» de la cuenta de gastos)")


def _retener(g: Expense, user_id: int, motivo: str | None, revision: date | None) -> None:
    motivo = " ".join((motivo or "").split())
    if len(motivo) < 3:
        bad_request("Indique por qué se retiene el pago")
    if not revision:
        bad_request("Indique la fecha en que se revisará el pago retenido")
    if revision < date.today():
        bad_request("La fecha de revisión no puede ser anterior a hoy")
    if g.pagado:
        bad_request("La factura ya está pagada")
    g.pago_retenido, g.pago_retenido_motivo, g.pago_retenido_revision = True, motivo, revision
    g.pago_retenido_user_id, g.pago_retenido_fecha = user_id, datetime.now()


def _nombres(db: Session, ids_unidades: set[int], ids_usuarios: set[int]) -> tuple[dict, dict]:
    unidades = dict(db.execute(select(Unit.id, Unit.codigo).where(Unit.id.in_(ids_unidades or {-1}))).all())
    usuarios = dict(db.execute(select(User.id, User.nombre).where(User.id.in_(ids_usuarios or {-1}))).all())
    return unidades, usuarios


def _gasto_out(g: Expense, unidades: dict, usuarios: dict, activos: dict) -> dict:
    d = g.to_dict()
    d["activo"] = activos.get(g.asset_id)
    d["unidad"] = unidades.get(g.unit_id)
    d["categoria_nombre"] = CATEGORIAS_GASTO.get(g.categoria, g.categoria)
    d["lugar"] = (f"Apartamento {d['unidad']}" if g.ambito == "apartamento"
                  else g.ambito_detalle or "Otro" if g.ambito == "otro" else "General")
    d["usuario"] = usuarios.get(g.user_id)
    d["liquido"] = round(float(g.total) - float(g.retencion or 0), 2)  # lo que se paga al proveedor
    d["retencion_nombre"] = RETENCIONES.get(g.retencion_tipo, ("",))[0] if g.retencion_tipo else None
    d["pago_retenido_por"] = usuarios.get(g.pago_retenido_user_id) if g.pago_retenido else None
    return d


def _doc_out(x: ReceivedDocument, gasto: Expense | None, unidades: dict, usuarios: dict, activos: dict) -> dict:
    d = x.to_dict(exclude=("fichero", "sha256"))
    d["activo"] = activos.get(x.asset_id)
    d["tipo_nombre"] = TIPOS_DOCUMENTO.get(x.tipo, x.tipo)
    d["unidad"] = unidades.get(x.unit_id)
    d["usuario"] = usuarios.get(x.user_id)
    d["gasto"] = _gasto_out(gasto, unidades, usuarios, activos) if gasto else None
    d["de_colaborador"] = x.revision is not None
    return d


def _activos(db: Session) -> dict:
    return dict(db.execute(select(Asset.id, Asset.nombre)).all())


# --------------------------------------------------------------------------- documentos recibidos (carpeta)
@router.get("/documentos-recibidos")
def list_documents(asset_id: int | None = None, tipo: str | None = None, desde: date | None = None,
                   hasta: date | None = None, q: str | None = None, revision: str | None = None,
                   scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    scope.require_any("documentos.ver")
    stmt = scoped(select(ReceivedDocument), ReceivedDocument.asset_id, scope.asset_ids("documentos.ver"))
    if asset_id:
        stmt = stmt.where(ReceivedDocument.asset_id == asset_id)
    if tipo:
        stmt = stmt.where(ReceivedDocument.tipo == tipo)
    if desde:
        stmt = stmt.where(ReceivedDocument.fecha >= desde)
    if hasta:
        stmt = stmt.where(ReceivedDocument.fecha <= hasta)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(ReceivedDocument.emisor.ilike(like), ReceivedDocument.referencia.ilike(like),
                              ReceivedDocument.descripcion.ilike(like), ReceivedDocument.nombre.ilike(like)))
    if revision:
        stmt = stmt.where(ReceivedDocument.revision == revision)
    docs = list(db.scalars(stmt.order_by(ReceivedDocument.fecha.desc(), ReceivedDocument.id.desc()).limit(3000)))
    # la documentación del personal de las subcontratas solo la ven dirección y Recepción 1 del activo
    reservado: dict[int, bool] = {}
    docs = [x for x in docs if x.tipo not in TIPOS_DOC_RESERVADOS or reservado.setdefault(
        x.asset_id, ausencias.gestiona(db, scope, db.get(Asset, x.asset_id)))]
    gastos = {g.documento_id: g for g in db.scalars(select(Expense).where(
        Expense.documento_id.in_([x.id for x in docs] or [-1])))}
    unidades, usuarios = _nombres(db, {x.unit_id for x in docs if x.unit_id} | {g.unit_id for g in gastos.values()
                                                                                  if g.unit_id},
                                  {x.user_id for x in docs if x.user_id})
    activos = _activos(db)
    return [_doc_out(x, gastos.get(x.id), unidades, usuarios, activos) for x in docs]


@router.post("/documentos-recibidos", status_code=201)
def upload_document(tareas: BackgroundTasks, ficheros: list[UploadFile] = File(...), asset_id: int = Form(...),
                    tipo: str = Form(...), fecha: date = Form(...), vencimiento: date | None = Form(None),
                    emisor: str | None = Form(None), referencia: str | None = Form(None),
                    descripcion: str | None = Form(None), unit_id: int | None = Form(None),
                    gasto: str | None = Form(None), confirmar_duplicado: bool = Form(False),
                    scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Sube un documento escaneado (foto o PDF) a la carpeta del activo. Con `gasto` (JSON con los datos del
    gasto) se anota además en la cuenta de gastos."""
    scope.require_asset("documentos.editar", asset_id)
    get_or_404(db, Asset, asset_id)
    if tipo not in TIPOS_DOCUMENTO:
        bad_request(f"Tipo de documento no válido. Opciones: {', '.join(TIPOS_DOCUMENTO)}")
    _unidad(db, asset_id, unit_id)
    if vencimiento and vencimiento < fecha:
        bad_request("El vencimiento no puede ser anterior a la fecha del documento")
    datos, mime, nombre = _fichero(ficheros)
    gasto_in = None
    if gasto:
        try:
            gasto_in = ExpenseIn(**json.loads(gasto))
        except (ValueError, ValidationError) as e:
            bad_request(f"Datos del gasto no válidos: {str(e)[:300]}")
    if gasto_in and gasto_in.pagado:
        bad_request(SIN_JUSTIFICANTE)
    valores = _valores_gasto(db, asset_id, gasto_in) if gasto_in else None
    confirmado = confirmar_duplicado or bool(gasto_in and gasto_in.confirmar_duplicado)
    huella = documentos.huella(datos)
    # duplicados, antes de guardar nada: la factura (proveedor + nº) y el propio fichero
    if valores:
        _comprobar(db, asset_id, valores, confirmado=confirmado)
    elif tipo == "factura" and emisor and referencia:
        for d in db.scalars(select(ReceivedDocument).where(ReceivedDocument.asset_id == asset_id,
                                                           ReceivedDocument.tipo == "factura")):
            if _prov(d.emisor) == _prov(emisor) and _num(d.referencia) == _num(referencia):
                raise HTTPException(409, f"Factura duplicada: {d.emisor} nº {d.referencia} ya está en la carpeta "
                                         f"(subida el {d.subido:%d/%m/%Y}). No se registra de nuevo.")
    igual = db.scalar(select(ReceivedDocument).where(ReceivedDocument.asset_id == asset_id,
                                                     ReceivedDocument.sha256 == huella).limit(1))
    if igual and not confirmado:
        raise HTTPException(409, f"Posible duplicado: ese mismo fichero ya se subió el {igual.subido:%d/%m/%Y} "
                                 f"({igual.emisor or igual.nombre}). Compruébelo antes de registrarlo.")
    x = ReceivedDocument(asset_id=asset_id, unit_id=unit_id, tipo=tipo, fecha=fecha, vencimiento=vencimiento,
                         emisor=" ".join((emisor or "").split()) or None, referencia=referencia or None,
                         descripcion=descripcion or None, nombre=nombre, fichero=documentos.guardar(datos), mime=mime,
                         tamano=len(datos), sha256=huella, user_id=scope.user.id)
    db.add(x)
    db.flush()
    g = None
    if gasto_in:
        g = Expense(asset_id=asset_id, documento_id=x.id, user_id=scope.user.id, **valores)
        if gasto_in.retener_pago:
            _retener(g, scope.user.id, gasto_in.retener_motivo, gasto_in.retener_revision)
        db.add(g)
        db.flush()
    audit(db, scope.user, "subir", "documento_recibido", x.id, {"tipo": tipo, "emisor": x.emisor,
                                                                 "gasto": g.id if g else None})
    db.commit()
    if g is not None and g.pago_retenido:
        tareas.add_task(avisos.pago_retenido, g.id, scope.user.id)
    unidades, usuarios = _nombres(db, {i for i in (x.unit_id, g.unit_id if g else None) if i}, {scope.user.id})
    return _doc_out(x, g, unidades, usuarios, _activos(db))


@router.post("/documentos-recibidos/leer")
def read_invoice(ficheros: list[UploadFile] = File(...), scope: Scope = Depends(get_scope),
                 db: Session = Depends(get_db)):
    """Lee la factura (PDF o fotos) y propone sus datos para el formulario. No guarda nada.
    Los datos «dudosos» (no encontrados o que no cuadran) se marcan para revisarlos."""
    scope.require_any("documentos.editar")
    from .. import lector_facturas, nif
    datos, _, _ = _fichero(ficheros)
    propios = {nif.normalizar(c) for c in db.scalars(select(Company.cif)) if c}
    proveedores = {nif.normalizar(s.nif): s.nombre
                   for s in db.scalars(select(Supplier).where(Supplier.nif.is_not(None)))}
    try:
        r = lector_facturas.leer(datos, propios, proveedores)
    except ValueError as e:
        bad_request(str(e))
    c = r["campos"]
    if c.get("nif") and c["nif"] in proveedores:
        r["proveedor_conocido"] = True
    return r


def _fichero(ficheros: list[UploadFile]) -> tuple[bytes, str, str]:
    """Un PDF o una o varias fotos. Varias fotos (páginas del mismo documento) se unen en un PDF."""
    if not ficheros:
        bad_request("Adjunte el documento escaneado")
    if len(ficheros) > 20:
        bad_request("Máximo 20 páginas por documento")
    partes = []
    for f in ficheros:
        try:
            partes.append(adjuntos.normalizar(f.file.read(adjuntos.TAM_MAX + 1), (f.filename or "documento")[-150:]))
        except ValueError as e:
            bad_request(f"{f.filename}: {e}")
    if len(partes) == 1:
        return partes[0]
    if any(m != "image/jpeg" for _, m, _ in partes):
        bad_request("Para varias páginas adjunte solo fotos, o un único PDF con todas")
    from PIL import Image
    paginas = [Image.open(BytesIO(d)).convert("RGB") for d, _, _ in partes]
    out = BytesIO()
    paginas[0].save(out, "PDF", save_all=True, append_images=paginas[1:], resolution=150)
    base = partes[0][2].rsplit(".", 1)[0]
    return out.getvalue(), "application/pdf", f"{base}.pdf"


@router.get("/documentos-recibidos/{did}/fichero")
def view_document(did: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, ReceivedDocument, did)
    _ver(scope, x.asset_id)
    if x.tipo in TIPOS_DOC_RESERVADOS and not ausencias.gestiona(db, scope, db.get(Asset, x.asset_id)):
        raise HTTPException(403, "La documentación del personal solo la ven Recepción 1 y dirección")
    audit(db, scope.user, "ver", "documento_recibido", did)
    db.commit()
    return Response(documentos.leer(x.fichero), media_type=x.mime, headers={
        "Content-Disposition": f'inline; filename="{x.nombre}"', "Cache-Control": "private, no-store"})


def _puede_borrar(scope: Scope, asset_id: int, autor: int | None, cuando: datetime) -> None:
    """Para que el control sea fiable, solo borra quien lo subió el mismo día, o la dirección."""
    scope.require_asset("documentos.editar", asset_id)
    if scope.company_level_ids("documentos.editar") is None:
        return
    a = scope.db.get(Asset, asset_id)
    if a.company_id in (scope.company_level_ids("documentos.editar") or set()):
        return
    if autor != scope.user.id or cuando.date() != date.today():
        raise HTTPException(403, "Solo la dirección, o quien lo registró el mismo día, puede borrarlo")


@router.delete("/documentos-recibidos/{did}")
def delete_document(did: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x = get_or_404(db, ReceivedDocument, did)
    _puede_borrar(scope, x.asset_id, x.user_id, x.subido)
    for g in db.scalars(select(Expense).where(Expense.documento_id == did)):
        g.documento_id = None  # el apunte del gasto se conserva
    for r in db.scalars(select(PresidencyReport).where(PresidencyReport.documento_id == did)):
        r.documento_id = None  # el registro del envío del informe se conserva
    documentos.borrar(x.fichero)
    audit(db, scope.user, "borrar", "documento_recibido", did, {"tipo": x.tipo, "emisor": x.emisor,
                                                                 "nombre": x.nombre})
    db.delete(x)
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------- cuenta de gastos
def _gastos_q(scope: Scope, asset_id, desde, hasta, categoria, unit_id, pagado):
    stmt = scoped(select(Expense), Expense.asset_id, scope.asset_ids("documentos.ver"))
    if asset_id:
        stmt = stmt.where(Expense.asset_id == asset_id)
    if desde:
        stmt = stmt.where(Expense.fecha >= desde)
    if hasta:
        stmt = stmt.where(Expense.fecha <= hasta)
    if categoria:
        stmt = stmt.where(Expense.categoria == categoria)
    if unit_id:
        stmt = stmt.where(Expense.unit_id == unit_id)
    if pagado is not None:
        stmt = stmt.where(Expense.pagado.is_(pagado))
    return stmt.order_by(Expense.fecha.desc(), Expense.id.desc())


def _gastos(db, scope, asset_id, desde, hasta, categoria, unit_id, pagado) -> list[dict]:
    scope.require_any("documentos.ver")
    filas = list(db.scalars(_gastos_q(scope, asset_id, desde, hasta, categoria, unit_id, pagado).limit(10000)))
    unidades, usuarios = _nombres(db, {g.unit_id for g in filas if g.unit_id},
                                  {i for g in filas for i in (g.user_id, g.pago_retenido_user_id) if i})
    activos = _activos(db)
    docs = dict(db.execute(select(ReceivedDocument.id, ReceivedDocument.tipo).where(
        ReceivedDocument.id.in_([g.documento_id for g in filas if g.documento_id] or [-1]))).all())
    out = []
    for g in filas:
        d = _gasto_out(g, unidades, usuarios, activos)
        d["documento_tipo"] = TIPOS_DOCUMENTO.get(docs.get(g.documento_id)) if g.documento_id else None
        out.append(d)
    return out


@router.get("/gastos")
def list_expenses(asset_id: int | None = None, desde: date | None = None, hasta: date | None = None,
                  categoria: str | None = None, unit_id: int | None = None, pagado: bool | None = None,
                  scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    filas = _gastos(db, scope, asset_id, desde, hasta, categoria, unit_id, pagado)
    por_cat = defaultdict(float)
    for g in filas:
        por_cat[g["categoria_nombre"]] += g["base"]
    return {"gastos": filas, "totales": {
        "base": round(sum(g["base"] for g in filas), 2), "cuota": round(sum(g["cuota"] for g in filas), 2),
        "total": round(sum(g["total"] for g in filas), 2),
        "retencion": round(sum(g["retencion"] for g in filas), 2),
        "opex": round(sum(g["base"] for g in filas if g["naturaleza"] == "OPEX"), 2),
        "capex": round(sum(g["base"] for g in filas if g["naturaleza"] == "CAPEX"), 2),
        "sin_naturaleza": sum(1 for g in filas if not g["naturaleza"]),
        "pendiente_pago": round(sum(g["liquido"] for g in filas if not g["pagado"]), 2),
        "retenidas": sum(1 for g in filas if g["pago_retenido"] and not g["pagado"]),
        "retenidas_importe": round(sum(g["liquido"] for g in filas if g["pago_retenido"] and not g["pagado"]), 2),
        "sin_documento": sum(1 for g in filas if not g["documento_id"]),
        "por_categoria": {k: round(v, 2) for k, v in sorted(por_cat.items(), key=lambda x: -x[1])}}}


@router.post("/gastos", status_code=201)
def create_expense(data: ExpenseIn, tareas: BackgroundTasks, scope: Scope = Depends(get_scope),
                   db: Session = Depends(get_db)):
    """Gasto sin documento escaneado (p. ej. un recibo domiciliado) o enlazado a un documento ya subido."""
    if data.documento_id:
        doc = get_or_404(db, ReceivedDocument, data.documento_id)
        asset_id = doc.asset_id
        if db.scalar(select(Expense.id).where(Expense.documento_id == doc.id)):
            bad_request("Ese documento ya tiene su gasto anotado: edítelo")
    elif data.asset_id:
        asset_id = data.asset_id
        get_or_404(db, Asset, asset_id)
    else:
        bad_request("Indique el activo del gasto")
    scope.require_asset("documentos.editar", asset_id)
    if data.pagado:
        bad_request(SIN_JUSTIFICANTE)
    valores = _valores_gasto(db, asset_id, data)
    _comprobar(db, asset_id, valores, confirmado=data.confirmar_duplicado)
    g = Expense(asset_id=asset_id, documento_id=data.documento_id, user_id=scope.user.id, **valores)
    if data.retener_pago:
        _retener(g, scope.user.id, data.retener_motivo, data.retener_revision)
    db.add(g)
    db.flush()
    audit(db, scope.user, "crear", "gasto", g.id, {"concepto": g.concepto, "total": float(g.total),
                                                    "pago_retenido": g.pago_retenido})
    db.commit()
    if g.pago_retenido:
        tareas.add_task(avisos.pago_retenido, g.id, scope.user.id)
    return _uno(db, scope, g.id)


def _uno(db: Session, scope: Scope, gid: int) -> dict:
    g = db.get(Expense, gid)
    unidades, usuarios = _nombres(db, {g.unit_id} if g.unit_id else set(),
                                  {i for i in (g.user_id, g.pago_retenido_user_id) if i})
    return _gasto_out(g, unidades, usuarios, _activos(db))


@router.put("/gastos/{gid}")
def update_expense(gid: int, data: ExpenseIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    g = get_or_404(db, Expense, gid)
    scope.require_asset("documentos.editar", g.asset_id)
    antes = {k: str(v) for k, v in g.to_dict().items()}
    if data.pagado and not g.pagado and g.pago_retenido:
        bad_request("El pago de esta factura está retenido: debe liberarlo quien se encarga de los pagos")
    if data.pagado and not g.pagado:
        bad_request(SIN_JUSTIFICANTE)
    valores = _valores_gasto(db, g.asset_id, data)
    _comprobar(db, g.asset_id, valores, excluir=gid, confirmado=True)  # al editar: solo el duplicado exacto
    if g.pagado and data.pagado:  # sigue pagada: se conserva la fecha de pago si no se cambia
        valores["fecha_pago"] = data.fecha_pago or g.fecha_pago
    if not data.pagado:  # desmarcar el pago: el justificante queda en la carpeta, sin enlazar
        g.justificante_id = None
    for k, v in valores.items():
        setattr(g, k, v)
    audit(db, scope.user, "editar", "gasto", gid,
          {k: [antes.get(k), str(v)] for k, v in g.to_dict().items() if antes.get(k) != str(v)})
    db.commit()
    return _uno(db, scope, gid)


@router.post("/gastos/{gid}/pagar")
def pay_expense(gid: int, ficheros: list[UploadFile] = File(...), fecha_pago: date | None = Form(None),
                forma_pago: str | None = Form(None), scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Marca la factura como pagada con su justificante de pago (obligatorio), que se guarda en la carpeta del
    activo como «Justificante de pago» enlazado al gasto."""
    g = get_or_404(db, Expense, gid)
    scope.require_asset("documentos.editar", g.asset_id)
    if g.pagado:
        bad_request("La factura ya está pagada")
    if g.pago_retenido:
        bad_request("El pago de esta factura está retenido: debe liberarlo quien se encarga de los pagos")
    if forma_pago:
        if forma_pago not in FORMAS_PAGO_GASTO:
            bad_request("Forma de pago no válida")
        g.forma_pago = forma_pago
    if not g.forma_pago:
        bad_request("Indique cómo se ha pagado: transferencia o cargo en cuenta")
    f = fecha_pago or date.today()
    if f > date.today():
        bad_request("La fecha de pago no puede ser futura")
    datos, mime, nombre = _fichero(ficheros)
    x = ReceivedDocument(asset_id=g.asset_id, unit_id=g.unit_id, tipo="justificante", fecha=f, emisor=g.proveedor,
                         referencia=g.numero_factura, descripcion=f"Justificante de pago · {g.concepto}"[:1000],
                         nombre=nombre, fichero=documentos.guardar(datos), mime=mime, tamano=len(datos),
                         sha256=documentos.huella(datos), user_id=scope.user.id)
    db.add(x)
    db.flush()
    g.pagado, g.fecha_pago, g.justificante_id = True, f, x.id
    audit(db, scope.user, "pagar", "gasto", gid, {"fecha_pago": f.isoformat(), "justificante": x.id,
                                                  "total": float(g.total)})
    db.commit()
    return _uno(db, scope, gid)


@router.post("/gastos/{gid}/retener-pago")
def hold_payment(gid: int, data: RetenerIn, tareas: BackgroundTasks, scope: Scope = Depends(get_scope),
                 db: Session = Depends(get_db)):
    """No pagar de momento: queda retenida y se avisa a quien paga (dirección, administración) hasta la revisión.
    Sirve también para cambiar el motivo o la fecha de revisión."""
    g = get_or_404(db, Expense, gid)
    scope.require_asset("documentos.editar", g.asset_id)
    _retener(g, scope.user.id, data.motivo, data.revision)
    audit(db, scope.user, "retener_pago", "gasto", gid, {"motivo": g.pago_retenido_motivo,
                                                         "revision": str(g.pago_retenido_revision)})
    db.commit()
    tareas.add_task(avisos.pago_retenido, g.id, scope.user.id)
    return _uno(db, scope, gid)


@router.post("/gastos/{gid}/liberar-pago")
def release_payment(gid: int, data: LiberarIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Se puede pagar: lo decide quien paga (finanzas) o quien la retuvo."""
    g = get_or_404(db, Expense, gid)
    if not g.pago_retenido:
        bad_request("El pago de esta factura no está retenido")
    if not (scope.can_asset("finanzas.ver", g.asset_id) and scope.can_asset("documentos.editar", g.asset_id)) \
            and g.pago_retenido_user_id != scope.user.id:
        raise HTTPException(403, "Solo quien se encarga de los pagos (o quien la retuvo) puede liberarla")
    audit(db, scope.user, "liberar_pago", "gasto", gid, {"motivo_retencion": g.pago_retenido_motivo,
                                                         "nota": data.nota})
    g.pago_retenido, g.pago_retenido_revision = False, None
    db.commit()
    return _uno(db, scope, gid)


@router.delete("/gastos/{gid}")
def delete_expense(gid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    g = get_or_404(db, Expense, gid)
    _puede_borrar(scope, g.asset_id, g.user_id, g.creado)
    audit(db, scope.user, "borrar", "gasto", gid, {"concepto": g.concepto, "total": float(g.total)})
    db.delete(g)
    db.commit()
    return {"ok": True}


@router.get("/gastos/excel")
def expenses_excel(asset_id: int | None = None, desde: date | None = None, hasta: date | None = None,
                   categoria: str | None = None, unit_id: int | None = None, scope: Scope = Depends(get_scope),
                   db: Session = Depends(get_db)):
    """Cuenta de gastos en Excel: detalle, resumen por categoría y mes, y gasto por apartamento."""
    from openpyxl import Workbook

    from .informes import ENTERO, EUR, FECHA, _hoja
    filas = _gastos(db, scope, asset_id, desde, hasta, categoria, unit_id, None)
    wb = Workbook()
    wb.remove(wb.active)
    periodo = f"{desde:%d/%m/%Y} – {hasta:%d/%m/%Y}" if desde and hasta else "todo el periodo"
    _hoja(wb, "Gastos", ["Fecha factura", "Activo", "Ámbito", "Categoría", "OPEX/CAPEX", "Concepto", "Proveedor",
                         "Nº factura",
                         "Vencimiento", "Base", "IVA %", "Cuota IVA", "Total con IVA", "Tipo de retención",
                         "Retención %", "Importe retención", "A pagar", "Pagado", "Fecha pago", "Forma de pago", "Pago retenido", "Revisión",
                         "Documento", "Registrado por"],
          [[date.fromisoformat(g["fecha"]), g["activo"], g["lugar"], g["categoria_nombre"], g["naturaleza"] or "",
            g["concepto"],
            g["proveedor"] or "", g["numero_factura"] or "",
            date.fromisoformat(g["vencimiento"]) if g["vencimiento"] else None, g["base"], g["tipo_iva"], g["cuota"],
            g["total"], g["retencion_nombre"] or "", g["retencion_pct"] or None, g["retencion"] or None, g["liquido"],
            "Sí" if g["pagado"] else "No", date.fromisoformat(g["fecha_pago"]) if g["fecha_pago"] else None,
            FORMAS_PAGO_GASTO.get(g["forma_pago"], ""),
            f"Sí: {g['pago_retenido_motivo']}" if g["pago_retenido"] else "",
            date.fromisoformat(g["pago_retenido_revision"]) if g["pago_retenido_revision"] else None,
            g["documento_tipo"] or "SIN DOCUMENTO", g["usuario"] or ""]
           for g in sorted(filas, key=lambda x: (x["fecha"], x["id"]))],
          {8: FECHA, 9: EUR, 10: ENTERO, 11: EUR, 12: EUR, 14: ENTERO, 15: EUR, 16: EUR, 18: FECHA, 21: FECHA},
          totales=[9, 11, 12, 15, 16],
          nota=f"Cuenta de gastos · {periodo}. Fecha y vencimiento tal como figuran en la factura recibida.")
    res = defaultdict(lambda: defaultdict(float))
    for g in filas:
        res[(g["activo"], g["categoria_nombre"])][g["fecha"][:7]] += g["base"]
    meses = sorted({m for v in res.values() for m in v})
    _hoja(wb, "Por categoría", ["Activo", "Categoría", *meses, "Total"],
          [[a, c, *[round(v.get(m, 0), 2) for m in meses], round(sum(v.values()), 2)]
           for (a, c), v in sorted(res.items())],
          {i: EUR for i in range(2, len(meses) + 3)}, totales=list(range(2, len(meses) + 3)),
          nota="Base imponible (sin IVA) por categoría y mes")
    por_apto = defaultdict(float)
    for g in filas:
        if g["unidad"]:
            por_apto[(g["activo"], g["unidad"])] += g["base"]
    _hoja(wb, "Por apartamento", ["Activo", "Apartamento", "Gasto (base)"],
          [[a, u, round(v, 2)] for (a, u), v in sorted(por_apto.items())], {2: EUR}, totales=[2],
          nota="Gastos imputados a un apartamento concreto")
    out = BytesIO()
    wb.save(out)
    audit(db, scope.user, "exportar", "gastos", None, {"desde": str(desde), "hasta": str(hasta),
                                                       "asset_id": asset_id})
    db.commit()
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Cuenta_de_gastos.xlsx"'})


@router.get("/gastos/catalogos")
def catalogs(scope: Scope = Depends(get_scope)):
    return {"tipos_documento": TIPOS_DOCUMENTO, "categorias": CATEGORIAS_GASTO, "ambitos": AMBITOS_GASTO,
            "formas_pago": FORMAS_PAGO_GASTO, "ivas": IVAS,
            "retenciones": {k: {"nombre": n, "pct": p} for k, (n, p) in RETENCIONES.items()},
            "naturalezas": NATURALEZAS}
