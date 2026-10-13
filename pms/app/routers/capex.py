"""CAPEX (fase 1): definición del proyecto de inversión, presupuesto base, plantillas «CAPEX tipo», base de precios
FIEBDC-3, documentos, licitación ciega a contratistas, cuadro comparativo, adjudicación (completa o por capítulos) y
aprobación de la Dirección General por correo cuando se supera el umbral (30.000 € sin IVA por defecto).

Flujo de estados: definicion → licitacion → adjudicacion → (pendiente_aprobacion →) adjudicado. Cancelado en
cualquier momento. Las mediciones solo se cambian en definición, para que todas las ofertas sean sobre lo mismo."""
import hmac
import io
import secrets
import zipfile
from datetime import date, datetime, time, timedelta
from html import escape
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import HTMLResponse
from pydantic import BaseModel, Field
from sqlalchemy import delete, insert, or_, select
from sqlalchemy.orm import Session

from .. import adjuntos, avisos, bc3, capex, documentos
from ..config import BASE_DIR, settings
from ..database import get_db
from ..models import (ESTADOS_CAPEX, ETIQUETAS_DOC_CAPEX, Asset, CapexApproval, CapexAward, CapexDocument, CapexItem,
                      CapexProject, CapexTemplate, EmailLog, PriceBase, PriceItem, Supplier, SupplierBid,
                      TenderInvite, TenderProcess, Unit, User)
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404, scoped

router = APIRouter(prefix="/api/capex", tags=["CAPEX"])
publico = APIRouter(tags=["CAPEX (enlaces)"])
STATIC = Path(BASE_DIR) / "static"
OTROS_FORMATOS = {  # además de PDF y fotos: planos CAD y documentos de oficina
    "dwg": "application/acad", "dxf": "application/dxf", "dwf": "model/vnd.dwf", "docx":
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "doc": "application/msword",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xls": "application/vnd.ms-excel",
    "bc3": "text/plain", "zip": "application/zip", "ifc": "application/x-step"}
TAM_DOC = 40 * 1024 * 1024
TAM_BC3 = 120 * 1024 * 1024
ADJUNTOS_MAX = 15 * 1024 * 1024  # por encima, los documentos se descargan desde el enlace
EDITABLE = ("definicion",)


# --------------------------------------------------------------------------- utilidades
def _proyecto(db: Session, scope: Scope, pid: int, gestionar: bool = False) -> CapexProject:
    p = get_or_404(db, CapexProject, pid)
    scope.require_asset("capex.gestionar" if gestionar else "capex.ver", p.asset_id)
    return p


def _editable(p: CapexProject) -> None:
    if p.estado not in EDITABLE:
        bad_request("Las mediciones y precios base solo se cambian en definición (todas las ofertas deben ser sobre "
                    "lo mismo). Si hace falta, vuelva el proyecto a definición.")


def _tocado(p: CapexProject) -> None:
    p.actualizado = datetime.now()


def _base_url(request: Request | None) -> str:
    return (settings.url or (str(request.base_url) if request else "")).rstrip("/")


def _aprobador(scope: Scope) -> bool:
    return (scope.user.email or "").lower() in settings.capex_aprobadores


def _item(x: CapexItem) -> dict:
    d = x.to_dict()
    d["cantidad"], d["precio"] = float(x.cantidad), float(x.precio)
    d["importe"] = capex.r2(d["cantidad"] * d["precio"])
    return d


def _doc(d: CapexDocument) -> dict:
    return {**d.to_dict(exclude=("fichero", "sha256")), "etiqueta_nombre": ETIQUETAS_DOC_CAPEX.get(d.etiqueta, d.etiqueta)}


def _resumen(db: Session, p: CapexProject) -> dict:
    d = p.to_dict()
    for k in ("presupuesto_objetivo", "gg_pct", "bi_pct", "iva_pct"):
        d[k] = float(d[k]) if d[k] is not None else None
    a = db.get(Asset, p.asset_id)
    u = db.get(Unit, p.unit_id) if p.unit_id else None
    d.update(activo=a.nombre, unidad=u.codigo if u else None, estado_nombre=ESTADOS_CAPEX.get(p.estado, p.estado))
    t = capex.totales(p, capex.partidas(db, p.id))
    d.update(pem=t["pem"], pbl=t["pbl"], total_iva=t["total"], desviacion_objetivo=t["desviacion_objetivo"])
    adj = [x for x in db.scalars(select(CapexAward).where(CapexAward.project_id == p.id,
                                                          CapexAward.estado.in_(("propuesta", "aprobada"))))]
    d["adjudicado"] = capex.r2(sum(float(x.importe) for x in adj)) if adj else None
    return d


def _fichero_doc(f: UploadFile) -> tuple[bytes, str, str]:
    nombre = (f.filename or "documento")[-150:]
    datos = f.file.read(TAM_DOC + 1)
    if len(datos) > TAM_DOC:
        bad_request(f"{nombre}: supera los 40 MB")
    ext = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
    if ext in OTROS_FORMATOS:
        return datos, OTROS_FORMATOS[ext], nombre
    try:
        return adjuntos.normalizar(datos, nombre)
    except ValueError as e:
        bad_request(f"{nombre}: {e}. También se admiten planos DWG/DXF y documentos Word/Excel.")


def _guardar_doc(db: Session, datos: bytes, mime: str, nombre: str, user_id: int | None, **kw) -> CapexDocument:
    d = CapexDocument(nombre=nombre[:160], fichero=documentos.guardar(datos), mime=mime[:100], tamano=len(datos),
                      sha256=documentos.huella(datos), user_id=user_id, **kw)
    db.add(d)
    return d


def _copiar_doc(db: Session, d: CapexDocument, user_id: int | None, **kw) -> CapexDocument:
    return _guardar_doc(db, documentos.leer(d.fichero), d.mime, d.nombre, user_id, etiqueta=d.etiqueta, notas=d.notas,
                        licitacion=d.licitacion, **kw)


def _aplicar_plantilla(db: Session, p: CapexProject, t: CapexTemplate, user_id: int | None) -> int:
    orden = len(capex.partidas(db, p.id))
    for n, x in enumerate(t.partidas or [], orden + 1):
        db.add(CapexItem(project_id=p.id, capitulo=str(x.get("capitulo") or "01")[:20],
                         capitulo_nombre=str(x.get("capitulo_nombre") or "General")[:200], codigo=x.get("codigo"),
                         descripcion=x.get("descripcion") or "", unidad=x.get("unidad"),
                         cantidad=x.get("cantidad") or 0, precio=x.get("precio") or 0, orden=n, origen="plantilla"))
    for d in db.scalars(select(CapexDocument).where(CapexDocument.template_id == t.id)):
        _copiar_doc(db, d, user_id, project_id=p.id)
    p.plantilla_id = t.id
    return len(t.partidas or [])


# --------------------------------------------------------------------------- proyectos
class ProjectIn(BaseModel):
    asset_id: int
    titulo: str = Field(min_length=3, max_length=200)
    descripcion: str | None = None
    ubicacion: str | None = Field(default=None, max_length=200)
    unit_id: int | None = None
    presupuesto_objetivo: float | None = Field(default=None, ge=0)
    fecha_inicio: date | None = None
    fecha_fin: date | None = None
    gg_pct: float = Field(default=13, ge=0, le=30)
    bi_pct: float = Field(default=6, ge=0, le=30)
    iva_pct: float = Field(default=21, ge=0, le=21)
    plantilla_id: int | None = None


class ProjectUpdate(BaseModel):
    titulo: str | None = Field(default=None, min_length=3, max_length=200)
    descripcion: str | None = None
    ubicacion: str | None = Field(default=None, max_length=200)
    unit_id: int | None = None
    presupuesto_objetivo: float | None = Field(default=None, ge=0)
    fecha_inicio: date | None = None
    fecha_fin: date | None = None
    gg_pct: float | None = Field(default=None, ge=0, le=30)
    bi_pct: float | None = Field(default=None, ge=0, le=30)
    iva_pct: float | None = Field(default=None, ge=0, le=21)


@router.get("/catalogos")
def catalogs(scope: Scope = Depends(get_scope)):
    return {"estados": ESTADOS_CAPEX, "etiquetas": ETIQUETAS_DOC_CAPEX, "niveles": capex.NIVEL_TXT,
            "colores": capex.COLORES, "umbral": settings.capex_umbral, "aprobador": _aprobador(scope),
            "aprobadores": list(settings.capex_aprobadores)}


@router.get("")
def list_projects(asset_id: int | None = None, estado: str | None = None, scope: Scope = Depends(get_scope),
                  db: Session = Depends(get_db)):
    stmt = scoped(select(CapexProject), CapexProject.asset_id, scope.asset_ids("capex.ver"))
    if asset_id:
        stmt = stmt.where(CapexProject.asset_id == asset_id)
    if estado == "abiertos":
        stmt = stmt.where(CapexProject.estado.not_in(("adjudicado", "cancelado")))
    elif estado:
        stmt = stmt.where(CapexProject.estado == estado)
    return [_resumen(db, p) for p in db.scalars(stmt.order_by(CapexProject.creado.desc()))]


@router.post("", status_code=201)
def create_project(data: ProjectIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    scope.require_asset("capex.gestionar", data.asset_id)
    get_or_404(db, Asset, data.asset_id)
    if data.unit_id and get_or_404(db, Unit, data.unit_id).asset_id != data.asset_id:
        bad_request("La unidad no es de este activo")
    p = CapexProject(**data.model_dump(exclude={"plantilla_id"}), codigo=capex.codigo_nuevo(db),
                     creado_por=scope.user.id)
    db.add(p)
    db.flush()
    if data.plantilla_id:
        _aplicar_plantilla(db, p, get_or_404(db, CapexTemplate, data.plantilla_id), scope.user.id)
    audit(db, scope.user, "crear", "capex", p.id, {"codigo": p.codigo, "plantilla": data.plantilla_id})
    db.commit()
    return _resumen(db, p)


@router.get("/{pid}")
def get_project(pid: int, request: Request, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid)
    items = capex.partidas(db, p.id)
    tenders = []
    for t in db.scalars(select(TenderProcess).where(TenderProcess.project_id == p.id).order_by(TenderProcess.id)):
        invs = []
        for i in db.scalars(select(TenderInvite).where(TenderInvite.tender_id == t.id).order_by(TenderInvite.id)):
            s = db.get(Supplier, i.supplier_id)
            b = db.scalar(select(SupplierBid).where(SupplierBid.invite_id == i.id))
            pem = sum(float(x.cantidad) * (capex.precio_oferta(b, x.id) or 0) for x in items) if b else None
            invs.append({**i.to_dict(exclude=("clave",)), "proveedor": s.nombre, "telefono": s.telefono,
                         "oferta": {"id": b.id, "estado": b.estado, "origen": b.origen, "pem": capex.r2(pem),
                                    "enviada_en": b.enviada_en.isoformat() if b.enviada_en else None,
                                    "faltan": sum(1 for x in items if capex.precio_oferta(b, x.id) is None)}
                         if b else None})
        tenders.append({**t.to_dict(), "invitaciones": invs})
    adj = []
    for x in db.scalars(select(CapexAward).where(CapexAward.project_id == p.id).order_by(CapexAward.id)):
        adj.append({**x.to_dict(), "pem": float(x.pem), "importe": float(x.importe),
                    "proveedor": db.get(Supplier, x.supplier_id).nombre})
    apr = [{**a.to_dict(exclude=("clave", "user_agent")), "importe": float(a.importe)}
           for a in db.scalars(select(CapexApproval).where(CapexApproval.project_id == p.id).order_by(CapexApproval.id))]
    return {"proyecto": _resumen(db, p), "totales": capex.totales(p, items), "partidas": [_item(x) for x in items],
            "documentos": [_doc(d) for d in db.scalars(select(CapexDocument).where(CapexDocument.project_id == p.id)
                                                      .order_by(CapexDocument.etiqueta, CapexDocument.id))],
            "licitaciones": tenders, "adjudicaciones": adj, "aprobaciones": apr,
            "gestionar": scope.can_asset("capex.gestionar", p.asset_id),
            "aprobar": _aprobador(scope) and p.estado == "pendiente_aprobacion"}


@router.put("/{pid}")
def update_project(pid: int, data: ProjectUpdate, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    cambios = data.model_dump(exclude_unset=True)
    if p.estado not in EDITABLE and {"gg_pct", "bi_pct"} & cambios.keys():
        bad_request("Los gastos generales y el beneficio industrial solo se cambian en definición")
    if cambios.get("unit_id") and get_or_404(db, Unit, cambios["unit_id"]).asset_id != p.asset_id:
        bad_request("La unidad no es de este activo")
    for k, v in cambios.items():
        setattr(p, k, v)
    _tocado(p)
    audit(db, scope.user, "editar", "capex", p.id, {k: str(v) for k, v in cambios.items()})
    db.commit()
    return _resumen(db, p)


class EstadoIn(BaseModel):
    estado: str
    motivo: str | None = Field(default=None, max_length=500)


@router.post("/{pid}/estado")
def change_state(pid: int, data: EstadoIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Cancelar, volver a definición (sin ofertas recibidas) o reabrir la licitación."""
    p = _proyecto(db, scope, pid, gestionar=True)
    enviadas = capex.ofertas(db, p.id)
    if data.estado == "cancelado":
        if p.estado == "adjudicado":
            bad_request("Un proyecto adjudicado no se cancela desde aquí")
        if not (data.motivo or "").strip():
            bad_request("Indique el motivo de la cancelación")
    elif data.estado == "definicion":
        if p.estado not in ("licitacion", "cancelado") or enviadas:
            bad_request("Solo vuelve a definición una licitación sin ofertas recibidas (o un proyecto cancelado)")
        for t in db.scalars(select(TenderProcess).where(TenderProcess.project_id == p.id, TenderProcess.estado == "abierta")):
            t.estado = "cerrada"
    elif data.estado == "licitacion":
        if p.estado not in ("adjudicacion",):
            bad_request("Solo se reabre la licitación desde el comparativo")
        for t in db.scalars(select(TenderProcess).where(TenderProcess.project_id == p.id)):
            t.estado = "abierta"
    else:
        bad_request("Cambio de estado no permitido")
    anterior, p.estado = p.estado, data.estado
    _tocado(p)
    audit(db, scope.user, "estado", "capex", p.id, {"de": anterior, "a": p.estado, "motivo": data.motivo})
    db.commit()
    return _resumen(db, p)


@router.delete("/{pid}")
def delete_project(pid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    if p.estado != "definicion" or db.scalar(select(TenderProcess.id).where(TenderProcess.project_id == p.id)):
        bad_request("Solo se borra un proyecto en definición y sin licitaciones: cancélelo")
    ficheros = [d.fichero for d in db.scalars(select(CapexDocument).where(CapexDocument.project_id == p.id))]
    db.execute(delete(CapexDocument).where(CapexDocument.project_id == p.id))
    db.execute(delete(CapexItem).where(CapexItem.project_id == p.id))
    audit(db, scope.user, "borrar", "capex", p.id, {"codigo": p.codigo})
    db.delete(p)
    db.commit()
    for f in ficheros:
        documentos.borrar(f)
    return {"ok": True}


# --------------------------------------------------------------------------- partidas (presupuesto)
class ItemIn(BaseModel):
    capitulo: str = Field(min_length=1, max_length=20)
    capitulo_nombre: str = Field(min_length=1, max_length=200)
    codigo: str | None = Field(default=None, max_length=40)
    descripcion: str = Field(min_length=1)
    unidad: str | None = Field(default=None, max_length=20)
    cantidad: float = Field(default=0, ge=0)
    precio: float = Field(default=0, ge=0)


class ItemUpdate(BaseModel):
    capitulo: str | None = Field(default=None, min_length=1, max_length=20)
    capitulo_nombre: str | None = Field(default=None, min_length=1, max_length=200)
    codigo: str | None = Field(default=None, max_length=40)
    descripcion: str | None = Field(default=None, min_length=1)
    unidad: str | None = Field(default=None, max_length=20)
    cantidad: float | None = Field(default=None, ge=0)
    precio: float | None = Field(default=None, ge=0)
    orden: int | None = None


@router.post("/{pid}/partidas", status_code=201)
def add_item(pid: int, data: ItemIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    _editable(p)
    x = CapexItem(project_id=p.id, **data.model_dump(), orden=len(capex.partidas(db, p.id)) + 1)
    db.add(x)
    _tocado(p)
    db.commit()
    return _item(x)


def _partida(db: Session, scope: Scope, iid: int) -> tuple[CapexItem, CapexProject]:
    x = get_or_404(db, CapexItem, iid)
    p = _proyecto(db, scope, x.project_id, gestionar=True)
    _editable(p)
    return x, p


@router.put("/partidas/{iid}")
def update_item(iid: int, data: ItemUpdate, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x, p = _partida(db, scope, iid)
    for k, v in data.model_dump(exclude_unset=True).items():
        if v is not None:
            setattr(x, k, v)
    _tocado(p)
    db.commit()
    return _item(x)


@router.delete("/partidas/{iid}")
def delete_item(iid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    x, p = _partida(db, scope, iid)
    db.delete(x)
    _tocado(p)
    db.commit()
    return {"ok": True}


class CapituloIn(BaseModel):
    capitulo: str = Field(min_length=1, max_length=20)
    nombre: str = Field(min_length=1, max_length=200)


@router.put("/{pid}/capitulos/{cap}")
def rename_chapter(pid: int, cap: str, data: CapituloIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    _editable(p)
    for x in db.scalars(select(CapexItem).where(CapexItem.project_id == p.id, CapexItem.capitulo == cap)):
        x.capitulo, x.capitulo_nombre = data.capitulo, data.nombre
    _tocado(p)
    db.commit()
    return {"ok": True}


class DesdeBaseIn(BaseModel):
    precio_ids: list[int] = Field(min_length=1, max_length=200)
    capitulo: str = Field(min_length=1, max_length=20)
    capitulo_nombre: str = Field(min_length=1, max_length=200)


@router.post("/{pid}/partidas/desde-base", status_code=201)
def add_from_base(pid: int, data: DesdeBaseIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Añade partidas de la base de precios a un capítulo (cantidad 0: se mide después)."""
    p = _proyecto(db, scope, pid, gestionar=True)
    _editable(p)
    orden = len(capex.partidas(db, p.id))
    nuevas = []
    for n, pi in enumerate(db.scalars(select(PriceItem).where(PriceItem.id.in_(data.precio_ids))), orden + 1):
        x = CapexItem(project_id=p.id, capitulo=data.capitulo, capitulo_nombre=data.capitulo_nombre, codigo=pi.codigo,
                      descripcion=pi.resumen + (f"\n{pi.texto}" if pi.texto and pi.texto != pi.resumen else ""),
                      unidad=pi.unidad, cantidad=0, precio=pi.precio, orden=n, origen="base")
        db.add(x)
        nuevas.append(x)
    _tocado(p)
    db.commit()
    return [_item(x) for x in nuevas]


def _leer_bc3(f: UploadFile) -> tuple[bytes, str]:
    datos = f.file.read(TAM_BC3 + 1)
    if len(datos) > TAM_BC3:
        bad_request("El fichero supera los 120 MB")
    nombre = f.filename or "base.bc3"
    if datos[:2] == b"PK":  # las bases públicas se descargan comprimidas
        try:
            z = zipfile.ZipFile(io.BytesIO(datos))
            dentro = [n for n in z.namelist() if n.lower().endswith(".bc3")]
            if not dentro:
                bad_request("El ZIP no contiene ningún fichero .bc3")
            nombre = max(dentro, key=lambda n: z.getinfo(n).file_size)
            if z.getinfo(nombre).file_size > TAM_BC3:
                bad_request("El .bc3 del ZIP supera los 120 MB")
            datos = z.read(nombre)
        except zipfile.BadZipFile:
            bad_request("El ZIP está dañado")
    return datos, Path(nombre).name


@router.post("/{pid}/importar-bc3")
def import_bc3(pid: int, fichero: UploadFile = File(...), reemplazar: bool = Form(False),
               scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Presupuesto con mediciones en FIEBDC-3 (Presto, Arquímedes, Menfis…): capítulos, partidas, cantidades y precios."""
    p = _proyecto(db, scope, pid, gestionar=True)
    _editable(p)
    datos, nombre = _leer_bc3(fichero)
    try:
        filas = bc3.presupuesto(bc3.leer(datos))
    except ValueError as e:
        bad_request(str(e))
    filas = [f for f in filas if f["cantidad"]]
    if not filas:
        bad_request("El fichero no trae mediciones: si es una base de precios, impórtela en «Base de precios»")
    if reemplazar:
        db.execute(delete(CapexItem).where(CapexItem.project_id == p.id))
    orden = 0 if reemplazar else len(capex.partidas(db, p.id))
    for n, f in enumerate(filas, orden + 1):
        db.add(CapexItem(project_id=p.id, capitulo=f["capitulo"][:20], capitulo_nombre=f["capitulo_nombre"][:200],
                         codigo=(f["codigo"] or "")[:40], descripcion=f["descripcion"], unidad=(f["unidad"] or "")[:20],
                         cantidad=f["cantidad"], precio=f["precio"], orden=n, origen="bc3"))
    _tocado(p)
    audit(db, scope.user, "importar_bc3", "capex", p.id, {"fichero": nombre, "partidas": len(filas)})
    db.commit()
    return {"partidas": len(filas), "capitulos": len({f["capitulo"] for f in filas})}


@router.get("/{pid}/presupuesto.xlsx")
def budget_excel(pid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    p = _proyecto(db, scope, pid)
    items = capex.partidas(db, p.id)
    t = capex.totales(p, items)
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws.append([f"{p.codigo} · {p.titulo} · {db.get(Asset, p.asset_id).nombre}"])
    ws.append([])
    ws.append(["Capítulo", "Código", "Descripción", "Ud", "Cantidad", "Precio €", "Importe €"])
    for c in ws[3]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
    for cap in t["capitulos"]:
        ws.append([f"{cap['capitulo']} {cap['nombre']}", "", "", "", "", "", cap["importe"]])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for x in items:
            if x.capitulo == cap["capitulo"]:
                ws.append(["", x.codigo, x.descripcion, x.unidad, float(x.cantidad), float(x.precio),
                           capex.r2(float(x.cantidad) * float(x.precio))])
    ws.append([])
    for k, v in (("Presupuesto de ejecución material", t["pem"]), (f"Gastos generales {float(p.gg_pct):g} %", t["gg"]),
                 (f"Beneficio industrial {float(p.bi_pct):g} %", t["bi"]), ("Presupuesto base de licitación", t["pbl"]),
                 (f"IVA {float(p.iva_pct):g} %", t["iva"]), ("Total con IVA", t["total"])):
        ws.append(["", "", k, "", "", "", v])
    for col, ancho in zip("ABCDEFG", (28, 12, 70, 6, 10, 12, 14)):
        ws.column_dimensions[col].width = ancho
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Presupuesto_{p.codigo}.xlsx"'})


# --------------------------------------------------------------------------- documentos
def _documento(db: Session, scope: Scope, did: int, gestionar: bool = False) -> CapexDocument:
    d = get_or_404(db, CapexDocument, did)
    if d.project_id:
        _proyecto(db, scope, d.project_id, gestionar=gestionar)
    elif not scope.has_any("capex.gestionar" if gestionar else "capex.ver"):
        raise HTTPException(403, "Sin permiso sobre CAPEX")
    return d


@router.post("/{pid}/documentos", status_code=201)
def upload_docs(pid: int, ficheros: list[UploadFile] = File(...), etiqueta: str = Form(...),
                notas: str | None = Form(None), licitacion: bool = Form(True), scope: Scope = Depends(get_scope),
                db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    if etiqueta not in ETIQUETAS_DOC_CAPEX:
        bad_request("Etiqueta no válida")
    if len(ficheros) > 30:
        bad_request("Máximo 30 ficheros a la vez")
    nuevos = [_guardar_doc(db, *_fichero_doc(f), scope.user.id, project_id=p.id, etiqueta=etiqueta,
                           notas=(notas or "").strip()[:300] or None, licitacion=licitacion) for f in ficheros]
    db.flush()
    audit(db, scope.user, "subir_documentos", "capex", p.id, {"etiqueta": etiqueta, "n": len(nuevos)})
    db.commit()
    return [_doc(d) for d in nuevos]


@router.get("/documentos/{did}")
def view_doc(did: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    d = _documento(db, scope, did)
    return Response(documentos.leer(d.fichero), media_type=d.mime, headers={
        "Content-Disposition": f'inline; filename="{_ascii(d.nombre)}"', "Cache-Control": "private, no-store"})


class DocUpdate(BaseModel):
    etiqueta: str | None = None
    notas: str | None = Field(default=None, max_length=300)
    licitacion: bool | None = None


@router.put("/documentos/{did}")
def update_doc(did: int, data: DocUpdate, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    d = _documento(db, scope, did, gestionar=True)
    if data.etiqueta is not None and data.etiqueta not in ETIQUETAS_DOC_CAPEX:
        bad_request("Etiqueta no válida")
    for k, v in data.model_dump(exclude_unset=True).items():
        if v is not None:
            setattr(d, k, v)
    db.commit()
    return _doc(d)


@router.delete("/documentos/{did}")
def delete_doc(did: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    d = _documento(db, scope, did, gestionar=True)
    f = d.fichero
    audit(db, scope.user, "borrar_documento", "capex", d.project_id or 0, {"documento": did, "nombre": d.nombre})
    db.delete(d)
    db.commit()
    documentos.borrar(f)
    return {"ok": True}


def _ascii(texto: str) -> str:
    from ..empleados import nombre_fichero
    return nombre_fichero(texto)


# --------------------------------------------------------------------------- plantillas «CAPEX tipo»
class TemplateIn(BaseModel):
    nombre: str = Field(min_length=3, max_length=160)
    descripcion: str | None = None
    project_id: int | None = None  # se crea a partir de un proyecto (partidas y documentos de plano y pliego)
    con_cantidades: bool = True


class TemplateUpdate(BaseModel):
    nombre: str | None = Field(default=None, min_length=3, max_length=160)
    descripcion: str | None = None
    activa: bool | None = None
    partidas: list[dict] | None = None


def _plantilla(db: Session, t: CapexTemplate) -> dict:
    docs = [_doc(d) for d in db.scalars(select(CapexDocument).where(CapexDocument.template_id == t.id))]
    pem = sum(float(x.get("cantidad") or 0) * float(x.get("precio") or 0) for x in t.partidas or [])
    return {**t.to_dict(), "n_partidas": len(t.partidas or []), "pem": capex.r2(pem), "documentos": docs,
            "capitulos": sorted({f"{x.get('capitulo')} {x.get('capitulo_nombre')}" for x in t.partidas or []})}


def _gestor(scope: Scope) -> None:
    if not scope.has_any("capex.gestionar"):
        raise HTTPException(403, "Las plantillas y bases de precios las gestiona la Dirección Técnica")


@router.get("/plantillas/lista")
def list_templates(scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    if not scope.has_any("capex.ver"):
        raise HTTPException(403, "Sin permiso sobre CAPEX")
    return [_plantilla(db, t) for t in db.scalars(select(CapexTemplate).order_by(CapexTemplate.nombre))]


@router.post("/plantillas/lista", status_code=201)
def create_template(data: TemplateIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _gestor(scope)
    if db.scalar(select(CapexTemplate.id).where(CapexTemplate.nombre == data.nombre.strip())):
        bad_request("Ya hay una plantilla con ese nombre")
    t = CapexTemplate(nombre=data.nombre.strip(), descripcion=data.descripcion, creado_por=scope.user.id, partidas=[])
    db.add(t)
    db.flush()
    if data.project_id:
        p = _proyecto(db, scope, data.project_id)
        t.partidas = [{"capitulo": x.capitulo, "capitulo_nombre": x.capitulo_nombre, "codigo": x.codigo,
                       "descripcion": x.descripcion, "unidad": x.unidad,
                       "cantidad": float(x.cantidad) if data.con_cantidades else 0, "precio": float(x.precio)}
                      for x in capex.partidas(db, p.id)]
        for d in db.scalars(select(CapexDocument).where(CapexDocument.project_id == p.id,
                                                        CapexDocument.etiqueta.in_(("plano", "pliego", "memoria")))):
            _copiar_doc(db, d, scope.user.id, template_id=t.id)
    audit(db, scope.user, "crear_plantilla", "capex", t.id, {"nombre": t.nombre, "proyecto": data.project_id})
    db.commit()
    return _plantilla(db, t)


@router.put("/plantillas/lista/{tid}")
def update_template(tid: int, data: TemplateUpdate, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _gestor(scope)
    t = get_or_404(db, CapexTemplate, tid)
    for k, v in data.model_dump(exclude_unset=True).items():
        if v is not None:
            setattr(t, k, v)
    db.commit()
    return _plantilla(db, t)


@router.delete("/plantillas/lista/{tid}")
def delete_template(tid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _gestor(scope)
    t = get_or_404(db, CapexTemplate, tid)
    for p in db.scalars(select(CapexProject).where(CapexProject.plantilla_id == t.id)):
        p.plantilla_id = None
    ficheros = [d.fichero for d in db.scalars(select(CapexDocument).where(CapexDocument.template_id == t.id))]
    db.execute(delete(CapexDocument).where(CapexDocument.template_id == t.id))
    db.delete(t)
    db.commit()
    for f in ficheros:
        documentos.borrar(f)
    return {"ok": True}


@router.post("/plantillas/lista/{tid}/documentos", status_code=201)
def upload_template_docs(tid: int, ficheros: list[UploadFile] = File(...), etiqueta: str = Form(...),
                         scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _gestor(scope)
    t = get_or_404(db, CapexTemplate, tid)
    if etiqueta not in ETIQUETAS_DOC_CAPEX:
        bad_request("Etiqueta no válida")
    nuevos = [_guardar_doc(db, *_fichero_doc(f), scope.user.id, template_id=t.id, etiqueta=etiqueta) for f in ficheros]
    db.commit()
    return [_doc(d) for d in nuevos]


class AplicarIn(BaseModel):
    plantilla_id: int


@router.post("/{pid}/plantilla")
def apply_template(pid: int, data: AplicarIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    _editable(p)
    n = _aplicar_plantilla(db, p, get_or_404(db, CapexTemplate, data.plantilla_id), scope.user.id)
    _tocado(p)
    db.commit()
    return {"partidas": n}


# --------------------------------------------------------------------------- base de precios (FIEBDC-3)
@router.get("/bases/lista")
def list_bases(scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    if not scope.has_any("capex.ver"):
        raise HTTPException(403, "Sin permiso sobre CAPEX")
    return [b.to_dict() for b in db.scalars(select(PriceBase).order_by(PriceBase.importado.desc()))]


@router.post("/bases/lista", status_code=201)
def import_base(fichero: UploadFile = File(...), nombre: str = Form(...), scope: Scope = Depends(get_scope),
                db: Session = Depends(get_db)):
    _gestor(scope)
    datos, archivo = _leer_bc3(fichero)
    try:
        filas = bc3.partidas_base(bc3.leer(datos))
    except ValueError as e:
        bad_request(str(e))
    if not filas:
        bad_request("No se han encontrado partidas en el fichero")
    b = PriceBase(nombre=nombre.strip()[:160] or archivo, fichero=archivo[:160], partidas=len(filas),
                  user_id=scope.user.id)
    db.add(b)
    db.flush()
    for i in range(0, len(filas), 2000):
        db.execute(insert(PriceItem), [{"base_id": b.id, "codigo": f["codigo"][:40], "unidad": (f["unidad"] or "")[:20],
                                        "resumen": f["resumen"][:300], "texto": f["texto"] or None,
                                        "precio": f["precio"], "capitulo": (f["capitulo"] or "")[:300] or None}
                                       for f in filas[i:i + 2000]])
    audit(db, scope.user, "importar_base_precios", "capex", b.id, {"fichero": archivo, "partidas": len(filas)})
    db.commit()
    return b.to_dict()


@router.delete("/bases/lista/{bid}")
def delete_base(bid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    _gestor(scope)
    b = get_or_404(db, PriceBase, bid)
    db.execute(delete(PriceItem).where(PriceItem.base_id == b.id))
    db.delete(b)
    db.commit()
    return {"ok": True}


@router.get("/precios/buscar")
def search_prices(q: str = "", base_id: int | None = None, limite: int = 60, scope: Scope = Depends(get_scope),
                  db: Session = Depends(get_db)):
    """Busca por código o por palabras de la descripción (todas deben aparecer)."""
    if not scope.has_any("capex.ver"):
        raise HTTPException(403, "Sin permiso sobre CAPEX")
    stmt = select(PriceItem)
    if base_id:
        stmt = stmt.where(PriceItem.base_id == base_id)
    for palabra in q.split()[:6]:
        like = f"%{palabra}%"
        stmt = stmt.where(or_(PriceItem.resumen.ilike(like), PriceItem.codigo.ilike(like), PriceItem.capitulo.ilike(like)))
    return [{**x.to_dict(), "precio": float(x.precio)}
            for x in db.scalars(stmt.order_by(PriceItem.codigo).limit(max(1, min(limite, 200))))]


# --------------------------------------------------------------------------- licitación ciega
class TenderIn(BaseModel):
    proveedor_ids: list[int] = Field(min_length=1, max_length=30)
    fecha_limite: date
    mensaje: str | None = Field(default=None, max_length=3000)
    documento_ids: list[int] | None = None  # None = los marcados para licitación


def _caducada(t: TenderProcess) -> bool:
    return datetime.now() > datetime.combine(t.fecha_limite, time(23, 59, 59))


def _enlace(base: str, inv: TenderInvite) -> str:
    return f"{base}/oferta/{inv.id}?k={inv.clave}"


def _enviar_invitacion(db: Session, p: CapexProject, t: TenderProcess, inv: TenderInvite, base: str,
                       docs: list[CapexDocument]) -> dict:
    s = db.get(Supplier, inv.supplier_id)
    a = db.get(Asset, p.asset_id)
    url = _enlace(base, inv)
    asunto = f"Petición de oferta {p.codigo}: {p.titulo} – {a.nombre}"
    texto = (f"Estimados señores:\n\nDesde el Grupo INVERSIETE les invitamos a presentar oferta para los trabajos "
             f"«{p.titulo}» en {a.nombre}{f' ({p.ubicacion})' if p.ubicacion else ''}.\n\n"
             f"{(t.mensaje.strip() + chr(10) + chr(10)) if t.mensaje else ''}"
             f"Pueden consultar las mediciones, descargar la documentación y escribir sus precios en este enlace "
             f"personal: {url}\nTambién pueden rellenar la plantilla Excel adjunta y devolverla por el mismo enlace.\n\n"
             f"Plazo de presentación: hasta el {t.fecha_limite:%d/%m/%Y} (incluido).\n"
             "Los precios unitarios se entienden de ejecución material, sin gastos generales, beneficio industrial "
             "ni IVA.\n\nQuedamos a su disposición para cualquier aclaración.\n\nAtentamente,\nDirección Técnica · "
             "Grupo INVERSIETE")
    adj = [(f"Oferta_{p.codigo}_{_ascii(s.nombre)}.xlsx", capex.excel_ciego(db, p, inv),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]
    peso = 0
    for d in docs:
        if peso + d.tamano > ADJUNTOS_MAX:
            texto += "\n\n(Parte de la documentación pesa demasiado para el correo: descárguela desde el enlace.)"
            break
        adj.append((d.nombre, documentos.leer(d.fichero), d.mime))
        peso += d.tamano
    html = avisos._html(asunto, "".join(f"<p>{escape(par)}</p>" for par in texto.split("\n\n")).replace(
        escape(url), f'<a href="{escape(url)}" style="font-weight:600">{escape(url)}</a>'))
    r = {"proveedor": s.nombre, "email": inv.email, "ok": False, "whatsapp": None}
    if inv.email and avisos.configurado():
        try:
            avisos.enviar(inv.email, asunto, texto, html, adj)
            inv.enviado, inv.envio_error, r["ok"] = datetime.now(), None, True
        except Exception as e:  # noqa: BLE001
            inv.envio_error = str(e)[:300]
            r["error"] = inv.envio_error
        db.add(EmailLog(clave=f"capex_licitacion:{inv.id}", tipo="capex", destinatario=inv.email,
                        asunto=asunto[:200], ok=r["ok"], error=inv.envio_error))
    else:
        r["error"] = "Sin correo del proveedor" if not inv.email else "Correo saliente no configurado"
    from ..firma_contrato import movil_whatsapp
    movil = movil_whatsapp(s.telefono or "") if s.telefono else None
    if movil:
        r["whatsapp"] = "https://wa.me/{}?text={}".format(movil, quote(
            f"Buenos días, desde el Grupo INVERSIETE les invitamos a ofertar «{p.titulo}» en {a.nombre}. "
            f"Mediciones, documentación y oferta en: {url} . Plazo: {t.fecha_limite:%d/%m/%Y}."))
    return r


def _docs_licitacion(db: Session, p: CapexProject, ids: list[int] | None) -> list[CapexDocument]:
    stmt = select(CapexDocument).where(CapexDocument.project_id == p.id)
    stmt = stmt.where(CapexDocument.id.in_(ids or [-1])) if ids is not None else stmt.where(CapexDocument.licitacion)
    return list(db.scalars(stmt.order_by(CapexDocument.etiqueta, CapexDocument.id)))


@router.post("/{pid}/licitacion", status_code=201)
def create_tender(pid: int, data: TenderIn, request: Request, scope: Scope = Depends(get_scope),
                  db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    if p.estado not in ("definicion", "licitacion"):
        bad_request("Solo se licita un proyecto en definición o en licitación")
    items = capex.partidas(db, p.id)
    if not items:
        bad_request("El proyecto no tiene partidas")
    if any(not float(x.cantidad) for x in items):
        bad_request("Hay partidas sin medición (cantidad 0): mídalas o bórrelas antes de licitar")
    if data.fecha_limite < date.today():
        bad_request("El plazo ya ha pasado")
    t = db.scalar(select(TenderProcess).where(TenderProcess.project_id == p.id, TenderProcess.estado == "abierta"))
    if t is None:
        t = TenderProcess(project_id=p.id, fecha_limite=data.fecha_limite, mensaje=data.mensaje,
                          creado_por=scope.user.id)
        db.add(t)
        db.flush()
    else:
        t.fecha_limite, t.mensaje = data.fecha_limite, data.mensaje or t.mensaje
    ya = {i.supplier_id for i in db.scalars(select(TenderInvite).where(TenderInvite.tender_id == t.id))}
    docs = _docs_licitacion(db, p, data.documento_ids)
    base = _base_url(request)
    res = []
    for sid in dict.fromkeys(data.proveedor_ids):
        s = get_or_404(db, Supplier, sid)
        if sid in ya:
            continue
        inv = TenderInvite(tender_id=t.id, supplier_id=s.id, email=(s.email or "").strip() or None,
                           clave=secrets.token_urlsafe(24))
        db.add(inv)
        db.flush()
        res.append(_enviar_invitacion(db, p, t, inv, base, docs))
    p.estado = "licitacion"
    _tocado(p)
    audit(db, scope.user, "licitar", "capex", p.id, {"licitacion": t.id, "proveedores": data.proveedor_ids,
                                                       "fecha_limite": t.fecha_limite.isoformat()})
    db.commit()
    return {"licitacion": t.id, "enviados": res}


@router.post("/invitaciones/{iid}/reenviar")
def resend_invite(iid: int, request: Request, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    inv = get_or_404(db, TenderInvite, iid)
    t = db.get(TenderProcess, inv.tender_id)
    p = _proyecto(db, scope, t.project_id, gestionar=True)
    s = db.get(Supplier, inv.supplier_id)
    inv.email = (s.email or "").strip() or inv.email
    r = _enviar_invitacion(db, p, t, inv, _base_url(request), _docs_licitacion(db, p, None))
    db.commit()
    return r


@router.get("/invitaciones/{iid}/excel")
def invite_excel(iid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    inv = get_or_404(db, TenderInvite, iid)
    p = _proyecto(db, scope, db.get(TenderProcess, inv.tender_id).project_id)
    s = db.get(Supplier, inv.supplier_id)
    return Response(capex.excel_ciego(db, p, inv),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Oferta_{p.codigo}_{_ascii(s.nombre)}.xlsx"'})


def _guardar_oferta(db: Session, inv: TenderInvite, precios: dict, origen: str, enviar: bool,
                    notas: str | None = None, plazo: int | None = None) -> SupplierBid:
    t = db.get(TenderProcess, inv.tender_id)
    validos = {str(x.id) for x in capex.partidas(db, t.project_id)}
    limpios = {}
    for k, v in (precios or {}).items():
        if str(k) not in validos or v in (None, ""):
            continue
        try:
            v = round(float(v), 2)
        except (TypeError, ValueError):
            bad_request(f"Precio no válido: «{v}»")
        if v < 0:
            bad_request("No se admiten precios negativos")
        limpios[str(k)] = v
    b = db.scalar(select(SupplierBid).where(SupplierBid.invite_id == inv.id))
    if b is None:
        b = SupplierBid(tender_id=t.id, invite_id=inv.id, supplier_id=inv.supplier_id)
        db.add(b)
    b.precios, b.origen, b.actualizado = limpios, origen, datetime.now()
    if notas is not None:
        b.notas = notas.strip()[:3000] or None
    if plazo is not None:
        b.plazo_dias = plazo
    if enviar:
        if not limpios:
            bad_request("La oferta no tiene ningún precio")
        b.estado, b.enviada_en = "enviada", datetime.now()
    db.flush()
    return b


@router.post("/invitaciones/{iid}/oferta-excel")
def upload_bid_excel(iid: int, fichero: UploadFile = File(...), scope: Scope = Depends(get_scope),
                     db: Session = Depends(get_db)):
    """La Dirección Técnica sube la plantilla que el contratista devolvió por correo."""
    inv = get_or_404(db, TenderInvite, iid)
    t = db.get(TenderProcess, inv.tender_id)
    p = _proyecto(db, scope, t.project_id, gestionar=True)
    if p.estado not in ("licitacion", "adjudicacion"):
        bad_request("La licitación ya está resuelta")
    try:
        precios = capex.leer_excel(db, p, inv, fichero.file.read(10 * 1024 * 1024))
    except ValueError as e:
        bad_request(str(e))
    b = _guardar_oferta(db, inv, precios, "recibida", True)
    audit(db, scope.user, "oferta_excel", "capex", p.id, {"invitacion": inv.id, "precios": len(precios)})
    db.commit()
    return {"id": b.id, "precios": len(precios)}


@router.post("/licitacion/{tid}/cerrar")
def close_tender(tid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Cierra la recepción de ofertas y pasa al comparativo."""
    t = get_or_404(db, TenderProcess, tid)
    p = _proyecto(db, scope, t.project_id, gestionar=True)
    if not capex.ofertas(db, p.id):
        bad_request("No hay ninguna oferta enviada todavía")
    t.estado = "cerrada"
    p.estado = "adjudicacion"
    _tocado(p)
    audit(db, scope.user, "cerrar_licitacion", "capex", p.id, {"licitacion": t.id})
    db.commit()
    return _resumen(db, p)


# --------------------------------------------------------------------------- comparativo y adjudicación
@router.get("/{pid}/comparativo")
def comparison(pid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    return capex.comparativo(db, _proyecto(db, scope, pid))


@router.get("/{pid}/comparativo.pdf")
def comparison_pdf(pid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid)
    pdf = capex.comparativo_pdf(db, p, capex.comparativo(db, p))
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="Comparativo_{p.codigo}.pdf"'})


class LoteIn(BaseModel):
    bid_id: int
    capitulos: list[str] = Field(default_factory=list)  # vacío = obra completa


class AwardIn(BaseModel):
    lotes: list[LoteIn] = Field(min_length=1, max_length=30)
    nota: str | None = Field(default=None, max_length=500)


@router.post("/{pid}/adjudicar")
def award(pid: int, data: AwardIn, request: Request, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    if p.estado not in ("licitacion", "adjudicacion"):
        bad_request("Solo se adjudica desde el comparativo (licitación con ofertas recibidas)")
    cmp = capex.comparativo(db, p)
    bids = {o["bid_id"]: o for o in cmp["ofertas"]}
    todos = [c["capitulo"] for c in cmp["capitulos"]]
    usados: set[str] = set()
    nuevos = []
    for lt in data.lotes:
        if lt.bid_id not in bids:
            bad_request("Esa oferta no es de este proyecto o no está enviada")
        caps = lt.capitulos or todos
        if lt.capitulos and len(data.lotes) == 1 and set(caps) == set(todos):
            caps = todos
        if not lt.capitulos and len(data.lotes) > 1:
            bad_request("Con varias contratas indique los capítulos de cada una")
        malos = set(caps) - set(todos)
        if malos:
            bad_request(f"Capítulos que no existen: {', '.join(sorted(malos))}")
        repetidos = usados & set(caps)
        if repetidos:
            bad_request(f"Capítulos adjudicados dos veces: {', '.join(sorted(repetidos))}")
        usados |= set(caps)
        incompletos = [c["capitulo"] for c in cmp["capitulos"] if c["capitulo"] in caps
                       and not c["ofertas"][lt.bid_id]["completo"]]
        if incompletos:
            bad_request(f"La oferta de {bids[lt.bid_id]['proveedor']} no tiene precio en todas las partidas de los "
                        f"capítulos {', '.join(incompletos)}")
        t = capex.lote(p, cmp, lt.bid_id, caps)
        nuevos.append((lt.bid_id, caps, t))
    for x in db.scalars(select(CapexAward).where(CapexAward.project_id == p.id, CapexAward.estado.in_(("propuesta", "rechazada")))):
        x.estado = "anulada"
    for a in db.scalars(select(CapexApproval).where(CapexApproval.project_id == p.id, CapexApproval.estado == "pendiente")):
        a.estado = "anulada"
    awards = []
    for bid_id, caps, t in nuevos:
        x = CapexAward(project_id=p.id, bid_id=bid_id, supplier_id=bids[bid_id]["supplier_id"],
                       capitulos=[] if caps == todos and len(nuevos) == 1 else caps, pem=t["pem"], importe=t["pbl"],
                       nota=data.nota, creado_por=scope.user.id)
        db.add(x)
        awards.append(x)
    for t in db.scalars(select(TenderProcess).where(TenderProcess.project_id == p.id)):
        t.estado = "cerrada"
    total = capex.r2(sum(t["pbl"] for _, _, t in nuevos))
    sin = [c for c in todos if c not in usados]
    db.flush()
    envios = []
    if total > settings.capex_umbral:
        p.estado = "pendiente_aprobacion"
        envios = _pedir_aprobacion(db, p, cmp, awards, total, sin, _base_url(request), scope.user)
    else:
        p.estado = "adjudicado"
        for x in awards:
            x.estado = "aprobada"
    _tocado(p)
    audit(db, scope.user, "adjudicar", "capex", p.id, {"total": total, "lotes": [(b, c) for b, c, _ in nuevos],
                                                         "sin_adjudicar": sin})
    db.commit()
    return {"proyecto": _resumen(db, p), "total": total, "requiere_aprobacion": total > settings.capex_umbral,
            "sin_adjudicar": sin, "aprobacion": envios}


def _pedir_aprobacion(db: Session, p: CapexProject, cmp: dict, awards: list[CapexAward], total: float,
                      sin: list[str], base: str, autor: User) -> list[dict]:
    a = db.get(Asset, p.asset_id)
    nombres = {c["capitulo"]: f"{c['capitulo']} {c['nombre']}" for c in cmp["capitulos"]}
    lotes = [{"proveedor": db.get(Supplier, x.supplier_id).nombre,
              "capitulos": [nombres[c] for c in x.capitulos] or ["Obra completa"], "importe": float(x.importe)}
             for x in awards]
    resumen = {"codigo": p.codigo, "titulo": p.titulo, "activo": a.nombre, "ubicacion": p.ubicacion,
               "pbl": cmp["base"]["pbl"], "objetivo": float(p.presupuesto_objetivo) if p.presupuesto_objetivo else None,
               "total": total, "iva_pct": float(p.iva_pct), "lotes": lotes, "sin_adjudicar": [nombres[c] for c in sin],
               "ofertas": [{"proveedor": o["proveedor"], "pbl": o["pbl"], "desviacion": o["desviacion"],
                            "completa": o["completa"], "plazo_dias": o["plazo_dias"]} for o in cmp["ofertas"]],
               "propone": autor.nombre, "nota": awards[0].nota if awards else None}
    pdf = capex.comparativo_pdf(db, p, cmp)
    out = []
    for email in settings.capex_aprobadores:
        ap = CapexApproval(project_id=p.id, email=email, clave=secrets.token_urlsafe(24), importe=total, resumen=resumen)
        db.add(ap)
        db.flush()
        out.append(_enviar_aprobacion(db, ap, pdf, base))
    return out


def _eur(v) -> str:
    return "—" if v is None else f"{v:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _enviar_aprobacion(db: Session, ap: CapexApproval, pdf: bytes, base: str) -> dict:
    r = ap.resumen
    url = f"{base}/capex/aprobacion/{ap.id}?k={ap.clave}"
    asunto = f"Aprobación de adjudicación {r['codigo']}: {r['titulo']} – {_eur(r['total'])}"
    filas = "".join(f"<tr><td style='padding:4px 8px'>{escape(x['proveedor'])}</td><td style='padding:4px 8px'>"
                    f"{escape(', '.join(x['capitulos']))}</td><td style='padding:4px 8px;text-align:right'>"
                    f"{_eur(x['importe'])}</td></tr>" for x in r["lotes"])
    ofs = ""
    for o in r["ofertas"]:
        desv = "" if o["desviacion"] is None else f" ({o['desviacion']:+.1f} % sobre el base)"
        ofs += f"<li>{escape(o['proveedor'])}: {_eur(o['pbl'])}{desv}{'' if o['completa'] else ' · incompleta'}</li>"
    boton = ("display:inline-block;padding:12px 26px;border-radius:8px;color:#fff;text-decoration:none;"
             "font-weight:700;font-size:15px;margin-right:12px")
    lugar = escape(r["activo"]) + (f", {escape(r['ubicacion'])}" if r.get("ubicacion") else "")
    objetivo = f" · objetivo {_eur(r['objetivo'])}" if r.get("objetivo") else ""
    sin = f"<p>Sin adjudicar: {escape(', '.join(r['sin_adjudicar']))}</p>" if r["sin_adjudicar"] else ""
    nota = f"<p>Nota: {escape(r['nota'])}</p>" if r.get("nota") else ""
    cuerpo = (f"<p>La Dirección Técnica propone adjudicar <b>{escape(r['titulo'])}</b> ({lugar}) por "
              f"<b>{_eur(r['total'])}</b> sin IVA, por encima del umbral de {_eur(settings.capex_umbral)} que requiere "
              f"su aprobación.</p>"
              f"<table style='border-collapse:collapse;font-size:13px;width:100%'><tr style='background:#f7f3ea'>"
              f"<th style='text-align:left;padding:4px 8px'>Contratista</th><th style='text-align:left;padding:4px 8px'>"
              f"Alcance</th><th style='text-align:right;padding:4px 8px'>Importe</th></tr>{filas}</table>"
              f"<p>Presupuesto base de licitación: <b>{_eur(r['pbl'])}</b>{objetivo}.</p>"
              f"<p>Ofertas recibidas:</p><ul>{ofs}</ul>{sin}{nota}"
              f"<p>Se adjunta el cuadro comparativo en PDF.</p>"
              f"<p style='margin:22px 0'><a href='{escape(url)}&amp;d=aprobar' style='{boton};background:#2e7d4f'>"
              f"Aprobar</a><a href='{escape(url)}&amp;d=rechazar' style='{boton};background:#b3261e'>Rechazar</a></p>"
              f"<p style='color:#6f6a61;font-size:12px'>Al pulsar se abre una página para confirmar la decisión (y "
              f"añadir un comentario si lo desea), sin entrar en el PMS. El enlace es personal: no lo reenvíe.</p>")
    texto = (f"{asunto}\n\nPropone: {r['propone']}. Importe sin IVA: {_eur(r['total'])}. "
             f"PBL: {_eur(r['pbl'])}.\nPara aprobar o rechazar: {url}")
    res = {"email": ap.email, "ok": False}
    if avisos.configurado():
        try:
            avisos.enviar(ap.email, asunto, texto, avisos._html("Aprobación de adjudicación CAPEX", cuerpo),
                          [(f"Comparativo_{r['codigo']}.pdf", pdf, "application/pdf")], urgente=True)
            ap.enviada_en, ap.envio_error, res["ok"] = datetime.now(), None, True
        except Exception as e:  # noqa: BLE001
            ap.envio_error = res["error"] = str(e)[:300]
        db.add(EmailLog(clave=f"capex_aprobacion:{ap.id}", tipo="capex", destinatario=ap.email, asunto=asunto[:200],
                        ok=res["ok"], error=ap.envio_error))
    else:
        res["error"] = ap.envio_error = "Correo saliente no configurado"
    return res


@router.post("/{pid}/aprobacion/reenviar")
def resend_approval(pid: int, request: Request, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _proyecto(db, scope, pid, gestionar=True)
    pend = list(db.scalars(select(CapexApproval).where(CapexApproval.project_id == p.id,
                                                       CapexApproval.estado == "pendiente")))
    if not pend:
        bad_request("No hay ninguna aprobación pendiente")
    pdf = capex.comparativo_pdf(db, p, capex.comparativo(db, p))
    res = [_enviar_aprobacion(db, ap, pdf, _base_url(request)) for ap in pend]
    db.commit()
    return res


class DecisionIn(BaseModel):
    decision: str = Field(pattern="^(aprobar|rechazar)$")
    comentario: str | None = Field(default=None, max_length=1000)


def _decidir(db: Session, p: CapexProject, ap: CapexApproval, data: DecisionIn, canal: str, ip: str | None,
             agente: str | None) -> None:
    if p.estado != "pendiente_aprobacion" or ap.estado != "pendiente":
        bad_request("Esta aprobación ya está resuelta")
    if data.decision == "rechazar" and not (data.comentario or "").strip():
        bad_request("Indique el motivo del rechazo")
    ahora = datetime.now()
    ap.estado = "aprobada" if data.decision == "aprobar" else "rechazada"
    ap.decidido_en, ap.comentario, ap.canal, ap.ip, ap.user_agent = (
        ahora, (data.comentario or "").strip() or None, canal, ip, (agente or "")[:300] or None)
    for otra in db.scalars(select(CapexApproval).where(CapexApproval.project_id == p.id,
                                                       CapexApproval.estado == "pendiente", CapexApproval.id != ap.id)):
        otra.estado = "anulada"  # basta con una decisión
    for x in db.scalars(select(CapexAward).where(CapexAward.project_id == p.id, CapexAward.estado == "propuesta")):
        x.estado = "aprobada" if ap.estado == "aprobada" else "rechazada"
    p.estado = "adjudicado" if ap.estado == "aprobada" else "adjudicacion"
    _tocado(p)
    audit(db, None, f"capex_{ap.estado}", "capex", p.id, {"aprobacion": ap.id, "email": ap.email, "canal": canal,
                                                            "comentario": ap.comentario, "ip": ip})
    db.commit()
    _avisar_decision(db, p, ap)


def _avisar_decision(db: Session, p: CapexProject, ap: CapexApproval) -> None:
    u = db.get(User, p.creado_por) if p.creado_por else None
    destinos = {e for e in [u.email if u else None, settings.mantenimiento_email] if e}
    if not destinos or not avisos.configurado():
        return
    que = "APROBADA" if ap.estado == "aprobada" else "RECHAZADA"
    asunto = f"Adjudicación {que}: {p.codigo} {p.titulo}"
    texto = (f"La Dirección General ({ap.email}) ha {que.lower()} la adjudicación de {p.codigo} «{p.titulo}» por "
             f"{_eur(float(ap.importe))} el {ap.decidido_en:%d/%m/%Y a las %H:%M}."
             f"{f' Comentario: {ap.comentario}' if ap.comentario else ''}"
             f"{'' if ap.estado == 'aprobada' else ' El proyecto vuelve al comparativo para revisar la adjudicación.'}")
    for d in destinos:
        try:
            avisos.enviar(d, asunto, texto, avisos._html(asunto, f"<p>{escape(texto)}</p>"))
            ok, err = True, None
        except Exception as e:  # noqa: BLE001
            ok, err = False, str(e)[:300]
        db.add(EmailLog(clave=f"capex_decision:{ap.id}", tipo="capex", destinatario=d, asunto=asunto[:200], ok=ok, error=err))
    db.commit()


@router.post("/{pid}/aprobacion")
def decide_in_app(pid: int, data: DecisionIn, request: Request, scope: Scope = Depends(get_scope),
                  db: Session = Depends(get_db)):
    """La Dirección General también puede decidir desde el PMS."""
    p = _proyecto(db, scope, pid)
    if not _aprobador(scope):
        raise HTTPException(403, "La aprueba la Dirección General")
    ap = db.scalar(select(CapexApproval).where(CapexApproval.project_id == p.id, CapexApproval.estado == "pendiente",
                                               CapexApproval.email == scope.user.email.lower()))
    if ap is None:
        bad_request("No tiene ninguna aprobación pendiente en este proyecto")
    _decidir(db, p, ap, data, "pms", request.client.host if request.client else None,
             request.headers.get("user-agent"))
    return _resumen(db, p)


# --------------------------------------------------------------------------- enlaces públicos
def _pagina(nombre: str) -> HTMLResponse:
    return HTMLResponse((STATIC / nombre).read_text(encoding="utf-8"),
                        headers={"Cache-Control": "no-store", "Referrer-Policy": "no-referrer",
                                 "X-Robots-Tag": "noindex"})


@publico.get("/oferta/{iid}", include_in_schema=False)
def bid_page(iid: int):
    return _pagina("oferta.html")


def _invitacion(db: Session, iid: int, k: str) -> tuple[TenderInvite, TenderProcess, CapexProject]:
    inv = db.get(TenderInvite, iid)
    if inv is None or inv.revocada or not hmac.compare_digest(inv.clave, k or ""):
        raise HTTPException(404, "Enlace no válido. Pida a la Dirección Técnica que se lo reenvíe.")
    t = db.get(TenderProcess, inv.tender_id)
    if datetime.now() > datetime.combine(t.fecha_limite, time(23, 59, 59)) + timedelta(days=60):
        raise HTTPException(404, "Este enlace ha caducado")
    return inv, t, db.get(CapexProject, t.project_id)


def _abierta(t: TenderProcess, p: CapexProject) -> bool:
    return t.estado == "abierta" and p.estado == "licitacion" and not _caducada(t)


@publico.get("/api/publico/oferta/{iid}")
def bid_data(iid: int, k: str, db: Session = Depends(get_db)):
    inv, t, p = _invitacion(db, iid, k)
    ahora = datetime.now()
    inv.abierto_en = inv.abierto_en or ahora
    inv.ultimo_acceso, inv.aperturas = ahora, (inv.aperturas or 0) + 1
    db.commit()
    a = db.get(Asset, p.asset_id)
    b = db.scalar(select(SupplierBid).where(SupplierBid.invite_id == inv.id))
    return {"proyecto": {"codigo": p.codigo, "titulo": p.titulo, "descripcion": p.descripcion, "activo": a.nombre,
                         "direccion": a.direccion, "ubicacion": p.ubicacion,
                         "gg_pct": float(p.gg_pct), "bi_pct": float(p.bi_pct)},
            "proveedor": db.get(Supplier, inv.supplier_id).nombre, "fecha_limite": t.fecha_limite.isoformat(),
            "mensaje": t.mensaje, "abierta": _abierta(t, p),
            # licitación ciega: sin precio base
            "partidas": [{"id": x.id, "capitulo": x.capitulo, "capitulo_nombre": x.capitulo_nombre, "codigo": x.codigo,
                          "descripcion": x.descripcion, "unidad": x.unidad, "cantidad": float(x.cantidad)}
                         for x in capex.partidas(db, p.id)],
            "documentos": [{"id": d.id, "nombre": d.nombre, "etiqueta": ETIQUETAS_DOC_CAPEX.get(d.etiqueta, d.etiqueta),
                            "tamano": d.tamano} for d in _docs_licitacion(db, p, None)],
            "oferta": {"precios": b.precios, "notas": b.notas, "plazo_dias": b.plazo_dias, "estado": b.estado,
                       "enviada_en": b.enviada_en.isoformat() if b.enviada_en else None} if b else None}


class PublicBidIn(BaseModel):
    precios: dict[str, float | None] = Field(default_factory=dict)
    notas: str | None = Field(default=None, max_length=3000)
    plazo_dias: int | None = Field(default=None, ge=1, le=1000)
    enviar: bool = False


@publico.post("/api/publico/oferta/{iid}")
def bid_save(iid: int, k: str, data: PublicBidIn, request: Request, db: Session = Depends(get_db)):
    inv, t, p = _invitacion(db, iid, k)
    if not _abierta(t, p):
        bad_request("El plazo de presentación de ofertas está cerrado")
    b = _guardar_oferta(db, inv, data.precios, "portal", data.enviar, data.notas, data.plazo_dias)
    if data.enviar:
        audit(db, None, "oferta_enviada", "capex", p.id, {"invitacion": inv.id, "precios": len(b.precios),
                                                            "ip": request.client.host if request.client else None})
        _avisar_oferta(db, p, inv)
    db.commit()
    return {"estado": b.estado, "precios": len(b.precios), "enviada_en": b.enviada_en.isoformat() if b.enviada_en else None}


def _avisar_oferta(db: Session, p: CapexProject, inv: TenderInvite) -> None:
    u = db.get(User, p.creado_por) if p.creado_por else None
    if not (u and u.email and avisos.configurado()):
        return
    s = db.get(Supplier, inv.supplier_id)
    asunto = f"Oferta recibida {p.codigo}: {s.nombre}"
    texto = f"{s.nombre} ha enviado su oferta para «{p.titulo}». Véala en CAPEX → {p.codigo} → Comparativo."
    try:
        avisos.enviar(u.email, asunto, texto, avisos._html(asunto, f"<p>{escape(texto)}</p>"))
        ok, err = True, None
    except Exception as e:  # noqa: BLE001
        ok, err = False, str(e)[:300]
    db.add(EmailLog(clave=f"capex_oferta:{inv.id}", tipo="capex", destinatario=u.email, asunto=asunto[:200], ok=ok, error=err))


@publico.get("/api/publico/oferta/{iid}/excel")
def bid_excel(iid: int, k: str, db: Session = Depends(get_db)):
    inv, t, p = _invitacion(db, iid, k)
    inv.descargas = (inv.descargas or 0) + 1
    db.commit()
    return Response(capex.excel_ciego(db, p, inv),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Oferta_{p.codigo}.xlsx"'})


@publico.post("/api/publico/oferta/{iid}/excel")
def bid_excel_upload(iid: int, k: str, fichero: UploadFile = File(...), db: Session = Depends(get_db)):
    inv, t, p = _invitacion(db, iid, k)
    if not _abierta(t, p):
        bad_request("El plazo de presentación de ofertas está cerrado")
    try:
        precios = capex.leer_excel(db, p, inv, fichero.file.read(10 * 1024 * 1024))
    except ValueError as e:
        bad_request(str(e))
    b = _guardar_oferta(db, inv, precios, "excel", False)
    db.commit()
    return {"estado": b.estado, "precios": precios}


@publico.get("/api/publico/oferta/{iid}/documentos/{did}")
def bid_document(iid: int, did: int, k: str, db: Session = Depends(get_db)):
    inv, t, p = _invitacion(db, iid, k)
    d = db.get(CapexDocument, did)
    if d is None or d.project_id != p.id or not d.licitacion:
        raise HTTPException(404, "Documento no encontrado")
    inv.descargas = (inv.descargas or 0) + 1
    db.commit()
    return Response(documentos.leer(d.fichero), media_type=d.mime, headers={
        "Content-Disposition": f'attachment; filename="{_ascii(d.nombre)}"', "Cache-Control": "private, no-store"})


@publico.get("/capex/aprobacion/{aid}", include_in_schema=False)
def approval_page(aid: int):
    return _pagina("capex_aprobacion.html")


def _aprobacion(db: Session, aid: int, k: str) -> tuple[CapexApproval, CapexProject]:
    ap = db.get(CapexApproval, aid)
    if ap is None or not hmac.compare_digest(ap.clave, k or ""):
        raise HTTPException(404, "Enlace no válido")
    return ap, db.get(CapexProject, ap.project_id)


@publico.get("/api/publico/capex-aprobacion/{aid}")
def approval_data(aid: int, k: str, db: Session = Depends(get_db)):
    ap, p = _aprobacion(db, aid, k)
    return {"resumen": ap.resumen, "estado": ap.estado, "decidido_en": ap.decidido_en.isoformat() if ap.decidido_en else None,
            "comentario": ap.comentario, "email": ap.email, "umbral": settings.capex_umbral,
            "pendiente": ap.estado == "pendiente" and p.estado == "pendiente_aprobacion"}


@publico.post("/api/publico/capex-aprobacion/{aid}")
def approval_decide(aid: int, k: str, data: DecisionIn, request: Request, db: Session = Depends(get_db)):
    ap, p = _aprobacion(db, aid, k)
    _decidir(db, p, ap, data, "correo", request.client.host if request.client else None,
             request.headers.get("user-agent"))
    return {"estado": ap.estado, "decidido_en": ap.decidido_en.isoformat()}


@publico.get("/api/publico/capex-aprobacion/{aid}/comparativo.pdf")
def approval_pdf(aid: int, k: str, db: Session = Depends(get_db)):
    ap, p = _aprobacion(db, aid, k)
    return Response(capex.comparativo_pdf(db, p, capex.comparativo(db, p)), media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="Comparativo_{p.codigo}.pdf"'})
