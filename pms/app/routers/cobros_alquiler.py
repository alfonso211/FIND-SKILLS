"""Alquiler residencial → Cobros de otro programa: lectura de las facturas en PDF (o ZIP), alta de las fichas de los
inquilinos y registro de la renta cobrada mes a mes como ingreso del activo (cuenta en la producción). Un documento
por inquilino con sus cobros y un Excel con todos los ingresos para la contabilidad."""
import io
from calendar import monthrange
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import cobros_alquiler as ca
from ..database import get_db
from ..models import Asset, Contact, ExternalInvoice
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404
from .historico import Cruce

router = APIRouter(prefix="/api/cobros-alquiler", tags=["cobros de alquiler"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIPO, SERIE = "renta", "ALQ"


class FilaIn(BaseModel):
    numero: str = Field(max_length=30)
    fecha: date
    cliente: str = Field(min_length=3, max_length=200)
    nif: str | None = Field(default=None, max_length=20)
    direccion: str | None = Field(default=None, max_length=300)
    total: float = Field(gt=0, lt=100000)
    desde: str = Field(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")
    hasta: str = Field(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")


class ImportIn(BaseModel):
    asset_id: int
    filas: list[FilaIn] = Field(min_length=1, max_length=300)


def _activo(db: Session, scope: Scope, asset_id: int, editar: bool = False) -> Asset:
    scope.require_asset("alquiler.editar" if editar else "alquiler.ver", asset_id)
    a = get_or_404(db, Asset, asset_id)
    if a.modalidad != "alquiler_residencial":
        bad_request("Solo para activos de alquiler residencial")
    return a


@router.post("/leer")
def read(ficheros: list[UploadFile] = File(...), asset_id: int = Form(...), scope: Scope = Depends(get_scope),
         db: Session = Depends(get_db)):
    """Lee las facturas y propone el periodo (de enero al mes actual). No guarda nada."""
    _activo(db, scope, asset_id, editar=True)
    subidas = []
    for f in ficheros[:50]:
        datos = f.file.read(60 * 1024 * 1024 + 1)
        if len(datos) > 60 * 1024 * 1024:
            bad_request(f"{f.filename}: supera 60 MB")
        subidas.append((f.filename or "fichero", datos))
    try:
        facturas, avisos = ca.leer_subida(subidas)
    except ValueError as e:
        bad_request(str(e))
    hoy = date.today()
    cruce = Cruce(db, [asset_id])
    vistos: set[str] = set()
    for x in facturas:
        k = ca.clave_inquilino(x["nif"], x["cliente"])
        if k in vistos:
            avisos.append(f"Factura {x['numero']}: el inquilino aparece en otra factura; se usará solo una")
        vistos.add(k)
        x["ficha"] = cruce.cliente(x["nif"], x["cliente"])[0] is not None
        x["desde"], x["hasta"] = f"{hoy.year}-01", f"{hoy:%Y-%m}"
        if x["mes"] and x["mes"] > f"{hoy:%Y-%m}":
            avisos.append(f"Factura {x['numero']}: es de un mes futuro ({x['mes']})")
    meses = {x["mes"] for x in facturas if x["mes"]}
    if len(meses) > 1:
        comun = max(meses, key=lambda m: sum(1 for x in facturas if x["mes"] == m))
        for x in facturas:
            if x["mes"] and x["mes"] != comun:
                avisos.append(f"Factura {x['numero']}: es de {x['mes']} y el resto de {comun}. Revise desde qué mes paga "
                              "(puede ser un inquilino nuevo)")
    return {"facturas": facturas, "avisos": avisos}


def _fecha_cobro(mes: str, dia: int) -> date:
    y, m = map(int, mes.split("-"))
    return date(y, m, min(dia, monthrange(y, m)[1]))


def _partir_nombre(nombre: str) -> tuple[str, str | None]:
    p = nombre.split()
    return (p[0], " ".join(p[1:]) or None) if len(p) > 1 else (nombre, None)


@router.post("/importar")
def import_(data: ImportIn, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Abre las fichas de los inquilinos que falten y guarda un ingreso cobrado por mes. Repetirlo actualiza."""
    a = _activo(db, scope, data.asset_id, editar=True)
    cruce = Cruce(db, [a.id])
    previas = {x.numero: x for x in db.scalars(select(ExternalInvoice).where(
        ExternalInvoice.asset_id == a.id, ExternalInvoice.tipo == TIPO))}
    fichas = cobros = 0
    for f in data.filas:
        lista = ca.meses(f.desde, f.hasta)
        if not lista:
            bad_request(f"{f.numero}: el periodo no es válido")
        nif = (f.nif or "").replace(" ", "").upper() or None
        if cruce.cliente(nif, f.cliente)[0] is None:
            nombre, apellidos = _partir_nombre(" ".join(f.cliente.split()))
            c = Contact(company_id=a.company_id, asset_id=a.id, tipo="inquilino", nombre=nombre, apellidos=apellidos,
                        documento_num=nif, documento_tipo="DNI" if nif else None, direccion=f.direccion,
                        notas="Ficha abierta al importar sus cobros de otro programa: falta asignar el piso y el contrato.")
            db.add(c)
            db.flush()
            cruce = Cruce(db, [a.id])
            fichas += 1
        k = ca.clave_inquilino(nif, f.cliente)
        mes_factura = f"{f.fecha:%Y-%m}"
        for mes in lista:
            numero = f"{k}-{mes}"
            x = previas.get(numero)
            if x is None:
                x = ExternalInvoice(asset_id=a.id, tipo=TIPO, serie=SERIE, numero=numero)
                db.add(x)
                previas[numero] = x
            cobro = _fecha_cobro(mes, f.fecha.day)
            x.fecha, x.nif, x.cliente, x.localizador = cobro, nif, " ".join(f.cliente.split()), None
            x.base, x.tipo_iva, x.cuota, x.total, x.fianza = f.total, 0, 0, f.total, 0
            x.detalle = {"cobrado": True, "fecha_cobro": cobro.isoformat(), "mes": mes, "factura_modelo": f.numero,
                         "factura": f.numero if mes == mes_factura else None, "direccion": f.direccion}
            x.user_id = scope.user.id
            cobros += 1
    audit(db, scope.user, "importar_cobros_alquiler", "activo", a.id,
          {"inquilinos": len(data.filas), "cobros": cobros, "fichas_nuevas": fichas})
    db.commit()
    return {"inquilinos": len(data.filas), "cobros": cobros, "fichas_nuevas": fichas}


def _por_inquilino(db: Session, a: Asset, anio: int | None = None) -> dict[str, dict]:
    stmt = select(ExternalInvoice).where(ExternalInvoice.asset_id == a.id, ExternalInvoice.tipo == TIPO)
    if anio:
        stmt = stmt.where(ExternalInvoice.fecha >= date(anio, 1, 1), ExternalInvoice.fecha <= date(anio, 12, 31))
    out: dict[str, dict] = defaultdict(lambda: {"cobros": []})
    for x in db.scalars(stmt.order_by(ExternalInvoice.fecha)):
        k = x.numero.rsplit("-", 2)[0]
        d = out[k]
        d.update(clave=k, cliente=x.cliente, nif=x.nif, direccion=(x.detalle or {}).get("direccion"))
        d["cobros"].append(x)
    return out


@router.get("")
def summary(asset_id: int, anio: int | None = None, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    a = _activo(db, scope, asset_id)
    cruce = Cruce(db, [a.id])
    filas = []
    for k, d in sorted(_por_inquilino(db, a, anio).items(), key=lambda kv: kv[1]["cliente"] or ""):
        cs = d["cobros"]
        filas.append({"clave": k, "cliente": d["cliente"], "contact_id": cruce.cliente(d["nif"], d["cliente"])[0],
                      "renta": float(cs[-1].total), "meses": len(cs), "total": round(sum(float(c.total) for c in cs), 2),
                      "desde": f"{cs[0].fecha:%Y-%m}", "hasta": f"{cs[-1].fecha:%Y-%m}"})
    return {"inquilinos": filas, "total": round(sum(f["total"] for f in filas), 2),
            "renta_mensual": round(sum(f["renta"] for f in filas), 2)}


@router.delete("/{clave}")
def delete_tenant(clave: str, asset_id: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Borra los cobros importados de un inquilino (para corregir y volver a importar). La ficha se conserva."""
    a = _activo(db, scope, asset_id, editar=True)
    d = _por_inquilino(db, a).get(clave)
    if not d:
        raise HTTPException(404, "No hay cobros de ese inquilino")
    for x in d["cobros"]:
        db.delete(x)
    audit(db, scope.user, "borrar_cobros_alquiler", "activo", a.id, {"clave": clave, "cobros": len(d["cobros"])})
    db.commit()
    return {"ok": True, "borrados": len(d["cobros"])}


def _pdf(a: Asset, d: dict) -> bytes:
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer

    from ..partes_trabajo import _doc_pdf, _esc, _estilos, _tabla
    cs = d["cobros"]
    total = sum(float(c.total) for c in cs)
    titulo = f"Relación de cobros de alquiler · {a.nombre}"
    out, doc, ancho, el = _doc_pdf(titulo, a, a.company)
    p, h = _estilos()
    el += [Paragraph(titulo, h),
           Paragraph(f"Inquilino: <b>{_esc(d['cliente'])}</b>" + (f" · NIF {_esc(d['nif'])}" if d["nif"] else "")
                     + (f" · {_esc(d['direccion'])}" if d["direccion"] else ""), p),
           Paragraph(f"{len(cs)} mensualidad(es) cobrada(s) · total {_eur(total)} · vivienda pendiente de asignar en el PMS",
                     p), Spacer(1, 3 * mm)]
    datos = [["Mes", "Concepto", "Fecha de cobro", "Factura", "Importe", "Estado"]]
    for c in cs:
        det = c.detalle or {}
        datos.append([f"{ca.MESES[c.fecha.month - 1].capitalize()} {c.fecha.year}", "Alquiler de vivienda",
                      f"{c.fecha:%d/%m/%Y}", det.get("factura") or f"(según {det.get('factura_modelo')})",
                      _eur(float(c.total)), "Cobrado"])
    datos.append(["", Paragraph("<b>Total</b>", p), "", "", _eur(total), ""])
    el.append(_tabla(datos, [36 * mm, 70 * mm, 34 * mm, 50 * mm, 34 * mm, ancho - 224 * mm]))
    doc.build(el)
    return out.getvalue()


def _eur(x: float) -> str:
    return f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


@router.get("/{clave}/pdf")
def tenant_pdf(clave: str, asset_id: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    a = _activo(db, scope, asset_id)
    d = _por_inquilino(db, a).get(clave)
    if not d:
        raise HTTPException(404, "No hay cobros de ese inquilino")
    return Response(_pdf(a, d), media_type="application/pdf", headers={
        "Content-Disposition": f'inline; filename="Cobros_{a.codigo}_{clave}.pdf"'})


@router.get("/excel/ingresos")
def excel(asset_id: int, anio: int | None = None, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Todos los cobros como ingresos por arrendamiento (uno por fila) para la contabilidad."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    a = _activo(db, scope, asset_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos arrendamiento"
    ws.append(["Fecha de cobro", "Mes", "Inquilino", "NIF", "Concepto", "Factura", "Base", "IVA", "Total", "Estado"])
    filas = [c for d in _por_inquilino(db, a, anio).values() for c in d["cobros"]]
    for c in sorted(filas, key=lambda c: (c.fecha, c.cliente or "")):
        det = c.detalle or {}
        ws.append([c.fecha, f"{c.fecha:%Y-%m}", c.cliente, c.nif or "", f"Alquiler de vivienda {a.nombre}",
                   det.get("factura") or "", float(c.base), float(c.cuota), float(c.total), "Cobrado"])
    ws.append([])
    ws.append(["", "", "", "", "", "Total", f"=SUM(G2:G{len(filas) + 1})", f"=SUM(H2:H{len(filas) + 1})",
               f"=SUM(I2:I{len(filas) + 1})", ""])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
    for fila in ws.iter_rows(min_row=2):
        fila[0].number_format = "DD/MM/YYYY"
        for c in fila[6:9]:
            c.number_format = "#,##0.00 €"
    for col, ancho in zip("ABCDEFGHIJ", (14, 10, 34, 12, 34, 12, 12, 10, 12, 10)):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="Ingresos_arrendamiento_{a.codigo}{f"_{anio}" if anio else ""}.xlsx"'})
