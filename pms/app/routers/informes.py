"""Informes en Excel: ocupación, producción, morosidad y costes de mantenimiento, por activo y mes.
Cada usuario solo ve los activos de su ámbito."""
from calendar import monthrange
from collections import defaultdict
from datetime import date, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..facturacion import lineas_de
from ..models import (MODALIDADES_CONTRATO, MODALIDADES_RESERVA, Asset, Charge, Contact, Invoice, Lease,
                      Reservation, Unit, WorkOrder)
from ..security import Scope, audit, get_scope
from ..utils import bad_request

router = APIRouter(prefix="/api/informes", tags=["informes"])

EUR, PCT, FECHA, ENTERO = '#,##0.00 "€"', "0.0%", "DD/MM/YYYY", "#,##0"
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
CABECERA = PatternFill("solid", fgColor="13294B")
TOTAL = PatternFill("solid", fgColor="E3EDF9")
INFORMES = {"ocupacion": "Ocupación", "produccion": "Producción", "morosidad": "Morosidad",
            "mantenimiento": "Costes de mantenimiento"}


# --------------------------------------------------------------------------- utilidades
def _meses(desde: date, hasta: date):
    """(etiqueta, primer día, último día) de cada mes del rango, recortado al rango."""
    y, m = desde.year, desde.month
    while (y, m) <= (hasta.year, hasta.month):
        ini, fin = date(y, m, 1), date(y, m, monthrange(y, m)[1])
        yield f"{MESES[m - 1]}-{y}", max(ini, desde), min(fin, hasta)
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def _noches(ent: date, sal: date, ini: date, fin: date) -> int:
    """Noches de la estancia [ent, sal) que caen entre ini y fin (ambos incluidos)."""
    return max(0, (min(sal, fin + timedelta(days=1)) - max(ent, ini)).days)


def _hoja(wb: Workbook, titulo: str, cabecera: list[str], filas: list[list], formatos: dict[int, str] | None = None,
          totales: list[int] | None = None, nota: str | None = None):
    ws = wb.create_sheet(titulo[:31])
    formatos = formatos or {}
    ws.append(cabecera)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), CABECERA
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for f in filas:
        ws.append(f)
    n = len(filas)
    if totales and n:
        ws.append(["TOTAL"] + [None] * (len(cabecera) - 1))
        for i in totales:
            col = get_column_letter(i + 1)
            ws.cell(n + 2, i + 1, f"=SUM({col}2:{col}{n + 1})")
        for c in ws[n + 2]:
            c.font, c.fill = Font(bold=True), TOTAL
    for i, fmt in formatos.items():
        for fila in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
            fila[0].number_format = fmt
    for i, cab in enumerate(cabecera, 1):
        largo = max([len(str(cab))] + [len(str(f[i - 1])) for f in filas[:500] if f[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)
    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecera))}{n + 1}"
    if nota:
        ws.cell(n + 4 if totales else n + 3, 1, nota).font = Font(italic=True, color="6B7785")
    return ws


def _activos(db: Session, scope: Scope, perm: str, asset_id: int | None, modalidades=None) -> list[Asset]:
    ids = scope.asset_ids(perm)
    stmt = select(Asset).order_by(Asset.codigo)
    if ids is not None:
        stmt = stmt.where(Asset.id.in_(ids or {-1}))
    if asset_id:
        stmt = stmt.where(Asset.id == asset_id)
    if modalidades:
        stmt = stmt.where(Asset.modalidad.in_(modalidades))
    return list(db.scalars(stmt))


def _reservas(db: Session, ids: list[int], desde: date, hasta: date):
    return db.execute(select(Reservation, Unit.asset_id).join(Unit, Unit.id == Reservation.unit_id).where(
        Unit.asset_id.in_(ids or [-1]), Reservation.estado.in_(("confirmada", "checkin", "checkout")),
        Reservation.fecha_entrada <= hasta, Reservation.fecha_salida > desde)).all()


# --------------------------------------------------------------------------- informes
def _ocupacion(db, scope, wb, desde, hasta, asset_id):
    turisticos = _activos(db, scope, "reservas.ver", asset_id, MODALIDADES_RESERVA)
    residenciales = _activos(db, scope, "alquiler.ver", asset_id, MODALIDADES_CONTRATO)
    if not (turisticos or residenciales):
        raise HTTPException(403, "Sin permiso para ver la ocupación de ningún activo")
    if turisticos:
        n_unid = {a.id: db.scalar(select(func.count()).select_from(Unit).where(
            Unit.asset_id == a.id, Unit.estado != "fuera_servicio")) for a in turisticos}
        res = _reservas(db, [a.id for a in turisticos], desde, hasta)
        fact = _facturado(db, [a.id for a in turisticos], desde, hasta) if any(
            scope.can_asset("finanzas.ver", a.id) for a in turisticos) else {}
        filas = []
        for a in turisticos:
            unidades = dict(db.execute(select(Unit.id, Unit.uso).where(Unit.asset_id == a.id,
                                                                       Unit.estado != "fuera_servicio")).all())
            for uso in sorted(set(unidades.values()), key=lambda x: (x == "garaje", x)):  # apartamentos y garajes
                ids_uso = {i for i, x in unidades.items() if x == uso}
                propias = [r for r, aid in res if aid == a.id and r.unit_id in ids_uso]
                clave = "garaje" if uso == "garaje" else "alojamiento"
                for mes, ini, fin in _meses(desde, hasta):
                    disp = len(ids_uso) * ((fin - ini).days + 1)
                    ocup = sum(_noches(r.fecha_entrada, r.fecha_salida, ini, fin) for r in propias)
                    ingresos = (round(fact.get((a.id, f"{ini:%Y-%m}"), {}).get(clave, 0), 2)
                                if scope.can_asset("finanzas.ver", a.id) else None)
                    llegadas = sum(1 for r in propias if ini <= r.fecha_entrada <= fin)
                    filas.append([a.nombre, uso, mes, len(ids_uso), disp, ocup, ocup / disp if disp else 0, llegadas,
                                  ingresos, None if ingresos is None else round(ingresos / ocup, 2) if ocup else 0,
                                  None if ingresos is None else round(ingresos / disp, 2) if disp else 0])
        _hoja(wb, "Turísticos", ["Activo", "Uso", "Mes", "Unidades", "Noches disponibles", "Noches ocupadas",
                                 "% ocupación", "Llegadas", "Facturado (base)", "ADR (precio medio noche)", "RevPAR"],
              filas, {4: ENTERO, 5: ENTERO, 6: PCT, 8: EUR, 9: EUR, 10: EUR}, totales=[4, 5, 7, 8],
              nota="Noches: reservas confirmadas, alojadas o salidas (sin canceladas ni no presentadas); unidades fuera de "
                   "servicio excluidas. Importes: alojamiento facturado en el mes según fecha de factura, sin IVA "
                   "(solo con permiso de finanzas).")
    if residenciales:
        filas = []
        for a in residenciales:
            unidades = db.execute(select(Unit.id, Unit.uso).where(Unit.asset_id == a.id)).all()
            usos = defaultdict(set)
            for uid, uso in unidades:
                usos[uso].add(uid)
            contratos = db.execute(select(Lease.unit_id, Lease.fecha_inicio, Lease.fecha_fin).join(Unit).where(
                Unit.asset_id == a.id, Lease.estado != "borrador")).all()
            for mes, ini, fin in _meses(desde, hasta):
                alquiladas = {u for u, fi, ff in contratos if fi <= fin and (ff is None or ff >= ini)}
                for uso, ids in sorted(usos.items()):
                    n = len(ids & alquiladas)
                    filas.append([a.nombre, mes, uso, len(ids), n, n / len(ids) if ids else 0])
        _hoja(wb, "Residencial", ["Activo", "Mes", "Uso", "Unidades", "Alquiladas", "% ocupación"], filas,
              {3: ENTERO, 4: ENTERO, 5: PCT},
              nota="Unidad alquilada: con contrato vigente, finalizado o rescindido que cubre algún día del mes.")


def _facturado(db, ids: list[int], desde: date, hasta: date) -> dict[tuple[int, str], dict[str, float]]:
    """Base imponible facturada por activo, mes (fecha de factura) y tipo de línea; también IVA, total y nº."""
    out: dict[tuple[int, str], dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for f in db.scalars(select(Invoice).where(Invoice.asset_id.in_(ids or [-1]), Invoice.fecha_expedicion >= desde,
                                              Invoice.fecha_expedicion <= hasta)):
        acc = out[(f.asset_id, f"{f.fecha_expedicion:%Y-%m}")]
        for x in lineas_de(f):
            acc[x["tipo"]] += x["base"]
        acc["base"] += float(f.base_imponible)
        acc["iva"] += float(f.cuota_iva)
        acc["total"] += float(f.total)
        acc["n"] += 1
    return out


def _produccion(db, scope, wb, desde, hasta, asset_id):
    """La producción es lo facturado: cuenta la fecha de la factura, no la de la reserva ni la de entrada."""
    activos = _activos(db, scope, "finanzas.ver", asset_id)
    if not activos:
        raise HTTPException(403, "Sin permiso para ver la producción (finanzas)")
    fact = _facturado(db, [a.id for a in activos], desde, hasta)
    filas = []
    for a in activos:
        for mes, ini, _ in _meses(desde, hasta):
            v = fact.get((a.id, f"{ini:%Y-%m}"), {})
            filas.append([a.nombre, mes] + [round(v.get(k, 0), 2) for k in
                                            ("alojamiento", "garaje", "renta", "servicio", "base", "iva", "total")]
                         + [int(v.get("n", 0))])
    _hoja(wb, "Producción", ["Activo", "Mes", "Alojamiento (base)", "Plazas de garaje (base)", "Rentas (base)",
                             "Servicios (base)", "Producción (base imponible)", "IVA", "Total facturado", "Nº facturas"],
          filas, {2: EUR, 3: EUR, 4: EUR, 5: EUR, 6: EUR, 7: EUR, 8: EUR, 9: ENTERO},
          totales=[2, 3, 4, 5, 6, 7, 8, 9],
          nota="Producción = lo facturado en cada mes según la fecha de la factura, sin IVA. Incluye las "
               "rectificativas (en negativo).")


def _tramo(dias: int) -> str:
    return "0-30 días" if dias <= 30 else "31-60 días" if dias <= 60 else "61-90 días" if dias <= 90 else "Más de 90 días"


def _morosidad(db, scope, wb, desde, hasta, asset_id):
    alq = _activos(db, scope, "alquiler.ver", asset_id, MODALIDADES_CONTRATO)
    tur = _activos(db, scope, "reservas.ver", asset_id, MODALIDADES_RESERVA)
    if not (alq or tur):
        raise HTTPException(403, "Sin permiso para ver la morosidad de ningún activo")
    nombres = {a.id: a.nombre for a in alq + tur}
    resumen = defaultdict(lambda: defaultdict(float))
    filas = []
    for c, u, t in db.execute(
            select(Charge, Unit, Contact).join(Lease, Lease.id == Charge.lease_id).join(Unit, Unit.id == Lease.unit_id)
            .join(Contact, Contact.id == Lease.tenant_id).where(
                Unit.asset_id.in_([a.id for a in alq] or [-1]), Charge.estado.in_(("pendiente", "parcial")),
                Charge.fecha_vencimiento <= hasta).order_by(Unit.asset_id, Charge.fecha_vencimiento)):
        pend = round(float(c.importe) - float(c.importe_pagado), 2)
        dias = (hasta - c.fecha_vencimiento).days
        resumen[nombres[u.asset_id]][_tramo(dias)] += pend
        filas.append([nombres[u.asset_id], u.codigo, f"{t.nombre} {t.apellidos or ''}".strip(), t.documento_num,
                      t.telefono, t.email, c.periodo, c.concepto, c.fecha_vencimiento, dias, _tramo(dias),
                      float(c.importe), float(c.importe_pagado), pend])
    if alq:
        _hoja(wb, "Recibos impagados", ["Activo", "Unidad", "Inquilino", "Documento", "Teléfono", "Email", "Periodo",
                                        "Concepto", "Vencimiento", "Días de retraso", "Tramo", "Importe", "Cobrado",
                                        "Pendiente"], filas, {8: FECHA, 9: ENTERO, 11: EUR, 12: EUR, 13: EUR},
              totales=[11, 12, 13], nota=f"Recibos vencidos a {hasta:%d/%m/%Y} sin cobrar del todo.")
    filas = []
    for r, aid in db.execute(select(Reservation, Unit.asset_id).join(Unit, Unit.id == Reservation.unit_id).where(
            Unit.asset_id.in_([a.id for a in tur] or [-1]), Reservation.estado.in_(("checkin", "checkout")),
            Reservation.fecha_entrada <= hasta, Reservation.importe_total > Reservation.importe_pagado)
            .order_by(Unit.asset_id, Reservation.fecha_entrada)):
        g = r.guest
        pend = round(float(r.importe_total) - float(r.importe_pagado), 2)
        dias = max(0, (hasta - r.fecha_salida).days)
        resumen[nombres[aid]][_tramo(dias)] += pend
        filas.append([nombres[aid], r.unit.codigo, r.localizador, f"{g.nombre} {g.apellidos or ''}".strip(), g.telefono,
                      g.email, r.canal, r.fecha_entrada, r.fecha_salida, r.estado, float(r.importe_total),
                      float(r.importe_pagado), pend])
    if tur:
        _hoja(wb, "Reservas con saldo", ["Activo", "Unidad", "Localizador", "Huésped", "Teléfono", "Email", "Canal",
                                         "Entrada", "Salida", "Estado", "Importe", "Cobrado", "Pendiente"], filas,
              {7: FECHA, 8: FECHA, 10: EUR, 11: EUR, 12: EUR}, totales=[10, 11, 12],
              nota="Reservas con check-in o check-out con importe pendiente de cobro.")
    tramos = ["0-30 días", "31-60 días", "61-90 días", "Más de 90 días"]
    filas = [[a] + [round(v[t], 2) for t in tramos] + [round(sum(v.values()), 2)] for a, v in sorted(resumen.items())]
    ws = _hoja(wb, "Resumen", ["Activo"] + tramos + ["Total pendiente"], filas, {i: EUR for i in range(1, 6)},
               totales=[1, 2, 3, 4, 5], nota=f"Deuda pendiente a {hasta:%d/%m/%Y} por antigüedad.")
    wb.move_sheet(ws, offset=-len(wb.sheetnames) + 1)


def _mantenimiento(db, scope, wb, desde, hasta, asset_id):
    activos = _activos(db, scope, "mantenimiento.ver", asset_id)
    if not activos:
        raise HTTPException(403, "Sin permiso para ver el mantenimiento de ningún activo")
    nombres = {a.id: a.nombre for a in activos}
    ots = list(db.scalars(select(WorkOrder).where(WorkOrder.asset_id.in_(nombres), WorkOrder.fecha_apertura >= desde,
                                                  WorkOrder.fecha_apertura <= hasta).order_by(WorkOrder.id)))
    codigos = dict(db.execute(select(Unit.id, Unit.codigo).where(Unit.id.in_({w.unit_id for w in ots if w.unit_id} or {-1}))).all())
    coste = lambda w: float(w.coste_real if w.coste_real is not None else 0)  # noqa: E731
    filas = [[f"OT-{w.id:05d}", nombres[w.asset_id], codigos.get(w.unit_id, "Zonas comunes"), w.titulo, w.tipo,
              w.categoria, w.prioridad, w.estado, w.fecha_apertura, w.fecha_cierre,
              (w.fecha_cierre - w.fecha_apertura).days if w.fecha_cierre else None, w.proveedor, w.asignado_a,
              float(w.coste_estimado) if w.coste_estimado is not None else None, coste(w)] for w in ots]
    _hoja(wb, "Órdenes de trabajo", ["OT", "Activo", "Unidad", "Título", "Tipo", "Instalación", "Prioridad", "Estado",
                                     "Apertura", "Cierre", "Días resolución", "Proveedor", "Asignado", "Coste estimado",
                                     "Coste real"], filas, {8: FECHA, 9: FECHA, 10: ENTERO, 13: EUR, 14: EUR},
          totales=[13, 14], nota="OT abiertas en el periodo. Coste real: el indicado al confirmar el trabajo.")
    for titulo, clave in (("Por instalación", lambda w: (nombres[w.asset_id], w.categoria)),
                          ("Por tipo", lambda w: (nombres[w.asset_id], w.tipo)),
                          ("Por proveedor", lambda w: (nombres[w.asset_id], w.proveedor or "(sin proveedor)"))):
        grupos = defaultdict(list)
        for w in ots:
            grupos[clave(w)].append(w)
        filas = [[a, k, len(ws), sum(1 for w in ws if w.estado not in ("cerrada", "cancelada")),
                  sum(1 for w in ws if w.prioridad == "urgente"), round(sum(coste(w) for w in ws), 2)]
                 for (a, k), ws in sorted(grupos.items(), key=lambda x: (x[0][0], -sum(coste(w) for w in x[1])))]
        _hoja(wb, titulo, ["Activo", titulo.split(" ", 1)[1].capitalize(), "Nº OT", "Abiertas", "Urgentes",
                           "Coste real"], filas, {2: ENTERO, 3: ENTERO, 4: ENTERO, 5: EUR}, totales=[2, 3, 4, 5])


GENERADORES = {"ocupacion": _ocupacion, "produccion": _produccion, "morosidad": _morosidad,
               "mantenimiento": _mantenimiento}


@router.get("/{informe}")
def report(informe: str, desde: date | None = None, hasta: date | None = None, asset_id: int | None = None,
           scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    if informe not in GENERADORES:
        raise HTTPException(404, f"Informe no disponible. Opciones: {', '.join(GENERADORES)}")
    hasta = hasta or date.today()
    desde = desde or date(hasta.year, 1, 1)
    if desde > hasta:
        bad_request("La fecha inicial es posterior a la final")
    if (hasta - desde).days > 3 * 366:
        bad_request("El periodo máximo es de 3 años")
    wb = Workbook()
    wb.remove(wb.active)
    GENERADORES[informe](db, scope, wb, desde, hasta, asset_id)
    wb.properties.title = f"{INFORMES[informe]} {desde:%d/%m/%Y}-{hasta:%d/%m/%Y}"
    wb.properties.creator = "INVERPMS"
    out = BytesIO()
    wb.save(out)
    audit(db, scope.user, "informe", informe, None, {"desde": str(desde), "hasta": str(hasta), "asset_id": asset_id})
    db.commit()
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Informe_{informe}_{desde}_{hasta}.xlsx"'})
