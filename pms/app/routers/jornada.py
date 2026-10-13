"""Registro de jornada (art. 34.9 ET, RD-ley 8/2019): entrada y salida diaria del personal propio.

- Cada empleado ficha desde su usuario del PMS («Mi jornada») o, si no tiene usuario, con su enlace personal
  (/fichar/{id}?k=…) desde el móvil. Se guarda la hora del servidor y, si el móvil lo permite, la ubicación.
- Recepción 1 del activo y la dirección consultan los registros, los corrigen (queda el original, quién y por qué)
  y sacan el informe mensual en PDF para firmar y el Excel para la asesoría. Se conservan 4 años."""
import hmac
import io
from datetime import date, datetime, timedelta
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.responses import HTMLResponse
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .. import empleados
from ..config import BASE_DIR
from ..database import get_db
from ..models import AREAS_PERSONAL, Asset, StaffMember, TimeEntry, User
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404

router = APIRouter(prefix="/api/jornada", tags=["registro de jornada"])
publico = APIRouter(tags=["registro de jornada (enlace)"])
PAGINA = Path(BASE_DIR) / "static" / "fichar.html"


class Gps(BaseModel):
    lat: float | None = Field(default=None, ge=-90, le=90)
    lon: float | None = Field(default=None, ge=-180, le=180)


def _t(t: TimeEntry, ahora: datetime | None = None) -> dict:
    d = t.to_dict()
    for k in ("lat", "lon", "lat_salida", "lon_salida"):
        d[k] = float(d[k]) if d[k] is not None else None
    d["horas"] = empleados.horas(t, ahora)
    d["corregido"] = t.original is not None
    return d


def _estado(db: Session, p: StaffMember) -> dict:
    ahora = datetime.now()
    hoy = empleados.del_mes(db, [p.id], ahora.strftime("%Y-%m"))
    abierto = empleados.abierto(db, p)
    return {"id": p.id, "nombre": p.nombre, "puesto": p.puesto, "dentro": abierto is not None,
            "desde": abierto.entrada.isoformat() if abierto else None,
            "hoy": [_t(t, ahora) for t in hoy if t.entrada.date() == ahora.date()],
            "horas_mes": round(sum(empleados.horas(t, ahora) for t in hoy), 2),
            "jornada_semanal": float(p.jornada_semanal) if p.jornada_semanal is not None else None}


# --------------------------------------------------------------------------- el propio empleado
def _mia(db: Session, scope: Scope) -> StaffMember:
    p = empleados.ficha_de_usuario(db, scope.user.id)
    if p is None or not empleados.es_propio(p):
        raise HTTPException(404, "Su usuario no está enlazado a una ficha de empleado. Pídalo a Recepción 1 o a dirección.")
    return p


@router.get("/mia")
def my_day(scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    return _estado(db, _mia(db, scope))


@router.post("/mia/fichar")
def clock_me(gps: Gps, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _mia(db, scope)
    t = empleados.fichar(db, p, "pms", gps.lat, gps.lon)
    audit(db, scope.user, "fichar", "personal", p.id, {"fichaje": t.id, "salida": t.salida is not None})
    db.commit()
    return {"fichaje": _t(t), **_estado(db, p)}


# --------------------------------------------------------------------------- enlace personal (sin usuario)
def _por_enlace(db: Session, pid: int, k: str) -> StaffMember:
    p = db.get(StaffMember, pid)
    if p is None or not p.fichaje_clave or not hmac.compare_digest(p.fichaje_clave, k or "") or not p.activo:
        raise HTTPException(404, "Enlace no válido. Pida uno nuevo a Recepción 1.")
    return p


@publico.get("/fichar/{pid}", include_in_schema=False)
def page(pid: int):
    return HTMLResponse(PAGINA.read_text(encoding="utf-8"), headers={"Cache-Control": "no-store",
                                                                     "Referrer-Policy": "no-referrer"})


@publico.get("/api/publico/fichaje/{pid}")
def link_state(pid: int, k: str, db: Session = Depends(get_db)):
    return _estado(db, _por_enlace(db, pid, k))


@publico.post("/api/publico/fichaje/{pid}")
def link_clock(pid: int, k: str, gps: Gps, request: Request, db: Session = Depends(get_db)):
    p = _por_enlace(db, pid, k)
    try:
        t = empleados.fichar(db, p, "enlace", gps.lat, gps.lon)
    except ValueError as e:
        bad_request(str(e))
    audit(db, None, "fichar", "personal", p.id, {"fichaje": t.id, "salida": t.salida is not None,
                                                  "ip": request.client.host if request.client else None})
    db.commit()
    return {"fichaje": _t(t), **_estado(db, p)}


# --------------------------------------------------------------------------- consulta y corrección (gestión)
def _empleados(db: Session, scope: Scope, asset_id: int | None) -> list[StaffMember]:
    stmt = select(StaffMember).where(StaffMember.supplier_id.is_(None))
    if asset_id:
        stmt = stmt.where(or_(StaffMember.asset_id.is_(None), StaffMember.asset_id == asset_id))
    return [p for p in db.scalars(stmt.order_by(StaffMember.nombre))
            if empleados.es_propio(p) and empleados.puede_gestionar(db, scope, p)]


@router.get("")
def month(mes: str = Query(pattern=r"^\d{4}-\d{2}$"), asset_id: int | None = None, staff_id: int | None = None,
          scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Resumen del mes por empleado (horas, días, registros sin salida) y el detalle de cada fichaje."""
    gente = _empleados(db, scope, asset_id)
    if staff_id:
        gente = [p for p in gente if p.id == staff_id]
    if not gente and not scope.has_any("personal.autorizar"):
        raise HTTPException(403, "El registro de jornada lo consultan Recepción 1 y la dirección")
    ahora = datetime.now()
    filas = empleados.del_mes(db, [p.id for p in gente], mes)
    nombres = dict(db.execute(select(User.id, User.nombre).where(
        User.id.in_({t.editado_por for t in filas if t.editado_por} or {-1}))).all())
    out = []
    for p in gente:
        mias = [t for t in filas if t.staff_id == p.id]
        out.append({"id": p.id, "nombre": p.nombre, "puesto": p.puesto, "area": AREAS_PERSONAL.get(p.area, p.area),
                    "activo": db.get(Asset, p.asset_id).nombre if p.asset_id else "Todos los activos",
                    "dentro": any(t.salida is None and t.entrada.date() == ahora.date() for t in mias),
                    "horas": round(sum(empleados.horas(t, ahora) for t in mias), 2),
                    "dias": len({t.entrada.date() for t in mias}),
                    "sin_salida": sum(1 for t in mias if t.salida is None and t.entrada.date() < ahora.date()),
                    "fichajes": [{**_t(t, ahora), "editado_nombre": nombres.get(t.editado_por)} for t in mias],
                    "enlace": bool(p.fichaje_clave), "usuario": bool(p.user_id)})
    return out


class Correccion(BaseModel):
    entrada: datetime
    salida: datetime | None = None
    motivo: str = Field(min_length=3, max_length=300)


def _comprueba(c: Correccion) -> None:
    if c.salida and c.salida <= c.entrada:
        bad_request("La salida debe ser posterior a la entrada")
    if c.entrada > datetime.now() + timedelta(minutes=5):
        bad_request("No se pueden registrar horas futuras")
    if c.salida and (c.salida - c.entrada) > timedelta(hours=16):
        bad_request("Un registro no puede pasar de 16 horas: revise las horas")


@router.post("/{fid}/corregir")
def fix(fid: int, data: Correccion, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Corrección (p. ej. olvidó fichar la salida): se guarda el registro original, quién lo cambia y el motivo."""
    t = get_or_404(db, TimeEntry, fid)
    p = db.get(StaffMember, t.staff_id)
    if not empleados.puede_gestionar(db, scope, p):
        raise HTTPException(403, "Lo corrigen Recepción 1 y la dirección")
    _comprueba(data)
    if t.original is None:
        t.original = {"entrada": t.entrada.isoformat(), "salida": t.salida.isoformat() if t.salida else None}
    t.entrada, t.salida = data.entrada.replace(tzinfo=None), data.salida.replace(tzinfo=None) if data.salida else None
    t.motivo, t.editado_por, t.editado_en = data.motivo.strip(), scope.user.id, datetime.now()
    audit(db, scope.user, "corregir_fichaje", "personal", p.id, {"fichaje": t.id, "original": t.original,
                                                                  "motivo": t.motivo})
    db.commit()
    return _t(t)


class Manual(Correccion):
    staff_id: int


@router.post("/manual", status_code=201)
def manual(data: Manual, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Registro a mano (el empleado no pudo fichar), con el motivo."""
    p = get_or_404(db, StaffMember, data.staff_id)
    if not (empleados.es_propio(p) and empleados.puede_gestionar(db, scope, p)):
        raise HTTPException(403, "Lo registran Recepción 1 y la dirección, solo para el personal propio")
    _comprueba(data)
    t = TimeEntry(staff_id=p.id, entrada=data.entrada.replace(tzinfo=None),
                  salida=data.salida.replace(tzinfo=None) if data.salida else None, origen="manual",
                  motivo=data.motivo.strip(), editado_por=scope.user.id, editado_en=datetime.now())
    db.add(t)
    db.flush()
    audit(db, scope.user, "fichaje_manual", "personal", p.id, {"fichaje": t.id, "motivo": t.motivo})
    db.commit()
    return _t(t)


# --------------------------------------------------------------------------- informe mensual y Excel
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


@router.get("/informe.pdf")
def report(staff_id: int, mes: str = Query(pattern=r"^\d{4}-\d{2}$"), scope: Scope = Depends(get_scope),
           db: Session = Depends(get_db)):
    """Registro mensual de jornada de un empleado, con las firmas de empresa y trabajador."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    p = get_or_404(db, StaffMember, staff_id)
    if not empleados.puede_gestionar(db, scope, p):
        raise HTTPException(403, "El informe lo sacan Recepción 1 y la dirección")
    a = db.get(Asset, p.asset_id) if p.asset_id else None
    empresa = a.company if a else None
    y, m = map(int, mes.split("-"))
    filas = empleados.del_mes(db, [p.id], mes)
    ss = getSampleStyleSheet()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=f"Registro de jornada {p.nombre} {mes}")
    cab = [Paragraph(f"<b>REGISTRO DE JORNADA · {MESES[m - 1].upper()} {y}</b>", ss["Title"]),
           Paragraph(f"Empresa: <b>{empresa.nombre if empresa else '—'}</b>"
                     f"{f' · CIF {empresa.cif}' if empresa and empresa.cif else ''}"
                     f"{f' · Centro: {a.nombre}' if a else ''}", ss["Normal"]),
           Paragraph(f"Trabajador: <b>{p.nombre}</b>{f' · DNI {p.dni}' if p.dni else ''}"
                     f"{f' · {p.puesto}' if p.puesto else ''}"
                     f"{f' · Jornada: {float(p.jornada_semanal):g} h/semana' if p.jornada_semanal else ''}", ss["Normal"]),
           Spacer(1, 4 * mm)]
    datos = [["Día", "Entrada", "Salida", "Horas", "Observaciones"]]
    total = 0.0
    for t in filas:
        h = empleados.horas(t)
        total += h
        obs = []
        if t.origen == "manual":
            obs.append(f"Registro manual: {t.motivo}")
        elif t.original:
            obs.append(f"Corregido: {t.motivo}")
        if t.salida is None:
            obs.append("Sin salida")
        datos.append([t.entrada.strftime("%d/%m/%Y"), t.entrada.strftime("%H:%M"),
                      t.salida.strftime("%H:%M") if t.salida else "—", f"{h:.2f}".replace(".", ","),
                      Paragraph("; ".join(obs), ss["BodyText"]) if obs else ""])
    datos.append(["", "", "Total", f"{total:.2f}".replace(".", ","), ""])
    tabla = Table(datos, colWidths=[26 * mm, 22 * mm, 22 * mm, 20 * mm, 88 * mm], repeatRows=1)
    tabla.setStyle(TableStyle([("GRID", (0, 0), (-1, -2), 0.4, colors.HexColor("#d6cfc0")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a1c")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    firmas = Table([["Firma de la empresa", "Firma del trabajador"], ["", ""]], colWidths=[89 * mm, 89 * mm],
                   rowHeights=[7 * mm, 26 * mm])
    firmas.setStyle(TableStyle([("BOX", (0, 0), (0, -1), 0.4, colors.grey), ("BOX", (1, 0), (1, -1), 0.4, colors.grey),
                                ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    doc.build(cab + [tabla, Spacer(1, 8 * mm), firmas, Spacer(1, 3 * mm),
                     Paragraph("Registro diario de jornada conforme al art. 34.9 del Estatuto de los Trabajadores. "
                               "Se conserva durante cuatro años a disposición de la persona trabajadora, de sus "
                               "representantes legales y de la Inspección de Trabajo.", ss["Italic"])])
    return Response(out.getvalue(), media_type="application/pdf", headers={
        "Content-Disposition": f'inline; filename="{empleados.nombre_fichero(f"Registro_jornada_{p.nombre}_{mes}")}.pdf"'})


@router.get("/excel")
def excel(mes: str = Query(pattern=r"^\d{4}-\d{2}$"), asset_id: int | None = None, scope: Scope = Depends(get_scope),
          db: Session = Depends(get_db)):
    """Todos los fichajes del mes para la asesoría (una hoja de detalle y otra de totales)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    datos = month(mes, asset_id, None, scope, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fichajes"
    ws.append(["Empleado", "Puesto", "Centro", "Día", "Entrada", "Salida", "Horas", "Origen", "Corregido / motivo"])
    for e in datos:
        for t in e["fichajes"]:
            ent = datetime.fromisoformat(t["entrada"])
            sal = datetime.fromisoformat(t["salida"]) if t["salida"] else None
            ws.append([e["nombre"], e["puesto"] or "", e["activo"], ent.date(), ent.strftime("%H:%M"),
                       sal.strftime("%H:%M") if sal else "", t["horas"], t["origen"],
                       t["motivo"] or ""])
    ws2 = wb.create_sheet("Totales")
    ws2.append(["Empleado", "Puesto", "Centro", "Días con fichaje", "Horas", "Registros sin salida"])
    for e in datos:
        ws2.append([e["nombre"], e["puesto"] or "", e["activo"], e["dias"], e["horas"], e["sin_salida"]])
    for hoja in (ws, ws2):
        for c in hoja[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A1A1C")
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = max(10, min(40, max(len(str(c.value or "")) for c in col) + 2))
    for fila in ws.iter_rows(min_row=2):
        fila[3].number_format = "DD/MM/YYYY"
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Registro_jornada_{mes}.xlsx"'})


def personal_hoy(db: Session, scope: Scope) -> dict | None:
    """Panel de control (Recepción 1 y dirección): quién falta hoy, quién está dentro, olvidos de fichaje y documentos
    que vencen."""
    from .. import ausencias as au
    from ..models import Absence, StaffDocument
    ids = scope.asset_ids("activos.ver")
    activos = list(db.scalars(select(Asset).where(Asset.activo) if ids is None else select(Asset).where(Asset.id.in_(ids))))
    mios = [a for a in activos if au.gestiona(db, scope, a)]
    if not mios and not scope.is_group_level("personal.autorizar"):
        return None
    hoy = date.today()
    faltan = []
    for x in db.scalars(select(Absence).where(Absence.asset_id.in_([a.id for a in mios] or [-1]), Absence.estado == "aprobada",
                                             Absence.desde <= hoy, Absence.hasta >= hoy).order_by(Absence.persona)):
        faltan.append({"persona": x.persona, "asset_id": x.asset_id, "hasta": x.hasta.isoformat(),
                       "tipo": "Ausencia" if x.tipo == "baja_medica" else x.tipo.replace("_", " ")})
    gente = [p for p in db.scalars(select(StaffMember).where(StaffMember.activo, StaffMember.supplier_id.is_(None)))
             if empleados.puede_gestionar(db, scope, p)]
    gids = [p.id for p in gente]
    dentro = [p.nombre for p in gente if empleados.abierto(db, p) and empleados.abierto(db, p).entrada.date() == hoy]
    olvidos = sum(1 for t in empleados.olvidos(db, hoy) if t.staff_id in gids)
    docs = sum(1 for d in db.scalars(select(StaffDocument).where(StaffDocument.staff_id.in_(gids or [-1]),
                                                                  StaffDocument.vence.is_not(None),
                                                                  StaffDocument.vence <= hoy + timedelta(days=30))))
    return {"faltan": faltan, "dentro": dentro, "olvidos": olvidos, "documentos": docs}
