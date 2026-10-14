"""Control del preventivo y de las inspecciones OCA: semáforo de cada plan, revisiones con sus documentos,
deficiencias con su OT y registro en Excel (ver app/preventivo.py)."""
import io
import json
from datetime import date

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Response, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import adjuntos, avisos, documentos, preventivo
from ..database import get_db
from ..models import Asset, PreventivePlan, PreventiveRevision, RevisionAttachment, Supplier
from ..schemas import RevisionIn
from ..security import Scope, audit, get_scope
from ..utils import bad_request, get_or_404, scoped

router = APIRouter(prefix="/api/preventivo", tags=["preventivo"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _abiertas():
    from .mantenimiento import ABIERTAS
    return ABIERTAS


def _plan_out(db: Session, p: PreventivePlan, hoy: date) -> dict:
    d = p.to_dict()
    d["tipo_nombre"] = preventivo.TIPOS_PLAN.get(p.tipo, p.tipo)
    d["activo_nombre"] = db.scalar(select(Asset.nombre).where(Asset.id == p.asset_id))
    d["proveedor_nombre"] = db.scalar(select(Supplier.nombre).where(Supplier.id == p.proveedor_id)) \
        if p.proveedor_id else p.proveedor
    d["puntos"] = preventivo.checklist_de(p)
    d["resultados"] = [[k, preventivo.RESULTADOS[k][0], preventivo.RESULTADOS[k][1]]
                       for k in preventivo.RESULTADOS_TIPO[p.tipo]]
    d.update(preventivo.estado(db, p, hoy, _abiertas()))
    return d


@router.get("/control")
def control(asset_id: int | None = None, tipo: str | None = None, scope: Scope = Depends(get_scope),
            db: Session = Depends(get_db)):
    """Todos los planes con su semáforo y los totales para el cuadro de control (lo urgente primero)."""
    ids = scope.asset_ids("mantenimiento.ver")
    q = scoped(select(PreventivePlan), PreventivePlan.asset_id, ids)
    if asset_id:
        scope.require_asset("mantenimiento.ver", asset_id)
        q = q.where(PreventivePlan.asset_id == asset_id)
    if tipo in preventivo.TIPOS_PLAN:
        q = q.where(PreventivePlan.tipo == tipo)
    hoy = date.today()
    planes = sorted((_plan_out(db, p, hoy) for p in db.scalars(q)), key=lambda x: (x["orden"], x["vence"]))
    activos = [p for p in planes if p["activo"]]
    cuenta = lambda f: sum(1 for p in activos if f(p))  # noqa: E731
    totales = {
        "planes": len(activos),
        "oca": cuenta(lambda p: p["tipo"] == "oca"),
        "vencidos": cuenta(lambda p: p["estado"] == "vencido"),
        "proximos": cuenta(lambda p: p["estado"] == "proximo"),
        "deficiencias": sum(len(p["deficiencias_abiertas"]) for p in activos),
        "subsanacion_vencida": cuenta(lambda p: p["estado"] == "subsanacion_vencida"),
        "sin_documento": cuenta(lambda p: any("documento" in a for a in p["avisos"])),
        "sin_registro": cuenta(lambda p: p["ultima"] is None),
        "al_dia": cuenta(lambda p: p["estado"] == "al_dia"),
    }
    return {"planes": planes, "totales": totales, "gravedades": preventivo.GRAVEDADES,
            "tipos": preventivo.TIPOS_PLAN}


def _plan(db: Session, scope: Scope, pid: int, perm: str) -> PreventivePlan:
    p = get_or_404(db, PreventivePlan, pid)
    scope.require_asset(perm, p.asset_id)
    return p


@router.get("/planes/{pid}/revisiones")
def history(pid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    p = _plan(db, scope, pid, "mantenimiento.ver")
    revs = db.scalars(select(PreventiveRevision).where(PreventiveRevision.plan_id == p.id)
                      .order_by(PreventiveRevision.fecha.desc(), PreventiveRevision.id.desc()))
    return {"plan": _plan_out(db, p, date.today()), "revisiones": [preventivo.revision_out(db, r) for r in revs]}


def _guardar_documentos(db: Session, r: PreventiveRevision, ficheros: list[UploadFile], user_id: int) -> list[int]:
    if len(ficheros) > 10:
        bad_request("Máximo 10 documentos por envío")
    nuevos = []
    for f in ficheros:
        datos = f.file.read(adjuntos.TAM_MAX + 1)
        if not datos:
            continue
        try:
            datos, mime, nombre = adjuntos.normalizar(datos, (f.filename or "documento")[-150:])
        except ValueError as e:
            bad_request(f"{f.filename}: {e}")
        a = RevisionAttachment(revision_id=r.id, nombre=nombre, fichero=documentos.guardar(datos), mime=mime,
                               tamano=len(datos), sha256=documentos.huella(datos), user_id=user_id)
        db.add(a)
        db.flush()
        nuevos.append(a.id)
    return nuevos


@router.post("/planes/{pid}/revisiones", status_code=201)
def register_revision(pid: int, tareas: BackgroundTasks, datos: str = Form(...),
                      ficheros: list[UploadFile] = File(default=[]), scope: Scope = Depends(get_scope),
                      db: Session = Depends(get_db)):
    """Registra una revisión (datos en JSON + documentos de la empresa): genera las OT de las deficiencias, cierra
    la OT preventiva abierta y fija la próxima fecha del plan."""
    p = _plan(db, scope, pid, "mantenimiento.editar")
    try:
        d = RevisionIn.model_validate(json.loads(datos))
    except (ValueError, ValidationError) as e:
        raise HTTPException(422, f"Datos de la revisión no válidos: {e}")
    if d.proveedor_id and not db.get(Supplier, d.proveedor_id):
        bad_request("Empresa no encontrada")
    try:
        r, ots = preventivo.registrar(db, p, scope.user, d.model_dump())
    except ValueError as e:
        db.rollback()
        bad_request(str(e))
    docs = _guardar_documentos(db, r, ficheros, scope.user.id)
    audit(db, scope.user, "registrar_revision", "plan_preventivo", p.id,
          {"revision": r.id, "resultado": r.resultado, "numero": r.numero, "documentos": len(docs),
           "ot_deficiencias": [w.id for w in ots], "ot_preventiva": r.work_order_id})
    db.commit()
    for w in ots:
        if w.prioridad == "urgente":
            tareas.add_task(avisos.ot_urgente, w.id, scope.user.id)
    out = preventivo.revision_out(db, r)
    out["aviso"] = preventivo.desfase(p, r.fecha, r.proxima)
    out["ots"] = [{"id": w.id, "titulo": w.titulo, "prioridad": w.prioridad} for w in ots]
    return out


@router.post("/revisiones/{rid}/documentos", status_code=201)
def add_documents(rid: int, ficheros: list[UploadFile] = File(...), scope: Scope = Depends(get_scope),
                  db: Session = Depends(get_db)):
    """Documentos que llegan después (el certificado de la OCA suele tardar unos días)."""
    r = get_or_404(db, PreventiveRevision, rid)
    scope.require_asset("mantenimiento.editar", r.asset_id)
    docs = _guardar_documentos(db, r, ficheros, scope.user.id)
    audit(db, scope.user, "adjuntar_revision", "plan_preventivo", r.plan_id, {"revision": r.id, "documentos": docs})
    db.commit()
    return preventivo.revision_out(db, r)


@router.get("/documentos/{aid}")
def view_document(aid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    a = get_or_404(db, RevisionAttachment, aid)
    r = db.get(PreventiveRevision, a.revision_id)
    scope.require_asset("mantenimiento.ver", r.asset_id)
    return Response(documentos.leer(a.fichero), media_type=a.mime, headers={
        "Content-Disposition": f'inline; filename="{a.nombre}"', "Cache-Control": "private, max-age=3600"})


@router.delete("/documentos/{aid}")
def delete_document(aid: int, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    a = get_or_404(db, RevisionAttachment, aid)
    r = db.get(PreventiveRevision, a.revision_id)
    scope.require_asset("mantenimiento.editar", r.asset_id)
    documentos.borrar(a.fichero)
    audit(db, scope.user, "borrar_documento_revision", "plan_preventivo", r.plan_id, {"revision": r.id, "nombre": a.nombre})
    db.delete(a)
    db.commit()
    return {"ok": True}


@router.get("/excel")
def register_excel(asset_id: int | None = None, scope: Scope = Depends(get_scope), db: Session = Depends(get_db)):
    """Registro de revisiones e inspecciones: una hoja con el estado de cada plan y otra con todo el historial."""
    datos = control(asset_id, None, scope, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Control"
    cab = ["Activo", "Tipo", "Plan", "Instalación", "Normativa", "Cada (días)", "Empresa", "Estado", "Vence",
           "Última revisión", "Resultado", "Nº acta / informe", "Documentos", "Deficiencias abiertas", "Avisos"]
    ws.append(cab)
    for p in datos["planes"]:
        u = p["ultima"] or {}
        ws.append([p["activo_nombre"], p["tipo_nombre"], p["titulo"], p["categoria"], p["normativa"],
                   p["periodicidad_dias"], p["proveedor_nombre"], p["estado_nombre"], date.fromisoformat(p["vence"]),
                   date.fromisoformat(u["fecha"]) if u else None, u.get("resultado_nombre"), u.get("numero"),
                   len(u.get("documentos") or []), len(p["deficiencias_abiertas"]), " · ".join(p["avisos"])])
    hist = wb.create_sheet("Historial")
    hist.append(["Activo", "Plan", "Fecha", "Prevista", "Resultado", "Empresa", "Técnico / inspector", "Nº acta",
                 "Próxima", "Plazo subsanación", "Deficiencias", "Documentos"])
    ids = {p["id"] for p in datos["planes"]}
    nombres = {p["id"]: (p["activo_nombre"], p["titulo"]) for p in datos["planes"]}
    for r in db.scalars(select(PreventiveRevision).where(PreventiveRevision.plan_id.in_(ids or {-1}))
                        .order_by(PreventiveRevision.fecha.desc())):
        o = preventivo.revision_out(db, r, breve=True)
        hist.append([*nombres[r.plan_id], r.fecha, r.fecha_prevista, o["resultado_nombre"], r.empresa, r.tecnico,
                     r.numero, r.proxima, r.plazo_subsanacion,
                     "; ".join(f"{x['descripcion']} ({preventivo.GRAVEDADES.get(x['gravedad'], '')}, OT-{x.get('ot_id')})"
                               for x in r.deficiencias or []), len(o["documentos"])])
    for hoja in (ws, hist):
        for c in hoja[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="13294B")
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = 16
            for c in col[1:]:
                if isinstance(c.value, date):
                    c.number_format = "DD/MM/YYYY"
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="control_preventivo_{date.today():%Y%m%d}.xlsx"'})
