"""Informe mensual para la presidencia: producción del mes (alquileres por tipo de estancia y servicios), gastos por
proveedor (OPEX/CAPEX), resultado y su porcentaje; y el resumen mensual del año anterior y del año en curso.
Todo sin IVA. La producción es lo facturado en el mes (fecha de la factura), como en el resto de informes."""
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import marca
from .facturacion import bases_por_tipo, lineas_de
from .models import Asset, Expense, Invoice, Reservation, Unit
from .routers.historico import mensual as historico_mensual
from .routers.informes import _facturado, _noches

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]
# Estancias por su duración total (noches), y las renovaciones mensuales aparte
TIPOS_ESTANCIA = [("corta", "Contratos inferiores a 1 semana"), ("semana", "Contratos de 1 semana"),
                  ("dos_semanas", "Contratos de 2 semanas"), ("mes", "Contratos de 1 mes"),
                  ("renovacion", "Renovaciones mensuales")]
ESTADOS_VALIDOS = ("confirmada", "checkin", "checkout")
SEMAFORO = {"rojo": "#c62828", "naranja": "#ef7d00", "verde": "#2e7d32"}


def tipo_estancia(r: Reservation) -> str:
    if r.renueva_id:
        return "renovacion"
    n = (r.fecha_salida - r.fecha_entrada).days
    return "corta" if n < 7 else "semana" if n < 14 else "dos_semanas" if n < 28 else "mes"


def semaforo(pct: float | None) -> str:
    """0 % a 12 % (o negativo) rojo; 13 % a 25 % naranja; 26 % o más verde."""
    if pct is None or pct < 12.5:
        return "rojo"
    return "naranja" if pct < 25.5 else "verde"


def fin_mes(anio: int, mes: int) -> date:
    return (date(anio, mes, 28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def _clave_prov(s: str | None) -> str:
    s = unicodedata.normalize("NFKD", (s or "").lower()).encode("ascii", "ignore").decode()
    s = re.sub(r"\b(s\.?\s?l\.?u?|s\.?\s?a\.?u?)\b", " ", re.sub(r"[^a-z0-9. ]", " ", s))
    return re.sub(r"[^a-z0-9]", "", s)


def _r(x) -> float:
    return round(float(x or 0), 2)


# --------------------------------------------------------------------------- datos
def _apartamentos(db: Session, asset_id: int) -> dict[int, str]:
    return dict(db.execute(select(Unit.id, Unit.uso).where(Unit.asset_id == asset_id)).all())


def _reservas(db: Session, asset_id: int, desde: date, hasta: date) -> list[Reservation]:
    return list(db.scalars(select(Reservation).join(Unit, Unit.id == Reservation.unit_id).where(
        Unit.asset_id == asset_id, Reservation.estado.in_(ESTADOS_VALIDOS),
        Reservation.fecha_entrada <= hasta, Reservation.fecha_salida > desde)))


def produccion_mes(db: Session, a: Asset, anio: int, mes: int) -> dict:
    """Página 1: ingresos del mes por tipo de estancia y por servicio, gastos por proveedor y resultado."""
    ini, fin = date(anio, mes, 1), fin_mes(anio, mes)
    usos = _apartamentos(db, a.id)
    reservas = {r.id: r for r in _reservas(db, a.id, ini, fin)}
    # pernoctaciones y contratos del mes (apartamentos; las plazas de garaje no son pernoctaciones)
    filas = {k: {"clave": k, "concepto": t, "contratos": 0, "pernoctaciones": 0, "importe": 0.0}
             for k, t in TIPOS_ESTANCIA}
    for r in reservas.values():
        if usos.get(r.unit_id) == "garaje":
            continue
        f = filas[tipo_estancia(r)]
        f["contratos"] += 1
        f["pernoctaciones"] += _noches(r.fecha_entrada, r.fecha_salida, ini, fin)
    # importes: lo facturado en el mes, sin IVA, según la estancia de cada factura
    facturas = list(db.scalars(select(Invoice).where(Invoice.asset_id == a.id, Invoice.fecha_expedicion >= ini,
                                                     Invoice.fecha_expedicion <= fin)))
    faltan = {f.reservation_id for f in facturas if f.reservation_id and f.reservation_id not in reservas}
    if faltan:
        reservas.update({r.id: r for r in db.scalars(select(Reservation).where(Reservation.id.in_(faltan)))})
    otros_aloj = rentas = garajes = 0.0
    servicios: dict[str, float] = defaultdict(float)
    for f, bases in bases_por_tipo(db, facturas):
        r = reservas.get(f.reservation_id)
        for t, b in bases.items():
            if t == "alojamiento":
                if r is not None and usos.get(r.unit_id) == "garaje":
                    garajes += b
                elif r is not None:
                    filas[tipo_estancia(r)]["importe"] += b
                else:
                    otros_aloj += b
            elif t == "renta":
                rentas += b
            elif t == "garaje":
                garajes += b
        for x in lineas_de(f):
            if x["tipo"] not in ("alojamiento", "renta", "garaje"):
                servicios[" ".join((x.get("concepto") or "Servicio").split())[:80]] += x["base"]
    alquileres = [{**f, "importe": _r(f["importe"])} for f in filas.values()]
    ext = historico_mensual(db, [a.id], ini, fin).get((a.id, f"{ini:%Y-%m}"), {})
    if otros_aloj:
        alquileres.append({"clave": "otros", "concepto": "Otros alojamientos facturados", "contratos": None,
                           "pernoctaciones": None, "importe": _r(otros_aloj)})
    if ext.get("alojamiento"):
        alquileres.append({"clave": "syade", "concepto": "Alojamiento facturado en el programa anterior (SYADE)",
                           "contratos": None, "pernoctaciones": None, "importe": _r(ext["alojamiento"])})
    if rentas:
        alquileres.append({"clave": "rentas", "concepto": "Rentas de contratos de alquiler", "contratos": None,
                           "pernoctaciones": None, "importe": _r(rentas)})
    serv = [{"concepto": k, "importe": _r(v)} for k, v in sorted(servicios.items(), key=lambda x: -x[1]) if _r(v)]
    if garajes:
        serv.append({"concepto": "Plazas de garaje", "importe": _r(garajes)})
    if ext.get("servicio"):
        serv.append({"concepto": "Servicios facturados en el programa anterior (SYADE)", "importe": _r(ext["servicio"])})
    total_alq = _r(sum(f["importe"] for f in alquileres))
    total_serv = _r(sum(s["importe"] for s in serv))
    ingresos = _r(total_alq + total_serv)

    # gastos del mes (fecha de la factura recibida), sin IVA, un renglón por proveedor
    grupos: dict[str, dict] = {}
    for g in db.scalars(select(Expense).where(Expense.asset_id == a.id, Expense.fecha >= ini, Expense.fecha <= fin)
                        .order_by(Expense.fecha, Expense.id)):
        k = f"s{g.supplier_id}" if g.supplier_id else _clave_prov(g.proveedor) or f"g{g.id}"
        x = grupos.setdefault(k, {"proveedor": g.proveedor or g.concepto, "facturas": [], "importe": 0.0,
                                  "naturalezas": set()})
        x["facturas"].append(g.numero_factura or "s/n")
        x["importe"] += float(g.base)
        x["naturalezas"].add(g.naturaleza or "sin indicar")
    gastos = [{"proveedor": x["proveedor"], "facturas": x["facturas"], "importe": _r(x["importe"]),
               "naturaleza": next(iter(x["naturalezas"])) if len(x["naturalezas"]) == 1 else "Varios"}
              for x in sorted(grupos.values(), key=lambda x: -x["importe"])]
    total_gastos = _r(sum(x["importe"] for x in gastos))
    resultado = _r(ingresos - total_gastos)
    pct = round(resultado / ingresos * 100, 1) if ingresos else None
    return {"alquileres": alquileres, "total_pernoctaciones": sum(f["pernoctaciones"] or 0 for f in alquileres),
            "total_alquileres": total_alq, "servicios": serv, "total_servicios": total_serv,
            "total_ingresos": ingresos, "gastos": gastos, "total_gastos": total_gastos, "resultado": resultado,
            "porcentaje": pct, "semaforo": semaforo(pct)}


def _var(actual, anterior) -> float | None:
    if actual is None or anterior in (None, 0):
        return None
    return round((actual - anterior) / abs(anterior) * 100, 1)


def resumen_anual(db: Session, a: Asset, anio: int, hasta_mes: int = 12) -> dict:
    """Por mes: producción (facturado sin IVA), pernoctaciones y % de ocupación, con la variación respecto al mes
    anterior y al mismo mes del año anterior."""
    desde, hasta = date(anio - 1, 1, 1), fin_mes(anio, 12)
    fact = _facturado(db, [a.id], desde, hasta)
    usos = _apartamentos(db, a.id)
    aloj = {u for u, uso in usos.items() if uso != "garaje"}
    activos = len(aloj)
    res = [r for r in _reservas(db, a.id, desde, hasta) if r.unit_id in aloj]

    def mes(y, m):
        ini, fin = date(y, m, 1), fin_mes(y, m)
        v = fact.get((a.id, f"{ini:%Y-%m}"), {})
        noches = sum(_noches(r.fecha_entrada, r.fecha_salida, ini, fin) for r in res)
        hay_datos = noches or not v.get("externo")  # meses solo con datos de SYADE: sin pernoctaciones
        disp = activos * fin.day
        return {"produccion": _r(v.get("base", 0)), "pernoctaciones": noches if hay_datos else None,
                "ocupacion": round(noches / disp * 100, 1) if hay_datos and disp else None}

    out = []
    for m in range(1, hasta_mes + 1):
        x, previo = mes(anio, m), mes(anio - 1, m)
        ant = out[-1] if out else mes(anio - 1, 12)
        out.append({"mes": m, "nombre": MESES[m - 1], **x,
                    "var_mes_produccion": _var(x["produccion"], ant["produccion"]),
                    "var_mes_pernoctaciones": _var(x["pernoctaciones"], ant["pernoctaciones"]),
                    "var_anio_produccion": _var(x["produccion"], previo["produccion"]),
                    "var_anio_pernoctaciones": _var(x["pernoctaciones"], previo["pernoctaciones"])})
    total_p = _r(sum(x["produccion"] for x in out))
    noches = [x["pernoctaciones"] for x in out if x["pernoctaciones"] is not None]
    return {"anio": anio, "meses": out, "total_produccion": total_p, "total_pernoctaciones": sum(noches),
            "apartamentos": activos}


def datos(db: Session, a: Asset, anio: int, mes: int) -> dict:
    return {"activo": {"id": a.id, "codigo": a.codigo, "nombre": a.nombre, "sociedad": a.company.nombre,
                       "cif": a.company.cif},
            "anio": anio, "mes": mes, "periodo": f"{MESES[mes - 1]} {anio}",
            **produccion_mes(db, a, anio, mes),
            "anterior": resumen_anual(db, a, anio - 1), "actual": resumen_anual(db, a, anio, mes)}


# --------------------------------------------------------------------------- formatos
def eur(x) -> str:
    return "—" if x is None else f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def num(x) -> str:
    return "—" if x is None else f"{x:,}".replace(",", ".")


def pct(x, signo: bool = False) -> str:
    if x is None:
        return "—"
    return (f"{x:+.1f} %" if signo else f"{x:.1f} %").replace(".", ",")


def pdf(d: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base = getSampleStyleSheet()
    oro, negro, linea = colors.HexColor(marca.ORO), colors.HexColor(marca.NEGRO), colors.HexColor("#e2d9c4")
    st_t = ParagraphStyle("t", parent=base["Title"], fontSize=17, textColor=negro, spaceAfter=2)
    st_s = ParagraphStyle("s", parent=base["Normal"], fontSize=10, textColor=colors.HexColor("#6d6658"))
    st_h = ParagraphStyle("h", parent=base["Heading3"], fontSize=11.5, textColor=oro, spaceBefore=8, spaceAfter=4)
    st_c = ParagraphStyle("c", parent=base["Normal"], fontSize=8.5, leading=10.5)
    ancho = 174 * mm
    a = d["activo"]
    logos = marca.cabecera_pdf(marca.clave_sociedad(a.get("cif")), marca.clave_activo(a["codigo"]), ancho, 16 * mm)

    st_th = ParagraphStyle("th", parent=st_c, fontName="Helvetica-Bold", fontSize=7.5, leading=9,
                           textColor=colors.HexColor("#6d5310"), alignment=2)

    def tabla(filas, cols, total_filas=(), color_fila=None, cab=True, izquierda=(0,)):
        t = Table(filas, colWidths=[ancho * c for c in cols], repeatRows=1 if cab else 0)
        estilo = [("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                  ("LINEBELOW", (0, 0), (-1, -1), 0.4, linea), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                  *[("ALIGN", (c, 0), (c, -1), "LEFT") for c in izquierda],
                  ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if cab:
            estilo += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f6f1e6")), ("FONTNAME", (0, 0), (-1, 0),
                                                                                     "Helvetica-Bold"),
                       ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6d5310"))]
        for i in total_filas:
            estilo += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"), ("BACKGROUND", (0, i), (-1, i),
                                                                         colors.HexColor("#faf8f3"))]
        if color_fila:
            estilo += color_fila
        t.setStyle(TableStyle(estilo))
        return t

    def cabecera(titulo):
        out = [logos, Spacer(1, 4 * mm)] if logos else []
        return out + [Paragraph(f"{titulo} · {a['nombre']}", st_t),
                      Paragraph(f"Informe a la presidencia · {d['periodo'].capitalize()} · {a['sociedad']} · "
                                "importes sin IVA", st_s), Spacer(1, 4 * mm)]

    cuerpo = cabecera("Producción del mes")
    # ingresos por alquileres
    cuerpo.append(Paragraph("Ingresos por alquileres", st_h))
    filas = [["Concepto", "Contratos", "Pernoctaciones", "Importe"]]
    filas += [[Paragraph(f["concepto"], st_c), num(f["contratos"]), num(f["pernoctaciones"]), eur(f["importe"])]
              for f in d["alquileres"]]
    filas.append(["Total alquileres", "", num(d["total_pernoctaciones"]), eur(d["total_alquileres"])])
    cuerpo.append(tabla(filas, (0.52, 0.14, 0.16, 0.18), total_filas=(len(filas) - 1,)))
    # servicios
    cuerpo.append(Paragraph("Ingresos por servicios", st_h))
    filas = [["Servicio", "Importe"]] + [[Paragraph(s["concepto"], st_c), eur(s["importe"])] for s in d["servicios"]]
    if len(filas) == 1:
        filas.append(["Sin servicios facturados en el mes", eur(0)])
    filas.append(["Total servicios", eur(d["total_servicios"])])
    cuerpo.append(tabla(filas, (0.82, 0.18), total_filas=(len(filas) - 1,)))
    cuerpo.append(Spacer(1, 3 * mm))
    cuerpo.append(tabla([[f"TOTAL PRODUCCIÓN · {d['periodo'].upper()}", eur(d["total_ingresos"])]], (0.82, 0.18),
                        total_filas=(0,), cab=False, color_fila=[("FONTSIZE", (0, 0), (-1, -1), 10),
                                                                 ("TEXTCOLOR", (0, 0), (-1, -1), negro)]))
    # gastos
    cuerpo.append(Paragraph("Gastos del mes", st_h))
    filas = [["Proveedor", "Nº de factura", "Importe sin IVA", "CAPEX / OPEX"]]
    filas += [[Paragraph(g["proveedor"] or "", st_c), Paragraph(", ".join(g["facturas"]), st_c), eur(g["importe"]),
               g["naturaleza"]] for g in d["gastos"]]
    if len(filas) == 1:
        filas.append(["Sin gastos registrados en el mes", "", eur(0), ""])
    filas.append(["Total gastos", "", eur(d["total_gastos"]), ""])
    cuerpo.append(tabla(filas, (0.38, 0.30, 0.17, 0.15), total_filas=(len(filas) - 1,), izquierda=(0, 1)))
    # resultado
    color = colors.HexColor(SEMAFORO[d["semaforo"]])
    cuerpo.append(Spacer(1, 3 * mm))
    cuerpo.append(tabla([["Total producción (ingresos)", eur(d["total_ingresos"])],
                         ["Total gastos", eur(-d["total_gastos"])],
                         ["RESULTADO DE PRODUCCIÓN", eur(d["resultado"])],
                         ["Porcentaje de resultado sobre la producción", pct(d["porcentaje"])]],
                        (0.82, 0.18), total_filas=(2,), cab=False,
                        color_fila=[("BACKGROUND", (0, 3), (-1, 3), color), ("TEXTCOLOR", (0, 3), (-1, 3), colors.white),
                                    ("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"), ("FONTSIZE", (0, 2), (-1, 3), 10)]))
    cuerpo.append(Paragraph("Semáforo: rojo de 0 % a 12 % · naranja de 13 % a 25 % · verde de 26 % en adelante.", st_s))

    for clave in ("anterior", "actual"):
        r = d[clave]
        cuerpo += [PageBreak()] + cabecera(f"Resumen mensual {r['anio']}")
        filas = [["Mes"] + [Paragraph(t, st_th) for t in (
            "Producción", "Pernoct.", "Ocupación", "Δ producción<br/>vs mes anterior",
            "Δ pernoct.<br/>vs mes anterior", "Δ producción<br/>vs año anterior", "Δ pernoct.<br/>vs año anterior")]]
        estilo = []
        for i, m in enumerate(r["meses"], start=1):
            filas.append([m["nombre"].capitalize(), eur(m["produccion"]), num(m["pernoctaciones"]),
                          pct(m["ocupacion"]), pct(m["var_mes_produccion"], True),
                          pct(m["var_mes_pernoctaciones"], True), pct(m["var_anio_produccion"], True),
                          pct(m["var_anio_pernoctaciones"], True)])
            for j, k in enumerate(("var_mes_produccion", "var_mes_pernoctaciones", "var_anio_produccion",
                                   "var_anio_pernoctaciones"), start=4):
                if m[k] is not None:
                    estilo.append(("TEXTCOLOR", (j, i), (j, i), colors.HexColor("#2e7d32" if m[k] >= 0 else "#c62828")))
        filas.append(["Total", eur(r["total_produccion"]), num(r["total_pernoctaciones"]), "", "", "", "", ""])
        cuerpo.append(tabla(filas, (0.12, 0.15, 0.10, 0.11, 0.13, 0.13, 0.13, 0.13), total_filas=(len(filas) - 1,),
                            color_fila=estilo))
        cuerpo.append(Spacer(1, 3 * mm))
        cuerpo.append(Paragraph(
            f"Producción: facturado en el mes sin IVA (incluye lo importado del programa anterior). Pernoctaciones y "
            f"ocupación: noches de reservas en {r['apartamentos']} apartamentos (sin plazas de garaje). Δ: variación "
            "respecto al mes anterior y al mismo mes del año anterior. «—»: sin datos.", st_s))

    out = io.BytesIO()

    def pie(c, doc):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.grey)
        c.drawString(18 * mm, 10 * mm, f"INVERPMS · Informe a la presidencia · {a['nombre']} · {d['periodo']}")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Página {doc.page}")
        c.restoreState()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm,
                      bottomMargin=16 * mm, title=f"Informe a la presidencia {a['nombre']} {d['periodo']}",
                      author="INVERPMS").build(cuerpo, onFirstPage=pie, onLaterPages=pie)
    return out.getvalue()


def excel(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Producción"
    neg = Font(bold=True)
    euro = '#,##0.00 "€"'
    ws.append([f"Informe a la presidencia · {d['activo']['nombre']} · {d['periodo']} · sin IVA"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Ingresos por alquileres", "Contratos", "Pernoctaciones", "Importe"])
    for f in d["alquileres"]:
        ws.append([f["concepto"], f["contratos"], f["pernoctaciones"], f["importe"]])
    ws.append(["Total alquileres", None, d["total_pernoctaciones"], d["total_alquileres"]])
    ws.append([])
    ws.append(["Ingresos por servicios", None, None, "Importe"])
    for s in d["servicios"]:
        ws.append([s["concepto"], None, None, s["importe"]])
    ws.append(["Total servicios", None, None, d["total_servicios"]])
    ws.append([f"TOTAL PRODUCCIÓN {d['periodo'].upper()}", None, None, d["total_ingresos"]])
    ws.append([])
    ws.append(["Gastos: proveedor", "Nº de factura", "CAPEX / OPEX", "Importe sin IVA"])
    for g in d["gastos"]:
        ws.append([g["proveedor"], ", ".join(g["facturas"]), g["naturaleza"], g["importe"]])
    ws.append(["Total gastos", None, None, d["total_gastos"]])
    ws.append([])
    ws.append(["RESULTADO DE PRODUCCIÓN", None, None, d["resultado"]])
    ws.append(["Porcentaje de resultado", None, None, None if d["porcentaje"] is None else d["porcentaje"] / 100])
    ws.cell(ws.max_row, 4).number_format = "0.0%"
    ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=SEMAFORO[d["semaforo"]][1:])
    ws.cell(ws.max_row, 4).font = Font(bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=3):
        if row[0].value and str(row[0].value).startswith(("Total", "TOTAL", "RESULTADO", "Ingresos", "Gastos")):
            for c in row:
                c.font = neg
        if row[3].number_format != "0.0%" and isinstance(row[3].value, (int, float)):
            row[3].number_format = euro
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    for clave in ("anterior", "actual"):
        r = d[clave]
        h = wb.create_sheet(f"Resumen {r['anio']}")
        h.append(["Mes", "Producción", "Pernoctaciones", "Ocupación", "Δ mes anterior producción",
                  "Δ mes anterior pernoctaciones", "Δ año anterior producción", "Δ año anterior pernoctaciones"])
        for c in h[1]:
            c.font = neg
        for m in r["meses"]:
            h.append([m["nombre"], m["produccion"], m["pernoctaciones"],
                      None if m["ocupacion"] is None else m["ocupacion"] / 100,
                      *[None if m[k] is None else m[k] / 100 for k in ("var_mes_produccion", "var_mes_pernoctaciones",
                                                                       "var_anio_produccion",
                                                                       "var_anio_pernoctaciones")]])
        h.append(["Total", r["total_produccion"], r["total_pernoctaciones"]])
        for row in h.iter_rows(min_row=2):
            row[1].number_format = euro
            for c in row[3:]:
                c.number_format = "+0.0%;-0.0%;0.0%" if c.column > 4 else "0.0%"
        for col, w in zip("ABCDEFGH", (14, 16, 16, 12, 22, 24, 22, 24)):
            h.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
