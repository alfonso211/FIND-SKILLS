"""Firma del contrato de alojamiento en tablet (firma electrónica simple, art. 3.10 del Reglamento eIDAS).

1. El contrato relleno (.docx) se convierte a PDF y se enseña al cliente en la tablet.
2. El cliente lo lee, acepta expresamente las condiciones destacadas y la información de protección de datos y
   firma con el dedo o un puntero.
3. Se incrusta la firma en el contrato, se añade una página de evidencias (fecha y hora, dispositivo, IP,
   empleado, huellas SHA-256 del documento sin firmar y de la firma) y se calcula la huella del PDF final.
4. El PDF firmado se guarda cifrado (no hay copia en papel) y se envía al cliente por correo o WhatsApp.
"""
import base64
import io
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm

from . import marca
from .contratos import _parrafos

MADRID = ZoneInfo("Europe/Madrid")
VALIDEZ_ENLACE_DIAS = 7


def disponible() -> bool:
    return shutil.which("soffice") is not None or shutil.which("libreoffice") is not None


def a_pdf(docx_bytes: bytes) -> bytes:
    """Convierte el .docx a PDF con LibreOffice (en el propio servidor)."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("El servidor no tiene instalado el conversor a PDF (LibreOffice)")
    with tempfile.TemporaryDirectory(prefix="pms_pdf_") as tmp:
        d = Path(tmp)
        (d / "contrato.docx").write_bytes(docx_bytes)
        r = subprocess.run([exe, f"-env:UserInstallation=file://{d}/perfil", "--headless", "--norestore",
                            "--convert-to", "pdf", "--outdir", str(d), str(d / "contrato.docx")],
                           capture_output=True, timeout=120)
        pdf = d / "contrato.pdf"
        if r.returncode != 0 or not pdf.exists():
            raise RuntimeError("No se ha podido convertir el contrato a PDF")
        return pdf.read_bytes()


def imagen_firma(data_url: str) -> bytes:
    """PNG de la firma (data URL del lienzo de la tablet), validado y recortado al trazo."""
    from PIL import Image, ImageOps
    if not data_url.startswith("data:image/png;base64,"):
        raise ValueError("Firma no válida")
    try:
        img = Image.open(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1], validate=True)))
        img.load()
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Firma no válida") from exc
    if img.width * img.height > 4000 * 2000:
        raise ValueError("Firma demasiado grande")
    rgba = img.convert("RGBA")
    caja = rgba.getchannel("A").getbbox()
    if not caja or (caja[2] - caja[0]) < 40 or (caja[3] - caja[1]) < 15:
        raise ValueError("La firma está vacía o es demasiado pequeña: firme de nuevo")
    rgba = rgba.crop(caja)
    fondo = Image.new("RGBA", rgba.size, (255, 255, 255, 0))
    fondo.alpha_composite(rgba)
    out = io.BytesIO()
    ImageOps.contain(fondo, (900, 300)).save(out, "PNG")
    return out.getvalue()


def insertar_firma(docx_bytes: bytes, firma_png: bytes) -> bytes:
    """Coloca la firma del cliente sobre su línea «Fdo.:» (no en la de la empresa, «P.p. Fdo.:»)."""
    d = docx.Document(io.BytesIO(docx_bytes))
    destino = next((p for p in _parrafos(d) if p.text.strip().startswith("Fdo.:")), None)
    if destino is None:
        raise RuntimeError("La plantilla no tiene la línea de firma del cliente («Fdo.:»)")
    anterior = destino._p.getprevious()
    vacio = None
    if anterior is not None and anterior.tag.endswith("}p") and not "".join(anterior.itertext()).strip():
        from docx.text.paragraph import Paragraph
        vacio = Paragraph(anterior, destino._parent)
    p = vacio or destino.insert_paragraph_before()
    # la plantilla usa interlineado fijo: con él la imagen se recortaría
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = p.paragraph_format.space_after = 0
    p.add_run().add_picture(io.BytesIO(firma_png), height=Cm(1.6))
    fila = next((a for a in p._p.iterancestors() if a.tag == qn("w:tr")), None)
    if fila is not None:  # una fila de alto exacto también la recortaría
        for alto in fila.iter(qn("w:trHeight")):
            alto.set(qn("w:hRule"), "atLeast")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def paginas_png(pdf: bytes, escala: float = 1.6) -> list[bytes]:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf)
    out = []
    for i in range(len(doc)):
        b = io.BytesIO()
        doc[i].render(scale=escala).to_pil().save(b, "PNG", optimize=True)
        out.append(b.getvalue())
    return out


def pagina_evidencias(ev: dict, firma_png: bytes, asset_codigo: str | None, cif: str | None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title="Evidencias de firma electrónica", creator=marca.NOMBRE)
    ss = getSampleStyleSheet()
    n = ParagraphStyle("n", parent=ss["Normal"], fontName="Helvetica", fontSize=9, leading=12)
    tit = ParagraphStyle("t", parent=n, fontName="Helvetica-Bold", fontSize=14, leading=18)
    peq = ParagraphStyle("p", parent=n, fontSize=7.5, leading=10, textColor=colors.HexColor("#6b6b6b"))
    esc = lambda s: str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    filas = [
        ("Documento", ev["documento"]), ("Firmante", f"{ev['firmante']} · {ev['documento_firmante'] or ''}"),
        ("Fecha y hora de la firma", ev["fecha_local"] + " (hora de Madrid) · " + ev["fecha_utc"] + " UTC"),
        ("Aceptación", "Ha leído el contrato completo y acepta expresamente las condiciones destacadas en negrita y "
                       "la información sobre protección de datos"),
        ("Dispositivo", ev["dispositivo"]), ("Dirección IP", ev["ip"]),
        ("Presentado por (recepción)", ev["empleado"]),
        ("Huella SHA-256 del contrato sin firmar", ev["sha256_contrato"]),
        ("Huella SHA-256 de la imagen de la firma", ev["sha256_firma"]),
    ]
    t = Table([[Paragraph(f"<b>{esc(k)}</b>", n), Paragraph(esc(v), n)] for k, v in filas],
              colWidths=[52 * mm, 122 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ddd6c8")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f7f3ea")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    firma = Image(io.BytesIO(firma_png))
    k = min(70 * mm / firma.imageWidth, 25 * mm / firma.imageHeight)
    firma.drawWidth, firma.drawHeight = firma.imageWidth * k, firma.imageHeight * k
    cuerpo = []
    logos = marca.cabecera_pdf(marca.clave_sociedad(cif), marca.clave_activo(asset_codigo), 174 * mm, 18 * mm)
    if logos:
        cuerpo += [logos, Spacer(1, 6 * mm)]
    cuerpo += [Paragraph("Evidencias de la firma electrónica", tit), Spacer(1, 4 * mm), t, Spacer(1, 6 * mm),
               Paragraph("<b>Firma manuscrita digitalizada del cliente</b>", n), Spacer(1, 2 * mm), firma,
               Spacer(1, 8 * mm),
               Paragraph("Firma electrónica simple (art. 3.10 del Reglamento (UE) 910/2014, eIDAS) recogida en "
                         "tableta en el establecimiento. El documento firmado se conserva cifrado y se ha puesto a "
                         "disposición del cliente por correo electrónico o WhatsApp. Cualquier alteración del PDF "
                         "cambia su huella SHA-256, que consta en el registro de " + marca.NOMBRE + ".", peq)]
    doc.build(cuerpo)
    return buf.getvalue()


def unir(*pdfs: bytes) -> bytes:
    import pypdfium2 as pdfium
    final = pdfium.PdfDocument.new()
    for p in pdfs:
        final.import_pages(pdfium.PdfDocument(p))
    out = io.BytesIO()
    final.save(out)
    return out.getvalue()


def ahora() -> tuple[datetime, str, str]:
    utc = datetime.now(ZoneInfo("UTC"))
    return utc.replace(tzinfo=None), utc.astimezone(MADRID).strftime("%d/%m/%Y %H:%M:%S"), utc.strftime(
        "%Y-%m-%d %H:%M:%S")


def movil_whatsapp(movil: str) -> str | None:
    """Número en formato internacional sin «+» (wa.me). Los móviles españoles de 9 cifras llevan el 34."""
    d = "".join(c for c in movil if c.isdigit())
    if movil.strip().startswith("00"):
        d = d[2:]
    if len(d) == 9 and d[0] in "67":
        d = "34" + d
    return d if 10 <= len(d) <= 15 else None
